"""Tests for scripts/orchestrator.py: verdict parsing, the Maker-Checker
governance loop, per-iteration checkpointing, deterministic policy
enforcement, and export gating.

run_pipeline() reads agents/*.md, config/style_guide.md, and
templates/cam/*.md relative to the current working directory (same as
deal_export.py / state_manager.py's own base_dir=None convention) -- the
`project_root` fixture below builds a minimal fake project in tmp_path and
chdirs into it so these tests never touch the real repository's deals/ or
templates/ directories.

Since _apply_deterministic_policy_checks() now forces REJECTED whenever a
draft's trailing structured JSON block doesn't cover every required CP and
risk category, every MockClient "draft" response that a test expects to be
approved and exported must include a *compliant* block -- see
_compliant_draft() below. Drafts used only to exercise an LLM-originated (or
deliberately triggered code-enforced) REJECTED path don't need one.
"""
import json
import os
from types import SimpleNamespace

import docx
import pytest

from orchestrator import (
    MAX_REVIEW_ITERATIONS,
    REQUIRED_RISK_TAXONOMY,
    _apply_deterministic_policy_checks,
    _check_reported_figures,
    _compute_collateral_cover_pct,
    _ground_truth_figures,
    _load_multi_period_financials,
    _normalize_category,
    _values_match,
    evaluate_financial_model,
    parse_underwriter_output,
    parse_verdict,
    run_pipeline,
)
from state_manager import read_state, state_path


class MockClient:
    """Stand-in for anthropic.Anthropic: returns each response in `responses`,
    in order, from successive .messages.create() calls."""

    def __init__(self, responses):
        self.responses = list(responses)
        self.call_count = 0
        self.messages = SimpleNamespace(create=self._create)

    def _create(self, **kwargs):
        self.call_count += 1
        if not self.responses:
            raise AssertionError("MockClient received more calls than responses were queued")
        text = self.responses.pop(0)
        return SimpleNamespace(content=[SimpleNamespace(text=text)])


@pytest.fixture
def project_root(tmp_path, monkeypatch):
    agents_dir = tmp_path / "agents"
    agents_dir.mkdir()
    (agents_dir / "underwriter_agent.md").write_text("MAKER PROMPT", encoding="utf-8")
    (agents_dir / "risk_reviewer_agent.md").write_text("CHECKER PROMPT", encoding="utf-8")

    cam_dir = tmp_path / "templates" / "cam"
    cam_dir.mkdir(parents=True)
    (cam_dir / "corporate_credit_cam.md").write_text("TEMPLATE", encoding="utf-8")

    monkeypatch.chdir(tmp_path)
    return tmp_path


def _approved_json(notes=None):
    return '```json\n' + json.dumps({"verdict": "APPROVED", "notes": notes}) + '\n```'


def _rejected_json(notes):
    return '```json\n' + json.dumps({"verdict": "REJECTED", "notes": notes}) + '\n```'


# A deal with no covenants/security_package/guarantees in state.json only
# ever requires the two standard CPs, and the Underwriter is always free to
# self-report every canonical risk category as covered.
STANDARD_CP_IDS = ["KYC-AML", "FACILITY-EXECUTION"]
ALL_CATEGORIES_COVERED = {category: {"status": "covered"} for category in REQUIRED_RISK_TAXONOMY}


def _compliant_draft(body="# Draft CAM", cp_ids=STANDARD_CP_IDS,
                      risk_categories=None, reported_figures=None):
    """A draft whose trailing structured JSON block satisfies every
    deterministic policy check by default (see agents/underwriter_agent.md's
    Structured Output guideline) -- for tests where the draft is expected to
    actually be approved and exported.

    reported_figures defaults to empty: per the Structured Output guideline,
    omitting a metric entirely (rather than reporting an unverifiable guess)
    is always compliant, and most tests below have no financial data behind
    them for a reported figure to be checked against anyway."""
    payload = {
        "cp_ids_included": list(cp_ids),
        "risk_categories_covered": risk_categories if risk_categories is not None else ALL_CATEGORIES_COVERED,
        "reported_figures": reported_figures if reported_figures is not None else {},
    }
    return body + "\n\n```json\n" + json.dumps(payload) + "\n```"


# ---------------------------------------------------------------------------
# parse_verdict()
# ---------------------------------------------------------------------------

def test_parse_verdict_extracts_approved():
    text = "Looks solid overall.\n" + _approved_json()
    verdict, notes = parse_verdict(text)
    assert verdict == "APPROVED"
    assert notes is None


def test_parse_verdict_extracts_rejected_with_notes():
    text = _rejected_json("DSCR in section 11 doesn't match the spreading workbook.")
    verdict, notes = parse_verdict(text)
    assert verdict == "REJECTED"
    assert notes == "DSCR in section 11 doesn't match the spreading workbook."


def test_parse_verdict_is_case_insensitive_on_verdict_value():
    text = '```json\n{"verdict": "approved", "notes": null}\n```'
    verdict, _ = parse_verdict(text)
    assert verdict == "APPROVED"


def test_parse_verdict_falls_back_to_rejected_when_json_block_missing():
    text = "The draft looks fine overall but I have some concerns."
    verdict, notes = parse_verdict(text)
    assert verdict == "REJECTED"
    assert notes == text


def test_parse_verdict_falls_back_to_rejected_on_malformed_json():
    text = '```json\n{"verdict": "APPROVED", "notes": \n```'  # truncated/invalid JSON
    verdict, notes = parse_verdict(text)
    assert verdict == "REJECTED"
    assert notes == text


def test_parse_verdict_falls_back_to_rejected_on_unrecognized_verdict_value():
    text = '```json\n{"verdict": "MAYBE", "notes": "unsure"}\n```'
    verdict, notes = parse_verdict(text)
    assert verdict == "REJECTED"
    assert notes == text


def test_parse_verdict_falls_back_to_rejected_on_empty_response():
    verdict, notes = parse_verdict("")
    assert verdict == "REJECTED"
    assert notes == ""


def test_parse_verdict_ignores_an_earlier_echoed_json_block_and_finds_the_trailing_verdict():
    """The grounding context hands the reviewer its own ```json blocks
    (financials/ratios/collateral) to check the draft against -- if it
    quotes one back while explaining its reasoning, that earlier block must
    not be mistaken for the real, trailing verdict block."""
    text = (
        "Here are the ratios I'm checking the draft against:\n"
        '```json\n{"dscr": 5.1, "gross_leverage": 0.68}\n```\n'
        "Everything in the draft reconciles with the above.\n"
        + _approved_json()
    )
    verdict, notes = parse_verdict(text)
    assert verdict == "APPROVED"
    assert notes is None


def test_parse_verdict_skips_multiple_non_verdict_blocks_to_find_the_real_one():
    text = (
        '```json\n{"financials": {"FY-Current": {"revenue": 1000}}}\n```\n'
        '```json\n{"ratios": {"FY-Current": {"dscr": 5.1}}}\n```\n'
        + _rejected_json("Collateral value doesn't match the supplied data.")
    )
    verdict, notes = parse_verdict(text)
    assert verdict == "REJECTED"
    assert notes == "Collateral value doesn't match the supplied data."


# ---------------------------------------------------------------------------
# parse_underwriter_output()
# ---------------------------------------------------------------------------

def test_parse_underwriter_output_extracts_cp_ids_and_categories():
    draft = _compliant_draft(cp_ids=["KYC-AML", "FACILITY-EXECUTION", "SEC-PERFECT-AST-001"])
    result = parse_underwriter_output(draft)
    assert result["cp_ids_included"] == ["KYC-AML", "FACILITY-EXECUTION", "SEC-PERFECT-AST-001"]
    assert result["risk_categories_covered"] == ALL_CATEGORIES_COVERED


def test_parse_underwriter_output_defaults_to_empty_when_block_missing():
    result = parse_underwriter_output("# Draft CAM with no trailing JSON block")
    assert result == {"cp_ids_included": [], "risk_categories_covered": {}, "reported_figures": {}}


def test_parse_underwriter_output_defaults_to_empty_on_malformed_json():
    result = parse_underwriter_output('# Draft\n```json\n{"cp_ids_included": [\n```')
    assert result == {"cp_ids_included": [], "risk_categories_covered": {}, "reported_figures": {}}


def test_parse_underwriter_output_extracts_reported_figures():
    draft = _compliant_draft(cp_ids=["KYC-AML"], reported_figures={"dscr": 1.45, "ebitda": 950000})
    result = parse_underwriter_output(draft)
    assert result["reported_figures"] == {"dscr": 1.45, "ebitda": 950000}


def test_parse_underwriter_output_degrades_safely_when_fields_have_the_wrong_type():
    """A field present but the wrong JSON type (e.g. a list instead of an
    object) must degrade to the safe empty default, not propagate a
    malformed value that later crashes _apply_deterministic_policy_checks()."""
    draft = (
        "# Draft\n\n```json\n"
        + json.dumps({
            "cp_ids_included": "KYC-AML",  # should be a list, not a bare string
            "risk_categories_covered": ["Market"],  # should be a dict, not a list
            "reported_figures": ["dscr", 1.5],  # should be a dict, not a list
        })
        + "\n```"
    )
    result = parse_underwriter_output(draft)
    assert result == {"cp_ids_included": [], "risk_categories_covered": {}, "reported_figures": {}}


def test_apply_deterministic_policy_checks_does_not_crash_on_malformed_risk_categories_type():
    draft = (
        "# Draft\n\n```json\n"
        + json.dumps({"cp_ids_included": [], "risk_categories_covered": ["Market"], "reported_figures": {}})
        + "\n```"
    )
    verdict, notes = _apply_deterministic_policy_checks(
        "APPROVED", None, draft, _policy_state(required_cps=[]), {},
    )
    assert verdict == "REJECTED"
    assert "Missing or malformed Risk Category" in notes


def test_parse_underwriter_output_ignores_unrelated_earlier_json_blocks():
    draft = (
        "# Draft CAM\n"
        '```json\n{"some_other_data": 123}\n```\n'
        + _compliant_draft(body="", cp_ids=["KYC-AML"]).strip()
    )
    result = parse_underwriter_output(draft)
    assert result["cp_ids_included"] == ["KYC-AML"]


# ---------------------------------------------------------------------------
# _normalize_category(): fixed string transform, not semantic judgment.
# ---------------------------------------------------------------------------

def test_normalize_category_strips_trailing_risk_suffix():
    assert _normalize_category("Market Risk") == _normalize_category("Market") == "market"


def test_normalize_category_is_case_and_whitespace_insensitive():
    assert _normalize_category("  KEY MAN  ") == _normalize_category("Key Man Risk") == "key man"


# ---------------------------------------------------------------------------
# _apply_deterministic_policy_checks(): the code-enforced overlay that can
# only ever move a verdict from APPROVED to REJECTED, never the reverse.
# ---------------------------------------------------------------------------

def _policy_state(required_cps=None, covenant_results=None, security_gaps=None):
    return {
        "required_conditions_precedent": required_cps or [{"cp_id": "KYC-AML", "text": "KYC/AML clearance."}],
        "covenant_results": covenant_results or [],
        "security_gaps": security_gaps or [],
    }


def test_apply_deterministic_policy_checks_passes_through_a_fully_compliant_approval():
    draft = _compliant_draft(cp_ids=["KYC-AML"])
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, _policy_state(), {})
    assert verdict == "APPROVED"
    assert notes is None


def test_apply_deterministic_policy_checks_overrides_approval_on_missing_cp():
    """A missing required cp_id must force REJECTED even when the LLM itself said APPROVED."""
    draft = _compliant_draft(cp_ids=[])  # KYC-AML required but not reported as included
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, _policy_state(), {})
    assert verdict == "REJECTED"
    assert "Missing Required CP KYC-AML" in notes


def test_apply_deterministic_policy_checks_overrides_on_malformed_not_applicable_category():
    """not_applicable with an empty justification is malformed and must trip the override."""
    categories = dict(ALL_CATEGORIES_COVERED)
    categories["Operational"] = {"status": "not_applicable", "justification": "   "}
    draft = _compliant_draft(cp_ids=["KYC-AML"], risk_categories=categories)
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, _policy_state(), {})
    assert verdict == "REJECTED"
    assert "Missing or malformed Risk Category: Operational" in notes


def test_apply_deterministic_policy_checks_accepts_normalized_category_key():
    """"Market Risk" as a key must satisfy the canonical "Market" requirement."""
    categories = dict(ALL_CATEGORIES_COVERED)
    del categories["Market"]
    categories["Market Risk"] = {"status": "covered"}
    draft = _compliant_draft(cp_ids=["KYC-AML"], risk_categories=categories)
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, _policy_state(), {})
    assert verdict == "APPROVED"
    assert notes is None


def test_apply_deterministic_policy_checks_accepts_case_insensitive_status_value():
    """A status of "Covered" (capitalized, plausible LLM drift) must be
    recognized the same as "covered" -- only the category *key* was
    normalized before, not the *status value* itself."""
    categories = dict(ALL_CATEGORIES_COVERED)
    categories["Market"] = {"status": "Covered"}
    draft = _compliant_draft(cp_ids=["KYC-AML"], risk_categories=categories)
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, _policy_state(), {})
    assert verdict == "APPROVED"
    assert notes is None


def test_apply_deterministic_policy_checks_overrides_on_covenant_failure():
    draft = _compliant_draft(cp_ids=["KYC-AML"])
    policy_state = _policy_state(covenant_results=[
        {"metric": "dscr", "type": "minimum", "threshold": 1.25, "actual": 1.0, "status": "FAIL", "headroom_pct": -0.2},
    ])
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, policy_state, {})
    assert verdict == "REJECTED"
    assert "Covenant FAIL: dscr" in notes


def test_apply_deterministic_policy_checks_overrides_on_unresolvable_covenant():
    draft = _compliant_draft(cp_ids=["KYC-AML"])
    policy_state = _policy_state(covenant_results=[
        {"metric": "made_up_metric", "type": "minimum", "threshold": 1.0, "actual": None,
         "status": "UNRESOLVABLE", "headroom_pct": None},
    ])
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, policy_state, {})
    assert verdict == "REJECTED"
    assert "Covenant UNRESOLVABLE: made_up_metric" in notes


def test_apply_deterministic_policy_checks_overrides_on_security_gap():
    draft = _compliant_draft(cp_ids=["KYC-AML"])
    policy_state = _policy_state(security_gaps=[
        "Uncharged Asset: AST-002 has no corresponding security charge registered.",
    ])
    verdict, notes = _apply_deterministic_policy_checks("APPROVED", None, draft, policy_state, {})
    assert verdict == "REJECTED"
    assert "Uncharged Asset: AST-002" in notes


def test_apply_deterministic_policy_checks_combines_llm_notes_with_code_enforced_reasons():
    draft = _compliant_draft(cp_ids=[])  # missing KYC-AML
    verdict, notes = _apply_deterministic_policy_checks(
        "REJECTED", "The narrative is too generic.", draft, _policy_state(), {},
    )
    assert verdict == "REJECTED"
    assert "The narrative is too generic." in notes
    assert "Missing Required CP KYC-AML" in notes


def test_apply_deterministic_policy_checks_never_overrides_rejected_to_approved():
    draft = _compliant_draft(cp_ids=["KYC-AML"])  # fully compliant
    verdict, notes = _apply_deterministic_policy_checks("REJECTED", "Weak mitigants.", draft, _policy_state(), {})
    assert verdict == "REJECTED"
    assert notes == "Weak mitigants."


# ---------------------------------------------------------------------------
# Narrative-accuracy check: reported_figures vs. ground_truth_figures. This
# closes the gap the CP/taxonomy/covenant checks don't -- those validate the
# deal's actual computed figures and structure, never what the drafted
# prose *states* those figures to be.
# ---------------------------------------------------------------------------

GROUND_TRUTH = {"dscr": 1.05, "gross_leverage": 3.4, "ebitda": 950000}


def test_apply_deterministic_policy_checks_passes_a_matching_reported_figure():
    draft = _compliant_draft(cp_ids=["KYC-AML"], reported_figures={"dscr": 1.0501})  # within 0.5% tolerance
    verdict, notes = _apply_deterministic_policy_checks(
        "APPROVED", None, draft, _policy_state(), GROUND_TRUTH,
    )
    assert verdict == "APPROVED"
    assert notes is None


def test_apply_deterministic_policy_checks_rejects_a_mismatched_reported_figure():
    """A stated DSCR of 1.45 against a computed 1.05 is well outside tolerance."""
    draft = _compliant_draft(cp_ids=["KYC-AML"], reported_figures={"dscr": 1.45})
    verdict, notes = _apply_deterministic_policy_checks(
        "APPROVED", None, draft, _policy_state(), GROUND_TRUTH,
    )
    assert verdict == "REJECTED"
    assert "Narrative/Ground-Truth Mismatch: reported dscr 1.45 vs computed 1.05" in notes


def test_apply_deterministic_policy_checks_flags_unresolvable_reported_figure():
    """A reported metric with no corresponding computed value must be
    flagged, not silently ignored -- same fallback discipline as the
    covenant-metric check in policy_engine.py."""
    draft = _compliant_draft(cp_ids=["KYC-AML"], reported_figures={"made_up_metric": 42})
    verdict, notes = _apply_deterministic_policy_checks(
        "APPROVED", None, draft, _policy_state(), GROUND_TRUTH,
    )
    assert verdict == "REJECTED"
    assert "UNRESOLVABLE_REPORTED_FIGURE" in notes
    assert "made_up_metric" in notes


def test_apply_deterministic_policy_checks_reported_figures_uses_same_unified_rejection_path():
    """A narrative mismatch must combine with LLM notes exactly like every
    other code-enforced reason -- not a separate/parallel field."""
    draft = _compliant_draft(cp_ids=["KYC-AML"], reported_figures={"dscr": 1.45})
    verdict, notes = _apply_deterministic_policy_checks(
        "REJECTED", "The tone is too informal.", draft, _policy_state(), GROUND_TRUTH,
    )
    assert verdict == "REJECTED"
    assert "The tone is too informal." in notes
    assert "Narrative/Ground-Truth Mismatch: reported dscr 1.45 vs computed 1.05" in notes


def test_apply_deterministic_policy_checks_ignores_omitted_figures():
    """Omitting a metric entirely (rather than guessing) must never itself be a violation."""
    draft = _compliant_draft(cp_ids=["KYC-AML"], reported_figures={})
    verdict, notes = _apply_deterministic_policy_checks(
        "APPROVED", None, draft, _policy_state(), GROUND_TRUTH,
    )
    assert verdict == "APPROVED"
    assert notes is None


# ---------------------------------------------------------------------------
# _values_match(): a fixed 0.5% relative tolerance, applied consistently.
# ---------------------------------------------------------------------------

def test_values_match_within_tolerance():
    assert _values_match(1.0501, 1.05) is True  # ~0.01% off


def test_values_match_rejects_beyond_tolerance():
    assert _values_match(1.45, 1.05) is False


def test_values_match_boundary_exactly_at_tolerance():
    actual = 100.0
    reported = actual * 1.005  # exactly 0.5%
    assert _values_match(reported, actual) is True


def test_values_match_handles_zero_actual_without_dividing_by_zero():
    assert _values_match(0, 0) is True
    assert _values_match(0.001, 0) is False


def test_values_match_returns_false_for_non_numeric_input():
    assert _values_match("not a number", 1.05) is False


# ---------------------------------------------------------------------------
# _compute_collateral_cover_pct() / _ground_truth_figures()
# ---------------------------------------------------------------------------

def test_compute_collateral_cover_pct_aggregates_across_assets():
    collateral = [
        {"exposure": 100, "collateral_value": 80},
        {"exposure": 50, "collateral_value": 40},
    ]
    assert _compute_collateral_cover_pct(collateral) == pytest.approx(120 / 150 * 100)


def test_compute_collateral_cover_pct_returns_none_when_no_exposure():
    assert _compute_collateral_cover_pct([]) is None
    assert _compute_collateral_cover_pct([{"exposure": 0, "collateral_value": 0}]) is None


def test_ground_truth_figures_combines_ratios_financials_and_collateral():
    financials = {"FY-Current": {"raw": {"revenue": 1000}, "ebitda": 510, "tangible_net_worth": 250}}
    ratios = {"FY-Current": {"dscr": 1.5, "gross_leverage": 0.68}}
    collateral = [{"exposure": 100, "collateral_value": 90}]

    truth = _ground_truth_figures(financials, ratios, collateral)

    assert truth["dscr"] == 1.5
    assert truth["gross_leverage"] == 0.68
    assert truth["ebitda"] == 510
    assert truth["tangible_net_worth"] == 250
    assert truth["collateral_cover_pct"] == pytest.approx(90.0)
    assert "raw" not in truth  # the nested raw-input dict is not itself a figure


def test_ground_truth_figures_handles_completely_empty_input():
    assert _ground_truth_figures({}, {}, []) == {}


# ---------------------------------------------------------------------------
# _check_reported_figures()
# ---------------------------------------------------------------------------

def test_check_reported_figures_returns_no_reasons_when_all_match():
    assert _check_reported_figures({"dscr": 1.05}, {"dscr": 1.05}) == []


def test_check_reported_figures_flags_unresolvable_and_mismatch_independently():
    reasons = _check_reported_figures(
        {"dscr": 1.45, "made_up_metric": 42}, {"dscr": 1.05},
    )
    assert len(reasons) == 2
    assert any("Narrative/Ground-Truth Mismatch" in r for r in reasons)
    assert any("UNRESOLVABLE_REPORTED_FIGURE" in r for r in reasons)


# ---------------------------------------------------------------------------
# _load_multi_period_financials(): --spread takes precedence over --financials
# ---------------------------------------------------------------------------

def test_load_multi_period_financials_prefers_spread_over_financials(tmp_path):
    financials_path = tmp_path / "financials.json"
    spread_path = tmp_path / "spread.json"
    financials_path.write_text(json.dumps({"FY-Current": {"revenue": 1}}))
    spread_path.write_text(json.dumps({"FY-Current": {"revenue": 2}}))

    result = _load_multi_period_financials(str(financials_path), str(spread_path))
    assert result == {"FY-Current": {"revenue": 2}}


def test_load_multi_period_financials_falls_back_to_financials_when_no_spread(tmp_path):
    financials_path = tmp_path / "financials.json"
    financials_path.write_text(json.dumps({"FY-Current": {"revenue": 1}}))

    result = _load_multi_period_financials(str(financials_path), None)
    assert result == {"FY-Current": {"revenue": 1}}


def test_load_multi_period_financials_returns_none_when_neither_given():
    assert _load_multi_period_financials(None, None) is None


# ---------------------------------------------------------------------------
# evaluate_financial_model() -- re-exported from spreading_builder and used
# directly by run_pipeline(); see test_spreading_builder.py for full coverage.
# ---------------------------------------------------------------------------

def test_evaluate_financial_model_usable_from_orchestrator():
    result = evaluate_financial_model({"FY-Current": {"revenue": 100, "cost_of_sales": 40}})
    assert result["financials"]["FY-Current"]["gross_profit"] == 60


# ---------------------------------------------------------------------------
# Governance loop: export gating, per-iteration checkpointing.
# ---------------------------------------------------------------------------

def test_run_pipeline_exports_when_approved_on_first_iteration(project_root):
    client = MockClient([
        _compliant_draft(),      # [1/3] Underwriter draft
        _approved_json(),        # [2/3] Risk Reviewer audit, iteration 1
    ])

    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 client=client)

    deal_dir = os.path.dirname(state_path("Acme Corp", "Fleet Loan"))
    assert os.path.isfile(os.path.join(deal_dir, "draft_v1.md"))
    assert not os.path.exists(os.path.join(deal_dir, "draft_v2.md"))

    state = read_state("Acme Corp", "Fleet Loan")
    assert [entry["verdict"] for entry in state["review_trail"]] == ["APPROVED"]
    assert state["review_verdict"] == "APPROVED"
    assert "export" in state["steps_completed"]
    assert "policy_state" in state

    docx_path = os.path.join(deal_dir, "Acme Corp_Fleet Loan_CAM.docx")
    xlsx_path = os.path.join(deal_dir, "Acme Corp_Fleet Loan_Spreading.xlsx")
    assert os.path.isfile(docx_path)
    assert os.path.isfile(xlsx_path)


def test_run_pipeline_revises_and_exports_after_one_rejection(project_root):
    client = MockClient([
        _compliant_draft("# Draft v1"),                # initial draft
        _rejected_json("Fix the EBITDA figure."),      # iteration 1 audit
        _compliant_draft("# Draft v2 (revised)"),      # Maker revision
        _approved_json(),                                # iteration 2 audit
    ])

    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 client=client)

    deal_dir = os.path.dirname(state_path("Acme Corp", "Fleet Loan"))
    with open(os.path.join(deal_dir, "draft_v1.md"), encoding="utf-8") as f:
        assert f.read() == _compliant_draft("# Draft v1")
    with open(os.path.join(deal_dir, "draft_v2.md"), encoding="utf-8") as f:
        assert f.read() == _compliant_draft("# Draft v2 (revised)")

    state = read_state("Acme Corp", "Fleet Loan")
    assert [entry["verdict"] for entry in state["review_trail"]] == ["REJECTED", "APPROVED"]
    assert state["review_trail"][0]["notes"] == "Fix the EBITDA figure."
    assert state["review_verdict"] == "APPROVED"

    # The exported draft must be the revised, approved one -- not the rejected first draft.
    docx_path = os.path.join(deal_dir, "Acme Corp_Fleet Loan_CAM.docx")
    doc = docx.Document(docx_path)
    assert doc.paragraphs[0].text == "Draft v2 (revised)"


def test_run_pipeline_exits_nonzero_and_skips_export_after_max_rejections(project_root):
    client = MockClient([
        "# Draft",
        _rejected_json("still wrong 1"),
        "# Draft",
        _rejected_json("still wrong 2"),
        "# Draft",
        _rejected_json("still wrong 3"),
    ])

    with pytest.raises(SystemExit) as exc_info:
        run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                     client=client)
    assert exc_info.value.code == 1

    deal_dir = os.path.dirname(state_path("Acme Corp", "Fleet Loan"))
    for i in range(1, MAX_REVIEW_ITERATIONS + 1):
        assert os.path.isfile(os.path.join(deal_dir, f"draft_v{i}.md"))

    state = read_state("Acme Corp", "Fleet Loan")
    assert state["review_verdict"] == "REJECTED"
    assert len(state["review_trail"]) == MAX_REVIEW_ITERATIONS
    assert not os.path.exists(os.path.join(deal_dir, "Acme Corp_Fleet Loan_CAM.docx"))
    assert client.call_count == 6  # 1 initial draft + 3 audits + 2 revisions, no 3rd revision


def test_run_pipeline_stores_financials_and_ratios_on_state(project_root):
    client = MockClient([_compliant_draft(), _approved_json()])
    multi_period_financials = {"FY-Current": {"revenue": 1000, "cost_of_sales": 400}}

    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 multi_period_financials=multi_period_financials, client=client)

    state = read_state("Acme Corp", "Fleet Loan")
    assert state["financials"]["FY-Current"]["gross_profit"] == 600
    assert "ratios" in state and "FY-Current" in state["ratios"]
    assert "spread" in state["steps_completed"]


def test_run_pipeline_stores_collateral_data_on_state(project_root):
    client = MockClient([_compliant_draft(), _approved_json()])
    collateral_data = [{"asset_class": "HGV", "exposure": 100, "collateral_value": 80,
                         "perfection_status": "Registered"}]

    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 collateral_data=collateral_data, client=client)

    state = read_state("Acme Corp", "Fleet Loan")
    assert state["collateral"] == collateral_data


# ---------------------------------------------------------------------------
# State preservation across re-runs: write_state()'s shallow merge means a
# fresh `financials={}` replaces the whole dict, so run_pipeline() must read
# existing state first and only replace financials/ratios/collateral when
# actually given new data for them, and only ever *add* to steps_completed.
# ---------------------------------------------------------------------------

def test_run_pipeline_preserves_existing_financials_when_rerun_without_new_data(project_root):
    multi_period_financials = {"FY-Current": {"revenue": 1000, "cost_of_sales": 400}}
    client1 = MockClient([_compliant_draft(), _approved_json()])
    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 multi_period_financials=multi_period_financials, client=client1)

    state_after_first = read_state("Acme Corp", "Fleet Loan")
    assert state_after_first["financials"]["FY-Current"]["gross_profit"] == 600

    # Second run for the same deal, no --financials/--spread this time --
    # must not wipe what the first run already checkpointed.
    client2 = MockClient([_compliant_draft("# Draft CAM v2"), _approved_json()])
    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 client=client2)

    state_after_second = read_state("Acme Corp", "Fleet Loan")
    assert state_after_second["financials"]["FY-Current"]["gross_profit"] == 600


def test_run_pipeline_preserves_existing_collateral_when_rerun_without_new_data(project_root):
    collateral_data = [{"asset_class": "HGV", "exposure": 100, "collateral_value": 80,
                         "perfection_status": "Registered"}]
    client1 = MockClient([_compliant_draft(), _approved_json()])
    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 collateral_data=collateral_data, client=client1)

    client2 = MockClient([_compliant_draft("# Draft CAM v2"), _approved_json()])
    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 client=client2)

    state = read_state("Acme Corp", "Fleet Loan")
    assert state["collateral"] == collateral_data


def test_run_pipeline_accumulates_steps_completed_without_duplicating_across_reruns(project_root):
    client1 = MockClient([_compliant_draft(), _approved_json()])
    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 client=client1)
    state1 = read_state("Acme Corp", "Fleet Loan")
    assert state1["steps_completed"].count("draft") == 1
    assert state1["steps_completed"].count("audit") == 1
    assert state1["steps_completed"].count("export") == 1

    client2 = MockClient([_compliant_draft("# Draft CAM v2"), _approved_json()])
    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 client=client2)
    state2 = read_state("Acme Corp", "Fleet Loan")
    assert state2["steps_completed"].count("draft") == 1
    assert state2["steps_completed"].count("audit") == 1
    assert state2["steps_completed"].count("export") == 1


# ---------------------------------------------------------------------------
# End-to-end: a deterministic policy failure (not an LLM-originated one)
# must still drive the governance loop through the exact same revision path,
# and produce a review_trail entry structurally identical to an
# LLM-originated one.
# ---------------------------------------------------------------------------

def test_run_pipeline_code_enforced_rejection_triggers_revision_and_matches_review_trail_shape(project_root):
    # The Reviewer says APPROVED both times; the *first* draft omits a
    # required CP, so the code-enforced layer must override it to REJECTED
    # and drive a revision cycle identical to an LLM-originated rejection.
    client = MockClient([
        _compliant_draft("# Draft v1", cp_ids=[]),   # missing KYC-AML/FACILITY-EXECUTION
        _approved_json(),                              # Reviewer wrongly says APPROVED
        _compliant_draft("# Draft v2 (fixed)"),        # Maker revision, now compliant
        _approved_json(),
    ])

    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 client=client)

    state = read_state("Acme Corp", "Fleet Loan")
    trail = state["review_trail"]
    assert [entry["verdict"] for entry in trail] == ["REJECTED", "APPROVED"]
    assert "Missing Required CP KYC-AML" in trail[0]["notes"]

    # Structurally identical to an LLM-originated entry: same keys, same types.
    llm_style_entry_keys = {"iteration", "verdict", "notes", "timestamp"}
    assert set(trail[0].keys()) == llm_style_entry_keys
    assert set(trail[1].keys()) == llm_style_entry_keys

    # The revision path actually ran: a second draft + audit call were made.
    assert client.call_count == 4

    docx_path = os.path.join(os.path.dirname(state_path("Acme Corp", "Fleet Loan")),
                              "Acme Corp_Fleet Loan_CAM.docx")
    doc = docx.Document(docx_path)
    assert doc.paragraphs[0].text == "Draft v2 (fixed)"


def test_run_pipeline_narrative_mismatch_triggers_revision_and_matches_review_trail_shape(project_root):
    """A drafted CAM can state figures that don't match what state.json
    actually computed even while every CP/taxonomy box is ticked -- the
    narrative-accuracy check must catch that independently, drive the same
    revision cycle, and produce a review_trail entry with the same shape as
    every other rejection reason."""
    client = MockClient([
        _compliant_draft("# Draft v1", reported_figures={"dscr": 5.0}),  # wildly wrong DSCR
        _approved_json(),                                                 # Reviewer wrongly says APPROVED
        _compliant_draft("# Draft v2 (fixed)", reported_figures={"dscr": 5.1}),  # now matches
        _approved_json(),
    ])
    multi_period_financials = {
        "FY-Current": {
            "revenue": 1000, "cost_of_sales": 400, "admin_expenses": 100,
            "depreciation": 50, "amortisation": 20, "other_income": 10,
            "interest_paid": 30, "scheduled_principal": 70,
        }
    }  # EBITDA=510, DSCR=510/(30+70)=5.1

    run_pipeline("Acme Corp", "Fleet Loan", "0.20%", "LGD 3 (15%)", "corporate_credit",
                 multi_period_financials=multi_period_financials, client=client)

    state = read_state("Acme Corp", "Fleet Loan")
    trail = state["review_trail"]
    assert [entry["verdict"] for entry in trail] == ["REJECTED", "APPROVED"]
    assert "Narrative/Ground-Truth Mismatch: reported dscr 5.0 vs computed 5.1" in trail[0]["notes"]

    llm_style_entry_keys = {"iteration", "verdict", "notes", "timestamp"}
    assert set(trail[0].keys()) == llm_style_entry_keys
    assert set(trail[1].keys()) == llm_style_entry_keys

    assert client.call_count == 4  # revision path actually ran

    docx_path = os.path.join(os.path.dirname(state_path("Acme Corp", "Fleet Loan")),
                              "Acme Corp_Fleet Loan_CAM.docx")
    doc = docx.Document(docx_path)
    assert doc.paragraphs[0].text == "Draft v2 (fixed)"
