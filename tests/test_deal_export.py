import json
import os
from pathlib import Path

import docx
import openpyxl
import pytest

from deal_export import (
    _downside_financial_data_from_state,
    _financial_data_from_state,
    _supplied_row_values,
    export_deal,
    main,
    workbook_inputs_from_state,
)
from state_manager import read_state, resolve_date_str, write_state


def _write(path, content="content"):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    with open(path, "w", encoding="utf-8") as f:
        f.write(content)


def test_a_bare_call_without_date_str_finds_the_state_of_an_earlier_dated_folder(tmp_path):
    """The state is read BEFORE today's output folder is created. Before, makedirs ran first and the auto-discovery
    (latest dated folder wins) then found the new, empty folder, so the workbook came out blank (#171 review)."""
    base = str(tmp_path)
    earlier = os.path.join(base, "deals", "Acme Corp", "Fleet Loan_2026-01-01")
    os.makedirs(earlier)
    with open(os.path.join(earlier, "state.json"), "w", encoding="utf-8") as f:
        json.dump({"financials": {"FY-Current": {"raw": {"revenue": 4321}}}}, f)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))

    output_dir = export_deal("Acme Corp", "Fleet Loan", "asset_finance", "# Draft", base_dir=base)   # no date_str

    assert os.path.basename(output_dir) != "Fleet Loan_2026-01-01"           # a new dated folder for today...
    sheet = openpyxl.load_workbook(os.path.join(output_dir, "Acme Corp_Fleet Loan_Spreading.xlsx"))["Financial Spreading"]
    revenue_row = next(r for r in sheet.iter_rows(values_only=True) if r[0] == "Revenue")
    assert 4321 in revenue_row                                              # ...but the earlier folder's figures


def test_creates_dated_output_folder(tmp_path):
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))

    output_dir = export_deal(
        "Acme Corp", "Fleet Loan", "asset_finance", "# Draft",
        date_str="2026-01-15", base_dir=base,
    )

    assert output_dir == os.path.join(base, "deals", "Acme Corp", "Fleet Loan_2026-01-15")
    assert os.path.isdir(output_dir)


def test_exports_docx_and_xlsx(tmp_path):
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))

    output_dir = export_deal(
        "Acme Corp", "Fleet Loan", "asset_finance",
        "# Acme Corp CAM\n\n| A | B |\n| --- | --- |\n| 1 | 2 |\n",
        date_str="2026-01-15", base_dir=base,
    )

    docx_path = os.path.join(output_dir, "Acme Corp_Fleet Loan_CAM.docx")
    xlsx_path = os.path.join(output_dir, "Acme Corp_Fleet Loan_Spreading.xlsx")
    assert os.path.isfile(docx_path)
    assert os.path.isfile(xlsx_path)

    doc = docx.Document(docx_path)
    assert doc.paragraphs[0].text == "Acme Corp CAM"
    assert len(doc.tables) == 1

    wb = openpyxl.load_workbook(xlsx_path)
    assert "Financial Spreading" in wb.sheetnames


def test_auto_saves_new_template_when_none_exists(tmp_path):
    base = str(tmp_path)
    # No templates/cam/ or templates/local/cam/ entry for this type at all.

    export_deal(
        "Acme Corp", "Fleet Loan", "brand_new_type", "# Brand New Draft",
        date_str="2026-01-15", base_dir=base,
    )

    new_template_path = os.path.join(base, "templates", "local", "cam", "brand_new_type_cam.md")
    assert os.path.isfile(new_template_path)
    with open(new_template_path, encoding="utf-8") as f:
        assert f.read() == "# Brand New Draft"


def test_does_not_overwrite_an_existing_local_override(tmp_path):
    base = str(tmp_path)
    override_path = os.path.join(base, "templates", "local", "cam", "asset_finance_cam.md")
    _write(override_path, "# Existing override -- must not change")

    export_deal(
        "Acme Corp", "Fleet Loan", "asset_finance", "# A totally different draft",
        date_str="2026-01-15", base_dir=base,
    )

    with open(override_path, encoding="utf-8") as f:
        assert f.read() == "# Existing override -- must not change"


def test_does_not_touch_local_override_when_shipped_default_exists(tmp_path):
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))

    export_deal(
        "Acme Corp", "Fleet Loan", "asset_finance", "# Draft",
        date_str="2026-01-15", base_dir=base,
    )

    local_override_path = os.path.join(base, "templates", "local", "cam", "asset_finance_cam.md")
    assert not os.path.exists(local_override_path)


def test_export_deal_rejects_unsafe_company(tmp_path):
    base = str(tmp_path)
    with pytest.raises(ValueError):
        export_deal("C:\\Windows\\Temp\\evil", "Fleet Loan", "asset_finance", "# Draft", base_dir=base)


def test_export_deal_rejects_unsafe_deal_type(tmp_path):
    """deal_type flows into template_resolver's path building
    (templates/local/cam/<deal_type>_cam.md) the same way company/proposal
    flow into the deals/ path -- it needs the same protection."""
    base = str(tmp_path)
    with pytest.raises(ValueError):
        export_deal("Acme Corp", "Fleet Loan", "../../evil", "# Draft", base_dir=base)


def test_export_deal_rejects_unsafe_date_str(tmp_path):
    base = str(tmp_path)
    with pytest.raises(ValueError):
        export_deal("Acme Corp", "Fleet Loan", "asset_finance", "# Draft",
                     date_str="../../evil", base_dir=base)


def test_financial_data_from_state_ignores_non_dict_period_value():
    """A period value that isn't a dict (e.g. malformed/legacy state.json)
    must not crash -- it degrades to an empty raw-figures dict for that
    period instead."""
    state = {"financials": {"FY-Current": "not a dict", "FY-1": {"raw": {"revenue": 100}}}}
    result = _financial_data_from_state(state)
    assert result == {"FY-Current": {}, "FY-1": {"revenue": 100}}


# ---------------------------------------------------------------------------
# _financial_data_from_state()'s historical-period rule (issue #121): raw
# inputs for analyst-supplied historical periods come from their own
# analyst_supplied_financials store. Issue #249 adds the separate per-period
# rule for framework-computed forward periods in the same deal.
# ---------------------------------------------------------------------------

def test_financial_data_from_state_reads_analyst_supplied_store_when_flagged():
    """Rule: R-EXPORT-001"""
    state = {
        "financials_source": "analyst-supplied",
        "analyst_supplied_financials": {"FY-Current": {"revenue": 500}},
        # Deliberately also present, to prove it's ignored for this deal --
        # an analyst-supplied deal's financials[period] never has a "raw"
        # sub-key populated in practice, but this guards against a future
        # bug that blends the two if it somehow were.
        "financials": {"FY-Current": {"raw": {"revenue": 999}}},
    }
    assert _financial_data_from_state(state) == {"FY-Current": {"revenue": 500}}


def test_financial_data_from_state_reads_framework_computed_store_by_default():
    """financials_source absent, or anything other than "analyst-supplied",
    must keep reading the pre-existing financials[period]["raw"] path --
    the framework-computed behavior this function already had before #121,
    unchanged.
    Rule: R-EXPORT-001"""
    state = {"financials": {"FY-Current": {"raw": {"revenue": 500}}}}
    assert _financial_data_from_state(state) == {"FY-Current": {"revenue": 500}}

    state["financials_source"] = "framework-computed"
    assert _financial_data_from_state(state) == {"FY-Current": {"revenue": 500}}


def test_financial_data_from_state_analyst_supplied_with_no_raw_breakdown_is_empty():
    """An analyst-supplied deal where no raw breakdown was ever given (the
    common case -- it's optional) must export blank raw-input cells rather
    than falling back to financials[period]["raw"], which this mode never
    populates anyway."""
    state = {"financials_source": "analyst-supplied", "financials": {"FY-Current": {}}}
    assert _financial_data_from_state(state) == {}


# ---------------------------------------------------------------------------
# _downside_financial_data_from_state(): the exported workbook's FY+1/FY+2/
# FY+3 "(Downside)" columns (see issue #38 / spreading_builder.py's
# COL_TO_PERIOD_KEY) are populated from this, mirroring how
# _financial_data_from_state() feeds the historical/forward base-case
# columns.
# ---------------------------------------------------------------------------

def test_downside_financial_data_from_state_reduces_downside_case_financials():
    state = {
        "downside_case": {
            "financials": {"FY+1": {"raw": {"revenue": 950}}, "FY+2": {"raw": {"revenue": 1000}}},
            "ratios": {"FY+1": {"dscr": 1.02}},
        },
    }
    result = _downside_financial_data_from_state(state)
    assert result == {"FY+1": {"revenue": 950}, "FY+2": {"revenue": 1000}}


def test_downside_financial_data_from_state_empty_when_no_downside_case():
    """No stress_assumptions were ever supplied for this deal -- state.json
    has no "downside_case" key at all (or it's {}) -- must degrade to an
    empty dict, not raise."""
    assert _downside_financial_data_from_state({}) == {}
    assert _downside_financial_data_from_state({"downside_case": {}}) == {}


def test_downside_financial_data_from_state_ignores_non_dict_downside_case():
    """Malformed/legacy state.json where "downside_case" itself isn't a dict
    must not crash."""
    assert _downside_financial_data_from_state({"downside_case": "not a dict"}) == {}


def test_export_deal_populates_downside_columns_in_the_exported_workbook(tmp_path):
    """End-to-end: a deal whose state.json carries a downside_case actually
    gets its FY+1 (Downside) column populated in the exported .xlsx --
    proves the full state.json -> deal_export.py -> spreading_builder.py
    wiring, not just the extraction function in isolation."""
    from state_manager import write_state

    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))
    write_state(
        "Acme Corp", "Fleet Loan", date_str="2026-01-15", base_dir=base,
        financials={"FY-Current": {"raw": {"revenue": 1000}}},
        downside_case={"financials": {"FY+1": {"raw": {"revenue": 950}}}},
    )

    output_dir = export_deal(
        "Acme Corp", "Fleet Loan", "asset_finance", "# Draft",
        date_str="2026-01-15", base_dir=base,
    )

    wb = openpyxl.load_workbook(os.path.join(output_dir, "Acme Corp_Fleet Loan_Spreading.xlsx"))
    ws = wb["Financial Spreading"]
    label_to_row = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
    assert ws.cell(row=label_to_row["Revenue"], column=8).value == 950  # H = FY+1 (Downside)


def test_analyst_supplied_historical_export_uses_recorded_values_and_labels_the_column(tmp_path):
    from state_manager import write_state

    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))
    write_state(
        "Acme Corp", "Fleet Loan", date_str="2026-01-15", base_dir=base,
        financials_source="analyst-supplied",
        analyst_supplied_financials={
            "FY-Current": {"revenue": 5000, "cost_of_sales": 3300, "admin_expenses": 1000},
        },
        financials={
            "FY-Current": {
                "gross_profit": 1700, "ebitda": 1300, "total_equity": 2300,
            },
        },
        ratios={
            "FY-Current": {
                "gross_leverage": 1.54, "dscr": 2.6, "gearing": 0.87,
            },
        },
    )

    output_dir = export_deal(
        "Acme Corp", "Fleet Loan", "asset_finance", "# Draft",
        date_str="2026-01-15", base_dir=base,
    )

    wb = openpyxl.load_workbook(os.path.join(output_dir, "Acme Corp_Fleet Loan_Spreading.xlsx"))
    ws = wb["Financial Spreading"]
    assert ws.cell(row=1, column=4).value == "FY-Current (analyst-supplied)"
    label_to_row = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}

    expected_values = {
        "Revenue": 5000,
        "Gross Profit": 1700,
        "EBITDA": 1300,
        "Total Equity": 2300,
        "Gross Leverage": 1.54,
        "DSCR": 2.6,
        "Gearing % (Interest-Bearing Debt / Equity)": 0.87,
    }
    for label, expected in expected_values.items():
        value = ws.cell(row=label_to_row[label], column=4).value
        assert value == expected, label
        assert not (isinstance(value, str) and value.startswith("=")), label

    # No recorded subtotal was supplied for Net Profit; do not let Excel
    # recompute one from raw cells that were not the basis of the supplied figure.
    assert ws.cell(row=label_to_row["Net Profit"], column=4).value is None


# ---------------------------------------------------------------------------
# CLI (/assemble's Bash step): which dated folder does it export into? Issue
# #97 -- a multi-day deal's output must land next to its own state.json, not
# in a new folder dated whenever the export happened to run.
# ---------------------------------------------------------------------------

def _cli_draft(tmp_path):
    draft = tmp_path / "draft.md"
    draft.write_text("# Draft", encoding="utf-8")
    return str(draft)


def test_cli_exports_into_the_deals_existing_dated_folder(tmp_path, monkeypatch):
    from datetime import datetime
    monkeypatch.chdir(tmp_path)
    existing = tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-01-10"
    existing.mkdir(parents=True)
    (existing / "state.json").write_text("{}", encoding="utf-8")

    main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance",
          "--draft", _cli_draft(tmp_path)])

    assert (existing / "Acme Corp_Fleet Loan_CAM.docx").exists()
    assert (existing / "Acme Corp_Fleet Loan_Spreading.xlsx").exists()
    today = datetime.now().strftime("%Y-%m-%d")
    assert not (tmp_path / "deals" / "Acme Corp" / f"Fleet Loan_{today}").exists()


def test_cli_picks_the_most_recent_of_several_existing_dated_folders(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    for date in ("2025-01-10", "2026-01-10"):
        (tmp_path / "deals" / "Acme Corp" / f"Fleet Loan_{date}").mkdir(parents=True)

    main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance",
          "--draft", _cli_draft(tmp_path)])

    assert (tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-01-10" / "Acme Corp_Fleet Loan_CAM.docx").exists()
    assert not (tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2025-01-10" / "Acme Corp_Fleet Loan_CAM.docx").exists()


def test_cli_date_str_flag_creates_export_only_folder_for_a_deal_with_no_state_without_hijacking_discovery(
        tmp_path, monkeypatch):
    from datetime import datetime
    monkeypatch.chdir(tmp_path)

    main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance",
          "--draft", _cli_draft(tmp_path), "--date-str", "2026-03-01"])

    export_dir = tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-03-01"
    assert (export_dir / "Acme Corp_Fleet Loan_CAM.docx").exists()
    assert not (export_dir / "state.json").exists()
    assert resolve_date_str("Acme Corp", "Fleet Loan") == datetime.now().strftime("%Y-%m-%d")


def test_cli_explicit_date_with_no_state_is_an_error_when_the_deal_has_state_in_another_folder(
        tmp_path, monkeypatch, capsys):
    """Decision D1 (#248): a named folder with no state.json must not silently export a blank workbook."""
    monkeypatch.chdir(tmp_path)
    write_state("Acme Corp", "Fleet Loan", date_str="2026-01-10", note="existing deal state")
    draft = _cli_draft(tmp_path)
    before = sorted(str(p.relative_to(tmp_path)) for p in tmp_path.rglob("*"))

    with pytest.raises(SystemExit) as caught:
        main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance",
              "--draft", draft, "--date-str", "2026-03-01"])

    message = str(caught.value.code)
    assert message.startswith("error: ") and "\n" not in message
    assert "2026-03-01" in message and "Fleet Loan_2026-01-10" in message
    assert capsys.readouterr().err == ""
    assert sorted(str(p.relative_to(tmp_path)) for p in tmp_path.rglob("*")) == before  # nothing written
    assert not (tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-03-01").exists()
    assert not (tmp_path / "templates").exists()


def test_export_deal_with_an_explicit_date_that_has_no_state_raises_when_state_exists_elsewhere(tmp_path):
    from state_manager import StateError
    write_state("Acme Corp", "Fleet Loan", date_str="2026-01-10", base_dir=str(tmp_path), note="existing deal state")

    with pytest.raises(StateError, match=r"2026-03-01.*Fleet Loan_2026-01-10"):
        export_deal("Acme Corp", "Fleet Loan", "asset_finance", "# Draft", date_str="2026-03-01", base_dir=str(tmp_path))

    assert not (tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-03-01").exists()


def test_export_deal_with_an_explicit_date_that_has_state_or_no_deal_state_anywhere_still_exports(tmp_path):
    write_state("Acme Corp", "Fleet Loan", date_str="2026-01-10", base_dir=str(tmp_path), note="existing deal state")
    own = export_deal("Acme Corp", "Fleet Loan", "asset_finance", "# Draft", date_str="2026-01-10", base_dir=str(tmp_path))
    new = export_deal("Acme Corp", "Brand New", "asset_finance", "# Draft", date_str="2026-03-01", base_dir=str(tmp_path))
    assert os.path.isfile(os.path.join(own, "Acme Corp_Fleet Loan_CAM.docx"))
    assert os.path.isfile(os.path.join(new, "Acme Corp_Brand New_CAM.docx"))


def test_cli_brand_new_deal_gets_a_folder_dated_today(tmp_path, monkeypatch):
    from datetime import datetime
    monkeypatch.chdir(tmp_path)

    main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance",
          "--draft", _cli_draft(tmp_path)])

    today = datetime.now().strftime("%Y-%m-%d")
    assert (tmp_path / "deals" / "Acme Corp" / f"Fleet Loan_{today}" / "Acme Corp_Fleet Loan_CAM.docx").exists()


def test_cli_rejects_an_unsafe_company_before_touching_the_deals_tree(tmp_path, monkeypatch):
    import deal_export
    monkeypatch.chdir(tmp_path)

    def must_not_resolve(*args, **kwargs):
        pytest.fail("resolve_export_date_str() (the deals/ glob) ran before sanitization")

    monkeypatch.setattr(deal_export, "resolve_export_date_str", must_not_resolve)
    with pytest.raises(SystemExit, match="error: company"):
        main(["--company", "../evil", "--proposal", "Fleet Loan", "--type", "asset_finance",
              "--draft", _cli_draft(tmp_path)])
    assert not (tmp_path / "deals").exists()


@pytest.mark.parametrize("bad", ["foo", "2026-1-5", "2026-01-05-extra", "", "2026-02-31"])
def test_cli_rejects_a_malformed_date_str(tmp_path, monkeypatch, bad):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(SystemExit):
        main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance",
              "--draft", _cli_draft(tmp_path), "--date-str", bad])
    assert not (tmp_path / "deals").exists()


def test_cli_does_not_export_into_another_deals_prefix_sharing_folder(tmp_path, monkeypatch):
    """Proposal "Fleet" must not claim "Fleet_Q2"'s folder (issue #97 review)."""
    monkeypatch.chdir(tmp_path)
    (tmp_path / "deals" / "Acme" / "Fleet_2026-01-10").mkdir(parents=True)
    (tmp_path / "deals" / "Acme" / "Fleet_Q2_2026-05-01").mkdir(parents=True)

    main(["--company", "Acme", "--proposal", "Fleet", "--type", "asset_finance",
          "--draft", _cli_draft(tmp_path)])

    assert (tmp_path / "deals" / "Acme" / "Fleet_2026-01-10" / "Acme_Fleet_CAM.docx").exists()
    assert not (tmp_path / "deals" / "Acme" / "Fleet_Q2_2026-05-01" / "Acme_Fleet_CAM.docx").exists()


# ---------------------------------------------------------------------------
# _supplied_row_values: what an analyst-supplied period writes into the workbook (issue #219)
# ---------------------------------------------------------------------------

def test_supplied_row_values_keeps_numbers_writes_na_for_a_null_ratio_and_leaves_everything_else_blank():
    """The docstring's contract, row by row. A number (zero included) is written; a ratio recorded as null is N/A, as the
    workbook writes it; a null subtotal, text, a boolean and a key with no workbook row are left out, so the cell stays
    blank. A defect here is a wrong number or a stray "N/A" in the exported workbook."""
    values = _supplied_row_values({
        "ebitda": 1000, "total_assets": 0, "dscr": 2.5,
        "gross_leverage": None,                 # a ratio recorded as null: N/A
        "net_profit": None,                     # a subtotal recorded as null: blank, not "N/A"
        "current_ratio": "n/m",                 # text: blank
        "gearing": True,                        # a boolean is not a number: blank
        "not_a_workbook_row": 5,                # no row for it: ignored
    })
    assert values == {"EBITDA": 1000, "Total Assets": 0, "DSCR": 2.5, "Gross Leverage": "N/A"}


def test_supplied_row_values_takes_the_later_period_dict_and_ignores_a_malformed_one():
    """The forecast columns pass the subtotals and then the ratios (or a base and an override): a later value wins for
    the same row, and anything that is not a dict is skipped rather than failing the export."""
    assert _supplied_row_values({"ebitda": 1}, {"ebitda": 2}) == {"EBITDA": 2}
    assert _supplied_row_values(None, "text", ["ebitda"], {"dscr": 1.5}) == {"DSCR": 1.5}
    assert _supplied_row_values() == {}


# ---------------------------------------------------------------------------
# main(): what the command line prints, defaults to and refuses, in-process (issue #182)
# ---------------------------------------------------------------------------

def _docx_text(path):
    return "\n".join(p.text for p in docx.Document(str(path)).paragraphs)


def test_cli_prints_where_the_files_went_and_they_are_there(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    existing = tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-01-10"
    existing.mkdir(parents=True)

    main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance", "--draft", _cli_draft(tmp_path)])

    out = capsys.readouterr()
    assert out.err == ""
    *notices, line = out.out.splitlines()
    assert len(notices) == 1 and notices[0].startswith("[Auto-Template] No template found for 'asset_finance'")
    assert line.startswith("Done! Files generated in ") and line.endswith("Fleet Loan_2026-01-10")
    assert os.path.samefile(line.removeprefix("Done! Files generated in "), existing)
    assert sorted(p.name for p in existing.iterdir()) == ["Acme Corp_Fleet Loan_CAM.docx", "Acme Corp_Fleet Loan_Spreading.xlsx"]


def test_cli_type_defaults_to_corporate_credit(tmp_path, monkeypatch):
    """With no --type the deal type is corporate_credit: with no template for a type, export_deal auto-saves one named for it."""
    monkeypatch.chdir(tmp_path)
    main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--draft", _cli_draft(tmp_path)])
    saved = sorted(p.name for p in (tmp_path / "templates").rglob("*.md"))
    assert saved == ["corporate_credit_cam.md"]


def test_cli_exports_the_draft_it_was_given_read_as_utf8(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "draft.md"
    draft.write_text("# Credit memo\n\nFacility of £1m \u2014 caf\u00e9 terms.\n", encoding="utf-8")

    main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance", "--draft", str(draft),
          "--date-str", "2026-03-01"])

    folder = tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-03-01"
    text = _docx_text(folder / "Acme Corp_Fleet Loan_CAM.docx")
    assert "Credit memo" in text and "Facility of £1m \u2014 caf\u00e9 terms." in text
    assert capsys.readouterr().out.strip().endswith("Fleet Loan_2026-03-01")


def test_cli_with_a_missing_draft_fails_before_creating_a_deal_folder(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(SystemExit, match="error: "):
        main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--draft", str(tmp_path / "nope.md")])
    assert not (tmp_path / "deals").exists() and not (tmp_path / "templates").exists()


def test_cli_turns_a_state_error_into_one_error_line_and_exports_nothing(tmp_path, monkeypatch, capsys):
    """A state the export cannot use is one `error:` exit message, and nothing is created in the deal folder."""
    monkeypatch.chdir(tmp_path)
    folder = tmp_path / "deals" / "Acme Corp" / "Fleet Loan_2026-01-10"
    folder.mkdir(parents=True)
    (folder / "state.json").write_text(json.dumps({"financials": [1, 2]}), encoding="utf-8")

    with pytest.raises(SystemExit) as caught:
        main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--type", "asset_finance", "--draft", _cli_draft(tmp_path)])

    message = str(caught.value.code)
    assert message.startswith("error: ") and '"financials" must be an object' in message and "\n" not in message
    assert capsys.readouterr().out == ""
    assert sorted(p.name for p in folder.iterdir()) == ["state.json"]
    assert not (tmp_path / "templates").exists()


@pytest.mark.parametrize(("bad", "message"), [
    ("2026-1-5", "--date-str '2026-1-5' must be a YYYY-MM-DD date"),
    ("2026-02-30", "--date-str '2026-02-30' is not a valid calendar date"),
])
def test_cli_names_the_bad_date_as_one_error_line_and_exits_1(tmp_path, monkeypatch, capsys, bad, message):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(SystemExit) as caught:
        main(["--company", "Acme Corp", "--proposal", "Fleet Loan", "--draft", _cli_draft(tmp_path), "--date-str", bad])
    assert caught.value.code == f"error: {message}"
    assert capsys.readouterr().err == ""  # sys.exit(str) prints the one line; no argparse usage block


def test_cli_rejects_an_unsafe_proposal_before_the_deals_tree_is_searched_as_it_does_a_company(tmp_path, monkeypatch):
    import deal_export
    monkeypatch.chdir(tmp_path)

    def must_not_resolve(*args, **kwargs):
        pytest.fail("resolve_export_date_str() (the deals/ glob) ran before the proposal was sanitized")

    monkeypatch.setattr(deal_export, "resolve_export_date_str", must_not_resolve)
    with pytest.raises(SystemExit, match="error: proposal"):
        main(["--company", "Acme Corp", "--proposal", "../evil", "--draft", _cli_draft(tmp_path)])
    assert not (tmp_path / "deals").exists()


@pytest.mark.parametrize("argv, missing", [
    ([], "--company"),
    (["--company", "A"], "--proposal"),
    (["--company", "A", "--proposal", "P"], "--draft"),
])
def test_cli_requires_company_proposal_and_draft(capsys, argv, missing):
    with pytest.raises(SystemExit) as caught:
        main(argv)
    assert caught.value.code == 2
    err = capsys.readouterr().err
    assert "usage:" in err and missing in err


def test_cli_help_lists_every_flag_and_exits_zero(capsys):
    with pytest.raises(SystemExit) as helped:
        main(["--help"])
    assert helped.value.code == 0
    text = capsys.readouterr().out
    for flag in ("--company", "--proposal", "--type", "--draft", "--date-str"):
        assert flag in text, flag


def _write_dated_state(base, date_str, revenue):
    folder = os.path.join(base, "deals", "Synthetic Co", f"Credit Facility_{date_str}")
    os.makedirs(folder, exist_ok=True)
    with open(os.path.join(folder, "state.json"), "w", encoding="utf-8") as state_file:
        json.dump({
            "financials": {"FY-Current": {"raw": {"revenue": revenue}}},
            "financials_source": "framework-computed",
        }, state_file)


def _exported_revenue(base, folder_date):
    workbook_path = os.path.join(
        base, "deals", "Synthetic Co", f"Credit Facility_{folder_date}",
        "Synthetic Co_Credit Facility_Spreading.xlsx",
    )
    sheet = openpyxl.load_workbook(workbook_path)["Financial Spreading"]
    revenue_row = next(row for row in sheet.iter_rows(values_only=True) if row[0] == "Revenue")
    return revenue_row


def test_explicit_date_exports_figures_from_that_dated_state(tmp_path):
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))
    _write_dated_state(base, "2026-01-10", 1000)
    _write_dated_state(base, "2026-03-12", 9999)

    output_dir = export_deal(
        "Synthetic Co", "Credit Facility", "asset_finance", "# Draft",
        date_str="2026-01-10", base_dir=base,
    )

    assert output_dir == os.path.join(base, "deals", "Synthetic Co", "Credit Facility_2026-01-10")
    assert 1000 in _exported_revenue(base, "2026-01-10")
    assert 9999 not in _exported_revenue(base, "2026-01-10")


def test_cli_date_str_exports_figures_from_the_requested_dated_state(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    _write(os.path.join(str(tmp_path), "templates", "cam", "asset_finance_cam.md"))
    _write_dated_state(str(tmp_path), "2026-01-10", 1000)
    _write_dated_state(str(tmp_path), "2026-03-12", 9999)

    main([
        "--company", "Synthetic Co", "--proposal", "Credit Facility", "--type", "asset_finance",
        "--draft", _cli_draft(tmp_path), "--date-str", "2026-01-10",
    ])

    assert 1000 in _exported_revenue(str(tmp_path), "2026-01-10")
    assert 9999 not in _exported_revenue(str(tmp_path), "2026-01-10")


def test_bare_export_still_uses_auto_discovered_state(tmp_path):
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))
    _write_dated_state(base, "2026-01-10", 1000)
    _write_dated_state(base, "2026-03-12", 9999)

    output_dir = export_deal(
        "Synthetic Co", "Credit Facility", "asset_finance", "# Draft",
        base_dir=base,
    )

    folder_date = os.path.basename(output_dir).removeprefix("Credit Facility_")
    assert folder_date != "2026-01-10"
    assert 9999 in _exported_revenue(base, folder_date)
    assert 1000 not in _exported_revenue(base, folder_date)



def test_mixed_basis_export_uses_framework_raw_for_computed_forward_periods(tmp_path, monkeypatch):
    """A computed /project forecast must use its own raw store even when history is analyst-supplied (#249).
    Rule: R-EXPORT-002"""
    from spreading_check import main as spreading_check_main

    base = str(tmp_path)
    date_str = "2026-10-04"
    monkeypatch.chdir(base)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))
    write_state(
        "Synthetic Co", "Credit Facility", date_str=date_str, base_dir=base,
        financials_source="analyst-supplied",
        analyst_supplied_financials={
            "FY-Current": {"revenue": 5000, "cost_of_sales": 3300, "admin_expenses": 1000},
        },
        financials={
            "FY-Current": {"gross_profit": 1700, "ebitda": 1300, "total_equity": 2300},
        },
        ratios={"FY-Current": {"gross_leverage": 1.54, "dscr": 2.6, "gearing": 0.87}},
    )

    repo_root = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    examples = os.path.join(repo_root, "docs", "examples", "synthetic_co")
    spreading_check_main([
        "--company", "Synthetic Co", "--proposal", "Credit Facility",
        "--financials", os.path.join(examples, "forward_input.json"),
        "--stress-assumptions", os.path.join(examples, "stress_input.json"),
        "--no-update-financials-source",
    ])
    state = read_state("Synthetic Co", "Credit Facility", date_str=date_str, base_dir=base)

    assert state["financials_source"] == "analyst-supplied"
    assert state.get("forecast_source") != "analyst-supplied"
    expected_forward_raw = state["financials"]["FY+1"]["raw"]
    downside_financials = state["downside_case"]["financials"]

    output_dir = export_deal(
        "Synthetic Co", "Credit Facility", "asset_finance", "# Synthetic Co CAM",
        date_str=date_str, base_dir=base,
    )
    wb = openpyxl.load_workbook(os.path.join(output_dir, "Synthetic Co_Credit Facility_Spreading.xlsx"))
    ws = wb["Financial Spreading"]
    label_to_row = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}

    assert ws.cell(row=1, column=4).value == "FY-Current (analyst-supplied)"
    assert ws.cell(row=1, column=5).value == "FY+1"
    assert ws.cell(row=1, column=8).value == "FY+1 (Downside)"
    assert ws.cell(row=1, column=9).value == "FY+2 (Downside)"
    assert ws.cell(row=1, column=10).value == "FY+3 (Downside)"
    assert ws.cell(row=label_to_row["Revenue"], column=4).value == 5000
    assert ws.cell(row=label_to_row["EBITDA"], column=4).value == 1300
    from spreading_builder import FIELD_LABELS, FORWARD_PERIOD_KEYS

    # Every recorded raw input must survive in its own base-case period column.
    assert set(expected_forward_raw).issubset(FIELD_LABELS)
    for raw_key, expected in expected_forward_raw.items():
        row_label = FIELD_LABELS[raw_key]
        assert ws.cell(row=label_to_row[row_label], column=5).value == expected, row_label

    # The source-selection fix must not change any downside raw cells: compare
    # every period already recorded in state.json, not only FY+1.
    assert "FY+1" in downside_financials
    for offset, period in enumerate(FORWARD_PERIOD_KEYS):
        period_data = downside_financials.get(period)
        if not isinstance(period_data, dict):
            continue
        for raw_key, expected in period_data.get("raw", {}).items():
            if raw_key in FIELD_LABELS:
                row_label = FIELD_LABELS[raw_key]
                assert ws.cell(row=label_to_row[row_label], column=8 + offset).value == expected, (
                    period, row_label
                )



def test_financial_data_from_state_keeps_supplied_forecast_periods_in_supplied_store():
    """Rule: R-EXPORT-002"""
    state = {
        "financials_source": "analyst-supplied",
        "forecast_source": "analyst-supplied",
        "analyst_supplied_financials": {
            "FY-Current": {"revenue": 5000},
            "FY+1": {"revenue": 4600},
        },
        # Even if another source has a raw record, the supplied forecast path
        # stays in the supplied store and is not replaced by computed figures.
        "financials": {"FY+1": {"raw": {"revenue": 9999}}},
    }

    assert _financial_data_from_state(state) == {
        "FY-Current": {"revenue": 5000},
        "FY+1": {"revenue": 4600},
    }


def test_financial_data_from_state_drops_supplied_forward_period_without_computed_counterpart():
    """A forward period only in the supplied store is dropped, never used as a fallback (PX354-2, #249).
    Partial check only: this covers the computed-selected forward direction, not the analyst-supplied historical direction.
    Rule: R-EXPORT-003"""
    state = {
        "financials_source": "analyst-supplied",
        "analyst_supplied_financials": {
            "FY-Current": {"revenue": 5000},
            "FY+1": {"revenue": 4600},
            "FY+2": {"revenue": 4700},
        },
        "financials": {"FY+2": {"raw": {"revenue": 4800}}},
    }

    assert _financial_data_from_state(state) == {
        "FY-Current": {"revenue": 5000},
        "FY+2": {"revenue": 4800},
    }

# ---------------------------------------------------------------------------
# Mutation-survivor contracts for raw-store selection and analyst annotations.
# Kept in this mutation-selected file even where broader integration coverage
# lives elsewhere, because the mutation job only executes this fixed file.
# ---------------------------------------------------------------------------

def test_dx_raw_financials_projects_only_raw_values_and_keeps_empty_periods():
    from deal_export import _raw_financials

    assert _raw_financials({
        "FY-Current": {
            "raw": {"revenue": 0, "cost_of_sales": None},
            "ebitda": 999,
            "ratios": {"dscr": 2.0},
        },
        "FY-1": {"ebitda": 50},
        "FY-2": None,
        "FY+1": "malformed period record",
    }) == {
        "FY-Current": {"revenue": 0, "cost_of_sales": None},
        "FY-1": {},
        "FY-2": {},
        "FY+1": {},
    }
    assert _raw_financials(None) == {}


def test_dx_analyst_forecast_columns_returns_four_none_values_when_neither_basis_is_supplied():
    from deal_export import _analyst_forecast_columns

    assert _analyst_forecast_columns({
        "financials": {"FY-Current": {"ebitda": 100}},
        "ratios": {"FY-Current": {"dscr": 2.0}},
        "downside_case": {"basis": {"FY+1": "analyst-supplied"}},
    }) == (None, None, None, None)


def test_dx_analyst_forecast_columns_selects_only_historical_supplied_periods_without_downside_outputs():
    from deal_export import _analyst_forecast_columns

    supplied, downside, unavailable, derived = _analyst_forecast_columns({
        "financials_source": "analyst-supplied",
        "forecast_source": "framework-computed",
        "financials": {
            "FY-2": {"ebitda": 12, "net_profit": None},
            "FY-1": "not a period object",
            "FY-Current": {},
            "FY+1": {"ebitda": 999},
        },
        "ratios": {
            "FY-1": {"dscr": None},
            "FY+2": {"dscr": 2.0},
        },
        "downside_case": {
            "basis": {"FY+1": "analyst-supplied"},
            "financials": {"FY+1": {"ebitda": 50}},
        },
    })

    assert list(supplied) == ["FY-2", "FY-1", "FY-Current"]
    assert supplied == {
        "FY-2": {"EBITDA": 12},
        "FY-1": {"DSCR": "N/A"},
        "FY-Current": {},
    }
    assert downside is None
    assert unavailable is None
    assert derived is None


def test_dx_analyst_forecast_columns_filters_downside_basis_availability_and_numeric_raw_fields():
    from deal_export import _analyst_forecast_columns

    supplied, supplied_downside, unavailable, derived = _analyst_forecast_columns({
        "financials_source": "analyst-supplied",
        "forecast_source": "analyst-supplied",
        "financials": {
            "FY+1": {"ebitda": 101, "net_profit": None},
            "FY+2": "malformed financials",
            "FY+3": {},
            "FY-1": {"ebitda": 777},
        },
        "ratios": {
            "FY+1": {"dscr": 1.7},
            "FY+2": {"current_ratio": None},
            "FY+3": "malformed ratios",
        },
        "downside_case": {
            "basis": {
                "FY+1": "analyst-supplied",
                "FY+2": "framework-computed",
                "FY+3": "framework-computed",
                "FY-1": "analyst-supplied",
                "FY+4": "analyst-supplied",
            },
            "financials": {
                "FY+1": {"ebitda": 90},
                "FY-1": {"ebitda": 700},
            },
            "ratios": {"FY+1": {"dscr": None}},
            "unavailable": {
                "FY+2": "no positive denominator",
                "FY-1": "not a forward period",
                "FY+4": "outside the model horizon",
            },
        },
        "analyst_supplied_financials": {
            "FY+1": {"revenue": 111, "cash": 10},
            "FY+2": {"revenue": 222, "flag": True},
            "FY+3": {"revenue": 333, "cash": 0, "flag": False, "note": "not numeric"},
            "FY+4": {"revenue": 444},
        },
    })

    assert supplied == {
        "FY-1": {"EBITDA": 777},
        "FY+1": {"EBITDA": 101, "DSCR": 1.7},
        "FY+2": {"Current Ratio": "N/A"},
        "FY+3": {},
    }
    assert supplied_downside == {"FY+1": {"EBITDA": 90, "DSCR": "N/A"}}
    assert unavailable == {"FY+2"}
    assert derived == {"FY+3": {"revenue", "cash"}}


def test_dx_analyst_forecast_columns_normalizes_malformed_nested_downside_and_raw_stores():
    from deal_export import _analyst_forecast_columns

    assert _analyst_forecast_columns({
        "forecast_source": "analyst-supplied",
        "financials": {"FY+1": []},
        "ratios": [],
        "downside_case": {
            "basis": [],
            "financials": "malformed",
            "ratios": None,
            "unavailable": ["FY+2"],
        },
        "analyst_supplied_financials": "malformed",
    }) == (None, None, set(), {})


def test_dx_export_preserves_supplied_forecast_values_and_distinguishes_downside_labels(tmp_path):
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "asset_finance_cam.md"))
    write_state(
        "Synthetic Co", "Credit Facility", date_str="2026-10-04", base_dir=base,
        financials_source="analyst-supplied",
        forecast_source="analyst-supplied",
        analyst_supplied_financials={
            "FY-Current": {"revenue": 500, "cost_of_sales": 250},
            "FY+1": {"revenue": 888, "cost_of_sales": 444},
        },
        financials={
            "FY-Current": {"ebitda": 120, "raw": {"revenue": 999}},
            "FY+1": {"ebitda": 600, "raw": {"revenue": 999}},
        },
        ratios={
            "FY+1": {"dscr": 2.2},
        },
        downside_case={
            "basis": {
                "FY+1": "analyst-supplied",
                "FY+2": "framework-computed",
                "FY+3": "framework-computed",
            },
            "financials": {
                "FY+1": {"ebitda": 123, "raw": {"revenue": 777}},
                "FY+3": {"raw": {"revenue": 333, "cost_of_sales": 100}},
            },
            "ratios": {"FY+1": {"dscr": None}},
            "unavailable": {"FY+2": "denominator not positive"},
        },
    )

    output_dir = export_deal(
        "Synthetic Co", "Credit Facility", "asset_finance", "# Synthetic CAM",
        date_str="2026-10-04", base_dir=base,
    )
    workbook = openpyxl.load_workbook(
        os.path.join(output_dir, "Synthetic Co_Credit Facility_Spreading.xlsx"),
    )
    sheet = workbook["Financial Spreading"]
    row_by_label = {
        sheet.cell(row=row, column=1).value: row
        for row in range(1, sheet.max_row + 1)
    }

    assert sheet.cell(row=1, column=4).value == "FY-Current (analyst-supplied)"
    assert sheet.cell(row=1, column=5).value == "FY+1 (analyst-supplied)"
    assert sheet.cell(row=1, column=8).value == "FY+1 (Downside, analyst-supplied)"
    assert sheet.cell(row=1, column=9).value == "FY+2 (Downside, unavailable)"
    assert sheet.cell(row=row_by_label["Revenue"], column=5).value == 888
    assert sheet.cell(row=row_by_label["Revenue"], column=5).value != 999
    assert sheet.cell(row=row_by_label["EBITDA"], column=5).value == 600
    assert sheet.cell(row=row_by_label["DSCR"], column=5).value == 2.2
    assert sheet.cell(row=row_by_label["EBITDA"], column=8).value == 123
    assert sheet.cell(row=row_by_label["DSCR"], column=8).value == "N/A"
    assert sheet.cell(row=row_by_label["Revenue"], column=9).value is None


def test_dx_cli_rejects_non_utf8_draft_as_one_error_without_creating_outputs(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "draft.md"
    draft.write_bytes(b"\xffnot UTF-8")

    with pytest.raises(SystemExit) as caught:
        main([
            "--company", "Synthetic Co", "--proposal", "Credit Facility",
            "--draft", str(draft),
        ])

    assert str(caught.value.code) == (
        "error: 'utf-8' codec can't decode byte 0xff in position 0: invalid start byte"
    )
    output = capsys.readouterr()
    assert output.out == ""
    assert output.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()



def _group_basis_for_export_tests():
    return {
        "basis": "supplied_consolidated",
        "reporting_entity": "Synthetic Group Ltd",
        "entities_covered": ["Synthetic Group Ltd", "Synthetic Group OpCo"],
        "source": {
            "document": "Borrower-supplied consolidated accounts",
            "accounting_basis": "UK GAAP, consolidated",
            "date": "2026-08-31",
        },
        "period_end": "2026-03-31",
        "periods": ["FY-Current", "FY-1"],
        "figures_basis": "supplied",
        "perimeter_confirmed": True,
        "perimeter_confirmed_on": "2026-10-21",
    }


def test_export_deal_adds_group_basis_disclosure_only_when_metadata_is_present(tmp_path):
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "corporate_credit_cam.md"))
    date_str = "2026-01-01"
    write_state(
        "Synthetic Group Ltd", "Synthetic Group Facility", date_str=date_str, base_dir=base,
        financials_source="analyst-supplied",
        analyst_supplied_financials={"FY-Current": {"revenue": 500}},
        group_basis=_group_basis_for_export_tests(),
    )
    output = export_deal(
        "Synthetic Group Ltd", "Synthetic Group Facility", "corporate_credit", "# Synthetic CAM",
        date_str=date_str, base_dir=base,
    )
    workbook = openpyxl.load_workbook(os.path.join(output, "Synthetic Group Ltd_Synthetic Group Facility_Spreading.xlsx"))
    assert workbook.sheetnames == ["Financial Spreading", "Collateral & Exposure", "Group basis"]
    sheet = workbook["Group basis"]
    assert sheet["A1"].value == (
        "Credit aggregation of consolidated accounts supplied by the borrower; "
        "not a statutory consolidation prepared by this framework."
    )
    rows = {sheet.cell(row=i, column=1).value: sheet.cell(row=i, column=2).value
            for i in range(2, sheet.max_row + 1)}
    assert rows["Reporting entity"] == "Synthetic Group Ltd"
    assert rows["Entities covered"] == "Synthetic Group Ltd; Synthetic Group OpCo"
    assert rows["Source document"] == "Borrower-supplied consolidated accounts"
    assert rows["Accounting basis"] == "UK GAAP, consolidated"
    assert rows["Figures basis"] == "supplied"
    assert rows["Perimeter confirmed"] == "Yes"
    assert rows["Perimeter confirmed on"] == "2026-10-21"


def test_export_deal_refuses_malformed_group_basis_before_creating_exports(tmp_path):
    base = str(tmp_path)
    state_file = os.path.join(base, "deals", "Synthetic Group Ltd", "Synthetic Group Facility_2026-01-01", "state.json")
    _write(state_file, json.dumps({"group_basis": {"basis": "not-supported"}}))
    _write(os.path.join(base, "templates", "cam", "corporate_credit_cam.md"))
    with pytest.raises(ValueError, match="group_basis"):
        export_deal(
            "Synthetic Group Ltd", "Synthetic Group Facility", "corporate_credit", "# Synthetic CAM",
            date_str="2026-01-01", base_dir=base,
        )
    assert not os.path.exists(os.path.join(base, "deals", "Synthetic Group Ltd",
                                           "Synthetic Group Facility_2026-01-01",
                                           "Synthetic Group Ltd_Synthetic Group Facility_Spreading.xlsx"))


def test_export_deal_refuses_supplied_group_basis_when_the_financials_source_is_not_recorded(tmp_path):
    """Absent financials_source is read as framework-computed by the exporter, so 'supplied' would mislabel the
    Financial Spreading figures (review CH4)."""
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "corporate_credit_cam.md"))
    write_state(
        "Synthetic Group Ltd", "Synthetic Group Facility", date_str="2026-01-01", base_dir=base,
        financials={"FY-Current": {"raw": {"revenue": 100}}}, group_basis=_group_basis_for_export_tests(),
    )
    with pytest.raises(ValueError, match="not recorded"):
        export_deal(
            "Synthetic Group Ltd", "Synthetic Group Facility", "corporate_credit", "# Synthetic CAM",
            date_str="2026-01-01", base_dir=base,
        )
    assert not os.path.exists(os.path.join(base, "deals", "Synthetic Group Ltd", "Synthetic Group Facility_2026-01-01",
                                           "Synthetic Group Ltd_Synthetic Group Facility_Spreading.xlsx"))


# ---------------------------------------------------------------------------
# Issue #182 mutation survivors (2026-10-21): deal_export boundaries and main()
# ---------------------------------------------------------------------------

@pytest.mark.parametrize(
    ("bad_path", "expected"),
    [
        (
            "financials",
            (
                {"FY+1": {"Current Ratio": 1.5}},
                {"FY+1": {"EBITDA": 18, "DSCR": 2.0}},
                set(),
                {"FY+2": {"revenue", "cash"}},
            ),
        ),
        (
            "ratios",
            (
                {"FY+1": {"EBITDA": 20}},
                {"FY+1": {"EBITDA": 18, "DSCR": 2.0}},
                set(),
                {"FY+2": {"revenue", "cash"}},
            ),
        ),
        (
            "downside_case",
            (
                {"FY+1": {"EBITDA": 20, "Current Ratio": 1.5}},
                None,
                set(),
                {},
            ),
        ),
        (
            "downside_case.basis",
            (
                {"FY+1": {"EBITDA": 20, "Current Ratio": 1.5}},
                None,
                set(),
                {},
            ),
        ),
        (
            "downside_case.financials",
            (
                {"FY+1": {"EBITDA": 20, "Current Ratio": 1.5}},
                {"FY+1": {"DSCR": 2.0}},
                set(),
                {"FY+2": {"revenue", "cash"}},
            ),
        ),
        (
            "downside_case.ratios",
            (
                {"FY+1": {"EBITDA": 20, "Current Ratio": 1.5}},
                {"FY+1": {"EBITDA": 18}},
                set(),
                {"FY+2": {"revenue", "cash"}},
            ),
        ),
        (
            "downside_case.unavailable",
            (
                {"FY+1": {"EBITDA": 20, "Current Ratio": 1.5}},
                {"FY+1": {"EBITDA": 18, "DSCR": 2.0}},
                set(),
                {"FY+2": {"revenue", "cash"}},
            ),
        ),
        (
            "analyst_supplied_financials",
            (
                {"FY+1": {"EBITDA": 20, "Current Ratio": 1.5}},
                {"FY+1": {"EBITDA": 18, "DSCR": 2.0}},
                set(),
                {"FY+2": set()},
            ),
        ),
    ],
)
def test_dx_analyst_forecast_columns_normalizes_each_non_dict_store_to_its_exact_result(bad_path, expected):
    """Every dict guard is observable: malformed stores are empty, not a traceback or unintended data."""
    from deal_export import _analyst_forecast_columns

    state = {
        "financials_source": "analyst-supplied",
        "forecast_source": "analyst-supplied",
        "financials": {"FY+1": {"ebitda": 20}},
        "ratios": {"FY+1": {"current_ratio": 1.5}},
        "downside_case": {
            "basis": {"FY+1": "analyst-supplied", "FY+2": "framework-computed"},
            "financials": {"FY+1": {"ebitda": 18}},
            "ratios": {"FY+1": {"dscr": 2.0}},
            "unavailable": {},
        },
        "analyst_supplied_financials": {
            "FY+2": {"revenue": 222, "cash": 0, "flag": True, "note": "not numeric"},
        },
    }
    if "." in bad_path:
        parent, key = bad_path.split(".", 1)
        state[parent][key] = 17
    else:
        state[bad_path] = 17

    assert _analyst_forecast_columns(state) == expected


def test_dx_export_deal_passes_each_component_to_the_sanitizer_with_its_exact_label(tmp_path, monkeypatch):
    import deal_export

    original_sanitizer = deal_export.sanitize_path_component
    sanitizer_calls = []

    def recording_sanitizer(value, field_name):
        sanitizer_calls.append((value, field_name))
        return original_sanitizer(value, field_name)

    monkeypatch.setattr(deal_export, "sanitize_path_component", recording_sanitizer)
    monkeypatch.setattr(deal_export, "read_state", lambda *args, **kwargs: {})
    monkeypatch.setattr(deal_export, "existing_state_date_str", lambda *args, **kwargs: None)
    monkeypatch.setattr(deal_export, "cam_template_path", lambda *args, **kwargs: "synthetic-template.md")
    monkeypatch.setattr(deal_export, "export_to_docx", lambda *args, **kwargs: None)
    monkeypatch.setattr(deal_export, "export_to_xlsx", lambda *args, **kwargs: None)

    returned = export_deal(
        "Synthetic Co", "Credit Facility", "asset_finance", "# Synthetic draft",
        date_str="2026-03-12", base_dir=str(tmp_path),
    )

    assert sanitizer_calls == [
        ("Synthetic Co", "company"),
        ("Credit Facility", "proposal"),
        ("asset_finance", "deal_type"),
        ("2026-03-12", "date_str"),
    ]
    expected_folder = tmp_path / "deals" / "Synthetic Co" / "Credit Facility_2026-03-12"
    assert Path(returned) == expected_folder
    assert expected_folder.is_dir()
    assert not (tmp_path / "templates").exists()


def test_dx_export_deal_validates_state_before_creating_any_folder(tmp_path, monkeypatch):
    """Rule: R-EXPORT-005"""
    import deal_export

    monkeypatch.setattr(deal_export, "read_state", lambda *args, **kwargs: {"financials": {}})

    def reject_state(state, *, keys, path):
        assert state == {"financials": {}}
        assert keys == ("group_basis",)
        raise ValueError("Synthetic Co state validation failed")

    monkeypatch.setattr(deal_export, "validate_state", reject_state)

    with pytest.raises(ValueError) as caught:
        export_deal(
            "Synthetic Co", "Credit Facility", "asset_finance", "# Synthetic draft",
            base_dir=str(tmp_path),
        )

    assert str(caught.value) == "Synthetic Co state validation failed"
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()


def test_dx_export_deal_passes_exact_state_data_to_writers_and_returns_the_exact_folder(tmp_path, monkeypatch):
    import deal_export

    group_basis = {"basis": "supplied_consolidated", "reporting_entity": "Synthetic Co"}
    state = {"group_basis": group_basis}
    financial_data = {"FY-Current": {"raw": {"revenue": 1100}}}
    downside_financial_data = {"FY+1": {"raw": {"revenue": 900}}}
    collateral_data = {"security": "synthetic security"}
    supplied_values = {"FY-Current": {"EBITDA": 120}}
    supplied_downside_values = {"FY+1": {"EBITDA": 90}}
    blank_downside_periods = {"FY+2"}
    derived_downside_fields = {"FY+3": {"revenue"}}
    events = []
    reads = []
    validations = []
    docx_calls = []
    xlsx_calls = []
    expected_folder = tmp_path / "deals" / "Synthetic Co" / "Credit Facility_2026-03-12"

    def fake_read_state(*args, **kwargs):
        events.append("read_state")
        reads.append((args, kwargs))
        return state

    def fake_validate_state(*args, **kwargs):
        events.append("validate_state")
        validations.append((args, kwargs))

    monkeypatch.setattr(deal_export, "read_state", fake_read_state)
    monkeypatch.setattr(deal_export, "validate_state", fake_validate_state)
    monkeypatch.setattr(deal_export, "existing_state_date_str", lambda *args, **kwargs: None)
    monkeypatch.setattr(deal_export, "cam_template_path", lambda *args, **kwargs: "synthetic-template.md")

    def get_financials(received_state):
        events.append("financial_data")
        assert received_state is state
        return financial_data

    def get_downside(received_state):
        events.append("downside_financial_data")
        assert received_state is state
        return downside_financial_data

    def get_collateral(received_state):
        events.append("collateral_data")
        assert received_state is state
        return collateral_data

    def get_forecast_columns(received_state):
        events.append("forecast_columns")
        assert received_state is state
        return supplied_values, supplied_downside_values, blank_downside_periods, derived_downside_fields

    def fake_docx_writer(draft, path):
        events.append("docx_writer")
        docx_calls.append((draft, path, Path(path).parent.is_dir()))

    def fake_xlsx_writer(company, path, **kwargs):
        events.append("xlsx_writer")
        xlsx_calls.append((company, path, kwargs))

    monkeypatch.setattr(deal_export, "_financial_data_from_state", get_financials)
    monkeypatch.setattr(deal_export, "_downside_financial_data_from_state", get_downside)
    monkeypatch.setattr(deal_export, "_collateral_data_from_state", get_collateral)
    monkeypatch.setattr(deal_export, "_analyst_forecast_columns", get_forecast_columns)
    monkeypatch.setattr(deal_export, "export_to_docx", fake_docx_writer)
    monkeypatch.setattr(deal_export, "export_to_xlsx", fake_xlsx_writer)

    returned = export_deal(
        "Synthetic Co", "Credit Facility", "asset_finance", "# Synthetic draft",
        date_str="2026-03-12", base_dir=str(tmp_path),
    )

    expected_state_path = deal_export.state_path(
        "Synthetic Co", "Credit Facility", date_str="2026-03-12", base_dir=str(tmp_path)
    )
    assert reads == [
        (
            ("Synthetic Co", "Credit Facility"),
            {"date_str": "2026-03-12", "base_dir": str(tmp_path), "keys": deal_export.STATE_KEYS_READ},
        )
    ]
    assert validations == [
        ((state,), {"keys": ("group_basis",), "path": expected_state_path})
    ]
    assert events == [
        "read_state", "validate_state", "forecast_columns", "financial_data",
        "collateral_data", "downside_financial_data", "docx_writer", "xlsx_writer",
    ]
    expected_docx = expected_folder / "Synthetic Co_Credit Facility_CAM.docx"
    expected_xlsx = expected_folder / "Synthetic Co_Credit Facility_Spreading.xlsx"
    assert len(docx_calls) == 1
    assert docx_calls[0][0] == "# Synthetic draft"
    assert Path(docx_calls[0][1]) == expected_docx
    assert docx_calls[0][2] is True
    assert len(xlsx_calls) == 1
    xlsx_company, xlsx_path, xlsx_kwargs = xlsx_calls[0]
    assert xlsx_company == "Synthetic Co"
    assert Path(xlsx_path) == expected_xlsx
    assert xlsx_kwargs == {
        "financial_data": financial_data,
        "collateral_data": collateral_data,
        "downside_financial_data": downside_financial_data,
        "supplied_values": supplied_values,
        "supplied_downside_values": supplied_downside_values,
        "blank_downside_periods": blank_downside_periods,
        "derived_downside_fields": derived_downside_fields,
        "group_basis": group_basis,
    }
    assert Path(returned) == expected_folder
    assert expected_folder.is_dir()
    assert not (tmp_path / "templates").exists()


def test_dx_main_reads_the_draft_as_utf8_and_passes_exact_resolved_arguments(tmp_path, monkeypatch, capsys):
    import builtins
    import deal_export

    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "synthetic-draft.md"
    draft_text = "# Synthetic Co\n\nFacility fee: £1m — café terms.\n"
    draft.write_text(draft_text, encoding="utf-8")
    real_open = builtins.open
    draft_opens = []
    validations = []
    resolutions = []
    exports = []
    resolved_date = "2026-04-03"
    output_folder = tmp_path / "deals" / "Synthetic Co" / "Credit Facility_2026-04-03"

    def recording_open(file, *args, **kwargs):
        if os.fspath(file) == str(draft):
            draft_opens.append((os.fspath(file), args, dict(kwargs)))
        return real_open(file, *args, **kwargs)

    def record_date_validation(*args, **kwargs):
        validations.append((args, kwargs))

    def record_date_resolution(company, proposal, date_str=None):
        resolutions.append((company, proposal, date_str))
        return resolved_date

    def record_export(*args, **kwargs):
        exports.append((args, kwargs))
        return str(output_folder)

    monkeypatch.setattr(builtins, "open", recording_open)
    monkeypatch.setattr(deal_export, "validate_iso_date_str", record_date_validation)
    monkeypatch.setattr(deal_export, "resolve_export_date_str", record_date_resolution)
    monkeypatch.setattr(deal_export, "export_deal", record_export)

    main([
        "--company", "Synthetic Co", "--proposal", "Credit Facility",
        "--draft", str(draft),
    ])

    assert draft_opens == [(str(draft), (), {"encoding": "utf-8"})]
    assert validations == []
    assert resolutions == [("Synthetic Co", "Credit Facility", None)]
    assert exports == [
        (
            ("Synthetic Co", "Credit Facility", "corporate_credit", draft_text),
            {"date_str": resolved_date},
        )
    ]
    captured = capsys.readouterr()
    assert captured.out == f"Done! Files generated in {output_folder}\n"
    assert captured.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()


def test_dx_main_converts_a_draft_read_oserror_to_the_exact_exit_message(tmp_path, monkeypatch, capsys):
    import builtins
    import deal_export

    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "unreadable-synthetic-draft.md"
    real_open = builtins.open

    def fail_draft_read(file, *args, **kwargs):
        if os.fspath(file) == str(draft):
            raise OSError("synthetic draft read failure")
        return real_open(file, *args, **kwargs)

    def unexpected_resolution(*args, **kwargs):
        pytest.fail("date discovery ran after the draft read failed")

    def unexpected_export(*args, **kwargs):
        pytest.fail("export ran after the draft read failed")

    monkeypatch.setattr(builtins, "open", fail_draft_read)
    monkeypatch.setattr(deal_export, "resolve_export_date_str", unexpected_resolution)
    monkeypatch.setattr(deal_export, "export_deal", unexpected_export)

    with pytest.raises(SystemExit) as caught:
        main([
            "--company", "Synthetic Co", "--proposal", "Credit Facility",
            "--draft", str(draft),
        ])

    assert caught.value.code == "error: synthetic draft read failure"
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()


def test_dx_main_converts_date_validation_failure_to_the_exact_exit_message(tmp_path, monkeypatch, capsys):
    import deal_export

    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "synthetic-draft.md"
    draft.write_text("# Synthetic draft\n", encoding="utf-8")
    validations = []

    def reject_date(value, *, field_name):
        validations.append((value, field_name))
        raise ValueError("synthetic invalid calendar date")

    def unexpected_resolution(*args, **kwargs):
        pytest.fail("date discovery ran after date validation failed")

    def unexpected_export(*args, **kwargs):
        pytest.fail("export ran after date validation failed")

    monkeypatch.setattr(deal_export, "validate_iso_date_str", reject_date)
    monkeypatch.setattr(deal_export, "resolve_export_date_str", unexpected_resolution)
    monkeypatch.setattr(deal_export, "export_deal", unexpected_export)

    with pytest.raises(SystemExit) as caught:
        main([
            "--company", "Synthetic Co", "--proposal", "Credit Facility",
            "--draft", str(draft), "--date-str", "2026-02-30",
        ])

    assert validations == [("2026-02-30", "--date-str")]
    assert caught.value.code == "error: synthetic invalid calendar date"
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()


def test_dx_main_sanitizes_company_before_date_discovery_with_an_exact_error(tmp_path, monkeypatch, capsys):
    import deal_export

    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "synthetic-draft.md"
    draft.write_text("# Synthetic draft\n", encoding="utf-8")
    sanitizer_calls = []

    def reject_unsafe_company(value, field_name):
        sanitizer_calls.append((value, field_name))
        if field_name == "company":
            raise ValueError("synthetic company rejected")
        return value

    def unexpected_resolution(*args, **kwargs):
        pytest.fail("date discovery ran before company sanitization succeeded")

    monkeypatch.setattr(deal_export, "sanitize_path_component", reject_unsafe_company)
    monkeypatch.setattr(deal_export, "resolve_export_date_str", unexpected_resolution)

    with pytest.raises(SystemExit) as caught:
        main([
            "--company", "../synthetic-unsafe", "--proposal", "Credit Facility",
            "--draft", str(draft),
        ])

    assert sanitizer_calls == [("../synthetic-unsafe", "company")]
    assert caught.value.code == "error: synthetic company rejected"
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()



def test_dx_main_validates_an_explicit_date_before_opening_the_draft(tmp_path, monkeypatch, capsys):
    import builtins
    import deal_export

    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "missing-synthetic-draft.md"
    real_open = builtins.open
    open_attempts = []

    def reject_impossible_date(value, *, field_name):
        assert value == "2026-02-30"
        assert field_name == "--date-str"
        raise ValueError("--date-str '2026-02-30' is not a valid calendar date")

    def recording_open(file, *args, **kwargs):
        if os.fspath(file) == str(draft):
            open_attempts.append(os.fspath(file))
        return real_open(file, *args, **kwargs)

    def unexpected_resolution(*args, **kwargs):
        pytest.fail("date discovery ran after the explicit date had been rejected")

    monkeypatch.setattr(deal_export, "validate_iso_date_str", reject_impossible_date)
    monkeypatch.setattr(builtins, "open", recording_open)
    monkeypatch.setattr(deal_export, "resolve_export_date_str", unexpected_resolution)

    with pytest.raises(SystemExit) as caught:
        main([
            "--company", "Synthetic Co", "--proposal", "Credit Facility",
            "--draft", str(draft), "--date-str", "2026-02-30",
        ])

    assert caught.value.code == "error: --date-str '2026-02-30' is not a valid calendar date"
    assert open_attempts == []
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()



@pytest.mark.parametrize(
    ("field", "value"),
    [
        ("company", os.path.join("..", "synthetic-unsafe")),
        ("proposal", os.path.join("..", "synthetic-unsafe")),
        ("deal_type", os.path.join("..", "synthetic-unsafe")),
        ("date_str", os.path.join("..", "synthetic-unsafe")),
    ],
)
def test_dx_export_deal_rejects_each_unsafe_path_component_before_writing(tmp_path, field, value):
    arguments = {
        "company": "Synthetic Co",
        "proposal": "Credit Facility",
        "deal_type": "asset_finance",
        "date_str": "2026-03-12",
    }
    arguments[field] = value
    expected = (
        f"{field} {value!r} isn't safe to use as a filesystem path "
        f"component (no {chr(92)} / : * ? \" < > | characters, and not \".\" or \"..\")."
    )

    with pytest.raises(ValueError) as caught:
        export_deal(
            arguments["company"], arguments["proposal"], arguments["deal_type"],
            "# Synthetic draft", date_str=arguments["date_str"], base_dir=str(tmp_path),
        )

    assert str(caught.value) == expected
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()



@pytest.mark.parametrize(
    ("failure_type", "failure_message"),
    [
        (OSError, "synthetic export write failure"),
        (ValueError, "synthetic export validation failure"),
    ],
)
def test_dx_main_converts_export_errors_to_exact_error_exit_messages(
    tmp_path, monkeypatch, capsys, failure_type, failure_message
):
    import deal_export

    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "synthetic-draft.md"
    draft.write_text("# Synthetic draft\n", encoding="utf-8")
    monkeypatch.setattr(
        deal_export, "resolve_export_date_str",
        lambda company, proposal, date_str=None: "2026-04-03",
    )

    def fail_export(*args, **kwargs):
        raise failure_type(failure_message)

    monkeypatch.setattr(deal_export, "export_deal", fail_export)

    with pytest.raises(SystemExit) as caught:
        main([
            "--company", "Synthetic Co", "--proposal", "Credit Facility",
            "--draft", str(draft),
        ])

    assert caught.value.code == f"error: {failure_message}"
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()



def test_dx_main_sanitizes_proposal_before_date_discovery_with_an_exact_error(tmp_path, monkeypatch, capsys):
    import deal_export

    monkeypatch.chdir(tmp_path)
    draft = tmp_path / "synthetic-draft.md"
    draft.write_text("# Synthetic draft\n", encoding="utf-8")
    sanitizer_calls = []

    def sanitize(value, field_name):
        sanitizer_calls.append((value, field_name))
        if field_name == "proposal":
            raise ValueError("proposal 'Synthetic Proposal' isn't safe")
        return value

    def unexpected_resolution(*args, **kwargs):
        pytest.fail("date discovery ran before proposal sanitization succeeded")

    monkeypatch.setattr(deal_export, "sanitize_path_component", sanitize)
    monkeypatch.setattr(deal_export, "resolve_export_date_str", unexpected_resolution)

    with pytest.raises(SystemExit) as caught:
        main([
            "--company", "Synthetic Co", "--proposal", "Synthetic Proposal",
            "--draft", str(draft),
        ])

    assert sanitizer_calls == [
        ("Synthetic Co", "company"),
        ("Synthetic Proposal", "proposal"),
    ]
    assert caught.value.code == "error: proposal 'Synthetic Proposal' isn't safe"
    captured = capsys.readouterr()
    assert captured.out == ""
    assert captured.err == ""
    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()


# ---------------------------------------------------------------------------
# Review fixes (CL6, CL7): the mutants a test can still kill. `required=True` on the three flags was invisible to the
# older test above because argparse's usage line names every flag, so `missing in err` is true whatever is missing.
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("argv, required", [
    ([], "--company, --proposal, --draft"),
    (["--proposal", "P", "--draft", "d.md"], "--company"),
    (["--company", "A", "--draft", "d.md"], "--proposal"),
    (["--company", "A", "--proposal", "P"], "--draft"),
])
def test_cli_names_exactly_the_flags_that_are_missing(capsys, argv, required):
    with pytest.raises(SystemExit) as caught:
        main(argv)

    assert caught.value.code == 2
    assert capsys.readouterr().err.rstrip().endswith(f"error: the following arguments are required: {required}")


def test_a_brand_new_deal_is_exported_into_a_folder_dated_year_month_day_by_the_clock(tmp_path, monkeypatch):
    from datetime import datetime

    import deal_export

    class FrozenDatetime:
        @staticmethod
        def now():
            return datetime(2026, 3, 9, 13, 5, 7)

    monkeypatch.setattr(deal_export, "datetime", FrozenDatetime)
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "corporate_credit_cam.md"))

    folder = export_deal("Synthetic Co", "Synthetic Fleet Loan", "corporate_credit", "# memo\n", base_dir=base)

    assert Path(folder) == tmp_path / "deals" / "Synthetic Co" / "Synthetic Fleet Loan_2026-03-09"
    assert (Path(folder) / "Synthetic Co_Synthetic Fleet Loan_CAM.docx").is_file()


def test_the_workbook_is_built_for_the_sanitised_company_and_written_beside_the_document(tmp_path, monkeypatch):
    seen = {}
    monkeypatch.setattr("deal_export.export_to_xlsx", lambda company, path, **kwargs: seen.update(company=company, path=path))
    monkeypatch.setattr("deal_export.export_to_docx", lambda text, path: seen.update(docx=path))
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "corporate_credit_cam.md"))

    folder = export_deal("Synthetic Co", "Synthetic Fleet Loan", "corporate_credit", "# memo\n",
                         date_str="2026-01-05", base_dir=base)

    assert seen["company"] == "Synthetic Co"
    assert Path(seen["path"]) == Path(folder) / "Synthetic Co_Synthetic Fleet Loan_Spreading.xlsx"
    assert Path(seen["docx"]) == Path(folder) / "Synthetic Co_Synthetic Fleet Loan_CAM.docx"


def test_a_second_new_deal_type_is_saved_even_though_the_local_template_folder_already_exists(tmp_path):
    base = str(tmp_path)

    export_deal("Synthetic Co", "Facility One", "first_new_type", "# first draft\n", date_str="2026-01-05",
                base_dir=base)
    export_deal("Synthetic Co", "Facility Two", "second_new_type", "# second draft\n", date_str="2026-01-05",
                base_dir=base)

    local = tmp_path / "templates" / "local" / "cam"
    assert (local / "first_new_type_cam.md").read_text(encoding="utf-8") == "# first draft\n"
    assert (local / "second_new_type_cam.md").read_text(encoding="utf-8") == "# second draft\n"


def test_an_auto_saved_template_keeps_non_ascii_text_even_where_the_default_encoding_is_cp1252(
    tmp_path, cp1252_default_open,
):
    draft = "# Synthetic Co\n\nFacility fee: £1m — café terms, 東京 office.\n"

    export_deal("Synthetic Co", "Facility One", "brand_new_type", draft, date_str="2026-01-05", base_dir=str(tmp_path))

    saved = tmp_path / "templates" / "local" / "cam" / "brand_new_type_cam.md"
    assert saved.read_text(encoding="utf-8") == draft


def test_workbook_inputs_from_state_selects_one_raw_store_per_period_and_passes_the_rest_through():
    group_basis = {"basis": "supplied_consolidated"}
    state = {
        "financials_source": "analyst-supplied",
        "analyst_supplied_financials": {"FY-Current": {"revenue": 500}, "FY+1": {"revenue": 999}},
        "financials": {"FY+1": {"raw": {"revenue": 600}}, "FY-Current": {"revenue": 500, "ebitda": 90}},
        "ratios": {"FY-Current": {"dscr": 1.4}},
        "collateral": [{"asset": "Synthetic asset"}],
        "group_basis": group_basis,
    }

    inputs = workbook_inputs_from_state(state)

    assert set(inputs) == {
        "financial_data", "collateral_data", "downside_financial_data", "supplied_values",
        "supplied_downside_values", "blank_downside_periods", "derived_downside_fields", "group_basis"}
    assert inputs["financial_data"] == {"FY-Current": {"revenue": 500}, "FY+1": {"revenue": 600}}
    assert inputs["collateral_data"] == [{"asset": "Synthetic asset"}]
    assert inputs["group_basis"] is group_basis
    assert inputs["supplied_values"]["FY-Current"] and inputs["downside_financial_data"] is None


def test_workbook_inputs_from_state_for_an_empty_state_is_all_none():
    assert workbook_inputs_from_state({}) == {
        "financial_data": None, "collateral_data": None, "downside_financial_data": None, "supplied_values": None,
        "supplied_downside_values": None, "blank_downside_periods": None, "derived_downside_fields": None,
        "group_basis": None}


def test_the_workbook_export_receives_exactly_the_shared_inputs(tmp_path, monkeypatch):
    seen = {}
    monkeypatch.setattr("deal_export.export_to_xlsx", lambda company, path, **kwargs: seen.update(kwargs))
    monkeypatch.setattr("deal_export.export_to_docx", lambda text, path: None)
    base = str(tmp_path)
    _write(os.path.join(base, "templates", "cam", "corporate_credit_cam.md"))
    write_state("Synthetic Co", "Synthetic Facility", date_str="2026-01-01", base_dir=base,
                financials_source="framework-computed", financials={"FY-Current": {"raw": {"revenue": 10}}})

    export_deal("Synthetic Co", "Synthetic Facility", "corporate_credit", "# memo", date_str="2026-01-01",
                base_dir=base)

    state = read_state("Synthetic Co", "Synthetic Facility", date_str="2026-01-01", base_dir=base)
    assert seen == workbook_inputs_from_state(state)


def test_an_unexportable_draft_creates_no_folder_and_saves_no_template(tmp_path):
    """An unterminated HTML comment is refused before the folder is created or the draft is saved as a template."""
    base = str(tmp_path)
    draft = "# Synthetic CAM\n\n<!-- Guidance: never closed\n\nBody.\n"

    with pytest.raises(ValueError, match="Unterminated HTML comment"):
        export_deal("Acme Corp", "Fleet Loan", "brand_new_type", draft, date_str="2026-01-15", base_dir=base)

    assert not (tmp_path / "deals").exists()
    assert not (tmp_path / "templates").exists()


def test_financial_data_from_state_does_not_fall_back_to_computed_store_for_missing_supplied_historical_period():
    """A missing historical period in the selected supplied store is not filled from computed raw figures.
    Rule: R-EXPORT-003"""
    state = {
        "financials_source": "analyst-supplied",
        "analyst_supplied_financials": {
            "FY-Current": {"revenue": 5000},
        },
        # FY-1 exists only in the other store; a historical fallback would wrongly export it.
        "financials": {"FY-1": {"raw": {"revenue": 4600}}},
    }

    assert _financial_data_from_state(state) == {
        "FY-Current": {"revenue": 5000},
    }


def test_financial_data_from_state_does_not_merge_raw_fields_from_both_stores_within_a_period():
    """The selected historical store supplies the whole period; fields from the other store are not added.
    Rule: R-EXPORT-004"""
    state = {
        "financials_source": "analyst-supplied",
        "analyst_supplied_financials": {
            "FY-Current": {"revenue": 5000},
        },
        # A field absent from the selected store must not be filled from the computed store.
        "financials": {"FY-Current": {"raw": {"cost_of_sales": 3300}}},
    }

    assert _financial_data_from_state(state) == {
        "FY-Current": {"revenue": 5000},
    }
