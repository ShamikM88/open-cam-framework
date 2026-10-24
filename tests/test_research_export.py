import os

import docx
import pytest

import research_export
from research_export import export_research_brief
from state_manager import read_state, write_state


VALID_BRIEF = (
    "# Synthetic Research Brief\n\n## Go/No-Go\n\nGO -- synthetic example.\n\n"
    "## Sources\n\n- Synthetic registry record: https://example.invalid/registry\n"
)
NO_SOURCES_BRIEF = (
    "# Synthetic Research Brief\n\n## Go/No-Go\n\n"
    "The word sources in this paragraph is not a sources section.\n"
)


VALID_BRIEF = (
    "# Synthetic Research Brief\n\n## Go/No-Go\n\nGO -- synthetic example.\n\n"
    "## Sources\n\n- Synthetic registry record: https://example.invalid/registry\n"
)
NO_SOURCES_BRIEF = (
    "# Synthetic Research Brief\n\n## Go/No-Go\n\n"
    "The word sources in this paragraph is not a sources section.\n"
)


def test_creates_dated_output_folder_when_none_exists_yet(tmp_path):
    base = str(tmp_path)
    output_path = export_research_brief(
        "Acme Corp", "Fleet Loan", VALID_BRIEF, date_str="2026-01-15", base_dir=base,
    )
    expected_dir = os.path.join(base, "deals", "Acme Corp", "Fleet Loan_2026-01-15")
    assert output_path == os.path.join(expected_dir, "Acme Corp_Fleet Loan_Research_Brief.docx")
    assert os.path.isfile(output_path)


def test_reuses_an_existing_dated_folder_when_date_str_not_given(tmp_path):
    """The brief belongs alongside whatever state.json /research (or
    /triage + /commercial) already wrote for this deal -- it must not
    create a second, differently-dated folder of its own."""
    base = str(tmp_path)
    write_state("Acme Corp", "Fleet Loan", date_str="2025-06-01", base_dir=base, triage={"go_no_go": "GO"})
    output_path = export_research_brief("Acme Corp", "Fleet Loan", VALID_BRIEF, base_dir=base)
    assert "Fleet Loan_2025-06-01" in output_path
    assert os.path.isfile(output_path)


def test_exports_readable_docx_content(tmp_path):
    output_path = export_research_brief(
        "Acme Corp", "Fleet Loan", VALID_BRIEF, date_str="2026-01-15", base_dir=str(tmp_path),
    )
    doc = docx.Document(output_path)
    assert doc.paragraphs[0].text == "Synthetic Research Brief"


def test_missing_sources_section_is_rejected_before_creating_any_output(tmp_path):
    with pytest.raises(ValueError, match="populated Sources or References section"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", NO_SOURCES_BRIEF, date_str="2026-01-15", base_dir=str(tmp_path),
        )
    assert not (tmp_path / "deals").exists()


def test_empty_sources_section_is_rejected_before_creating_any_output(tmp_path):
    brief = "# Synthetic brief\n\n## Sources\n\n## Findings\n\nNo source list here."
    with pytest.raises(ValueError, match="populated Sources or References section"):
        export_research_brief("Acme Co", "Synthetic Lease", brief, date_str="2026-01-15", base_dir=str(tmp_path))
    assert not (tmp_path / "deals").exists()


@pytest.mark.parametrize("heading", ["# SOURCES", "### rEfErEnCeS ###", "###### Sources  ####"])
def test_populated_sources_heading_is_case_insensitive_any_level_and_allows_closing_hashes(tmp_path, heading):
    brief = f"# Synthetic brief\n\n{heading}\n\n1) Synthetic reference\n"
    output = export_research_brief("Acme Co", "Synthetic Lease", brief, date_str="2026-01-15",
                                   base_dir=str(tmp_path))
    assert os.path.isfile(output)


def test_sources_section_can_be_populated_by_a_url_without_a_list_item(tmp_path):
    brief = "# Synthetic brief\n\n## References\n\nhttps://example.invalid/source\n"
    output = export_research_brief("Acme Co", "Synthetic Lease", brief, date_str="2026-01-15",
                                   base_dir=str(tmp_path))
    assert os.path.isfile(output)


def test_word_sources_in_a_paragraph_does_not_count_as_a_section(tmp_path):
    with pytest.raises(ValueError, match="populated Sources or References section"):
        export_research_brief("Acme Co", "Synthetic Lease", NO_SOURCES_BRIEF, base_dir=str(tmp_path))
    assert not (tmp_path / "deals").exists()


def test_sources_headings_and_entries_inside_fenced_code_are_ignored(tmp_path):
    brief = (
        "# Synthetic brief\n\n```markdown\n## Sources\n- Fake source: https://example.invalid/fake\n```\n"
        "\n## Findings\nNo real source section.\n"
    )
    with pytest.raises(ValueError, match="populated Sources or References section"):
        export_research_brief("Acme Co", "Synthetic Lease", brief, base_dir=str(tmp_path))
    assert not (tmp_path / "deals").exists()


def test_sources_section_stops_at_next_heading_of_same_or_higher_level(tmp_path):
    brief = "# Synthetic brief\n\n## Sources\n\n### Still in this section\n\n## Findings\n- Not a source entry\n"
    with pytest.raises(ValueError, match="populated Sources or References section"):
        export_research_brief("Acme Co", "Synthetic Lease", brief, base_dir=str(tmp_path))
    assert not (tmp_path / "deals").exists()


def test_override_adds_the_disclosed_footnote_as_the_last_docx_paragraph(tmp_path, monkeypatch):
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    reason = "  analyst  explicitly\n authorised\t omission  "
    output = export_research_brief(
        "Acme Co", "Synthetic Lease", NO_SOURCES_BRIEF, date_str="2026-01-15", base_dir=str(tmp_path),
        analyst_override_no_sources=reason,
    )
    doc = docx.Document(output)
    assert doc.paragraphs[-1].text == (
        "Footnote: No sources list is included in this brief. Exported under analyst override on "
        "2026-10-14: analyst explicitly authorised omission."
    )


@pytest.mark.parametrize("reason", ["", "   ", "\n \t"])
def test_override_with_empty_or_whitespace_reason_is_rejected_before_writing(tmp_path, reason):
    with pytest.raises(ValueError, match="reason must not be empty"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", NO_SOURCES_BRIEF, base_dir=str(tmp_path),
            analyst_override_no_sources=reason,
        )
    assert not (tmp_path / "deals").exists()


def test_override_reason_is_required_even_when_sources_exist(tmp_path):
    with pytest.raises(ValueError, match="reason must not be empty"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", VALID_BRIEF, base_dir=str(tmp_path),
            analyst_override_no_sources="  ",
        )
    assert not (tmp_path / "deals").exists()


def test_override_has_no_effect_when_sources_are_already_populated(tmp_path, monkeypatch):
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    output = export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, base_dir=str(tmp_path),
        analyst_override_no_sources="Explicitly authorised synthetic exception",
    )
    doc = docx.Document(output)
    assert not any(p.text.startswith("Footnote: No sources list") for p in doc.paragraphs)


@pytest.mark.parametrize(
    "company,proposal,date_str",
    [
        ("C:\\\\Windows\\\\Temp\\\\evil", "Fleet Loan", None),
        ("Acme Corp", "../../evil", None),
        ("Acme Corp", "Fleet Loan", "../../evil"),
    ],
)
def test_rejects_unsafe_path_components(company, proposal, date_str, tmp_path):
    with pytest.raises(ValueError):
        export_research_brief(company, proposal, VALID_BRIEF, date_str=date_str, base_dir=str(tmp_path))


def test_research_command_requires_a_flat_sources_list_and_explicit_override():
    from pathlib import Path

    command = (Path(__file__).resolve().parents[1] / ".claude" / "commands" / "research.md").read_text(
        encoding="utf-8"
    )
    assert "Mandatory final Sources section" in command
    assert "flat" in command and "docx_builder.py" in command and "nested bullets" in command
    assert "every document the brief relies on" in command
    assert "URL wherever one exists" in command
    assert "citation key" in command and "WRG note 14" in command
    assert 'explicitly authorises exporting without that section, in chat' in command
    assert "record that reason" in command
    assert "Silence, an ambiguous answer, or “just export it” is not explicit authorisation." in command
    assert "Never pass the flag on your own initiative." in command
    assert '--analyst-override-no-sources "<reason>"' in command



def test_fenced_source_items_do_not_populate_a_real_sources_section(tmp_path):
    brief = (
        "# Synthetic brief\n\n## Sources\n\n"
        "```text\n- Fake source: https://example.invalid/fake\n```\n\n"
        "## Findings\nOther synthetic content.\n"
    )
    with pytest.raises(ValueError, match="populated Sources or References section"):
        export_research_brief("Acme Co", "Synthetic Lease", brief, base_dir=str(tmp_path))
    assert not (tmp_path / "deals").exists()


@pytest.mark.parametrize("reason", [123, []])
def test_override_reason_must_be_text(tmp_path, reason):
    with pytest.raises(ValueError, match="reason must be text"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", NO_SOURCES_BRIEF, base_dir=str(tmp_path),
            analyst_override_no_sources=reason,
        )
    assert not (tmp_path / "deals").exists()


def _state_with_trail(base_dir, trail, date_str="2026-10-14"):
    return write_state(
        "Acme Co", "Synthetic Lease", date_str=date_str, base_dir=str(base_dir), review_trail=trail,
    )


def test_unapproved_release_discloses_trailing_rejection_count_near_document_top_and_records_state(
    tmp_path, monkeypatch,
):
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    trail = [
        {"iteration": 1, "verdict": "REJECTED", "notes": "old finding"},
        {"iteration": 2, "verdict": "APPROVED", "notes": "previous review cycle"},
        {"iteration": 3, "verdict": "REJECTED", "notes": "synthetic finding one"},
        {"iteration": 4, "verdict": "REJECTED", "notes": "synthetic finding two"},
        {"iteration": 5, "verdict": "REJECTED", "notes": "synthetic finding three"},
    ]
    _state_with_trail(tmp_path, trail)
    output_path = export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
        analyst_release_unapproved="  analyst requested\n manual review ",
    )
    doc = docx.Document(output_path)
    assert doc.paragraphs[0].text == "Synthetic Research Brief"
    assert doc.paragraphs[1].text == (
        "Review status: Not approved by the independent Risk Reviewer. 3 review passes returned REJECTED. "
        "Released for manual analyst review at the analyst's request on 2026-10-14: analyst requested manual review."
    )
    state = read_state("Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path))
    assert state["review_stopped_by_analyst"] == [
        {"date": "2026-10-14", "passes": 3, "reason": "analyst requested manual review"},
    ]
    assert state["review_trail"] == trail


@pytest.mark.parametrize("reason", ["", "   ", "\n \t"])
def test_unapproved_release_requires_nonempty_reason_before_writing(tmp_path, reason):
    with pytest.raises(ValueError, match="Analyst release reason must not be empty"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", VALID_BRIEF, base_dir=str(tmp_path),
            analyst_release_unapproved=reason,
        )
    assert not (tmp_path / "deals").exists()


def test_unapproved_release_state_is_written_only_after_docx_export(tmp_path, monkeypatch):
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    _state_with_trail(tmp_path, [{"verdict": "REJECTED", "notes": "synthetic finding"}])
    real_export = research_export.export_to_docx
    observed = {}

    def assert_export_then_return(markdown, output_path):
        real_export(markdown, output_path)
        observed["docx_exists"] = os.path.isfile(output_path)
        state = read_state("Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path))
        observed["release_record_present_during_export"] = "review_stopped_by_analyst" in state

    monkeypatch.setattr(research_export, "export_to_docx", assert_export_then_return)
    export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
        analyst_release_unapproved="explicit synthetic release",
    )
    assert observed == {"docx_exists": True, "release_record_present_during_export": False}
    state = read_state("Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path))
    assert state["review_stopped_by_analyst"][0]["passes"] == 1


def test_source_override_and_unapproved_release_combine_status_near_top_and_sources_footnote_last(
    tmp_path, monkeypatch,
):
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    _state_with_trail(tmp_path, [
        {"verdict": "REJECTED", "notes": "synthetic finding one"},
        {"verdict": "REJECTED", "notes": "synthetic finding two"},
        {"verdict": "REJECTED", "notes": "synthetic finding three"},
    ])
    output_path = export_research_brief(
        "Acme Co", "Synthetic Lease", NO_SOURCES_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
        analyst_override_no_sources="analyst authorised missing source list",
        analyst_release_unapproved="analyst explicitly stopped review",
    )
    doc = docx.Document(output_path)
    assert doc.paragraphs[0].text == "Synthetic Research Brief"
    assert doc.paragraphs[1].text.startswith(
        "Review status: Not approved by the independent Risk Reviewer. 3 review passes returned REJECTED."
    )
    assert doc.paragraphs[-1].text == (
        "Footnote: No sources list is included in this brief. Exported under analyst override on "
        "2026-10-14: analyst authorised missing source list."
    )


def test_trailing_rejection_counter_stops_at_latest_non_rejected_entry():
    assert research_export._trailing_rejected_count([
        {"verdict": "REJECTED"}, {"verdict": "APPROVED"},
        {"verdict": "REJECTED"}, {"verdict": "REJECTED"},
    ]) == 2
    assert research_export._trailing_rejected_count([{"verdict": "APPROVED"}]) == 0


def test_unapproved_release_is_refused_when_the_deal_has_no_review_trail(tmp_path):
    with pytest.raises(ValueError, match="no trailing REJECTED entry"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
            analyst_release_unapproved="analyst asked",
        )
    assert not (tmp_path / "deals").exists()


@pytest.mark.parametrize("trail", [
    [{"verdict": "APPROVED", "notes": "clean"}],
    [{"verdict": "REJECTED", "notes": "old"}, {"verdict": "APPROVED", "notes": "latest is approved"}],
])
def test_unapproved_release_is_refused_when_the_latest_verdict_is_approved(tmp_path, trail):
    _state_with_trail(tmp_path, trail)
    state_file = tmp_path / "deals" / "Acme Co" / "Synthetic Lease_2026-10-14" / "state.json"
    before = state_file.read_bytes()
    with pytest.raises(ValueError, match="no trailing REJECTED entry"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
            analyst_release_unapproved="analyst asked",
        )
    assert state_file.read_bytes() == before, "the refusal must not touch state.json"
    assert not list(state_file.parent.glob("*.docx")), "the refusal must not write a document"


@pytest.mark.parametrize(("count", "wording"), [(1, "1 review pass returned REJECTED."),
                                                (2, "2 review passes returned REJECTED.")])
def test_unapproved_release_status_uses_the_singular_for_one_pass(tmp_path, monkeypatch, count, wording):
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    _state_with_trail(tmp_path, [{"verdict": "REJECTED", "notes": "synthetic"}] * count)
    output = export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
        analyst_release_unapproved="analyst asked",
    )
    assert wording in docx.Document(output).paragraphs[1].text


def test_two_releases_keep_both_records_in_order(tmp_path, monkeypatch):
    _state_with_trail(tmp_path, [{"verdict": "REJECTED", "notes": "one"}] * 3)
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
        analyst_release_unapproved="first release",
    )
    write_state(
        "Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path),
        review_trail=[{"verdict": "REJECTED", "notes": "one"}] * 4,
    )
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-16")
    export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
        analyst_release_unapproved="second release",
    )
    state = read_state("Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path))
    assert state["review_stopped_by_analyst"] == [
        {"date": "2026-10-14", "passes": 3, "reason": "first release"},
        {"date": "2026-10-16", "passes": 4, "reason": "second release"},
    ]


def test_an_approved_export_after_a_release_leaves_the_records_and_the_verdict_alone(tmp_path, monkeypatch):
    monkeypatch.setattr(research_export, "_today_date", lambda: "2026-10-14")
    _state_with_trail(tmp_path, [{"verdict": "REJECTED", "notes": "one"}] * 3)
    export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
        analyst_release_unapproved="manual review please",
    )
    write_state(
        "Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path),
        review_verdict="APPROVED",
        review_trail=[{"verdict": "REJECTED", "notes": "one"}] * 3 + [{"verdict": "APPROVED", "notes": "ok"}],
    )
    output = export_research_brief(
        "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
    )
    assert not any(p.text.startswith("Review status:") for p in docx.Document(output).paragraphs)
    state = read_state("Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path))
    assert state["review_stopped_by_analyst"] == [
        {"date": "2026-10-14", "passes": 3, "reason": "manual review please"},
    ]
    assert state["review_verdict"] == "APPROVED"


def test_a_release_refuses_a_state_whose_record_is_not_a_list_and_changes_nothing(tmp_path):
    _state_with_trail(tmp_path, [{"verdict": "REJECTED", "notes": "one"}])
    write_state(
        "Acme Co", "Synthetic Lease", date_str="2026-10-14", base_dir=str(tmp_path),
        review_stopped_by_analyst={"passes": 1},
    )
    state_file = tmp_path / "deals" / "Acme Co" / "Synthetic Lease_2026-10-14" / "state.json"
    before = state_file.read_bytes()
    with pytest.raises(ValueError, match="must be a list"):
        export_research_brief(
            "Acme Co", "Synthetic Lease", VALID_BRIEF, date_str="2026-10-14", base_dir=str(tmp_path),
            analyst_release_unapproved="again",
        )
    assert state_file.read_bytes() == before and not list(state_file.parent.glob("*.docx"))


def test_research_export_rejects_framework_vocabulary_before_creating_output(tmp_path):
    brief = VALID_BRIEF.replace("GO -- synthetic example.", "GO -- this framework approves the opportunity.")
    with pytest.raises(ValueError, match="Rule A: forbidden framework vocabulary.*this framework"):
        export_research_brief(
            "Synthetic Co", "Fleet Loan", brief, date_str="2026-10-24", base_dir=str(tmp_path),
        )
    assert not (tmp_path / "deals").exists()


def test_research_export_rejects_unfilled_placeholder_before_creating_output(tmp_path):
    brief = VALID_BRIEF.replace("GO -- synthetic example.", "GO -- subject to [Amount].")
    with pytest.raises(ValueError, match="Rule B: unfilled template placeholder.*Amount"):
        export_research_brief(
            "Synthetic Co", "Fleet Loan", brief, date_str="2026-10-24", base_dir=str(tmp_path),
        )
    assert not (tmp_path / "deals").exists()


def test_an_unterminated_html_comment_is_rejected_before_creating_any_output(tmp_path):
    brief = VALID_BRIEF + "\n<!-- never closed\n"
    with pytest.raises(ValueError, match="Unterminated HTML comment"):
        export_research_brief("Synthetic Co", "Fleet Loan", brief, date_str="2026-10-24", base_dir=str(tmp_path))
    assert not (tmp_path / "deals").exists()
