"""The export read-back verifier (issue #196): what it accepts, what it reports, and that it never writes to a deal."""
import base64
import json
import tempfile
from pathlib import Path

import docx
import openpyxl
import pytest

import export_readback
from deal_export import export_deal
from state_manager import StateError, state_path, write_state

COMPANY = "Synthetic Co"
PROPOSAL = "Synthetic Fleet Loan"
DATE = "2026-01-05"
FOLDER = Path("deals") / COMPANY / f"{PROPOSAL}_{DATE}"
WORKBOOK = FOLDER / f"{COMPANY}_{PROPOSAL}_Spreading.xlsx"
DOCUMENT = FOLDER / f"{COMPANY}_{PROPOSAL}_CAM.docx"
DRAFT = FOLDER / "draft_approved.md"

PNG_1X1 = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")
OTHER_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg==")

DRAFT_MARKDOWN = """# Synthetic Co credit memo

The facility is a [synthetic link](https://example.invalid/page) with **bold** wording.

- First bullet
- Second bullet

| Metric | Value |
| :--- | ---: |
| Revenue | 5,000 |
| Debt | 2,400 |

![Chart alt text](chart.png)
"""

RAW = {
    "revenue": 5000, "cost_of_sales": 3000, "admin_expenses": 1200, "interest_paid": 50,
    "current_debt": 0, "overdraft": 0, "long_term_debt": 2400, "loan_notes": 0, "scheduled_principal": 200,
}


def build_deal(**state_fields):
    """Write a state and export it, in the current working directory, as `/assemble` would leave it."""
    fields = {"financials_source": "framework-computed", "financials": {"FY-Current": {"raw": RAW}}}
    fields.update(state_fields)
    Path("chart.png").write_bytes(PNG_1X1)
    write_state(COMPANY, PROPOSAL, date_str=DATE, **fields)
    export_deal(COMPANY, PROPOSAL, "corporate_credit", DRAFT_MARKDOWN, date_str=DATE)
    DRAFT.write_text(DRAFT_MARKDOWN, encoding="utf-8")
    write_state(COMPANY, PROPOSAL, date_str=DATE, approved_draft_path=DRAFT.as_posix())


@pytest.fixture(autouse=True)
def private_scratch_root(tmp_path_factory, monkeypatch):
    """Point the temporary directory at one this test owns, so a parallel worker's scratch files cannot be mistaken
    for a leftover of ours."""
    root = tmp_path_factory.mktemp("scratch_root")
    monkeypatch.setattr(tempfile, "tempdir", str(root))
    return root


@pytest.fixture
def deal(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    build_deal()
    return tmp_path


def reexport():
    """Write the unedited artifacts again."""
    export_deal(COMPANY, PROPOSAL, "corporate_credit", DRAFT_MARKDOWN, date_str=DATE)


def verify(**kwargs):
    return export_readback.verify_deal(COMPANY, PROPOSAL, **kwargs)


def edit_workbook(change):
    workbook = openpyxl.load_workbook(WORKBOOK)
    change(workbook)
    workbook.save(WORKBOOK)


def edit_document(change):
    document = docx.Document(DOCUMENT)
    change(document)
    document.save(DOCUMENT)


def cell_holding(workbook, value, sheet="Financial Spreading"):
    for row in workbook[sheet].iter_rows():
        for cell in row:
            if cell.value == value:
                return cell
    raise AssertionError(f"no cell holds {value!r}")


def folder_bytes():
    return {p.as_posix(): p.read_bytes() for p in sorted(Path("deals").rglob("*")) if p.is_file()}


def leftover_scratch_dirs():
    return sorted(p.name for p in Path(tempfile.gettempdir()).iterdir())


# ---------------------------------------------------------------------------
# A faithful export matches
# ---------------------------------------------------------------------------

def test_an_untouched_export_matches_in_both_artifacts(deal):
    report = verify()

    assert report["status"] == "match"
    assert report["mismatch_count"] == 0 and report["findings"] == [] and report["truncated"] is False
    assert report["artifacts"]["workbook"]["status"] == "match"
    assert report["artifacts"]["document"]["status"] == "match"
    assert Path(report["selected_folder"]) == FOLDER
    assert Path(report["artifacts"]["approved_draft"]) == DRAFT
    assert report["limits"] and all(isinstance(line, str) for line in report["limits"])


def test_an_analyst_supplied_export_matches(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    build_deal(financials_source="analyst-supplied", financials={"FY-Current": {"revenue": 500, "ebitda": 90}},
               analyst_supplied_financials={"FY-Current": {"revenue": 500}}, ratios={"FY-Current": {"dscr": 1.4}})

    assert verify()["status"] == "match"


def test_a_workbook_with_a_group_basis_sheet_matches(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    build_deal(
        financials_source="analyst-supplied", analyst_supplied_financials={"FY-Current": {"revenue": 500}},
        financials={"FY-Current": {"revenue": 500}},
        group_basis={
            "basis": "supplied_consolidated", "reporting_entity": "Synthetic Group Ltd",
            "entities_covered": ["Synthetic Group Ltd", "Synthetic Group OpCo"],
            "source": {"document": "Borrower-supplied consolidated accounts", "accounting_basis": "UK GAAP",
                       "date": "2026-08-31"},
            "period_end": "2026-03-31", "periods": ["FY-Current"], "figures_basis": "supplied",
            "perimeter_confirmed": True, "perimeter_confirmed_on": "2026-10-21"})

    assert verify()["status"] == "match"
    assert openpyxl.load_workbook(WORKBOOK).sheetnames[-1] == "Group basis"


def test_opening_and_saving_the_files_unchanged_is_not_reported_as_an_edit(deal):
    edit_workbook(lambda workbook: None)
    edit_document(lambda document: None)

    assert verify()["status"] == "match"


def test_a_state_changed_after_export_makes_the_workbook_stale(deal):
    write_state(COMPANY, PROPOSAL, date_str=DATE, financials={"FY-Current": {"raw": {**RAW, "revenue": 6000}}})

    report = verify()

    assert report["status"] == "mismatch"
    assert report["artifacts"]["workbook"]["status"] == "mismatch"
    assert report["artifacts"]["document"]["status"] == "match"
    assert {f["kind"] for f in report["findings"]} == {"value"}


# ---------------------------------------------------------------------------
# Workbook differences
# ---------------------------------------------------------------------------

def test_an_edited_input_cell_is_reported_with_its_location_and_both_values(deal):
    edit_workbook(lambda workbook: setattr(cell_holding(workbook, 5000), "value", 5001))

    report = verify()

    assert report["status"] == "mismatch" and report["artifacts"]["document"]["status"] == "match"
    assert [(f["artifact"], f["kind"], f["expected"], f["actual"]) for f in report["findings"]] == [
        ("workbook", "value", 5000, 5001)]
    assert report["findings"][0]["location"].startswith("Financial Spreading!")


def test_an_overwritten_formula_is_reported_as_the_formula_text(deal):
    holder = {}

    def overwrite(workbook):
        formula_cell = next(c for row in workbook["Financial Spreading"].iter_rows() for c in row
                            if isinstance(c.value, str) and c.value.startswith("="))
        holder["formula"], holder["coordinate"] = formula_cell.value, formula_cell.coordinate
        formula_cell.value = 123

    edit_workbook(overwrite)

    finding, = verify()["findings"]

    assert finding["location"] == f"Financial Spreading!{holder['coordinate']}"
    assert (finding["kind"], finding["expected"], finding["actual"]) == ("value", holder["formula"], 123)


def test_a_cleared_cell_is_a_difference(deal):
    edit_workbook(lambda workbook: setattr(cell_holding(workbook, 5000), "value", None))

    finding, = verify()["findings"]

    assert (finding["expected"], finding["actual"]) == (5000, None)


def test_a_changed_number_format_bold_or_italic_is_reported_by_kind(deal):
    def restyle(workbook):
        cell = cell_holding(workbook, 5000)
        cell.number_format = "0.000"
        cell.font = openpyxl.styles.Font(
            name=cell.font.name, sz=cell.font.sz, color=cell.font.color,
            bold=not cell.font.bold, italic=not cell.font.italic)

    edit_workbook(restyle)

    assert sorted(f["kind"] for f in verify()["findings"]) == ["bold", "italic", "number_format"]


def test_a_removed_or_added_sheet_is_reported(deal):
    edit_workbook(lambda workbook: workbook.remove(workbook["Collateral & Exposure"]))
    finding, = verify()["findings"]
    assert (finding["location"], finding["kind"]) == ("sheet names", "sheets")
    assert finding["expected"] == ["Financial Spreading", "Collateral & Exposure"]
    assert finding["actual"] == ["Financial Spreading"]

    edit_workbook(lambda workbook: workbook.create_sheet("Notes"))
    assert verify()["findings"][0]["actual"] == ["Financial Spreading", "Notes"]


def test_column_width_merged_cells_and_freeze_panes_are_part_of_the_comparison(deal):
    def relayout(workbook):
        sheet = workbook["Financial Spreading"]
        sheet.column_dimensions["B"].width = 77
        sheet.merge_cells("AA60:AB60")
        sheet.freeze_panes = "C3"

    edit_workbook(relayout)

    assert sorted(f["kind"] for f in verify()["findings"]) == ["column_width", "freeze_panes", "merged_cells"]


def test_numbers_are_compared_with_a_tight_tolerance(deal):
    coordinate = cell_holding(openpyxl.load_workbook(WORKBOOK), 5000).coordinate

    def set_value(value):
        edit_workbook(lambda workbook: setattr(workbook["Financial Spreading"][coordinate], "value", value))

    set_value(5000 + 1e-12)
    assert verify()["status"] == "match"

    set_value(5000.001)
    assert verify()["status"] == "mismatch"


@pytest.mark.parametrize("expected, actual, equal", [
    (1, 1.0, True), (0.1 + 0.2, 0.3, True), (5, 5.001, False), (True, 1, False), (1, True, False),
    ("a", "a", True), ("a", "A", False), ("1", 1, False), (None, None, True), (None, "", False), (None, 0, False),
])
def test_value_equality_rules(expected, actual, equal):
    assert export_readback._values_equal(expected, actual) is equal


# ---------------------------------------------------------------------------
# Document differences
# ---------------------------------------------------------------------------

def cell_texts(row):
    """The plain text of each cell of a table row as the verifier records it (a list of paragraph blocks per cell)."""
    return ["".join(part[1] for block in cell["paragraphs"] for part in block["parts"]) for cell in row]


def paragraph_with(document, text):
    return next(p for p in document.paragraphs if text in p.text)


def test_an_edited_paragraph_is_reported_by_block_with_both_versions(deal):
    def reword(document):
        run = paragraph_with(document, "First bullet").runs[0]
        run.text = "Changed bullet"

    edit_document(reword)

    report = verify()
    finding, = report["findings"]

    assert report["artifacts"]["workbook"]["status"] == "match"
    assert (finding["artifact"], finding["kind"]) == ("document", "paragraph")
    assert finding["location"].startswith("block ")
    assert "First bullet" in json.dumps(finding["expected"]) and "Changed bullet" in json.dumps(finding["actual"])


def test_an_edited_table_cell_is_reported(deal):
    edit_document(lambda document: setattr(document.tables[0].cell(1, 1), "text", "9,999"))

    finding, = verify()["findings"]

    assert finding["kind"] == "table"
    assert cell_texts(finding["expected"]["rows"][1]) == ["Revenue", "5,000"]
    assert cell_texts(finding["actual"]["rows"][1]) == ["Revenue", "9,999"]


def test_an_added_or_removed_block_is_reported_as_a_count_change(deal):
    expected_blocks = len(export_readback._document_blocks(docx.Document(DOCUMENT)))
    edit_document(lambda document: document.add_paragraph("An unauthorised extra paragraph"))

    assert verify()["findings"] == [{
        "artifact": "document", "location": "block count", "kind": "blocks",
        "expected": expected_blocks, "actual": expected_blocks + 1}]

    reexport()
    edit_document(lambda document: paragraph_with(document, "First bullet")._p.getparent().remove(
        paragraph_with(document, "First bullet")._p))

    assert verify()["findings"][0]["actual"] == expected_blocks - 1


def test_a_changed_style_or_run_formatting_is_reported(deal):
    def restyle(document):
        paragraph = paragraph_with(document, "Second bullet")
        paragraph.style = document.styles["Normal"]
        paragraph.runs[0].bold = True

    edit_document(restyle)

    finding, = verify()["findings"]

    assert finding["expected"]["style"] != finding["actual"]["style"]


def test_a_retargeted_hyperlink_is_reported(deal):
    def retarget(document):
        relationships = document.part.rels
        link = next(rel for rel in relationships.values() if rel.is_external and "example.invalid" in rel.target_ref)
        link._target = "https://example.invalid/other"

    edit_document(retarget)

    finding, = verify()["findings"]

    assert "https://example.invalid/page" in json.dumps(finding["expected"])
    assert "https://example.invalid/other" in json.dumps(finding["actual"])


def test_a_swapped_embedded_image_or_changed_alt_text_is_reported(deal):
    def swap(document):
        image_part = next(rel.target_part for rel in document.part.rels.values() if "image" in rel.reltype)
        image_part._blob = OTHER_PNG

    edit_document(swap)
    finding, = verify()["findings"]
    assert finding["expected"]["image_hashes"] != finding["actual"]["image_hashes"]

    reexport()

    def rename_alt(document):
        for element in document.element.body.iter():
            if element.tag.endswith("}docPr"):
                element.set("descr", "A different description")

    edit_document(rename_alt)
    finding, = verify()["findings"]
    assert finding["expected"]["image_alt_text"] == ["Chart alt text"]
    assert finding["actual"]["image_alt_text"] == ["A different description"]


def test_a_draft_edited_after_export_makes_the_document_stale(deal):
    DRAFT.write_text(DRAFT_MARKDOWN.replace("Second bullet", "Third bullet"), encoding="utf-8")

    report = verify()

    assert report["artifacts"]["document"]["status"] == "mismatch"
    assert report["artifacts"]["workbook"]["status"] == "match"


# ---------------------------------------------------------------------------
# The report
# ---------------------------------------------------------------------------

def test_findings_are_capped_but_the_count_is_exact(deal):
    def scribble(workbook):
        sheet = workbook["Financial Spreading"]
        for row in range(1, 41):
            for column in range(20, 27):
                sheet.cell(row=row, column=column).value = f"edit {row}-{column}"

    edit_workbook(scribble)

    report = verify()

    assert report["mismatch_count"] == 280 > export_readback.MAX_FINDINGS
    assert len(report["findings"]) == export_readback.MAX_FINDINGS and report["truncated"] is True
    assert report["artifacts"]["workbook"]["mismatch_count"] == 280


# ---------------------------------------------------------------------------
# Unusable input: one clear error, nothing compared
# ---------------------------------------------------------------------------

def test_a_deal_with_no_state_is_refused(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    with pytest.raises(StateError, match="no existing state.json"):
        verify()


def test_an_explicit_date_with_no_state_is_refused_and_a_bad_date_is_refused(deal):
    with pytest.raises(StateError, match="no state.json in the dated folder"):
        verify(date_str="2025-12-31")
    with pytest.raises(ValueError, match="--date-str"):
        verify(date_str="05/01/2026")


def test_the_named_date_selects_that_folder_rather_than_the_newest(deal):
    write_state(COMPANY, PROPOSAL, date_str="2026-06-01", financials_source="framework-computed")

    assert Path(verify(date_str=DATE)["selected_folder"]) == FOLDER
    with pytest.raises(StateError, match="does not exist"):
        verify()   # the newest dated folder has a state but no exports


@pytest.mark.parametrize("missing, label", [(WORKBOOK, "workbook"), (DOCUMENT, "document")])
def test_a_missing_artifact_is_refused_by_name(deal, missing, label):
    missing.unlink()

    with pytest.raises(StateError, match=f"the {label} .* does not exist"):
        verify()


def test_a_corrupt_workbook_or_document_is_one_error_not_a_traceback(deal):
    WORKBOOK.write_bytes(b"not a workbook")
    with pytest.raises(ValueError, match="cannot read workbook"):
        verify()

    reexport()
    DOCUMENT.write_bytes(b"not a document")
    with pytest.raises(ValueError, match="cannot read document"):
        verify()


@pytest.mark.parametrize("recorded, message", [
    (None, "records no approved_draft_path"),
    ("   ", "records no approved_draft_path"),
    (5, "records no approved_draft_path"),
    ("/etc/passwd", "relative path inside the deal's dated folder"),
    ("C:\\evil\\draft.md", "relative path inside the deal's dated folder"),
    ("deals/../../outside.md", "relative path inside the deal's dated folder"),
    ("deals\\..\\outside.md", "relative path inside the deal's dated folder"),
    ("elsewhere/draft.md", "is not inside the deal's dated folder"),
    (str(FOLDER / "missing.md"), "does not exist"),
])
def test_an_unusable_approved_draft_path_is_refused(deal, recorded, message):
    Path("elsewhere").mkdir()
    Path("elsewhere/draft.md").write_text("# other", encoding="utf-8")
    path = state_path(COMPANY, PROPOSAL, date_str=DATE)
    state = json.loads(Path(path).read_text(encoding="utf-8"))
    state["approved_draft_path"] = recorded
    Path(path).write_text(json.dumps(state), encoding="utf-8")

    with pytest.raises(StateError, match=message):
        verify()


def test_an_approved_draft_symlinked_out_of_the_folder_is_refused(deal):
    outside = Path("outside.md")
    outside.write_text("# outside", encoding="utf-8")
    link = FOLDER / "linked.md"
    try:
        link.symlink_to(outside.resolve())
    except (OSError, NotImplementedError):
        pytest.skip("symbolic links are not available here")
    write_state(COMPANY, PROPOSAL, date_str=DATE, approved_draft_path=link.as_posix())

    with pytest.raises(StateError, match="is not inside the deal's dated folder"):
        verify()


def test_a_malformed_group_basis_or_source_in_the_state_is_refused(deal):
    path = state_path(COMPANY, PROPOSAL, date_str=DATE)
    state = json.loads(Path(path).read_text(encoding="utf-8"))
    state["group_basis"] = {"basis": "nonsense"}
    Path(path).write_text(json.dumps(state), encoding="utf-8")
    with pytest.raises(StateError, match="group_basis"):
        verify()

    del state["group_basis"]
    state["financials_source"] = 5
    Path(path).write_text(json.dumps(state), encoding="utf-8")
    with pytest.raises(StateError, match="financials_source"):
        verify()


# ---------------------------------------------------------------------------
# Read-only for the deal, and no scratch left behind
# ---------------------------------------------------------------------------

def test_verifying_changes_nothing_in_the_deal_and_leaves_no_scratch_files(deal):
    before_scratch = leftover_scratch_dirs()
    before = folder_bytes()
    verify()
    edit_workbook(lambda workbook: setattr(cell_holding(workbook, 5000), "value", 1))
    changed = folder_bytes()
    assert verify()["status"] == "mismatch"

    assert folder_bytes() == changed and changed != before
    assert leftover_scratch_dirs() == before_scratch


def test_scratch_files_are_removed_even_when_a_comparison_fails(deal):
    before_scratch = leftover_scratch_dirs()
    WORKBOOK.write_bytes(b"broken")

    with pytest.raises(ValueError):
        verify()

    assert leftover_scratch_dirs() == before_scratch


# ---------------------------------------------------------------------------
# The command line, in process
# ---------------------------------------------------------------------------

def test_main_prints_the_report_and_exits_zero_on_a_match(deal, capsys):
    assert export_readback.main(["--company", COMPANY, "--proposal", PROPOSAL]) == 0

    out = capsys.readouterr()
    assert json.loads(out.out)["status"] == "match" and out.err == ""


def test_main_exits_one_with_the_findings_on_a_mismatch(deal, capsys):
    edit_workbook(lambda workbook: setattr(cell_holding(workbook, 5000), "value", 4999))

    assert export_readback.main(["--company", COMPANY, "--proposal", PROPOSAL, "--date-str", DATE]) == 1

    report = json.loads(capsys.readouterr().out)
    assert report["status"] == "mismatch" and report["mismatch_count"] == 1


def test_main_turns_an_unusable_deal_into_one_error_line(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    with pytest.raises(SystemExit) as caught:
        export_readback.main(["--company", COMPANY, "--proposal", PROPOSAL])

    assert caught.value.code == (
        f"error: no existing state.json for company '{COMPANY}' and proposal '{PROPOSAL}'")


def test_main_refuses_missing_required_flags_with_a_usage_error(capsys):
    with pytest.raises(SystemExit) as caught:
        export_readback.main(["--company", COMPANY])

    assert caught.value.code == 2 and "--proposal" in capsys.readouterr().err


# ---------------------------------------------------------------------------
# Review fixes (CL2, CL3, CH1, CH2): hidden or recoloured data, the far-corner cell, run formatting
# ---------------------------------------------------------------------------

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.shared import Inches, Pt, RGBColor
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule


def font_like(cell, **changes):
    font = cell.font
    values = {"name": font.name, "sz": font.sz, "bold": font.bold, "italic": font.italic, "color": font.color,
              "underline": font.underline, "strike": font.strike}
    values.update(changes)
    return Font(**values)


CELL_EDITS = {
    "fill": lambda c: setattr(c, "fill", PatternFill("solid", fgColor="FFFF00")),
    "font_color": lambda c: setattr(c, "font", font_like(c, color="FFFFFFFF")),
    "font_size": lambda c: setattr(c, "font", font_like(c, sz=(c.font.sz or 11) + 5)),
    "font_name": lambda c: setattr(c, "font", font_like(c, name="Courier New")),
    "underline": lambda c: setattr(c, "font", font_like(c, underline="single")),
    "strike": lambda c: setattr(c, "font", font_like(c, strike=True)),
    "border": lambda c: setattr(c, "border", Border(left=Side(style="thick"))),
    "alignment": lambda c: setattr(c, "alignment", Alignment(horizontal="right", wrap_text=True)),
}


@pytest.mark.parametrize("kind", sorted(CELL_EDITS))
def test_a_change_to_a_cells_look_is_reported_by_its_own_kind_and_location(deal, kind):
    holder = {}

    def change(workbook):
        cell = cell_holding(workbook, 5000)
        holder["coordinate"] = cell.coordinate
        CELL_EDITS[kind](cell)

    edit_workbook(change)

    report = verify()

    assert report["status"] == "mismatch"
    assert {f["location"] for f in report["findings"]} == {f"Financial Spreading!{holder['coordinate']}"}
    assert kind in {f["kind"] for f in report["findings"]}


def test_a_hidden_sheet_row_or_column_is_reported(deal):
    def hide(workbook):
        workbook["Collateral & Exposure"].sheet_state = "hidden"
        sheet = workbook["Financial Spreading"]
        sheet.row_dimensions[5].hidden = True
        sheet.column_dimensions["C"].hidden = True

    edit_workbook(hide)

    findings = {(f["location"], f["kind"], f["expected"], f["actual"]) for f in verify()["findings"]}

    assert ("Collateral & Exposure!sheet", "sheet_state", "visible", "hidden") in findings
    assert ("Financial Spreading!row 5", "row_hidden", False, True) in findings
    assert ("Financial Spreading!column C", "column_hidden", False, True) in findings


def test_a_changed_row_height_protection_defined_name_validation_or_conditional_format_is_reported(deal):
    def tweak(workbook):
        sheet = workbook["Financial Spreading"]
        sheet.row_dimensions[3].height = 60
        sheet.protection.sheet = True
        workbook.defined_names["Hidden_Total"] = DefinedName("Hidden_Total", attr_text="'Financial Spreading'!$B$3")
        validation = DataValidation(type="whole", operator="between", formula1="0", formula2="9")
        sheet.add_data_validation(validation)
        validation.add("B3")
        sheet.conditional_formatting.add("B3:B4", CellIsRule(operator="greaterThan", formula=["1"],
                                                             fill=PatternFill("solid", bgColor="FF0000")))

    edit_workbook(tweak)

    kinds = {f["kind"] for f in verify()["findings"]}

    assert {"row_height", "protection", "defined_names", "data_validations", "conditional_formatting"} <= kinds


def test_a_cell_in_the_far_corner_is_refused_at_once_instead_of_comparing_billions_of_positions(deal):
    edit_workbook(lambda workbook: setattr(workbook["Financial Spreading"]["XFD1048576"], "value", "x"))

    with pytest.raises(ValueError, match=r"worksheet 'Financial Spreading' uses 1048576 rows by 16384 columns"):
        verify()


def test_the_far_corner_refusal_is_one_error_line_on_the_command_line(deal):
    edit_workbook(lambda workbook: setattr(workbook["Financial Spreading"]["XFD1048576"], "value", "x"))

    with pytest.raises(SystemExit) as caught:
        export_readback.main(["--company", COMPANY, "--proposal", PROPOSAL])

    assert str(caught.value.code).startswith("error: worksheet 'Financial Spreading' uses 1048576 rows")


def first_bullet_run(document):
    return paragraph_with(document, "First bullet").runs[0]


RUN_EDITS = {
    "size": lambda run: setattr(run.font, "size", Pt(1)),
    "colour": lambda run: setattr(run.font.color, "rgb", RGBColor(255, 255, 255)),
    "underline": lambda run: setattr(run.font, "underline", True),
    "strike": lambda run: setattr(run.font, "strike", True),
    "highlight": lambda run: setattr(run.font, "highlight_color", WD_COLOR_INDEX.YELLOW),
}


@pytest.mark.parametrize("name", sorted(RUN_EDITS))
def test_a_change_to_a_runs_size_colour_underline_strike_or_highlight_is_reported(deal, name):
    edit_document(lambda document: RUN_EDITS[name](first_bullet_run(document)))

    finding, = verify()["findings"]

    assert finding["kind"] == "paragraph" and finding["expected"]["parts"] != finding["actual"]["parts"]
    assert "First bullet" in json.dumps(finding["actual"]["parts"])


def test_a_change_to_paragraph_alignment_or_spacing_is_reported(deal):
    edit_document(lambda document: setattr(paragraph_with(document, "First bullet"), "alignment",
                                           WD_ALIGN_PARAGRAPH.CENTER))
    finding, = verify()["findings"]
    assert finding["expected"]["layout"] != finding["actual"]["layout"]

    reexport()
    edit_document(lambda document: setattr(paragraph_with(document, "First bullet").paragraph_format,
                                           "space_after", Pt(40)))
    finding, = verify()["findings"]
    assert finding["expected"]["layout"] != finding["actual"]["layout"]


def test_a_change_to_the_formatting_inside_a_table_cell_is_reported(deal):
    edit_document(lambda document: setattr(document.tables[0].cell(1, 1).paragraphs[0].runs[0], "bold", True))

    finding, = verify()["findings"]

    assert finding["kind"] == "table" and finding["expected"]["rows"] != finding["actual"]["rows"]


def test_a_changed_page_margin_header_or_footer_is_reported_as_a_section_difference(deal):
    edit_document(lambda document: setattr(document.sections[0], "left_margin", Inches(3)))
    finding, = verify()["findings"]
    assert (finding["location"], finding["kind"]) == ("section 1", "section")
    assert finding["expected"]["margins"] != finding["actual"]["margins"]

    reexport()

    def edit_footer(document):
        document.sections[0].footer.paragraphs[0].text = "Edited footer"

    edit_document(edit_footer)
    finding, = verify()["findings"]
    assert finding["location"] == "section 1" and finding["expected"]["footer"] == {"linked": True}
    footer_blocks = finding["actual"]["footer"]["blocks"]
    assert finding["actual"]["footer"]["linked"] is False
    assert "Edited footer" in json.dumps(footer_blocks)


def test_the_limits_name_what_is_still_not_compared(deal):
    limits = " ".join(verify()["limits"])

    for phrase in ("cached formula results", "not by colour or font", "images in a workbook",
                   "rules inside data validations", "comments, footnotes", "text boxes", "file properties"):
        assert phrase in limits, phrase


def test_an_added_section_is_reported_as_a_section_count_difference(deal):
    edit_document(lambda document: document.add_section())

    findings = verify()["findings"]

    assert {"artifact": "document", "location": "section count", "kind": "sections", "expected": 1, "actual": 2} \
        in findings


def test_a_deleted_moved_or_repointed_chart_is_reported(deal):
    def delete_one(workbook):
        sheet = workbook["Financial Spreading"]
        sheet._charts.remove(sorted(sheet._charts, key=lambda chart: chart.anchor._from.row)[0])

    edit_workbook(delete_one)
    finding, = verify()["findings"]
    assert finding["kind"] == "charts" and len(finding["expected"]) == len(finding["actual"]) + 1

    reexport()

    def repoint(workbook):
        chart = workbook["Financial Spreading"]._charts[0]
        chart.series[0].val.numRef.f = "'Financial Spreading'!$B$9:$G$9"

    edit_workbook(repoint)
    finding, = verify()["findings"]
    assert finding["kind"] == "charts" and "$B$9:$G$9" in json.dumps(finding["actual"])

    reexport()

    def move(workbook):
        workbook["Financial Spreading"]._charts[0].anchor._from.row += 3

    edit_workbook(move)
    assert [f["kind"] for f in verify()["findings"]] == ["charts"]


def test_a_chart_part_without_a_range_reads_as_none():
    from types import SimpleNamespace

    assert export_readback._reference(None) is None
    assert export_readback._reference(SimpleNamespace(numRef=None, strRef=None)) is None
    assert export_readback._reference(SimpleNamespace(numRef=None, strRef=SimpleNamespace(f="A1:B1"))) == "A1:B1"
    assert export_readback._reference(SimpleNamespace(v="a literal name")) is None


# ---------------------------------------------------------------------------
# Second-pass review fixes (CH5, CH6, CH7): header and footer content, superscript and hyperlink formatting, and
# direct table layout
# ---------------------------------------------------------------------------

from docx.enum.table import WD_ALIGN_VERTICAL  # noqa: E402
from docx.oxml import parse_xml  # noqa: E402
from docx.oxml.ns import nsdecls  # noqa: E402


def document_with_header_image(png_bytes, tmp_path, name):
    picture = tmp_path / f"{name}.png"
    picture.write_bytes(png_bytes)
    document = docx.Document()
    header = document.sections[0].header
    header.is_linked_to_previous = False
    header.paragraphs[0].add_run().add_picture(str(picture))
    return document


def test_a_swapped_header_image_is_a_section_difference(tmp_path):
    """CH5: header and footer content used to be read as text only, so a different logo read as a match."""
    original = document_with_header_image(PNG_1X1, tmp_path, "one")
    swapped = document_with_header_image(OTHER_PNG, tmp_path, "two")

    assert export_readback._document_sections(original) != export_readback._document_sections(swapped)
    header = export_readback._document_sections(original)[0]["header"]
    assert header["linked"] is False and header["blocks"][0]["image_hashes"][0] is not None


def test_a_header_table_or_linked_status_is_part_of_the_section_signature(tmp_path):
    plain, with_table = docx.Document(), docx.Document()
    story = with_table.sections[0].footer
    story.is_linked_to_previous = False
    story.add_table(1, 1, 5000).cell(0, 0).text = "Total"

    assert export_readback._document_sections(plain)[0]["footer"] == {"linked": True}
    assert "Total" in json.dumps(export_readback._document_sections(with_table)[0]["footer"])


def test_a_run_changed_to_superscript_is_reported(deal):
    """CH6: superscript and subscript were not part of the run signature."""
    edit_document(lambda document: setattr(first_bullet_run(document).font, "superscript", True))

    finding, = verify()["findings"]

    assert finding["kind"] == "paragraph" and finding["expected"]["parts"] != finding["actual"]["parts"]


def test_a_hyperlink_whose_visible_formatting_changed_is_reported(deal):
    """CH6: a hyperlink kept only its target and text, so white one-point link text read as a match."""
    def hide(document):
        paragraph = paragraph_with(document, "synthetic link")
        link = next(item for item in paragraph.iter_inner_content() if hasattr(item, "url"))
        link.runs[0].font.size = Pt(1)
        link.runs[0].font.color.rgb = RGBColor(255, 255, 255)

    edit_document(hide)

    finding, = verify()["findings"]

    assert finding["kind"] == "paragraph"
    assert "https://example.invalid/page" in json.dumps(finding["actual"]["parts"])
    assert finding["expected"]["parts"] != finding["actual"]["parts"]


def shade_cell(document):
    cell = document.tables[0].cell(1, 1)
    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:val="clear" w:fill="FFFF00"/>'))


TABLE_EDITS = {
    "shading": shade_cell,
    "width": lambda document: setattr(document.tables[0].cell(1, 1), "width", Inches(2)),
    "vertical alignment": lambda document: setattr(document.tables[0].cell(1, 1), "vertical_alignment",
                                                   WD_ALIGN_VERTICAL.CENTER),
    "table borders": lambda document: document.tables[0]._tbl.tblPr.append(parse_xml(
        f'<w:tblBorders {nsdecls("w")}><w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/></w:tblBorders>')),
}


@pytest.mark.parametrize("name", sorted(TABLE_EDITS))
def test_a_change_to_direct_table_or_cell_layout_is_reported(deal, name):
    """CH7: shading, borders, widths and alignment applied directly were invisible to the comparison."""
    edit_document(TABLE_EDITS[name])

    finding, = verify()["findings"]

    assert finding["kind"] == "table" and finding["expected"] != finding["actual"]


# ---------------------------------------------------------------------------
# Third-pass review fixes (CH10, CH11, CH12)
# ---------------------------------------------------------------------------

def test_a_changed_table_column_width_is_reported(deal):
    """CH10: the grid column widths were not read."""
    def narrow(document):
        grid = document.tables[0]._tbl.tblGrid
        grid.gridCol_lst[0].w = Inches(1)

    edit_document(narrow)

    finding, = verify()["findings"]

    assert finding["kind"] == "table" and finding["expected"]["layout"] != finding["actual"]["layout"]


def test_a_restyled_style_definition_is_reported_even_though_no_paragraph_changed(deal):
    """CH11: only the style's name was compared, so redefining 'List Bullet' as one-point text read as a match."""
    edit_document(lambda document: setattr(document.styles["List Bullet"].font, "size", Pt(1)))

    finding, = verify()["findings"]

    assert (finding["location"], finding["kind"]) == ("style List Bullet", "style")
    assert finding["expected"][4] != finding["actual"][4] and finding["actual"][4] == 1.0


def test_a_resized_body_image_is_reported(deal):
    """CH12: an image kept its hash and alt text but could be shrunk to a dot."""
    edit_document(lambda document: setattr(document.inline_shapes[0], "width", Inches(0.1)))

    finding, = verify()["findings"]

    assert finding["kind"] == "paragraph"
    assert finding["expected"]["image_extents"] != finding["actual"]["image_extents"]


def test_a_resized_header_image_is_a_section_difference(tmp_path):
    original = document_with_header_image(PNG_1X1, tmp_path, "size_one")
    resized = document_with_header_image(PNG_1X1, tmp_path, "size_two")
    resized.sections[0].header.paragraphs[0].runs[0].element.xpath(".//wp:extent")[0].set("cx", "12345")

    assert export_readback._document_sections(original) != export_readback._document_sections(resized)


def test_a_style_only_one_side_uses_is_left_to_the_block_comparison(deal):
    edit_document(lambda document: setattr(document.sections[0].footer, "is_linked_to_previous", False))

    kinds = {f["kind"] for f in verify()["findings"]}

    assert "style" not in kinds


def test_a_character_style_has_a_signature_without_paragraph_layout():
    document = docx.Document()
    character = next(style for style in document.styles if style.type == 2)   # WD_STYLE_TYPE.CHARACTER

    signature = export_readback._style_signature(character)

    assert signature[0] == character.name and len(signature) == 10


# ---------------------------------------------------------------------------
# #408 mutation-survivor slice: fast helper contracts and exact verifier output
# ---------------------------------------------------------------------------

def test_408_numeric_comparison_uses_relative_and_absolute_tolerances():
    """Mutants #10/#11: each tolerance matters in a different numeric regime."""
    assert export_readback._values_equal(1_000_000_000.0, 1_000_000_000.5)
    assert export_readback._values_equal(0.0, 0.5e-9)
    assert not export_readback._values_equal(1_000_000_000.0, 1_000_000_002.0)


def test_408_color_signature_keeps_kind_value_and_zero_tint():
    """Mutants #3-#14: every supported color kind and falsey tint retains its real signature."""
    from types import SimpleNamespace

    colours = {
        "rgb": "00AA11",
        "theme": 4,
        "indexed": 7,
        "auto": True,
    }
    for kind, value in colours.items():
        color = SimpleNamespace(type=kind, rgb=colours["rgb"], theme=colours["theme"],
                                indexed=colours["indexed"], auto=colours["auto"], tint=0)
        assert export_readback._color(color) == [kind, value, 0]
    assert export_readback._color(None) is None
    tint = SimpleNamespace(type="rgb", rgb="00AA11", theme=0, indexed=0, auto=False, tint=0.25)
    assert export_readback._color(tint) == ["rgb", "00AA11", 0.25]


def test_408_border_and_fill_signatures_keep_side_colours_and_empty_fill():
    """Mutants #18 (border), #3-#5 (fill): foreground/background and border colours are compared."""
    from types import SimpleNamespace

    color = SimpleNamespace(type="rgb", rgb="00AA11", theme=None, indexed=None, auto=None, tint=0)
    side = SimpleNamespace(style="thin", color=color)
    cell = SimpleNamespace(border=SimpleNamespace(left=side, right=None, top=None, bottom=None))
    assert export_readback._border(cell) == {
        "left": ["thin", ["rgb", "00AA11", 0]], "right": None, "top": None, "bottom": None,
    }

    empty = SimpleNamespace(fill=SimpleNamespace(fill_type=None, fgColor=color, bgColor=color))
    assert export_readback._fill(empty) is None
    filled = SimpleNamespace(fill=SimpleNamespace(fill_type="solid", fgColor=color, bgColor=color))
    assert export_readback._fill(filled) == ["solid", ["rgb", "00AA11", 0], ["rgb", "00AA11", 0]]


def test_408_workbook_comparison_pins_sheet_and_defined_name_findings(monkeypatch):
    """Mutants #7/#17/#18/#29-#45: findings retain exact artifact, location, kind and both values."""
    from types import SimpleNamespace

    class DefinedNames:
        def __init__(self, entries):
            self.entries = entries

        def items(self):
            return self.entries.items()

    expected = SimpleNamespace(
        sheetnames=["Expected"],
        defined_names=DefinedNames({"SyntheticName": SimpleNamespace(attr_text="'Expected'!$A$1")}),
    )
    actual = SimpleNamespace(
        sheetnames=["Actual"],
        defined_names=DefinedNames({"SyntheticName": SimpleNamespace(attr_text="'Actual'!$A$1")}),
    )
    monkeypatch.setattr(export_readback, "_load_workbook",
                        lambda path: expected if path == "expected" else actual)
    assert export_readback.compare_workbooks("expected", "actual") == [
        {"artifact": "workbook", "location": "sheet names", "kind": "sheets",
         "expected": ["Expected"], "actual": ["Actual"]},
        {"artifact": "workbook", "location": "defined names", "kind": "defined_names",
         "expected": [("SyntheticName", "'Expected'!$A$1")],
         "actual": [("SyntheticName", "'Actual'!$A$1")]},
    ]


def test_408_workbook_comparison_continues_after_a_missing_expected_sheet(monkeypatch):
    """Mutant #46: skipping a sheet absent from the actual workbook must not stop later sheet comparisons."""
    from types import SimpleNamespace

    empty = SimpleNamespace(items=lambda: [])
    expected = SimpleNamespace(sheetnames=["Removed", "Retained"], defined_names=empty)
    actual = SimpleNamespace(sheetnames=["Retained"], defined_names=empty)
    compared = []
    monkeypatch.setattr(export_readback, "_load_workbook",
                        lambda path: expected if path == "expected" else actual)
    monkeypatch.setattr(export_readback, "_compare_sheets",
                        lambda name, want, have: compared.append(name) or [])
    findings = export_readback.compare_workbooks("expected", "actual")
    assert compared == ["Retained"]
    assert findings == [{"artifact": "workbook", "location": "sheet names", "kind": "sheets",
                         "expected": ["Removed", "Retained"], "actual": ["Retained"]}]


def test_408_compare_sheets_checks_cell_limit_and_reports_the_exact_refusal(monkeypatch):
    """Mutants #13/#15/#16: exactly the limit is allowed; an over-limit refusal keeps its exact message."""
    from types import SimpleNamespace

    class Sheet:
        def __init__(self, title, rows, columns):
            self.title, self.max_row, self.max_column = title, rows, columns

        def cell(self, row, column):
            return SimpleNamespace(value=None, coordinate=f"A{row}")

    monkeypatch.setattr(export_readback, "CELL_PROPERTIES", ())
    monkeypatch.setattr(export_readback, "_compare_sheet_layout", lambda name, want, have: [])
    monkeypatch.setattr(export_readback, "MAX_COMPARED_CELLS", 1)
    assert export_readback._compare_sheets("Synthetic", Sheet("w", 1, 1), Sheet("h", 1, 1)) == []

    with pytest.raises(ValueError) as refused:
        export_readback._compare_sheets("Synthetic", Sheet("w", 2, 1), Sheet("h", 2, 1))
    assert str(refused.value) == (
        "worksheet 'Synthetic' uses 2 rows by 1 columns, more than the 1 positions "
        "a read-back compares; an export is not that large, so the file was edited"
    )


def test_408_compare_sheets_checks_last_row_and_first_column_and_stops_at_bounds(monkeypatch):
    """Mutants #22/#23/#28/#30: all and only used cell coordinates are compared."""
    from types import SimpleNamespace

    class SparseSheet:
        def __init__(self, title, rows=2, columns=2, values=None):
            self.title, self.max_row, self.max_column = title, rows, columns
            self.values = values or {}

        def cell(self, row, column):
            value = self.values.get((row, column))
            return SimpleNamespace(value=value, coordinate=f"{chr(64 + column)}{row}")

    monkeypatch.setattr(export_readback, "CELL_PROPERTIES", ())
    monkeypatch.setattr(export_readback, "_compare_sheet_layout", lambda name, want, have: [])
    expected = SparseSheet("expected", values={(1, 1): "same", (2, 1): "tail", (1, 2): "same"})
    actual = SparseSheet("actual", values={(1, 1): "same", (2, 1): "changed", (1, 2): "same"})
    findings = export_readback._compare_sheets("Synthetic", expected, actual)
    assert [(item["location"], item["expected"], item["actual"]) for item in findings] == [
        ("Synthetic!A2", "tail", "changed")
    ]

    # A cell the worksheet claims is outside its used range is deliberately non-blank in this fake. The helper must
    # not inspect it: real openpyxl worksheets' cells beyond max_row/max_column are outside the comparison contract.
    bounded_expected = SparseSheet("expected", rows=1, columns=1, values={(2, 1): "outside"})
    bounded_actual = SparseSheet("actual", rows=1, columns=1, values={(2, 1): "different"})
    assert export_readback._compare_sheets("Synthetic", bounded_expected, bounded_actual) == []
    extra_column_expected = SparseSheet("expected", rows=1, columns=1, values={(1, 2): "outside"})
    extra_column_actual = SparseSheet("actual", rows=1, columns=1, values={(1, 2): "different"})
    assert export_readback._compare_sheets("Synthetic", extra_column_expected, extra_column_actual) == []


def test_408_compare_sheets_finding_preserves_artifact_and_both_property_values(monkeypatch):
    """Mutants #58/#59/#60/#61/#62: a changed cell property has an exact finding record."""
    from types import SimpleNamespace

    class Sheet:
        max_row = 1
        max_column = 1

        def __init__(self, title, marker):
            self.title, self.marker = title, marker

        def cell(self, row, column):
            return SimpleNamespace(value="same", coordinate="A1", marker=self.marker)

    monkeypatch.setattr(export_readback, "CELL_PROPERTIES", (("marker", lambda cell: cell.marker),))
    monkeypatch.setattr(export_readback, "_compare_sheet_layout", lambda name, want, have: [])
    assert export_readback._compare_sheets("Synthetic", Sheet("want", "expected"),
                                           Sheet("have", "actual")) == [
        {"artifact": "workbook", "location": "Synthetic!A1", "kind": "marker",
         "expected": "expected", "actual": "actual"}
    ]


def test_408_chart_signature_keeps_references_and_chart_properties():
    """Mutants #3/#5/#6/#10-#13/#17-#19: chart series refs and exact grouping/type are part of the signature."""
    from types import SimpleNamespace

    ref = lambda value: SimpleNamespace(numRef=SimpleNamespace(f=value), strRef=None)
    chart = SimpleNamespace(
        tagname="barChart",
        anchor=SimpleNamespace(_from=SimpleNamespace(row=2, col=3)),
        grouping="stacked",
        type="bar",
        series=[SimpleNamespace(tx=ref("Sheet!$A$1"), val=ref("Sheet!$B$1"), cat=ref("Sheet!$C$1"))],
    )
    assert export_readback._chart_signature(chart) == [
        "barChart", 2, 3, "stacked", "bar",
        [["Sheet!$A$1", "Sheet!$B$1", "Sheet!$C$1"]],
    ]
    chart_without_optional_attributes = SimpleNamespace(
        tagname="pieChart", anchor=SimpleNamespace(_from=SimpleNamespace(row=0, col=0)), series=[]
    )
    assert export_readback._chart_signature(chart_without_optional_attributes)[3:5] == [None, None]


def test_408_sheet_charts_are_sorted_only_by_anchor_row_and_column(monkeypatch):
    """Mutants #2/#4/#7/#8: chart order is based on anchor, preserving input order for tied anchors."""
    from types import SimpleNamespace

    row_zero = SimpleNamespace(sig=["zChart", 0, 0, "bar"])
    row_one = SimpleNamespace(sig=["aChart", 1, 9, "pie"])
    monkeypatch.setattr(export_readback, "_chart_signature", lambda chart: chart.sig)
    assert export_readback._sheet_charts(SimpleNamespace(_charts=[row_one, row_zero])) == [
        ["zChart", 0, 0, "bar"], ["aChart", 1, 9, "pie"]
    ]

    row_first = SimpleNamespace(sig=["chart", 0, 9, "bar"])
    column_first = SimpleNamespace(sig=["chart", 1, 0, "pie"])
    assert export_readback._sheet_charts(SimpleNamespace(_charts=[column_first, row_first])) == [
        ["chart", 0, 9, "bar"], ["chart", 1, 0, "pie"]
    ]

    tied_z = SimpleNamespace(sig=["chart", 2, 4, "z"])
    tied_a = SimpleNamespace(sig=["chart", 2, 4, "a"])
    assert export_readback._sheet_charts(SimpleNamespace(_charts=[tied_z, tied_a])) == [
        ["chart", 2, 4, "z"], ["chart", 2, 4, "a"]
    ]


def test_408_sheet_layout_comparison_reports_every_layout_difference(monkeypatch):
    """Mutants #32-#177: the labels and expected/actual values of layout findings are contractual."""
    from openpyxl import Workbook
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.workbook.defined_name import DefinedName

    book_w, book_h = Workbook(), Workbook()
    want, have = book_w.active, book_h.active
    want.title, have.title = "want", "have"
    want.sheet_state, have.sheet_state = "hidden", "visible"
    want.protection.sheet = True
    have.protection.sheet = False
    want.row_dimensions[1].hidden, have.row_dimensions[1].hidden = True, False
    want.row_dimensions[1].height, have.row_dimensions[1].height = 21, 28
    want.row_dimensions[4].height = 18
    have.row_dimensions[5].height = 19
    want.column_dimensions["A"].width, have.column_dimensions["A"].width = 9, 15
    want.column_dimensions["A"].hidden, have.column_dimensions["A"].hidden = True, False
    want.column_dimensions["B"].width = 21
    have.column_dimensions["C"].width = 22
    want.merge_cells("A1:B1")
    have.merge_cells("B1:C1")
    want.freeze_panes, have.freeze_panes = "B2", "C3"
    want.data_validations.add(DataValidation(type="whole", operator="greaterThan", formula1="0"))
    want.conditional_formatting.add("A1", CellIsRule(operator="greaterThan", formula=["0"]))
    want.defined_names.add(DefinedName("SyntheticLocal", attr_text="'want'!$A$1"))
    have.defined_names.add(DefinedName("SyntheticLocal", attr_text="'have'!$A$1"))
    monkeypatch.setattr(export_readback, "_sheet_charts",
                        lambda sheet: [["chart", 1, 1, "bar"]] if sheet.title == "want" else [])

    findings = export_readback._compare_sheet_layout("Synthetic", want, have)
    actual = {(item["location"], item["kind"]): (item["artifact"], item["expected"], item["actual"])
              for item in findings}
    expected_locations = {
        ("Synthetic!sheet", "sheet_state"),
        ("Synthetic!sheet", "protection"),
        ("Synthetic!row 1", "row_hidden"), ("Synthetic!row 1", "row_height"),
        ("Synthetic!row 4", "row_height"), ("Synthetic!row 5", "row_height"),
        ("Synthetic!column A", "column_width"), ("Synthetic!column A", "column_hidden"),
        ("Synthetic!column B", "column_width"), ("Synthetic!column C", "column_width"),
        ("Synthetic!merged cells", "merged_cells"), ("Synthetic!freeze panes", "freeze_panes"),
        ("Synthetic!data validations", "data_validations"),
        ("Synthetic!conditional formats", "conditional_formatting"),
        ("Synthetic!charts", "charts"), ("Synthetic!defined names", "sheet_defined_names"),
    }
    assert set(actual) == expected_locations
    assert actual[("Synthetic!sheet", "protection")] == ("workbook", True, False)
    assert actual[("Synthetic!row 1", "row_hidden")] == ("workbook", True, False)
    assert actual[("Synthetic!row 1", "row_height")] == ("workbook", 21.0, 28.0)
    assert actual[("Synthetic!row 4", "row_height")] == ("workbook", 18.0, None)
    assert actual[("Synthetic!row 5", "row_height")] == ("workbook", None, 19.0)
    assert actual[("Synthetic!column A", "column_width")] == ("workbook", 9.0, 15.0)
    assert actual[("Synthetic!column A", "column_hidden")] == ("workbook", True, False)
    assert actual[("Synthetic!column B", "column_width")] == ("workbook", 21.0, have.column_dimensions["B"].width)
    assert actual[("Synthetic!column C", "column_width")] == ("workbook", want.column_dimensions["C"].width, 22.0)
    assert actual[("Synthetic!merged cells", "merged_cells")] == ("workbook", ["A1:B1"], ["B1:C1"])
    assert actual[("Synthetic!freeze panes", "freeze_panes")] == ("workbook", "B2", "C3")
    assert actual[("Synthetic!data validations", "data_validations")] == ("workbook", 1, 0)
    assert actual[("Synthetic!conditional formats", "conditional_formatting")] == ("workbook", 1, 0)
    assert actual[("Synthetic!charts", "charts")] == ("workbook", [["chart", 1, 1, "bar"]], [])
    assert actual[("Synthetic!defined names", "sheet_defined_names")] == (
        "workbook", [("SyntheticLocal", "'want'!$A$1")], [("SyntheticLocal", "'have'!$A$1")]
    )


def test_408_text_value_none_and_non_none_are_distinct():
    """Mutants #1/#3/#4: None remains None while other values are stringified."""
    assert export_readback._text_value(None) is None
    assert export_readback._text_value(0) == "0"
    assert export_readback._text_value("synthetic") == "synthetic"


def test_408_run_signature_pins_font_flags_and_optional_colour_fields():
    """Mutants #3/#4/#9-#22 in _run_signature: each font slot has its own meaning."""
    from types import SimpleNamespace

    font = SimpleNamespace(
        underline="single", strike=True, name="Synthetic Font", size=SimpleNamespace(pt=11.0),
        color=SimpleNamespace(type=None, rgb=None, theme_color=None), highlight_color="yellow",
        superscript=True, subscript=False,
        all_caps=True, small_caps=False, hidden=True,
    )
    run = SimpleNamespace(text="Synthetic run", font=font, bold=True, italic=False)
    signature = export_readback._run_signature(run)
    assert signature == [
        "run", "Synthetic run", True, False, "single", True, "Synthetic Font", 11.0,
        None, None, "yellow", "True", "False", True, False, True,
    ]

    font.color = SimpleNamespace(type="rgb", rgb="00AA11", theme_color="theme-value")
    font.superscript = False
    font.subscript = True
    run.bold = False
    signature = export_readback._run_signature(run)
    assert signature[2:4] == [False, False]
    assert signature[8:10] == ["00AA11", "theme-value"]
    assert signature[11:16] == ["False", "True", True, False, True]

    font.color = SimpleNamespace(type=None, rgb=None, theme_color=None)
    signature = export_readback._run_signature(run)
    assert signature[8:10] == [None, None]


def test_408_paragraph_block_pins_hyperlinks_image_extents_and_layout(tmp_path):
    """Mutants #3/#4/#16/#30/#31/#53/#54/#56: preserve exact block markers and paragraph properties."""
    from docx.shared import Inches, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.opc.constants import RELATIONSHIP_TYPE as RT

    document = docx.Document()
    paragraph = document.add_paragraph()
    paragraph.alignment = 2
    paragraph.paragraph_format.left_indent = Pt(7)
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(9)
    paragraph.paragraph_format.line_spacing = 1.5
    run = paragraph.add_run("synthetic link")
    run.bold = True
    relationship_id = document.part.relate_to("https://example.invalid/synthetic", RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink_run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = "linked"
    hyperlink_run.append(text)
    hyperlink.append(hyperlink_run)
    paragraph._p.append(hyperlink)

    signature = export_readback._paragraph_block(document.part, paragraph)
    link, = [part for part in signature["parts"] if part[0] == "link"]
    assert link[0] == "link" and link[1] == "https://example.invalid/synthetic"
    assert signature["layout"][0] == export_readback._text_value(paragraph.alignment)
    assert signature["layout"][1:] == [7.0, 8.0, 9.0, "1.5"]

    picture_path = tmp_path / "synthetic.png"
    picture_path.write_bytes(PNG_1X1)
    with_image = document.add_paragraph()
    with_image.add_run().add_picture(str(picture_path), width=Inches(1), height=Inches(1))
    image_signature = export_readback._paragraph_block(document.part, with_image)
    assert image_signature["image_hashes"] == [export_readback.hashlib.sha256(PNG_1X1).hexdigest()]
    blip_id, = with_image._p.xpath(".//a:blip/@r:embed")
    document.part.related_parts.pop(blip_id)
    missing_relation = export_readback._paragraph_block(document.part, with_image)
    assert missing_relation["image_hashes"] == [None]
    assert len(image_signature["image_extents"]) == 1
    assert image_signature["image_extents"][0][0] is not None
    assert image_signature["image_extents"][0][1] is not None


def test_408_layout_properties_uses_xml_local_names_and_namespaced_attributes():
    """Mutants #1/#5/#6/#9/#11/#15/#16/#19: exact XML local names and attribute pairs matter."""
    from types import SimpleNamespace

    node = SimpleNamespace(tag="{urn:synthetic}item",
                           attrib={"{urn:synthetic}name": "value", "plain": "literal"})
    element = SimpleNamespace(xpath=lambda expression: [node])
    assert export_readback._layout_properties(element, ".//*") == [
        ["item", [("name", "value"), ("plain", "literal")]]
    ]


def test_408_table_block_keeps_style_nested_layout_and_picture_relationship(tmp_path):
    """Mutants #5-#9/#18/#19/#26: table keys and nested paragraphs include their real content."""
    from types import SimpleNamespace

    document = docx.Document()
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    picture_path = tmp_path / "cell-picture.png"
    picture_path.write_bytes(PNG_1X1)
    cell_paragraph = table.cell(0, 0).paragraphs[0]
    cell_paragraph.add_run().add_picture(str(picture_path))
    result = export_readback._table_block(document.part, table)
    assert result["type"] == "table" and result["style"] == "Table Grid"
    cell = result["rows"][0][0]
    assert "layout" in cell and "paragraphs" in cell
    assert cell["paragraphs"][0]["image_hashes"] == [export_readback.hashlib.sha256(PNG_1X1).hexdigest()]

    no_style = SimpleNamespace(style=None, _tbl=table._tbl, rows=[])
    assert export_readback._table_block(document.part, no_style)["style"] is None


def test_408_blocks_keep_paragraph_and_table_order_and_tag_types():
    """Mutants #6/#7/#9/#10 in _blocks: body children are identified and emitted in document order."""
    document = docx.Document()
    document.add_paragraph("before")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "inside table"
    document.add_paragraph("after")
    blocks = export_readback._document_blocks(document)
    assert [block["type"] for block in blocks] == ["paragraph", "table", "paragraph"]
    assert blocks[0]["parts"][0][1] == "before"
    assert blocks[1]["rows"][0][0]["paragraphs"][0]["parts"][0][1] == "inside table"
    assert blocks[2]["parts"][0][1] == "after"


def test_408_section_block_pins_page_margins_and_all_six_header_footer_slots():
    """Mutants #1-#18/#25/#26/#28/#29/#31/#32/#34/#35 in _section_block."""
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches

    document = docx.Document()
    section = document.sections[0]
    section.page_width = Inches(8)
    section.page_height = Inches(10)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.6)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.8)
    section.different_first_page_header_footer = True
    for story_name in ("header", "footer", "first_page_header", "first_page_footer", "even_page_header",
                       "even_page_footer"):
        story = getattr(section, story_name)
        story.is_linked_to_previous = False
        story.paragraphs[0].text = "synthetic " + story_name

    result = export_readback._section_block(section)
    assert result["type"] == "section"
    assert result["page"] == [576.0, 720.0, export_readback._text_value(section.orientation)]
    assert result["margins"] == [36.0, 43.2, 50.4, 57.6]
    assert result["different_first_page"] is True
    assert set(result) == {
        "type", "page", "margins", "different_first_page", "header", "footer",
        "first_page_header", "first_page_footer", "even_page_header", "even_page_footer",
    }
    for story_name in ("header", "footer", "first_page_header", "first_page_footer", "even_page_header",
                       "even_page_footer"):
        assert result[story_name]["linked"] is False
        assert ("synthetic " + story_name) in str(result[story_name])


def test_408_style_signature_keeps_base_style_font_colour_and_paragraph_layout():
    """Mutants #2/#7-#19/#23/#27-#31: compare every field in a directly read style definition."""
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    document = docx.Document()
    style = document.styles["Heading 1"]
    style.font.name = "Synthetic Font"
    style.font.size = Pt(17)
    style.font.bold = True
    style.font.italic = True
    style.font.underline = True
    style.font.color.rgb = RGBColor(1, 2, 3)
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    style.paragraph_format.left_indent = Pt(6)
    style.paragraph_format.space_before = Pt(7)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.25
    font, layout = style.font, style.paragraph_format
    expected = [
        style.name, export_readback._text_value(style.type),
        style.base_style.name if style.base_style is not None else None,
        font.name, export_readback._points(font.size), font.bold, font.italic,
        export_readback._text_value(font.underline), export_readback._text_value(font.color.rgb),
        export_readback._text_value(font.color.theme_color),
        export_readback._text_value(layout.alignment), export_readback._points(layout.left_indent),
        export_readback._points(layout.space_before), export_readback._points(layout.space_after),
        export_readback._text_value(layout.line_spacing),
    ]
    assert export_readback._style_signature(style) == expected
    plain_style = document.styles["Normal"]
    assert export_readback._style_signature(plain_style)[8:10] == [None, None]


def test_408_used_style_names_recurses_into_nested_dicts_and_lists():
    """Mutant #4: style names inside nested table cells and section blocks must also be found."""
    value = {"style": "Synthetic root", "nested": [{"style": "Synthetic cell"}, ["literal", {"style": "Synthetic story"}]]}
    assert export_readback._used_style_names(value, set()) == {
        "Synthetic root", "Synthetic cell", "Synthetic story"
    }


def test_408_compare_documents_pins_block_number_style_and_section_findings(monkeypatch):
    """Mutants #46/#47/#54/#77/#87/#88/#118/#128/#129: locations, shared styles, and exact artifact are observable."""
    from docx.shared import Pt
    expected_doc, actual_doc = docx.Document(), docx.Document()
    expected_doc.add_paragraph("expected block")
    actual_doc.add_paragraph("actual block")
    monkeypatch.setattr(export_readback, "_load_document",
                        lambda path: expected_doc if path == "expected" else actual_doc)
    findings = export_readback.compare_documents("expected", "actual")
    assert len(findings) == 1
    assert (findings[0]["artifact"], findings[0]["location"], findings[0]["kind"]) == (
        "document", "block 1", "paragraph"
    )

    # A style used by both documents remains checked even when paragraph text and styles do not change.
    actual_doc.paragraphs[0].text = expected_doc.paragraphs[0].text
    expected_doc.styles["Normal"].font.size = None
    actual_doc.styles["Normal"].font.size = None
    actual_doc.styles["Normal"].font.name = "Different Synthetic Font"
    findings = export_readback.compare_documents("expected", "actual")
    style_findings = [f for f in findings if f["kind"] == "style"]
    assert len(style_findings) == 1
    assert (style_findings[0]["artifact"], style_findings[0]["location"], style_findings[0]["kind"]) == (
        "document", "style Normal", "style"
    )

    # When a style is used on one side only, it is already a block difference; it is not a shared style.
    expected_doc.paragraphs[0].style = "Heading 1"
    actual_doc.paragraphs[0].style = "Normal"
    expected_doc.styles["Heading 1"].font.size = None
    actual_doc.styles["Heading 1"].font.size = Pt(1)
    findings = export_readback.compare_documents("expected", "actual")
    assert [(f["location"], f["kind"]) for f in findings] == [("block 1", "paragraph")]


def test_408_compare_documents_section_finding_has_exact_location_and_values(monkeypatch):
    """Mutants #118/#128/#129 in section comparison: section differences stay labelled document/section N."""
    from docx.shared import Inches

    expected_doc, actual_doc = docx.Document(), docx.Document()
    actual_doc.sections[0].left_margin = Inches(3)
    monkeypatch.setattr(export_readback, "_load_document",
                        lambda path: expected_doc if path == "expected" else actual_doc)
    findings = export_readback.compare_documents("expected", "actual")
    assert len(findings) == 1
    assert (findings[0]["artifact"], findings[0]["location"], findings[0]["kind"]) == (
        "document", "section 1", "section"
    )
    assert findings[0]["expected"]["margins"] != findings[0]["actual"]["margins"]


def test_408_approved_draft_error_and_base_dir_resolution_are_exact(tmp_path, monkeypatch):
    """Mutants #9/#11/#12/#29: preserve exact diagnostics and resolve repository-relative paths against base_dir."""
    from pathlib import Path
    from state_manager import StateError

    with pytest.raises(StateError) as missing:
        export_readback._approved_draft({}, str(tmp_path / "deal"), None)
    assert str(missing.value) == (
        "state.json records no approved_draft_path, so the document cannot be checked; "
        "/assemble records it when it exports"
    )

    folder = tmp_path / "deals" / "Synthetic Co" / f"Synthetic Fleet Loan_{DATE}"
    folder.mkdir(parents=True)
    draft = folder / "draft.md"
    draft.write_text("synthetic draft", encoding="utf-8")
    working = tmp_path / "working"
    working.mkdir()
    monkeypatch.chdir(working)
    recorded = f"deals/Synthetic Co/Synthetic Fleet Loan_{DATE}/draft.md"
    assert export_readback._approved_draft({"approved_draft_path": recorded}, str(folder), str(tmp_path)) == str(draft)


def test_408_verify_deal_preserves_call_arguments_report_shape_and_truncation_boundary(tmp_path, monkeypatch):
    """Mutants #3-#138/#175-#211: path plumbing, UTF-8 read, temp names, report keys and exact cap are pinned."""
    import builtins
    import tempfile
    from pathlib import Path

    base_dir = str(tmp_path)
    folder = tmp_path / "deals" / "Synthetic Co" / f"Synthetic Fleet Loan_{DATE}"
    folder.mkdir(parents=True)
    state_file = folder / "state.json"
    state_file.write_text("{}", encoding="utf-8")
    draft = folder / "draft.md"
    draft.write_text("synthetic draft", encoding="utf-8")

    calls = {"sanitize": [], "date": [], "state_path": [], "find_deal": [], "read_state": [],
             "validate_state": [], "approved": [], "open": [], "temp": [], "xlsx": [], "docx": []}

    def sanitize(value, field_name):
        calls["sanitize"].append((value, field_name))
        return value

    def validate_date(value, field_name):
        calls["date"].append((value, field_name))

    def state_path_for(company, proposal, **kwargs):
        calls["state_path"].append((company, proposal, kwargs.copy()))
        return str(state_file)

    def find_deal(company, proposal, **kwargs):
        calls["find_deal"].append((company, proposal, kwargs.copy()))
        return {"record_found": True, "date": DATE}

    def read_state(company, proposal, **kwargs):
        calls["read_state"].append((company, proposal, kwargs.copy()))
        return {"synthetic": "state"}

    def validate_state(state, **kwargs):
        calls["validate_state"].append((state, kwargs.copy()))

    def approved(state, deal_folder, passed_base_dir):
        calls["approved"].append((state, deal_folder, passed_base_dir))
        return str(draft)

    real_open = builtins.open
    def recording_open(*args, **kwargs):
        calls["open"].append(kwargs.copy())
        return real_open(*args, **kwargs)

    real_temp_directory = tempfile.TemporaryDirectory
    def recording_temp_directory(*args, **kwargs):
        temporary = real_temp_directory(*args, **kwargs)
        calls["temp"].append({"prefix": kwargs.get("prefix"), "path": temporary.name})
        return temporary

    findings_xlsx = [
        {"artifact": "workbook", "location": "Synthetic!A1", "kind": "value", "expected": 1, "actual": 2},
        {"artifact": "workbook", "location": "Synthetic!B1", "kind": "value", "expected": 3, "actual": 4},
    ]
    findings_docx = [
        {"artifact": "document", "location": "block 1", "kind": "paragraph", "expected": "old", "actual": "new"}
    ]

    monkeypatch.setattr(export_readback, "sanitize_path_component", sanitize)
    monkeypatch.setattr(export_readback, "validate_iso_date_str", validate_date)
    monkeypatch.setattr(export_readback, "state_path", state_path_for)
    monkeypatch.setattr(export_readback, "find_deal", find_deal)
    monkeypatch.setattr(export_readback, "read_state", read_state)
    monkeypatch.setattr(export_readback, "validate_state", validate_state)
    monkeypatch.setattr(export_readback, "_require_file", lambda path, what: None)
    monkeypatch.setattr(export_readback, "_approved_draft", approved)
    monkeypatch.setattr(export_readback, "workbook_inputs_from_state", lambda state: {})
    monkeypatch.setattr(export_readback, "export_to_xlsx",
                        lambda company, path, **kwargs: calls["xlsx"].append((company, path, kwargs.copy())))
    monkeypatch.setattr(export_readback, "export_to_docx",
                        lambda markdown, path: calls["docx"].append((markdown, path)))
    monkeypatch.setattr(export_readback, "compare_workbooks", lambda expected, actual: findings_xlsx)
    monkeypatch.setattr(export_readback, "compare_documents", lambda expected, actual: findings_docx)
    monkeypatch.setattr(builtins, "open", recording_open)
    monkeypatch.setattr(tempfile, "TemporaryDirectory", recording_temp_directory)
    monkeypatch.setattr(export_readback, "MAX_FINDINGS", 3)

    report = export_readback.verify_deal(
        "Synthetic Co", "Synthetic Fleet Loan", date_str=DATE, base_dir=base_dir
    )

    assert calls["sanitize"] == [("Synthetic Co", "company"), ("Synthetic Fleet Loan", "proposal")]
    assert calls["date"] == [(DATE, "--date-str")]
    assert calls["state_path"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
    ]
    assert calls["read_state"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir,
                                                   "keys": export_readback.STATE_KEYS_READ})
    ]
    assert calls["validate_state"] == [
        ({"synthetic": "state"}, {"keys": ("group_basis",), "path": str(state_file)})
    ]
    assert calls["approved"] == [({"synthetic": "state"}, str(folder), base_dir)]
    assert len(calls["open"]) == 1 and str(calls["open"][0].get("encoding", "")).lower() == "utf-8"
    assert len(calls["temp"]) == 1 and calls["temp"][0]["prefix"] == "export_readback_"
    scratch = Path(calls["temp"][0]["path"])
    assert calls["xlsx"] == [("Synthetic Co", str(scratch / "expected.xlsx"), {})]
    assert calls["docx"] == [("synthetic draft", str(scratch / "expected.docx"))]
    assert report == {
        "status": "mismatch",
        "selected_folder": str(folder),
        "artifacts": {
            "workbook": {"path": str(folder / "Synthetic Co_Synthetic Fleet Loan_Spreading.xlsx"),
                         "status": "mismatch", "mismatch_count": 2},
            "document": {"path": str(folder / "Synthetic Co_Synthetic Fleet Loan_CAM.docx"),
                         "status": "mismatch", "mismatch_count": 1},
            "approved_draft": str(draft),
        },
        "mismatch_count": 3,
        "findings": findings_xlsx + findings_docx,
        "truncated": False,
        "limits": list(export_readback.LIMITS),
    }

    # Also cover the discovery branch: the selected folder date must come from find_deal.
    calls.update({"state_path": [], "find_deal": [], "read_state": [], "validate_state": [], "approved": [],
                  "open": [], "temp": [], "xlsx": [], "docx": []})
    report = export_readback.verify_deal("Synthetic Co", "Synthetic Fleet Loan", base_dir=base_dir)
    assert calls["find_deal"] == [("Synthetic Co", "Synthetic Fleet Loan", {"base_dir": base_dir})]
    assert calls["state_path"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
    ]
    assert calls["read_state"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir,
                                                   "keys": export_readback.STATE_KEYS_READ})
    ]
    assert report["selected_folder"] == str(folder) and report["truncated"] is False


def test_408_main_help_required_flags_date_forwarding_and_json_shape(capsys, monkeypatch):
    """Pin help invocation, required company, date forwarding and exact JSON output shape."""
    with pytest.raises(SystemExit) as help_exit:
        export_readback.main(["--help"])
    assert help_exit.value.code == 0
    capsys.readouterr()  # Help wording is prose; only the CLI contract is pinned here.

    with pytest.raises(SystemExit) as required_exit:
        export_readback.main(["--proposal", "Synthetic Fleet Loan"])
    assert required_exit.value.code == 2
    assert "the following arguments are required: --company" in capsys.readouterr().err

    calls = []
    monkeypatch.setattr(export_readback, "verify_deal", lambda company, proposal, **kwargs: (
        calls.append((company, proposal, kwargs.copy())) or {"status": "match"}
    ))
    assert export_readback.main([
        "--company", "Synthetic Co", "--proposal", "Synthetic Fleet Loan", "--date-str", DATE
    ]) == 0
    assert calls == [("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE})]
    assert capsys.readouterr().out == '{\n  "status": "match"\n}\n'


def test_408_named_equivalent_text_and_path_spelling_cases():
    """Explicit aliases/redundancies that have no distinct observable behavior are recorded here."""
    # rsplit with one split and rsplit with no maxsplit both return the same final component.
    assert "tag".rsplit("}", 1)[-1] == "tag".rsplit("}")[-1]
    # UTF-8 codec names are case-insensitive aliases.
    assert "café".encode("utf-8") == "café".encode("UTF-8")
