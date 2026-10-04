"""Normalisers and the compare/update helper for the golden-output regression tests (issue #145).

A .docx or .xlsx is a zip of XML that embeds timestamps and ordering details, so it can never be compared byte for
byte. These functions reduce a generated file to the *content a reader would see* as plain text -- one line per
block or cell -- so a snapshot is a readable, reviewable diff:

* docx: body blocks in order -- paragraphs as `P[<style>] <text>` with `**bold**`, `*italic*`, `` `monospace` `` and
  `<link: url | text>` markers (adjacent runs with the same formatting are merged, so how the builder splits runs
  is not part of the snapshot), tables as `TABLE` plus one `| cell | cell |` line per row (a cell's alignment
  and bold are shown), an embedded image as `<image alt="description">` (its caption is an ordinary italic
  paragraph; an image that could not be embedded is the visible placeholder paragraph). Core properties (created/modified timestamps, author) are never read.
* xlsx: for each sheet its name, used range, column widths and merged ranges, then every non-empty cell as
  `<coordinate>: <value or formula>` plus its number format and bold/italic when not default. Formulas are kept
  as text (the file is never opened by Excel, so there are no cached values).

Nothing here is imported by production code.
"""
import difflib
import re
from pathlib import Path

import docx
import openpyxl
from docx.oxml.ns import qn

SNAPSHOT_DIR = Path(__file__).resolve().parent / "snapshots"
MONOSPACE_FONTS = {"Consolas", "Courier New"}
BODY_PARAGRAPH, BODY_TABLE = qn("w:p"), qn("w:tbl")


def _toggle_on(properties, tag):
    """A WordprocessingML toggle (<w:b/>) is on unless it says <w:b w:val="0"/> (or "false"/"off")."""
    element = properties.find(qn(tag))
    return element is not None and element.get(qn("w:val"), "1").lower() not in ("0", "false", "off")


def _run_format(run_element):
    properties = run_element.find(qn("w:rPr"))
    bold = italic = mono = False
    if properties is not None:
        bold = _toggle_on(properties, "w:b")
        italic = _toggle_on(properties, "w:i")
        fonts = properties.find(qn("w:rFonts"))
        mono = fonts is not None and fonts.get(qn("w:ascii")) in MONOSPACE_FONTS
    return bold, italic, mono


def _run_text(run_element):
    parts = []
    for child in run_element:
        if child.tag == qn("w:t"):
            parts.append(child.text or "")
        elif child.tag == qn("w:tab"):
            parts.append("\t")
        elif child.tag == qn("w:br"):
            parts.append("\n")
        elif child.tag == qn("w:drawing"):
            # The accessibility description is part of what a reader (or a screen reader) gets, so a lost one shows.
            properties = child.find(".//" + qn("wp:docPr"))
            alt = properties.get("descr") if properties is not None else None
            parts.append(f'<image alt="{alt}">' if alt else "<image>")
    return "".join(parts)


def _wrap(text, bold, italic, mono):
    if not text:
        return ""
    if mono:
        text = f"`{text}`"
    if italic:
        text = f"*{text}*"
    if bold:
        text = f"**{text}**"
    return text


def _paragraph_text(paragraph_element, document):
    """Merge adjacent runs of identical formatting, then render with markers."""
    segments = []   # [(format-key, text)]
    for child in paragraph_element:
        if child.tag == qn("w:r"):
            key, text = _run_format(child), _run_text(child)
        elif child.tag == qn("w:hyperlink"):
            rel = child.get(qn("r:id"))
            url = document.part.rels[rel].target_ref if rel in document.part.rels else "?"
            key = ("link", url)
            text = "".join(_run_text(r) for r in child.findall(qn("w:r")))
        else:
            continue
        if segments and segments[-1][0] == key:
            segments[-1] = (key, segments[-1][1] + text)
        else:
            segments.append((key, text))
    out = []
    for key, text in segments:
        if key and key[0] == "link":
            out.append(f"<link: {key[1]} | {text}>")
        else:
            out.append(_wrap(text, *key))
    return "".join(out)


def _style_name(paragraph_element, document):
    properties = paragraph_element.find(qn("w:pPr"))
    style_id = None
    if properties is not None and properties.find(qn("w:pStyle")) is not None:
        style_id = properties.find(qn("w:pStyle")).get(qn("w:val"))
    if style_id is None:
        return "Normal"
    for style in document.styles:
        if style.style_id == style_id:
            return style.name
    return style_id


def _alignment(paragraph_element):
    properties = paragraph_element.find(qn("w:pPr"))
    if properties is not None and properties.find(qn("w:jc")) is not None:
        return properties.find(qn("w:jc")).get(qn("w:val"))
    return "left"


def _table_lines(table_element, document):
    rows = table_element.findall(qn("w:tr"))
    lines = [f"TABLE {len(rows)} rows"]
    for row in rows:
        cells = []
        for cell in row.findall(qn("w:tc")):
            paragraphs = cell.findall(qn("w:p"))
            text = " / ".join(_paragraph_text(p, document) for p in paragraphs)
            align = _alignment(paragraphs[0]) if paragraphs else "left"
            cells.append(text if align == "left" else f"{text} ({align})")
        lines.append("| " + " | ".join(cells) + " |")
    return lines


def normalize_docx(path):
    document = docx.Document(str(path))
    lines = []
    for element in document.element.body:
        if element.tag == BODY_PARAGRAPH:
            lines.append(f"P[{_style_name(element, document)}] {_paragraph_text(element, document)}".rstrip())
        elif element.tag == BODY_TABLE:
            lines.extend(_table_lines(element, document))
    return "\n".join(lines) + "\n"


def _cell_attrs(cell):
    notes = []
    if cell.number_format != "General":
        notes.append(f"fmt={cell.number_format}")
    if cell.font is not None and cell.font.bold:
        notes.append("bold")
    if cell.font is not None and cell.font.italic:
        notes.append("italic")
    return f"  [{', '.join(notes)}]" if notes else ""


def normalize_xlsx(path):
    workbook = openpyxl.load_workbook(str(path))   # not data_only: keep formulas as text
    lines = []
    for sheet in workbook.worksheets:
        lines.append(f"== sheet {sheet.title!r} dims={sheet.dimensions}")
        widths = {k: round(v.width, 1) for k, v in sorted(sheet.column_dimensions.items()) if v.width}
        if widths:
            lines.append(f"widths: {widths}")
        if sheet.merged_cells.ranges:
            lines.append(f"merged: {sorted(str(r) for r in sheet.merged_cells.ranges)}")
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is not None:
                    lines.append(f"{cell.coordinate}: {cell.value!r}{_cell_attrs(cell)}")
    return "\n".join(lines) + "\n"


def compare_to_snapshot(name, actual, update):
    """Assert `actual` equals tests/snapshots/<name>; with `update`, (re)write it instead.

    A missing snapshot is a failure, never silently created: a snapshot that appeared by itself would be a test that
    can never fail the first time it runs. Line endings are normalised so a Windows checkout compares equal.
    """
    path = SNAPSHOT_DIR / name
    if update:
        path.parent.mkdir(parents=True, exist_ok=True)
        with open(path, "w", encoding="utf-8", newline="\n") as f:
            f.write(actual)
        return
    assert path.exists(), (f"no snapshot {path.name}; review the output below, then create it with "
                           f"`pytest --update-snapshots`:\n{actual}")
    expected = path.read_text(encoding="utf-8").replace("\r\n", "\n")
    if expected != actual:
        diff = "".join(difflib.unified_diff(expected.splitlines(True), actual.splitlines(True),
                                            f"snapshots/{name} (expected)", "generated output"))
        raise AssertionError(f"output differs from snapshots/{name}. If the change is intended, review this diff "
                             f"and run `pytest --update-snapshots`:\n{diff}")


FORBIDDEN_PATTERNS = [
    (r"```", "a raw code fence leaked into the document"),
    (r"cp_ids_included|risk_categories_covered|reported_figures", "the structured-output JSON leaked into the document"),
    (r"(?m)^\s*\|?\s*-{3,}\s*\|", "a raw Markdown table separator leaked into the document"),
    (r"\*\*\S", "a raw Markdown bold marker leaked into the document"),
    (r"(?m)^#{1,6} ", "a raw Markdown heading marker leaked into the document"),
    (r"\bnan\b|\binf\b|\bNone\b", "a Python null/not-a-number value leaked into the document"),
    (r"Traceback \(most recent call last\)|\b[A-Z][A-Za-z]*Error: |\bException\b", "an error message leaked into the document"),
    (r"\[Image not embedded", "an image failed to embed"),
]


def forbidden_content(text):
    """Reasons `text` (a document's visible text) contains something that should never reach a client."""
    return [reason for pattern, reason in FORBIDDEN_PATTERNS if re.search(pattern, text)]


def visible_text_docx(path):
    """The text a reader sees (paragraphs and table cells), without this module's formatting markers."""
    document = docx.Document(str(path))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(parts)


def visible_text_xlsx(path):
    workbook = openpyxl.load_workbook(str(path))
    return "\n".join(str(cell.value) for sheet in workbook.worksheets for row in sheet.iter_rows()
                     for cell in row if cell.value is not None)
