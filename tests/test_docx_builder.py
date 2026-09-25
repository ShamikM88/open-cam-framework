import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from docx_builder import export_to_docx


def _build(tmp_path, markdown_text):
    out = tmp_path / "out.docx"
    export_to_docx(markdown_text, str(out))
    return docx.Document(str(out))


def test_headings(tmp_path):
    doc = _build(tmp_path, "# Title\n## Subtitle\n")
    assert doc.paragraphs[0].text == "Title"
    assert doc.paragraphs[0].style.name == "Heading 1"
    assert doc.paragraphs[1].text == "Subtitle"
    assert doc.paragraphs[1].style.name == "Heading 2"


def test_bullet_list(tmp_path):
    doc = _build(tmp_path, "- First\n- Second\n")
    assert [p.text for p in doc.paragraphs] == ["First", "Second"]
    assert doc.paragraphs[0].style.name == "List Bullet"


def test_plain_paragraph(tmp_path):
    doc = _build(tmp_path, "Just a sentence.\n")
    assert doc.paragraphs[0].text == "Just a sentence."


def test_blank_lines_are_skipped(tmp_path):
    doc = _build(tmp_path, "First.\n\n\nSecond.\n")
    assert [p.text for p in doc.paragraphs] == ["First.", "Second."]


MARKDOWN_TABLE = (
    "| Name | Amount | Notes |\n"
    "| :--- | ---: | :---: |\n"
    "| Alpha | 100 | ok |\n"
    "| Beta | 200 | **flag** |\n"
)


def test_table_is_rendered_as_a_real_table(tmp_path):
    doc = _build(tmp_path, MARKDOWN_TABLE)
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert len(table.rows) == 3  # header + 2 body rows
    assert len(table.columns) == 3
    assert [c.text for c in table.rows[0].cells] == ["Name", "Amount", "Notes"]
    assert [c.text for c in table.rows[1].cells] == ["Alpha", "100", "ok"]
    # inline "**bold**" markers are stripped from cell text
    assert [c.text for c in table.rows[2].cells] == ["Beta", "200", "flag"]


def test_table_header_is_bold_and_body_is_not(tmp_path):
    doc = _build(tmp_path, MARKDOWN_TABLE)
    table = doc.tables[0]
    header_run = table.rows[0].cells[0].paragraphs[0].runs[0]
    body_run = table.rows[1].cells[0].paragraphs[0].runs[0]
    assert header_run.bold is True
    assert body_run.bold in (False, None)


def test_table_column_alignment_follows_separator_row(tmp_path):
    doc = _build(tmp_path, MARKDOWN_TABLE)
    header_cells = doc.tables[0].rows[0].cells
    assert header_cells[0].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert header_cells[1].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert header_cells[2].paragraphs[0].alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_content_after_table_continues_as_a_paragraph(tmp_path):
    doc = _build(tmp_path, MARKDOWN_TABLE + "After the table.\n")
    assert len(doc.tables) == 1
    assert doc.paragraphs[-1].text == "After the table."


def test_ragged_table_rows_do_not_crash(tmp_path):
    markdown = (
        "| A | B |\n"
        "| --- | --- |\n"
        "| 1 |\n"          # missing a column -> padded with a blank cell
        "| 2 | 3 | 4 |\n"  # extra column -> silently dropped
    )
    doc = _build(tmp_path, markdown)
    table = doc.tables[0]
    assert [c.text for c in table.rows[1].cells] == ["1", ""]
    assert [c.text for c in table.rows[2].cells] == ["2", "3"]


def test_multiple_tables_in_one_document(tmp_path):
    markdown = MARKDOWN_TABLE + "\nSome text.\n\n" + MARKDOWN_TABLE
    doc = _build(tmp_path, markdown)
    assert len(doc.tables) == 2


def test_table_style_has_visible_borders(tmp_path):
    doc = _build(tmp_path, MARKDOWN_TABLE)
    assert doc.tables[0].style.name == "Table Grid"


# ---------------------------------------------------------------------------
# Inline markdown (bold/italic/code) in headings-adjacent body text, H3, and
# stripping content that has no place in a client-facing Word document (a
# horizontal rule, or a fenced code block like the Underwriter's trailing
# structured JSON block).
# ---------------------------------------------------------------------------

def test_h3_heading(tmp_path):
    doc = _build(tmp_path, "### Sub-subtitle\n")
    assert doc.paragraphs[0].text == "Sub-subtitle"
    assert doc.paragraphs[0].style.name == "Heading 3"


def test_bold_in_plain_paragraph_renders_as_a_bold_run(tmp_path):
    doc = _build(tmp_path, "**Verdict:** Approve.\n")
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "Verdict: Approve."
    assert paragraph.runs[0].text == "Verdict:"
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[1].bold in (False, None)


def test_italic_in_plain_paragraph_renders_as_an_italic_run(tmp_path):
    doc = _build(tmp_path, "*A note in italics.*\n")
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "A note in italics."
    assert paragraph.runs[0].italic is True


def test_bold_in_bullet_list_renders_as_a_bold_run(tmp_path):
    doc = _build(tmp_path, "- **Label:** detail\n")
    paragraph = doc.paragraphs[0]
    assert paragraph.style.name == "List Bullet"
    assert paragraph.text == "Label: detail"
    assert paragraph.runs[0].bold is True


def test_horizontal_rule_is_skipped(tmp_path):
    doc = _build(tmp_path, "First.\n\n---\n\nSecond.\n")
    assert [p.text for p in doc.paragraphs] == ["First.", "Second."]


def test_fenced_code_block_is_skipped_entirely(tmp_path):
    markdown = (
        "Before the block.\n"
        "```json\n"
        '{"verdict": "APPROVED", "notes": null}\n'
        "```\n"
        "After the block.\n"
    )
    doc = _build(tmp_path, markdown)
    assert [p.text for p in doc.paragraphs] == ["Before the block.", "After the block."]
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "verdict" not in full_text
    assert "```" not in full_text


def test_unterminated_fence_does_not_discard_the_rest_of_the_document(tmp_path):
    """A stray/odd ``` (e.g. from truncation) must not silently swallow
    every line through EOF -- only a genuinely closed fence gets skipped."""
    markdown = (
        "Before.\n"
        "```json\n"
        '{"unterminated": true\n'
        "## Section 2\n"
        "Real narrative content that must survive.\n"
    )
    doc = _build(tmp_path, markdown)
    texts = [p.text for p in doc.paragraphs]
    assert "Before." in texts
    assert "Section 2" in texts
    assert "Real narrative content that must survive." in texts


def test_footnote_style_trailing_asterisk_is_not_treated_as_italic(tmp_path):
    """A single `*` used as a footnote marker (common in financial
    narrative, e.g. "Net Income* is 5.2x... Note 1*") must not be
    misread as an emphasis delimiter -- that would silently eat both
    asterisks and italicize unrelated text in between."""
    doc = _build(tmp_path, "Net Income* is 5.2x Interest Expense, per Note 1*.\n")
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "Net Income* is 5.2x Interest Expense, per Note 1*."
    assert all(not r.italic for r in paragraph.runs)


def test_triple_asterisk_renders_as_a_single_bold_italic_run(tmp_path):
    doc = _build(tmp_path, "***Critical:*** breach detected\n")
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "Critical: breach detected"
    assert paragraph.runs[0].text == "Critical:"
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[0].italic is True
    assert "*" not in paragraph.text


# ---------------------------------------------------------------------------
# Soft-wrapped source lines must join into one continuous Word paragraph
# (matching normal Markdown semantics -- a single newline is a soft wrap,
# only a blank line or a new block starts a fresh paragraph), not fragment
# into one Word paragraph per source line.
# ---------------------------------------------------------------------------

def test_wrapped_plain_paragraph_joins_into_one_paragraph(tmp_path):
    markdown = "This sentence was wrapped\nacross three separate\nsource lines.\n"
    doc = _build(tmp_path, markdown)
    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "This sentence was wrapped across three separate source lines."


def test_blank_line_still_separates_wrapped_paragraphs(tmp_path):
    markdown = "First para line one\nfirst para line two.\n\nSecond para line one\nsecond para line two.\n"
    doc = _build(tmp_path, markdown)
    assert [p.text for p in doc.paragraphs] == [
        "First para line one first para line two.",
        "Second para line one second para line two.",
    ]


def test_heading_ends_a_wrapped_paragraph(tmp_path):
    markdown = "Body text that wraps\nonto a second line.\n## Next Section\n"
    doc = _build(tmp_path, markdown)
    assert doc.paragraphs[0].text == "Body text that wraps onto a second line."
    assert doc.paragraphs[1].text == "Next Section"
    assert doc.paragraphs[1].style.name == "Heading 2"


def test_wrapped_bullet_item_joins_into_one_bullet(tmp_path):
    markdown = "- A bullet whose text wraps\n  onto a continuation line.\n"
    doc = _build(tmp_path, markdown)
    assert len(doc.paragraphs) == 1
    assert doc.paragraphs[0].text == "A bullet whose text wraps onto a continuation line."
    assert doc.paragraphs[0].style.name == "List Bullet"


def test_two_wrapped_bullets_stay_separate(tmp_path):
    markdown = "- First bullet wraps\n  onto a second line.\n- Second bullet, single line.\n"
    doc = _build(tmp_path, markdown)
    assert [p.text for p in doc.paragraphs] == [
        "First bullet wraps onto a second line.",
        "Second bullet, single line.",
    ]


def test_numbered_list_items_are_not_merged_together(tmp_path):
    """A numbered-list marker ('1. ', '2. ', ...) must still end the
    previous item's paragraph even though it isn't a bullet/heading/table
    row -- otherwise every item in the list collapses into one paragraph."""
    markdown = (
        "1. First item, short.\n"
        "2. Second item wraps\n"
        "   onto a continuation line.\n"
        "3. Third item, short.\n"
    )
    doc = _build(tmp_path, markdown)
    assert [p.text for p in doc.paragraphs] == [
        "1. First item, short.",
        "2. Second item wraps onto a continuation line.",
        "3. Third item, short.",
    ]


def test_table_still_ends_a_preceding_wrapped_paragraph(tmp_path):
    markdown = "Intro text that wraps\nonto a second line.\n" + MARKDOWN_TABLE
    doc = _build(tmp_path, markdown)
    assert doc.paragraphs[0].text == "Intro text that wraps onto a second line."
    assert len(doc.tables) == 1


# ---------------------------------------------------------------------------
# Markdown links ([text](url)) must render as real, clickable Word
# hyperlinks -- not pass through as literal bracket/paren text.
# ---------------------------------------------------------------------------

def _hyperlink_targets(doc, paragraph):
    """The external relationship URL(s) referenced by every w:hyperlink
    element in `paragraph`, resolved via the document part's relationship
    table (python-docx has no high-level API for reading hyperlinks back)."""
    rels = paragraph.part.rels
    return [
        rels[el.get(qn("r:id"))].target_ref
        for el in paragraph._p.findall(qn("w:hyperlink"))
    ]


def test_markdown_link_renders_as_a_real_hyperlink(tmp_path):
    doc = _build(tmp_path, "See [Companies House](https://example.com/company/123) for details.\n")
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "See Companies House for details."
    assert "[" not in paragraph.text and "(" not in paragraph.text
    assert _hyperlink_targets(doc, paragraph) == ["https://example.com/company/123"]


def test_markdown_link_in_a_bullet_item_renders_as_a_hyperlink(tmp_path):
    doc = _build(tmp_path, "- [Source A](https://example.com/a)\n- [Source B](https://example.com/b)\n")
    assert [p.text for p in doc.paragraphs] == ["Source A", "Source B"]
    assert _hyperlink_targets(doc, doc.paragraphs[0]) == ["https://example.com/a"]
    assert _hyperlink_targets(doc, doc.paragraphs[1]) == ["https://example.com/b"]


def test_multiple_links_in_one_paragraph(tmp_path):
    doc = _build(tmp_path, "See [A](https://example.com/a) and [B](https://example.com/b).\n")
    paragraph = doc.paragraphs[0]
    assert paragraph.text == "See A and B."
    assert _hyperlink_targets(doc, paragraph) == ["https://example.com/a", "https://example.com/b"]
