"""Every command-line entry point, run for real as a subprocess (issue #144).

The slash commands call these scripts from a Bash step (`python scripts/<name>.py ...`), so the argparse
setup and the `if __name__ == "__main__":` block are the actual interface. Unit tests that call `main()` or an
internal function never run that code. Each script here is executed in a throwaway working directory (every
script resolves `deals/`, `config/`, `templates/` and `inputs/` against the working directory) with the
Anthropic credentials removed from its environment, and must

- print usage and exit 0 for `--help`,
- exit 2 with a usage message for an unknown flag,
- do its job on a representative minimal input (where that is possible without a model call),
- fail with a non-zero exit when given an input that cannot work.

`coverage` follows the child processes (`patch = ["subprocess"]` in pyproject.toml), so these runs count
toward the coverage of the `__main__` blocks.
"""
import ast
import atexit
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import docx
import pytest

from test_source_manifest import MALFORMED_SOURCE_STATES, VALID_SOURCE_STATES

REPO_ROOT = Path(__file__).resolve().parents[1]
SCRIPTS_DIR = REPO_ROOT / "scripts"

# Every script that defines a command-line entry point. A new one must be added here AND given tests below;
# test_the_list_matches_the_scripts_that_have_a_main_block fails until it is.
CLI_SCRIPTS = (
    "calibrate", "check_coverage", "check_model_freshness", "check_test_count", "conventions", "deal_export",
    "mutation_report", "orchestrator", "pii_scan", "policy_check", "research_export", "review_compare", "run_evals", "source_manifest",
    "spreading_check", "state_manager",
    "supplied_forecast", "transcription_check", "wiki_publish",
)


def has_main_block(path):
    """True if the module has a top-level `if __name__ == "__main__":` (parsed, so a mention in a docstring
    or a comment does not count)."""
    tree = ast.parse(path.read_text(encoding="utf-8"))
    for node in tree.body:
        if isinstance(node, ast.If) and isinstance(node.test, ast.Compare):
            left, comparators = node.test.left, node.test.comparators
            if (isinstance(left, ast.Name) and left.id == "__name__" and len(comparators) == 1
                    and isinstance(comparators[0], ast.Constant) and comparators[0].value == "__main__"):
                return True
    return False


# An empty directory standing in for the user's home and config directories in every child process. Removing the
# ANTHROPIC_* variables is not enough: the SDK also reads a credentials profile from the user's config directory
# (~/.config/anthropic, %APPDATA%\\Anthropic), so on a machine with one configured a child could authenticate and
# make a real, paid model call. Pointing the home variables at an empty directory removes that path as well.
_ISOLATED_HOME = tempfile.mkdtemp(prefix="cli-home-")
atexit.register(shutil.rmtree, _ISOLATED_HOME, ignore_errors=True)
HOME_VARIABLES = ("HOME", "USERPROFILE", "APPDATA", "LOCALAPPDATA", "XDG_CONFIG_HOME", "XDG_DATA_HOME")


def child_env():
    """The parent's environment minus anything that could reach a model or an account: no ANTHROPIC_* variable and
    no real home or config directory (so no on-disk SDK credentials profile either)."""
    env = {k: v for k, v in os.environ.items() if not k.upper().startswith("ANTHROPIC")}
    env.update(dict.fromkeys(HOME_VARIABLES, _ISOLATED_HOME))
    env["PYTHONIOENCODING"] = "utf-8"
    return env


@pytest.fixture
def cli(tmp_path):
    """run(script, *args) -> CompletedProcess, in an empty working directory."""
    workdir = tmp_path / "work"
    workdir.mkdir()

    def run(script, *args, cwd=None):
        return subprocess.run(
            [sys.executable, str(SCRIPTS_DIR / f"{script}.py"), *map(str, args)],
            cwd=cwd or workdir, env=child_env(), capture_output=True, text=True, encoding="utf-8", timeout=120,
            stdin=subprocess.DEVNULL)

    run.workdir = workdir
    return run


# ---------------------------------------------------------------------------
# The list of CLIs is exactly the scripts that have one
# ---------------------------------------------------------------------------

def test_the_list_matches_the_scripts_that_have_a_main_block():
    found = sorted(p.stem for p in SCRIPTS_DIR.glob("*.py") if has_main_block(p))
    assert found == sorted(CLI_SCRIPTS), (
        "scripts with an `if __name__ == \"__main__\":` block and CLI_SCRIPTS disagree; add the new script "
        "to CLI_SCRIPTS and give it --help / bad-flag / representative tests in this file")


@pytest.mark.parametrize("source, expected", [
    ('if __name__ == "__main__":\n    pass\n', True),
    ("if __name__ == '__main__':\n    pass\n", True),
    ('"""if __name__ == \\"__main__\\": in a docstring"""\nx = 1\n', False),
    ('# if __name__ == "__main__":\nx = 1\n', False),
    ('def main():\n    if __name__ == "__main__":\n        pass\n', False),   # not top level
])
def test_the_main_block_detector_reads_code_not_text(tmp_path, source, expected):
    path = tmp_path / "m.py"
    path.write_text(source, encoding="utf-8")
    assert has_main_block(path) is expected


# ---------------------------------------------------------------------------
# --help and a bad flag, for every script
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("script", CLI_SCRIPTS)
def test_help_prints_usage_and_exits_zero(cli, script):
    result = cli(script, "--help")
    assert result.returncode == 0, result.stderr
    assert result.stdout.startswith("usage:") and result.stderr == ""


@pytest.mark.parametrize("script", CLI_SCRIPTS)
def test_an_unknown_flag_exits_two_with_usage_on_stderr(cli, script):
    result = cli(script, "--no-such-flag-xyz")
    assert result.returncode == 2
    assert "usage:" in result.stderr and result.stdout == ""


# ---------------------------------------------------------------------------
# Representative real runs
# ---------------------------------------------------------------------------

def test_pii_scan_reports_findings_as_json(cli):
    clean = cli.workdir / "clean.md"
    clean.write_text("# Template\n\nBorrower: [Company Name]\n", encoding="utf-8")
    dirty = cli.workdir / "dirty.md"
    dirty.write_text("Contact jane.doe@example.com about GBP 1,250,000.\n", encoding="utf-8")
    assert json.loads(cli("pii_scan", clean).stdout) == {"findings": []}
    findings = json.loads(cli("pii_scan", dirty).stdout)["findings"]
    assert findings, "an email address and a currency figure must be reported"


def test_pii_scan_on_a_file_that_is_not_utf8_is_one_error_line(cli):
    binary = cli.workdir / "binary.md"
    binary.write_bytes(b"\xff\xfe not text")
    _assert_a_clean_state_error(cli("pii_scan", binary), "codec can't decode")


def test_pii_scan_on_a_missing_file_fails(cli):
    result = cli("pii_scan", "no-such-file.md")
    _assert_a_clean_state_error(result, "no-such-file.md")




def test_check_model_freshness_real_cli_reports_installed_sdk_shape_offline(cli, tmp_path):
    report_path = tmp_path / "model-freshness.json"
    result = cli("check_model_freshness", "--json", report_path, cwd=REPO_ROOT)
    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    assert "only as fresh as the installed Anthropic SDK" in result.stdout
    assert "does not query the API" in result.stdout
    with open(report_path, encoding="utf-8") as f:
        report = json.load(f)
    assert set(report) == {
        "report_version", "sdk_version", "freshness_note", "sdk_tables", "models", "findings", "summary"
    }
    assert report["report_version"] == 1
    assert report["models"] and report["summary"]["models_checked"] == len(report["models"])
    assert report["summary"]["finding_count"] == len(report["findings"])
    assert report["summary"]["has_findings"] is bool(report["findings"])
    assert set(report["sdk_tables"]) == {"known_models", "deprecated_models"}
    assert all(item["status"] in {"available", "SDK table not found"}
               for item in report["sdk_tables"].values())


def test_state_manager_check_steps_reports_missing_steps_and_always_exits_zero(cli):
    result = cli("state_manager", "--check-steps", "--company", "Acme", "--proposal", "Loan",
                 "--required", "spread, collateral")
    assert result.returncode == 0
    assert result.stdout == '{\n  "missing_steps": [\n    "spread",\n    "collateral"\n  ],\n  "ok": false\n}\n'
    assert json.loads(result.stdout) == {"missing_steps": ["spread", "collateral"], "ok": False}


def test_source_manifest_saves_a_source_then_reports_it(cli):
    source = cli.workdir / "filing.txt"
    source.write_text("the filing", encoding="utf-8")
    before = json.loads(cli("source_manifest", "--company", "Acme", "--proposal", "Loan", "--check-sources").stdout)
    assert before == {"missing_saved_sources": False, "sources_declared": "not_declared"}   # nothing declared, nothing missing
    saved = cli("source_manifest", "--company", "Acme", "--proposal", "Loan", "--step", "triage",
                "--claim", "legal identity", "--file", source, "--url", "https://example.invalid/filing")
    assert saved.returncode == 0, saved.stderr
    entry = json.loads(saved.stdout)
    assert entry["step"] == "triage" and entry["claim"] == "legal identity"
    manifests = list(cli.workdir.glob("deals/Acme/Loan_*/sources/manifest.json"))
    assert len(manifests) == 1 and json.loads(manifests[0].read_text(encoding="utf-8"))[0]["step"] == "triage"


def test_source_manifest_requires_its_save_arguments_unless_checking(cli):
    result = cli("source_manifest", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 2 and "--step" in result.stderr and "--claim" in result.stderr


def test_policy_check_computes_the_policy_state_for_a_new_deal(cli):
    result = cli("policy_check", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 0, result.stderr
    state = json.loads(result.stdout)["policy_state"]
    assert {c["cp_id"] for c in state["required_conditions_precedent"]} >= {"KYC-AML"}


def test_conventions_write_then_read_round_trip_for_a_company_and_the_enterprise(cli):
    written = cli("conventions", "--company", "Acme", "--proposal", "Loan", "--write",
                  "--financials-source", "analyst-supplied", "--note", "Depreciation sits inside COGS")
    assert written.returncode == 0, written.stderr
    record = json.loads(written.stdout)
    assert record["financials_source_default"] == "analyst-supplied"
    read = json.loads(cli("conventions", "--company", "Acme", "--read").stdout)
    assert read["found"] is True and read["convention"]["financials_source_note"] == "Depreciation sits inside COGS"
    assert json.loads(cli("conventions", "--enterprise", "--read").stdout) == {"found": False, "convention": None}
    assert cli("conventions", "--enterprise", "--write", "--financials-source", "framework-computed").returncode == 0
    assert json.loads(cli("conventions", "--enterprise", "--read").stdout)["found"] is True


@pytest.mark.parametrize("args, fragment", [
    (("--enterprise", "--write"), "--financials-source"),
    (("--proposal", "Loan", "--enterprise", "--read"), "--company"),
    (("--read",), "required"),
    (("--company", "A", "--enterprise", "--read"), "not allowed"),
])
def test_conventions_rejects_inconsistent_arguments(cli, args, fragment):
    result = cli("conventions", *args)
    assert result.returncode == 2 and fragment in result.stderr


def test_deal_export_writes_the_docx_and_xlsx_into_the_dated_folder(cli):
    draft = cli.workdir / "draft.md"
    draft.write_text("# Credit Assessment Memorandum\n\n- **Borrower:** [Company Name]\n\nBody text.\n",
                     encoding="utf-8")
    result = cli("deal_export", "--company", "Acme", "--proposal", "Loan", "--type", "corporate_credit",
                 "--draft", draft)
    assert result.returncode == 0, result.stderr
    assert "Done! Files generated in" in result.stdout
    folder = next(cli.workdir.glob("deals/Acme/Loan_*"))
    assert (folder / "Acme_Loan_CAM.docx").stat().st_size > 0 and (folder / "Acme_Loan_Spreading.xlsx").stat().st_size > 0


def test_deal_export_rejects_a_malformed_date_and_a_missing_draft(cli):
    draft = cli.workdir / "draft.md"
    draft.write_text("# T\n", encoding="utf-8")
    for bad in ("foo", "2026-02-30"):
        bad_date = cli("deal_export", "--company", "A", "--proposal", "P", "--draft", draft, "--date-str", bad)
        assert bad_date.returncode == 1 and bad_date.stderr.startswith("error: --date-str ")
        assert bad_date.stderr.count("\n") == 1 and "usage:" not in bad_date.stderr and "Traceback" not in bad_date.stderr
    missing = cli("deal_export", "--company", "A", "--proposal", "P", "--draft", "nope.md")
    _assert_a_clean_state_error(missing, "nope.md")
    assert not (cli.workdir / "deals").exists()  # nothing was created for either failure


def test_research_export_help_documents_the_analyst_override_flag(cli):
    result = cli("research_export", "--help")
    assert result.returncode == 0, result.stderr
    assert "--analyst-override-no-sources" in result.stdout and "REASON" in result.stdout
    assert "--analyst-release-unapproved" in result.stdout
    assert "explicit analyst authorisation" in result.stdout


def test_research_export_writes_the_brief_docx(cli):
    brief = cli.workdir / "brief.md"
    brief.write_text("# Research brief\n\n- **Verdict:** Go\n\nSector notes [placeholder].\n\n## Sources\n\n- Synthetic registry: https://example.invalid/registry\n", encoding="utf-8")
    result = cli("research_export", "--company", "Acme", "--proposal", "Brief", "--brief", brief)
    assert result.returncode == 0, result.stderr
    assert "Research brief exported to" in result.stdout
    assert list(cli.workdir.glob("deals/Acme/Brief_*/*Research_Brief.docx"))


def test_research_export_without_sources_fails_cleanly_without_writing(cli):
    brief = cli.workdir / "brief.md"
    brief.write_text("# Synthetic brief\n\nFindings refer to sources but give no list.\n", encoding="utf-8")
    result = cli("research_export", "--company", "Acme", "--proposal", "Brief", "--brief", brief)
    assert result.returncode == 1
    assert len(result.stderr.splitlines()) == 1 and result.stderr.startswith("error: ")
    assert "populated Sources or References section" in result.stderr
    assert "Traceback" not in result.stderr and result.stdout == ""
    assert not (cli.workdir / "deals").exists()


def test_research_export_override_succeeds_and_discloses_the_reason(cli):
    import docx

    brief = cli.workdir / "brief.md"
    brief.write_text("# Synthetic brief\n\nSynthetic findings only.\n", encoding="utf-8")
    result = cli(
        "research_export", "--company", "Acme", "--proposal", "Brief", "--brief", brief,
        "--analyst-override-no-sources", "Synthetic analyst-authorised reason",
    )
    assert result.returncode == 0, result.stderr
    output = next(cli.workdir.glob("deals/Acme/Brief_*/*Research_Brief.docx"))
    footnote = docx.Document(output).paragraphs[-1].text
    assert footnote.startswith("Footnote: No sources list is included in this brief. Exported under analyst override on ")
    assert footnote.endswith(": Synthetic analyst-authorised reason.")


def test_research_export_override_without_a_reason_fails_cleanly(cli):
    brief = cli.workdir / "brief.md"
    brief.write_text("# Synthetic brief\n\nSynthetic findings only.\n", encoding="utf-8")
    result = cli(
        "research_export", "--company", "Acme", "--proposal", "Brief", "--brief", brief,
        "--analyst-override-no-sources",
    )
    assert result.returncode == 1
    assert len(result.stderr.splitlines()) == 1 and result.stderr.startswith("error: ")
    assert "reason must not be empty" in result.stderr
    assert "Traceback" not in result.stderr and result.stdout == ""
    assert not (cli.workdir / "deals").exists()


def test_research_export_on_a_missing_brief_fails_cleanly(cli):
    result = cli("research_export", "--company", "A", "--proposal", "P", "--brief", "nope.md")
    _assert_a_clean_state_error(result, "nope.md")


def test_spreading_check_recomputes_and_checkpoints_the_figures(cli):
    financials = cli.workdir / "fin.json"
    financials.write_text(json.dumps({"FY-Current": {"revenue": 1000, "cost_of_sales": 600, "admin_expenses": 100}}),
                          encoding="utf-8")
    result = cli("spreading_check", "--company", "Acme", "--proposal", "Loan", "--financials", financials)
    assert result.returncode == 0, result.stderr
    out = json.loads(result.stdout)
    assert out["financials"]["FY-Current"]["gross_profit"] == 400 and out["financials_source"] == "framework-computed"
    state_files = list(cli.workdir.glob("deals/Acme/Loan_*/state.json"))
    assert len(state_files) == 1
    assert json.loads(state_files[0].read_text(encoding="utf-8"))["financials"]["FY-Current"]["ebitda"] == 300


def test_spreading_check_no_update_financials_source_leaves_the_flag_alone(cli):
    financials = cli.workdir / "fin.json"
    financials.write_text(json.dumps({"FY-Current": {"revenue": 10}}), encoding="utf-8")
    result = cli("spreading_check", "--company", "Acme", "--proposal", "Loan", "--financials", financials,
                 "--no-update-financials-source")
    assert result.returncode == 0, result.stderr
    assert json.loads(result.stdout)["financials_source"] is None


def test_spreading_check_on_a_missing_input_fails_cleanly(cli):
    result = cli("spreading_check", "--company", "A", "--proposal", "P", "--financials", "nope.json")
    _assert_a_clean_state_error(result, "nope.json")


@pytest.mark.parametrize("name", ["a/b", "..", "a\\b"])
def test_deal_export_reports_unsafe_names_as_one_error_line(cli, name):
    draft = cli.workdir / "draft.md"
    draft.write_text("# Draft\n", encoding="utf-8")
    for flag in ("--company", "--proposal"):
        args = ["--company", "Acme", "--proposal", "Loan", "--draft", draft]
        args[args.index(flag) + 1] = name
        result = cli("deal_export", *args)
        _assert_a_clean_state_error(result, "isn't safe to use as a filesystem path component")
    assert not (cli.workdir / "deals").exists()


def test_research_export_reports_unsafe_names_as_one_error_line(cli):
    brief = cli.workdir / "brief.md"
    # A populated Sources section, so the exporter's sources check passes and the unsafe name is what fails.
    brief.write_text("# Research brief\n\n## Sources\n\n- Synthetic source: https://example.invalid/a\n", encoding="utf-8")
    result = cli("research_export", "--company", "a/b", "--proposal", "Brief", "--brief", brief)
    _assert_a_clean_state_error(result, "isn't safe to use as a filesystem path component")
    assert not (cli.workdir / "deals").exists()


def test_policy_check_reports_missing_draft_as_one_error_line(cli):
    result = cli("policy_check", "--company", "Acme", "--proposal", "Loan", "--draft", "nope.md")
    _assert_a_clean_state_error(result, "nope.md")


def test_policy_check_reports_unsafe_names_as_one_error_line(cli):
    result = cli("policy_check", "--company", "a/b", "--proposal", "Loan")
    _assert_a_clean_state_error(result, "isn't safe to use as a filesystem path component")
    assert not (cli.workdir / "deals").exists()


def test_state_manager_reports_unsafe_names_as_one_error_line(cli):
    for flag, name in (("--company", "a/b"), ("--proposal", "a/b")):
        result = cli("state_manager", "--check-steps", "--company", "Acme", "--proposal", "Loan",
                     "--required", "spread", flag, name)
        _assert_a_clean_state_error(result, "isn't safe to use as a filesystem path component")
    assert not (cli.workdir / "deals").exists()


def test_conventions_reports_unsafe_company_as_one_error_line(cli):
    result = cli("conventions", "--company", "a/b", "--read")
    _assert_a_clean_state_error(result, "isn't safe to use as a filesystem path component")
    assert not (cli.workdir / "deals").exists()


def test_spreading_check_reports_unsafe_names_as_one_error_line(cli):
    result = cli("spreading_check", "--company", "a/b", "--proposal", "Loan")
    _assert_a_clean_state_error(result, "isn't safe to use as a filesystem path component")
    assert not (cli.workdir / "deals").exists()


def _staged_transcription(cli, revenue="1,000"):
    (cli.workdir / "shot.png").write_bytes(b"synthetic image bytes")
    path = cli.workdir / "t.json"
    path.write_text(json.dumps({
        "mode": "framework-computed",
        "source": {"file": "shot.png", "kind": "image", "description": "Synthetic Co accounts", "unit": "GBP thousands",
                   "cost_sign": "positive"},
        "periods": {"FY-Current": {"lines": {"revenue": revenue, "cost_of_sales": "600"},
                                   "subtotals": {"gross_profit": "400"}}}}), encoding="utf-8")
    return path


def test_transcription_check_reads_back_without_touching_a_deal_then_commits_the_confirmed_digest(cli):
    staged = _staged_transcription(cli)
    shown = cli("transcription_check", "--transcription", staged, "--json")
    assert shown.returncode == 0, shown.stderr
    digest = json.loads(shown.stdout)["digest"]
    assert not (cli.workdir / "deals").exists(), "a read-back must not create or change a deal"
    refused = cli("transcription_check", "--transcription", staged, "--commit", "--confirm", "0" * 16,
                  "--company", "Acme", "--proposal", "Loan")
    assert refused.returncode == 1 and refused.stderr.startswith("error: refused:") and refused.stdout == ""
    assert "Traceback" not in refused.stderr and not (cli.workdir / "deals").exists()
    done = cli("transcription_check", "--transcription", staged, "--commit", "--confirm", digest,
               "--company", "Acme", "--proposal", "Loan")
    assert done.returncode == 0, done.stderr
    assert json.loads(done.stdout)["committed"] is True
    (state_file,) = cli.workdir.glob("deals/Acme/Loan_*/state.json")
    state = json.loads(state_file.read_text(encoding="utf-8"))
    assert state["financials"]["FY-Current"]["gross_profit"] == 400 and state["steps_completed"] == ["spread"]
    assert state["financials_transcriptions"][0]["confirmed_digest"] == digest


def test_transcription_check_blocks_a_discrepancy_and_a_commit_without_its_arguments(cli):
    staged = _staged_transcription(cli, revenue="1,100")
    shown = cli("transcription_check", "--transcription", staged)
    assert shown.returncode == 0 and "MISMATCH" in shown.stdout and "STATUS: BLOCKED" in shown.stdout
    digest = json.loads(cli("transcription_check", "--transcription", staged, "--json").stdout)["digest"]
    blocked = cli("transcription_check", "--transcription", staged, "--commit", "--confirm", digest,
                  "--company", "Acme", "--proposal", "Loan")
    assert blocked.returncode == 1 and "unresolved cross-foot discrepancies" in blocked.stderr
    incomplete = cli("transcription_check", "--transcription", staged, "--commit")
    assert incomplete.returncode == 2 and "--company" in incomplete.stderr
    assert not (cli.workdir / "deals").exists()


def test_transcription_check_refuses_a_replaced_image_and_a_deal_in_the_other_mode_without_touching_anything(cli):
    staged = _staged_transcription(cli)
    digest = json.loads(cli("transcription_check", "--transcription", staged, "--json").stdout)["digest"]
    (cli.workdir / "shot.png").write_bytes(b"a different screenshot")
    stale = cli("transcription_check", "--transcription", staged, "--commit", "--confirm", digest,
                "--company", "Acme", "--proposal", "Loan")
    assert stale.returncode == 1 and "the image file itself" in stale.stderr and not (cli.workdir / "deals").exists()

    digest = json.loads(cli("transcription_check", "--transcription", staged, "--json").stdout)["digest"]
    state_dir = cli.workdir / "deals" / "Acme" / "Loan_2026-01-15"
    state_dir.mkdir(parents=True)
    state_file = state_dir / "state.json"
    state_file.write_text(json.dumps({"company": "Acme", "proposal": "Loan", "date": "2026-01-15",
                                      "financials_source": "analyst-supplied",
                                      "financials": {"FY-1": {"gross_profit": 1}}}), encoding="utf-8")
    before = state_file.read_bytes()
    mixed = cli("transcription_check", "--transcription", staged, "--commit", "--confirm", digest,
                "--company", "Acme", "--proposal", "Loan")
    assert mixed.returncode == 1 and len(mixed.stderr.strip().splitlines()) == 1
    assert "never mixes the two bases" in mixed.stderr and "Traceback" not in mixed.stderr
    assert state_file.read_bytes() == before and not (state_dir / "sources").exists()


def test_transcription_check_on_a_missing_or_unreadable_input_fails_with_one_error_line(cli):
    result = cli("transcription_check", "--transcription", "nope.json")
    assert result.returncode == 1 and result.stderr.startswith("error: cannot read nope.json")
    bad = cli.workdir / "bad.json"
    bad.write_text("{}", encoding="utf-8")
    result = cli("transcription_check", "--transcription", bad)
    assert result.returncode == 1 and "mode must be" in result.stderr and "Traceback" not in result.stderr


def _forecast_file(cli, data=None):
    path = cli.workdir / "forecast.json"
    path.write_text(json.dumps(data or {"forecast": {"FY+1": {"subtotals": {"ebitda": 300},
                                                              "ratios": {"dscr": 1.4}}}}), encoding="utf-8")
    return path


def test_supplied_forecast_records_the_figures_as_given_and_prints_a_summary(cli):
    stress = cli.workdir / "stress.json"
    stress.write_text(json.dumps({"revenue_haircut_pct": 10}), encoding="utf-8")
    result = cli("supplied_forecast", "--company", "Acme", "--proposal", "Loan", "--forecast", _forecast_file(cli),
                 "--stress-assumptions", stress, "--note", "Management budget convention")
    assert result.returncode == 0, result.stderr
    summary = json.loads(result.stdout)
    assert summary["forecast_source"] == "analyst-supplied" and summary["forecast_years"] == ["FY+1"]
    assert "FY+1" in summary["downside"]["unavailable"], "stress was asked for but there are no raw lines"
    (state_file,) = cli.workdir.glob("deals/Acme/Loan_*/state.json")
    state = json.loads(state_file.read_text(encoding="utf-8"))
    assert state["financials"]["FY+1"] == {"ebitda": 300} and state["ratios"]["FY+1"] == {"dscr": 1.4}
    assert "multi_period_financials" not in state


def test_supplied_forecast_refuses_bad_input_and_an_incompatible_deal_with_one_error_line(cli):
    missing_note = cli("supplied_forecast", "--company", "A", "--proposal", "P", "--forecast", _forecast_file(cli))
    assert missing_note.returncode == 2 and "--note" in missing_note.stderr
    bad = _forecast_file(cli, {"forecast": {"FY+1": {"subtotals": {"ebitda": "300"}}}})
    result = cli("supplied_forecast", "--company", "A", "--proposal", "P", "--forecast", bad, "--note", "n")
    assert result.returncode == 1 and result.stderr.startswith("error: forecast.FY+1.subtotals.ebitda")
    assert "Traceback" not in result.stderr and not (cli.workdir / "deals").exists()
    gone = cli("supplied_forecast", "--company", "A", "--proposal", "P", "--forecast", "nope.json", "--note", "n")
    assert gone.returncode == 1 and "cannot read forecast file" in gone.stderr
    framework = cli.workdir / "fin.json"
    framework.write_text(json.dumps({"FY-Current": {"revenue": 1}}), encoding="utf-8")
    assert cli("spreading_check", "--company", "A", "--proposal", "P", "--financials", framework).returncode == 0
    (state_file,) = cli.workdir.glob("deals/A/P_*/state.json")
    before = state_file.read_bytes()
    mixed = cli("supplied_forecast", "--company", "A", "--proposal", "P", "--forecast", _forecast_file(cli),
                "--note", "n")
    assert mixed.returncode == 1 and "framework-computed" in mixed.stderr and len(mixed.stderr.splitlines()) == 1
    assert state_file.read_bytes() == before


def test_check_coverage_passes_a_compliant_report_and_fails_a_weak_one(cli):
    config = cli.workdir / "pyproject.toml"
    config.write_text(
        '[tool.opencam.coverage]\noverall = 85\ncritical = 90\ncritical_modules = ["alpha"]\n', encoding="utf-8")
    report = cli.workdir / "coverage.json"

    def write(overall, alpha):
        report.write_text(json.dumps({"totals": {"percent_covered": overall},
                                      "files": {"scripts/alpha.py": {"summary": {"percent_covered": alpha}}}}),
                          encoding="utf-8")

    write(96, 96)
    ok = cli("check_coverage", report, "--config", config)
    assert ok.returncode == 0 and "Coverage floors met." in ok.stdout
    write(96, 50)
    weak = cli("check_coverage", report, "--config", config)
    assert weak.returncode == 1 and "FAIL: alpha coverage 50.00%" in weak.stdout
    assert cli("check_coverage", report, "--config", config, "--report-only").returncode == 0


def test_check_coverage_on_a_missing_report_exits_two_naming_the_file(cli):
    result = cli("check_coverage", "no-such-report.json")
    assert result.returncode == 2 and "no-such-report.json" in result.stderr


def test_mutation_report_prints_a_per_module_table_and_never_fails_over_a_score(cli):
    config = cli.workdir / "pyproject.toml"
    config.write_text('[tool.mutmut]\nonly_mutate = ["src/alpha.py", "src/beta.py"]\n'
                      '[tool.opencam.coverage]\ncritical_modules = ["alpha"]\n', encoding="utf-8")
    results = cli.workdir / "all.txt"
    results.write_text("    alpha.x_f__mutmut_1: killed\n    alpha.x_f__mutmut_2: survived\n"
                       "    beta.x_g__mutmut_1: killed\n", encoding="utf-8")
    out = cli.workdir / "report.json"
    result = cli("mutation_report", results, "--config", config, "--json", out)
    assert result.returncode == 0, result.stderr
    assert "| alpha | critical | 2 | 1 | 0 | 1 | 0 | 50.0% | 50.0% |" in result.stdout
    assert "| beta | supporting | 1 | 1 | 0 | 0 | 0 | 100.0% | 100.0% |" in result.stdout
    assert json.loads(out.read_text(encoding="utf-8"))["total"]["total"] == 3
    results.write_text("    alpha.x_f__mutmut_1: survived\n", encoding="utf-8")      # survivors only: warned, exit 0
    poor = cli("mutation_report", results, "--config", config)
    assert poor.returncode == 0 and "survivors-only" in poor.stdout


def test_mutation_report_on_a_missing_results_file_exits_two_naming_it(cli):
    result = cli("mutation_report", "no-such-results.txt")
    assert result.returncode == 2 and "no-such-results.txt" in result.stderr


def test_check_test_count_passes_on_a_match_and_fails_on_a_mismatch(cli):
    output = cli.workdir / "pytest_output.txt"
    badge = cli.workdir / "badge.json"
    badge.write_text('{"passed": 5}', encoding="utf-8")
    output.write_text("5 passed in 1s\n", encoding="utf-8")
    assert cli("check_test_count", output, "--badge-path", badge).returncode == 0
    output.write_text("6 passed in 1s\n", encoding="utf-8")
    mismatch = cli("check_test_count", output, "--badge-path", badge)
    assert mismatch.returncode == 1 and "mismatch" in mismatch.stdout
    assert json.loads(badge.read_text(encoding="utf-8")) == {"passed": 5}  # the plain check never edits the file


def test_calibrate_mock_writes_a_style_guide_and_a_template_from_a_sample_pdf(cli):
    from test_calibrate import _write_pdf
    samples = cli.workdir / "inputs" / "calibration_samples"
    samples.mkdir(parents=True)
    (cli.workdir / "config").mkdir()  # calibrate writes into an existing config/, as in a checkout
    _write_pdf(samples / "sample.pdf", "Synthetic Co", pages=1, lines_per_page=5)
    result = cli("calibrate", "--mock", "--type", "asset_finance", "--on-overflow", "split")
    assert result.returncode == 0, result.stderr
    from calibrate import MOCK_STYLE_GUIDE, MOCK_TEMPLATE
    assert (cli.workdir / "config" / "style_guide.md").read_text(encoding="utf-8") == MOCK_STYLE_GUIDE
    assert (cli.workdir / "templates" / "local" / "cam" / "asset_finance_cam.md").read_text(
        encoding="utf-8") == MOCK_TEMPLATE


def test_calibrate_without_the_mock_flag_and_without_a_key_falls_back_to_mock(cli):
    """The documented onboarding path, run for real: no flag, no ANTHROPIC_* variable (child_env strips them)."""
    from calibrate import MOCK_STYLE_GUIDE, MOCK_TEMPLATE
    from test_calibrate import _write_pdf
    samples = cli.workdir / "inputs" / "calibration_samples"
    samples.mkdir(parents=True)
    (cli.workdir / "config").mkdir()
    _write_pdf(samples / "sample.pdf", "Synthetic Co", pages=1, lines_per_page=5)
    result = cli("calibrate", "--type", "asset_finance")
    assert result.returncode == 0, result.stderr
    assert "[INFO] ANTHROPIC_API_KEY not found. Running calibrate.py in --mock mode." in result.stdout
    assert (cli.workdir / "config" / "style_guide.md").read_text(encoding="utf-8") == MOCK_STYLE_GUIDE
    assert (cli.workdir / "templates" / "local" / "cam" / "asset_finance_cam.md").read_text(
        encoding="utf-8") == MOCK_TEMPLATE


def test_calibrate_with_only_a_scanned_pdf_is_one_error_line_on_stderr(cli):
    from pypdf import PdfWriter

    samples = cli.workdir / "inputs" / "calibration_samples"
    samples.mkdir(parents=True)
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with (samples / "scanned.pdf").open("wb") as handle:
        writer.write(handle)
    result = cli("calibrate", "--mock", "--type", "corporate_credit")
    assert result.returncode == 1
    assert result.stderr.startswith("error: ") and len(result.stderr.strip().splitlines()) == 1, result.stderr
    assert "no extractable text" in result.stderr and "Traceback" not in result.stderr
    assert "scanned.pdf" in result.stdout, "the per-PDF warning names the file on stdout"
    assert not (cli.workdir / "config").exists() and not (cli.workdir / "templates").exists()


def test_calibrate_without_sample_pdfs_says_so_and_writes_nothing(cli):
    result = cli("calibrate", "--mock")
    assert "No sample PDFs found" in result.stdout
    assert not (cli.workdir / "config").exists() and not (cli.workdir / "templates").exists()


def test_calibrate_rejects_an_unknown_overflow_mode(cli):
    result = cli("calibrate", "--on-overflow", "shrug")
    assert result.returncode == 2 and "invalid choice" in result.stderr


def test_run_evals_validates_lists_and_dry_runs_the_shipped_dataset(cli, tmp_path):
    validate = cli("run_evals", "--validate")
    assert validate.returncode == 0 and "15 cases, all valid" in validate.stdout
    listing = cli("run_evals", "--list")
    assert listing.returncode == 0 and "chk-clean-control" in listing.stdout
    out = tmp_path / "evals-out"
    dry = cli("run_evals", "--dry-run", "--out", out)
    assert dry.returncode == 0, dry.stderr + dry.stdout
    assert "0 calls" in dry.stdout and len(list(out.glob("*/results.json"))) == 1


def test_run_evals_refuses_conflicting_modes_and_an_unaffordable_plan(cli, tmp_path):
    clash = cli("run_evals", "--live", "--dry-run")
    assert clash.returncode == 2 and "mutually exclusive" in clash.stderr
    over_cap = cli("run_evals", "--dry-run", "--repeats", "50", "--out", tmp_path / "unused")
    assert over_cap.returncode == 2 and "Refused" in over_cap.stderr
    assert not (tmp_path / "unused").exists()


def test_orchestrator_without_credentials_is_one_clear_error_before_anything_happens(tmp_path):
    """The one script that cannot be run end to end without a model (issue #187). In a copy of just the files it reads,
    with every Anthropic variable removed and no route to a credentials profile, it must exit 1 with one `error:` line
    and no traceback, before it prints a progress line, reads a deal or checkpoints anything: it never gets as far as a
    request, and it neither succeeds silently nor hangs."""
    env = child_env()
    assert not any(k.upper().startswith("ANTHROPIC") for k in env)
    for variable in HOME_VARIABLES:      # no route to a credentials profile in the real home/config directory
        assert env[variable] == _ISOLATED_HOME and os.listdir(_ISOLATED_HOME) == [], variable
    work = tmp_path / "orchestrator-cwd"
    shutil.copytree(REPO_ROOT / "agents", work / "agents")
    shutil.copytree(REPO_ROOT / "templates" / "cam", work / "templates" / "cam")
    (work / "config").mkdir()
    shutil.copy(REPO_ROOT / "config" / "settings.json", work / "config" / "settings.json")
    result = subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / "orchestrator.py"), "--company", "Acme", "--proposal", "Loan"],
        cwd=work, env=env, capture_output=True, text=True, encoding="utf-8", timeout=120, stdin=subprocess.DEVNULL)
    assert result.returncode == 1
    assert "Traceback" not in result.stderr and "TypeError" not in result.stderr
    assert result.stderr.startswith("error: no Anthropic credentials were found") and len(result.stderr.splitlines()) == 1
    assert "ANTHROPIC_API_KEY" in result.stderr and "docs/configuration.md#credentials" in result.stderr
    assert result.stdout == "", "stopped before the pipeline printed or did anything"
    assert not (work / "deals").exists(), "no deal folder or checkpoint was created"


def test_orchestrator_new_review_refusal_is_one_error_before_model_or_write(cli):
    """A refused fresh review must win over missing model credentials and leave every dated state untouched."""
    future = cli.workdir / "deals" / "Acme" / "Loan_2099-01-10"
    future.mkdir(parents=True)
    future_state = future / "state.json"
    future_state.write_text(json.dumps({
        "company": "Acme", "proposal": "Loan", "date": "2099-01-10",
        "steps_completed": [], "inputs": {"pd": "SENTINEL"},
    }), encoding="utf-8")
    before = future_state.read_bytes()

    result = cli("orchestrator", "--company", "Acme", "--proposal", "Loan", "--new-review")

    from datetime import datetime
    today = datetime.now().strftime("%Y-%m-%d")
    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error: ")
    assert len(result.stderr.strip().splitlines()) == 1, result.stderr
    assert "strictly later than the newest existing state-bearing review" in result.stderr
    assert "Traceback" not in result.stderr
    assert future_state.read_bytes() == before
    assert not (cli.workdir / "deals" / "Acme" / f"Loan_{today}").exists()


def test_orchestrator_on_a_missing_input_file_fails_before_starting(tmp_path):
    work = tmp_path / "orchestrator-cwd"
    work.mkdir()
    result = subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / "orchestrator.py"), "--company", "A", "--proposal", "P",
         "--financials", "nope.json"],
        cwd=work, env=child_env(), capture_output=True, text=True, encoding="utf-8", timeout=120)
    assert result.returncode != 0 and "nope.json" in result.stderr
    assert not (work / "deals").exists()


@pytest.mark.parametrize(("flag", "content"), [
    ("--financials", b"{not json"),
    ("--spread", b"\xff\xfe not text"),
    ("--collateral", b"{not json"),
    ("--stress-assumptions", b"\xff\xfe not text"),
])
def test_orchestrator_reports_an_unusable_input_file_as_one_error_line(tmp_path, flag, content):
    work = tmp_path / "orchestrator-cwd"
    work.mkdir()
    (work / "input.json").write_bytes(content)
    result = subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / "orchestrator.py"), "--company", "A", "--proposal", "P", flag, "input.json"],
        cwd=work, env=child_env(), capture_output=True, text=True, encoding="utf-8", timeout=120)
    assert result.returncode == 1 and result.stdout == ""
    assert result.stderr.startswith("error: ") and len(result.stderr.strip().splitlines()) == 1, result.stderr
    assert "Traceback" not in result.stderr
    assert not (work / "deals").exists()


def test_orchestrator_on_a_missing_input_file_is_one_error_line(tmp_path):
    work = tmp_path / "orchestrator-cwd"
    work.mkdir()
    result = subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / "orchestrator.py"), "--company", "A", "--proposal", "P",
         "--collateral", "nope.json"],
        cwd=work, env=child_env(), capture_output=True, text=True, encoding="utf-8", timeout=120)
    assert result.returncode == 1 and result.stderr.startswith("error: ") and "nope.json" in result.stderr
    assert "Traceback" not in result.stderr and len(result.stderr.strip().splitlines()) == 1


# ---------------------------------------------------------------------------
# A state.json the code cannot use is an error message and exit status 1, never a traceback (issues #170, #171)
# ---------------------------------------------------------------------------

def _write_state(workdir, content):
    folder = workdir / "deals" / "Acme" / "Loan_2026-01-01"
    folder.mkdir(parents=True, exist_ok=True)
    text = content if isinstance(content, str) else json.dumps(content)
    (folder / "state.json").write_text(text, encoding="utf-8")
    return folder / "state.json"


def _assert_a_clean_state_error(result, *fragments):
    assert result.returncode == 1, (result.returncode, result.stdout, result.stderr)
    assert "Traceback" not in result.stderr and result.stdout == ""
    assert result.stderr.startswith("error: ") and len(result.stderr.strip().splitlines()) == 1, result.stderr
    for fragment in fragments:
        assert fragment in result.stderr, (fragment, result.stderr)


def test_spreading_check_reports_a_malformed_state_cleanly(cli):
    state_file = _write_state(cli.workdir, {"multi_period_financials": [1]})
    before = state_file.read_bytes()
    (cli.workdir / "fin.json").write_text(json.dumps({"FY-Current": {"revenue": 100}}), encoding="utf-8")
    result = cli("spreading_check", "--company", "Acme", "--proposal", "Loan", "--financials", "fin.json")
    _assert_a_clean_state_error(result, '"multi_period_financials" must be an object keyed by period', "state.json")
    assert state_file.read_bytes() == before


def test_spreading_check_reports_a_newer_schema_cleanly(cli):
    state_file = _write_state(cli.workdir, {"schema_version": "9.9.9"})
    before = state_file.read_bytes()
    (cli.workdir / "fin.json").write_text(json.dumps({"FY-Current": {"revenue": 100}}), encoding="utf-8")
    result = cli("spreading_check", "--company", "Acme", "--proposal", "Loan", "--financials", "fin.json")
    _assert_a_clean_state_error(result, "9.9.9", "downgrade")
    assert state_file.read_bytes() == before


def test_policy_check_reports_a_malformed_state_cleanly(cli):
    _write_state(cli.workdir, {"financials": [1, 2], "covenants": "none"})
    result = cli("policy_check", "--company", "Acme", "--proposal", "Loan")
    _assert_a_clean_state_error(result, '"financials"', '"covenants"')


def test_policy_check_reports_a_corrupt_state_file_cleanly(cli):
    _write_state(cli.workdir, '{"financials": ')
    _assert_a_clean_state_error(cli("policy_check", "--company", "Acme", "--proposal", "Loan"), "is corrupted")


def test_policy_check_reports_a_utf16_state_file_cleanly(cli):
    """What PowerShell's `>` redirection writes: not UTF-8, not a traceback."""
    folder = cli.workdir / "deals" / "Acme" / "Loan_2026-01-01"
    folder.mkdir(parents=True)
    (folder / "state.json").write_bytes("{}".encode("utf-16"))
    _assert_a_clean_state_error(cli("policy_check", "--company", "Acme", "--proposal", "Loan"), "is corrupted")


def test_policy_check_reports_a_top_level_list_cleanly(cli):
    _write_state(cli.workdir, [1, 2, 3])
    _assert_a_clean_state_error(cli("policy_check", "--company", "Acme", "--proposal", "Loan"), "the top level must be")


def test_deal_export_reports_a_malformed_state_cleanly(cli):
    _write_state(cli.workdir, {"financials_source": 5})
    (cli.workdir / "draft.md").write_text("# Draft\n", encoding="utf-8")
    result = cli("deal_export", "--company", "Acme", "--proposal", "Loan", "--draft", "draft.md")
    _assert_a_clean_state_error(result, '"financials_source" must be "framework-computed" or "analyst-supplied"')





def test_state_manager_exports_reviewer_view_with_a_real_cli_subprocess(cli):
    source = _write_state(cli.workdir, {
        "company": "Acme", "proposal": "Loan", "date": "2026-01-01", "schema_version": "1.1.0",
        "steps_completed": ["triage"], "review_verdict": "REJECTED",
        "review_trail": [{"iteration": 1, "verdict": "REJECTED", "notes": "synthetic finding",
                          "timestamp": "2026-01-01T10:00:00"}],
    })
    before = source.read_bytes()
    result = cli("state_manager", "--export-reviewer-view", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 0, result.stderr
    paths = json.loads(result.stdout)
    assert set(paths) == {"state_for_review", "review_trail"}
    assert paths["state_for_review"].replace("\\", "/") == "deals/Acme/Loan_2026-01-01/review_view/state_for_review.json"
    assert paths["review_trail"].replace("\\", "/") == "deals/Acme/Loan_2026-01-01/review_view/review_trail.json"
    record = json.loads((cli.workdir / paths["state_for_review"]).read_text(encoding="utf-8"))
    trail = json.loads((cli.workdir / paths["review_trail"]).read_text(encoding="utf-8"))
    assert "review_trail" not in record and "review_verdict" not in record
    assert trail[0]["verdict"] == "REJECTED"
    assert source.read_bytes() == before


def test_state_manager_reviewer_view_refuses_bad_state_without_creating_view(cli):
    source = _write_state(cli.workdir, {"schema_version": "9.9.9", "review_trail": ["malformed"]})
    before = source.read_bytes()
    result = cli("state_manager", "--export-reviewer-view", "--company", "Acme", "--proposal", "Loan")
    _assert_a_clean_state_error(result, "9.9.9")
    assert source.read_bytes() == before
    assert not (source.parent / "review_view").exists()



def test_state_manager_check_steps_still_requires_required(cli):
    result = cli("state_manager", "--check-steps", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 2
    assert result.stdout == ""
    assert "error: --required is required with --check-steps" in result.stderr


def test_state_manager_check_steps_reports_a_malformed_state_cleanly(cli):
    _write_state(cli.workdir, {"steps_completed": "triage"})
    result = cli("state_manager", "--check-steps", "--company", "Acme", "--proposal", "Loan", "--required", "spread")
    _assert_a_clean_state_error(result, '"steps_completed" must be a list')


def test_state_manager_check_steps_still_exits_zero_for_a_missing_step(cli):
    """The documented 'always exits 0' is about RESULTS: a state that cannot be read is an error, a missing step is not."""
    _write_state(cli.workdir, {"steps_completed": ["triage"]})
    result = cli("state_manager", "--check-steps", "--company", "Acme", "--proposal", "Loan", "--required", "spread")
    assert result.returncode == 0 and json.loads(result.stdout) == {"missing_steps": ["spread"], "ok": False}


def test_source_manifest_check_sources_reports_a_malformed_state_cleanly(cli):
    _write_state(cli.workdir, [1])
    result = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    _assert_a_clean_state_error(result, "the top level must be")


def test_the_orchestrator_reports_a_malformed_state_cleanly_before_any_model_call(tmp_path):
    work = tmp_path / "orchestrator-cwd"
    shutil.copytree(REPO_ROOT / "agents", work / "agents")
    shutil.copytree(REPO_ROOT / "templates" / "cam", work / "templates" / "cam")
    (work / "config").mkdir()
    shutil.copy(REPO_ROOT / "config" / "settings.json", work / "config" / "settings.json")
    _write_state(work, {"steps_completed": "triage"})
    # The credential check comes first (issue #187), so give the child a placeholder key and an address nothing listens
    # on: the run must stop on the state before it makes any request, and if it ever did, the request could go nowhere.
    env = child_env()
    env.update(ANTHROPIC_API_KEY="synthetic-placeholder", ANTHROPIC_BASE_URL="http://127.0.0.1:9")
    result = subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / "orchestrator.py"), "--company", "Acme", "--proposal", "Loan"],
        cwd=work, env=env, capture_output=True, text=True, encoding="utf-8", timeout=120,
        stdin=subprocess.DEVNULL)
    assert result.returncode == 1 and "Traceback" not in result.stderr
    assert result.stderr.startswith("error: ") and '"steps_completed" must be a list' in result.stderr
    assert "Underwriter Agent drafting" not in result.stdout        # stopped before the pipeline did anything


def test_supplied_forecast_turns_a_malformed_stored_shock_into_one_error_line_and_changes_nothing(cli):
    first = cli("supplied_forecast", "--company", "A", "--proposal", "P", "--forecast", _forecast_file(cli),
                "--note", "n")
    assert first.returncode == 0, first.stderr
    (state_file,) = cli.workdir.glob("deals/A/P_*/state.json")
    state = json.loads(state_file.read_text(encoding="utf-8"))
    state["stress_assumptions"] = {"revenue_haircut_pct": "ten"}      # a shape-valid object, a malformed value
    state_file.write_text(json.dumps(state), encoding="utf-8")
    before = {p: p.read_bytes() for p in cli.workdir.glob("deals/**/*") if p.is_file()}
    result = cli("supplied_forecast", "--company", "A", "--proposal", "P", "--forecast", _forecast_file(cli),
                 "--note", "n")
    assert result.returncode == 1 and result.stdout == ""
    assert len(result.stderr.splitlines()) == 1 and result.stderr.startswith("error: ")
    assert "stress_assumptions" in result.stderr and "revenue_haircut_pct" in result.stderr
    assert "Traceback" not in result.stderr
    assert {p: p.read_bytes() for p in cli.workdir.glob("deals/**/*") if p.is_file()} == before


# ---------------------------------------------------------------------------
# source_manifest.py: an unusable manifest or an unsafe name is one error line, never a traceback (#207)
# ---------------------------------------------------------------------------

UNUSABLE_MANIFEST_BYTES = [b"{ not json", b"", b"\xff\xfe[]", b"{}", b"null", b'{"a": 1}', b'"text"', b"123", b"[1]",
                           b'[{"filename": "a.txt"}, "x"]']
UNSAFE_NAMES = ["a/b", "a\\b", "..", ".", "a:b", "a*b", "a?b", 'a"b', "a<b", "a>b", "a|b"]


def _deal_with_manifest(cli, content):
    state_file = _write_state(cli.workdir, {"triage": {"sources": ["https://example.invalid/registry"]}})
    sources = state_file.parent / "sources"
    sources.mkdir()
    (sources / "manifest.json").write_bytes(content)
    return sources


def _tree(root):
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in sorted((root / "deals").rglob("*")) if p.is_file()}


@pytest.mark.parametrize("content", UNUSABLE_MANIFEST_BYTES)
def test_source_manifest_check_sources_reports_an_unusable_manifest_as_one_error_line(cli, content):
    _deal_with_manifest(cli, content)
    before = _tree(cli.workdir)
    result = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    _assert_a_clean_state_error(result, "manifest.json", "the file was not modified")
    assert _tree(cli.workdir) == before


@pytest.mark.parametrize("content", UNUSABLE_MANIFEST_BYTES)
def test_source_manifest_save_reports_an_unusable_manifest_and_copies_nothing(cli, content):
    sources = _deal_with_manifest(cli, content)
    source = cli.workdir / "filing.txt"
    source.write_text("the filing", encoding="utf-8")
    before = _tree(cli.workdir)
    result = cli("source_manifest", "--company", "Acme", "--proposal", "Loan", "--step", "triage", "--claim",
                 "legal identity", "--file", source)
    _assert_a_clean_state_error(result, "manifest.json", "the file was not modified")
    assert sorted(p.name for p in sources.iterdir()) == ["manifest.json"], "no copy of the document was left behind"
    assert _tree(cli.workdir) == before


@pytest.mark.parametrize("name", UNSAFE_NAMES)
def test_source_manifest_reports_an_unsafe_company_or_proposal_as_one_error_line(cli, name):
    source = cli.workdir / "filing.txt"
    source.write_text("the filing", encoding="utf-8")
    for company, proposal in ((name, "Loan"), ("Acme", name)):
        check = cli("source_manifest", "--check-sources", "--company", company, "--proposal", proposal)
        _assert_a_clean_state_error(check, "isn't safe to use as a filesystem path component")
        save = cli("source_manifest", "--company", company, "--proposal", proposal, "--step", "triage", "--claim", "x",
                   "--file", source)
        _assert_a_clean_state_error(save, "isn't safe to use as a filesystem path component")
    assert not (cli.workdir / "deals").exists(), "nothing was created for an unsafe name"


def test_source_manifest_save_reports_a_missing_file_as_one_error_line(cli):
    result = cli("source_manifest", "--company", "Acme", "--proposal", "Loan", "--step", "triage", "--claim", "x",
                 "--file", cli.workdir / "no-such-file.txt")
    _assert_a_clean_state_error(result, "needs an existing file")
    assert not (cli.workdir / "deals").exists()


def test_source_manifest_check_sources_preserves_missing_flag_for_known_shapes(cli):
    _write_state(cli.workdir, {"triage": {"sources": ["https://example.invalid/registry"]}})
    absent = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    assert absent.returncode == 0 and json.loads(absent.stdout) == {"missing_saved_sources": True, "sources_declared": "declared"}
    sources = next(cli.workdir.glob("deals/Acme/Loan_*")) / "sources"
    sources.mkdir()
    (sources / "manifest.json").write_text("[]", encoding="utf-8")
    empty = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    assert empty.returncode == 0 and json.loads(empty.stdout) == {"missing_saved_sources": True, "sources_declared": "declared"}
    (sources / "manifest.json").write_text('[{"filename": "a.txt", "claim": "x"}]', encoding="utf-8")
    saved = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    assert saved.returncode == 0 and json.loads(saved.stdout) == {"missing_saved_sources": False, "sources_declared": "declared"}
    assert saved.stderr == ""



def test_source_manifest_check_sources_reports_unknown_shape_for_a_completed_step(cli):
    _write_state(cli.workdir, {
        "steps_completed": ["triage"],
        "triage": {"verdict": "Go", "rationale": "Synthetic rationale without a sources key"},
    })
    result = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 0 and result.stderr == ""
    assert json.loads(result.stdout) == {
        "missing_saved_sources": True,
        "sources_declared": "unknown",
    }


def test_source_manifest_check_sources_reports_known_sources_shape_for_an_uncompleted_deal(cli):
    help_result = cli("source_manifest", "--help")
    assert help_result.returncode == 0 and help_result.stderr == ""
    assert all(text in help_result.stdout for text in (
        "missing_saved_sources", "sources_declared", "list[str]", "Explicit [] is valid",
    ))
    _write_state(cli.workdir, {})
    result = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 0 and result.stderr == ""
    assert json.loads(result.stdout) == {
        "missing_saved_sources": False,
        "sources_declared": "not_declared",
    }


@pytest.mark.parametrize("state", MALFORMED_SOURCE_STATES)
@pytest.mark.parametrize("saved", [False, True])
def test_source_manifest_cli_fails_closed_for_malformed_consumed_metadata_without_writes(cli, state, saved):
    state_file = _write_state(cli.workdir, state)
    if saved:
        sources = state_file.parent / "sources"
        sources.mkdir()
        (sources / "saved.txt").write_text("Synthetic saved material", encoding="utf-8")
        (sources / "manifest.json").write_text(
            '[{"step": "research", "filename": "saved.txt"}]', encoding="utf-8",
        )
    before = {p.relative_to(cli.workdir): p.read_bytes() for p in cli.workdir.rglob("*") if p.is_file()}
    result = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 0 and result.stderr == ""
    assert json.loads(result.stdout) == {"missing_saved_sources": True, "sources_declared": "unknown"}
    assert {p.relative_to(cli.workdir): p.read_bytes() for p in cli.workdir.rglob("*") if p.is_file()} == before


@pytest.mark.parametrize("state, expected, has_citations", VALID_SOURCE_STATES)
@pytest.mark.parametrize("saved", [False, True])
def test_source_manifest_cli_preserves_valid_empty_and_combined_research_without_writes(
    cli, state, expected, has_citations, saved,
):
    state_file = _write_state(cli.workdir, state)
    if saved:
        source = cli.workdir / "synthetic.txt"
        source.write_text("Synthetic research evidence", encoding="utf-8")
        result = cli("source_manifest", "--company", "Acme", "--proposal", "Loan",
                     "--step", "research", "--claim", "Synthetic filing and market context", "--file", source)
        assert result.returncode == 0 and result.stderr == ""
    before = {p.relative_to(cli.workdir): p.read_bytes() for p in cli.workdir.rglob("*") if p.is_file()}
    result = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    assert result.returncode == 0 and result.stderr == ""
    assert json.loads(result.stdout) == {
        "missing_saved_sources": has_citations and not saved, "sources_declared": expected,
    }
    assert state_file.exists()
    assert {p.relative_to(cli.workdir): p.read_bytes() for p in cli.workdir.rglob("*") if p.is_file()} == before


def _deal_with_a_dangling_manifest(cli):
    from test_source_manifest import dangling_link
    state_file = _write_state(cli.workdir, {"triage": {"sources": ["https://example.invalid/registry"]}})
    sources = state_file.parent / "sources"
    sources.mkdir()
    dangling_link(sources / "manifest.json", cli.workdir / "nowhere")
    return sources


def test_source_manifest_check_sources_treats_a_dangling_link_manifest_as_unusable_not_as_nothing_saved(cli):
    sources = _deal_with_a_dangling_manifest(cli)
    result = cli("source_manifest", "--check-sources", "--company", "Acme", "--proposal", "Loan")
    _assert_a_clean_state_error(result, "manifest.json", "target does not exist", "the file was not modified")
    assert os.path.lexists(sources / "manifest.json")


def test_source_manifest_save_refuses_a_dangling_link_manifest_and_copies_nothing(cli):
    sources = _deal_with_a_dangling_manifest(cli)
    source = cli.workdir / "filing.txt"
    source.write_text("the filing", encoding="utf-8")
    result = cli("source_manifest", "--company", "Acme", "--proposal", "Loan", "--step", "triage", "--claim",
                 "legal identity", "--file", source)
    _assert_a_clean_state_error(result, "manifest.json", "target does not exist", "the file was not modified")
    assert [p.name for p in sources.iterdir()] == ["manifest.json"], "no copy of the document was left behind"


def test_spreading_check_default_mode_refuses_analyst_supplied_deal_without_writing(cli):
    deal_dir = cli.workdir / "deals" / "Acme" / "Loan_2026-01-01"
    deal_dir.mkdir(parents=True)
    state_path = deal_dir / "state.json"
    state_path.write_text(json.dumps({
        "company": "Acme",
        "proposal": "Loan",
        "date": "2026-01-01",
        "schema_version": "1.1.0",
        "financials_source": "analyst-supplied",
        "financials_source_note": "Synthetic institution convention.",
        "analyst_supplied_financials": {
            "FY-Current": {"revenue": 5000, "cost_of_sales": 3000},
        },
        "financials": {
            "FY-Current": {"raw": {"revenue": 5000, "cost_of_sales": 3000}},
        },
    }), encoding="utf-8")
    before = state_path.read_bytes()
    financials = cli.workdir / "fin.json"
    financials.write_text(json.dumps({
        "FY-Current": {"revenue": 5100, "cost_of_sales": 3050, "admin_expenses": 1000},
    }), encoding="utf-8")

    result = cli(
        "spreading_check", "--company", "Acme", "--proposal", "Loan", "--financials", financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error: ")
    assert result.stderr.count("\n") == 1
    assert "one deal never mixes the two modes" in result.stderr
    assert state_path.read_bytes() == before


def test_research_export_release_unapproved_is_real_cli_and_records_state(cli):
    brief = cli.workdir / "brief.md"
    brief.write_text(
        "# Synthetic research brief\n\n## Findings\nSynthetic analyst notes.\n\n"
        "## Sources\n\n- Synthetic source: https://example.invalid/research\n",
        encoding="utf-8",
    )
    state_dir = cli.workdir / "deals" / "Acme" / "Brief_2026-10-14"
    state_dir.mkdir(parents=True)
    state_path = state_dir / "state.json"
    state_path.write_text(json.dumps({
        "company": "Acme", "proposal": "Brief", "date": "2026-10-14",
        "review_trail": [
            {"iteration": 1, "verdict": "APPROVED", "notes": "previous cycle"},
            {"iteration": 2, "verdict": "REJECTED", "notes": "synthetic finding one"},
            {"iteration": 3, "verdict": "REJECTED", "notes": "synthetic finding two"},
            {"iteration": 4, "verdict": "REJECTED", "notes": "synthetic finding three"},
        ],
    }), encoding="utf-8")
    result = cli(
        "research_export", "--company", "Acme", "--proposal", "Brief", "--brief", brief,
        "--analyst-release-unapproved", "analyst requested manual review",
    )
    assert result.returncode == 0, result.stderr
    output_path = next(cli.workdir.glob("deals/Acme/Brief_*/*Research_Brief.docx"))
    paragraphs = [p.text for p in docx.Document(output_path).paragraphs]
    assert paragraphs[1].startswith(
        "Review status: Not approved by the independent Risk Reviewer. 3 review passes returned REJECTED."
    )
    # The CLI stamps the release with the day it runs; the date is read back from the document and the record must
    # agree with it, so the test depends on neither the day it runs nor a midnight rollover during the run.
    released_on = re.search(r"at the analyst's request on (\d{4}-\d{2}-\d{2}): ", paragraphs[1]).group(1)
    state = json.loads(state_path.read_text(encoding="utf-8"))
    assert state["review_stopped_by_analyst"] == [
        {"date": released_on, "passes": 3, "reason": "analyst requested manual review"},
    ]


def test_research_export_release_unapproved_cli_refuses_a_deal_without_a_rejection(cli):
    brief = cli.workdir / "brief.md"
    brief.write_text(
        "# Synthetic research brief\n\n## Sources\n\n- Synthetic source: https://example.invalid/source\n",
        encoding="utf-8",
    )
    state_dir = cli.workdir / "deals" / "Acme" / "Brief_2026-10-14"
    state_dir.mkdir(parents=True)
    (state_dir / "state.json").write_text(json.dumps({
        "company": "Acme", "proposal": "Brief", "date": "2026-10-14",
        "review_trail": [{"iteration": 1, "verdict": "APPROVED", "notes": "clean"}],
    }), encoding="utf-8")
    result = cli(
        "research_export", "--company", "Acme", "--proposal", "Brief", "--brief", brief,
        "--analyst-release-unapproved", "analyst asked",
    )
    assert result.returncode == 1 and result.stderr.startswith("error: ")
    assert result.stderr.count("\n") == 1 and "no trailing REJECTED entry" in result.stderr
    assert "Traceback" not in result.stderr
    assert not list(state_dir.glob("*.docx"))
    assert "review_stopped_by_analyst" not in (state_dir / "state.json").read_text(encoding="utf-8")


def test_research_export_release_unapproved_cli_rejects_blank_reason_without_output(cli):
    brief = cli.workdir / "brief.md"
    brief.write_text(
        "# Synthetic research brief\n\n## Sources\n\n- Synthetic source: https://example.invalid/source\n",
        encoding="utf-8",
    )
    result = cli(
        "research_export", "--company", "Acme", "--proposal", "Brief", "--brief", brief,
        "--analyst-release-unapproved", "   ",
    )
    assert result.returncode == 1 and result.stderr.startswith("error: ")
    assert "Analyst release reason must not be empty" in result.stderr and "Traceback" not in result.stderr
    assert not (cli.workdir / "deals").exists()


def test_orchestrator_refuses_forward_only_run_on_analyst_supplied_deal_before_credentials_or_writing(cli):
    from state_manager import state_path, write_state

    company, proposal, date_str = "Synthetic Co", "Synthetic Facility", "2026-01-05"
    write_state(
        company, proposal, date_str=date_str, base_dir=str(cli.workdir),
        financials_source="analyst-supplied",
        financials_source_note="Synthetic analyst convention.",
        analyst_supplied_financials={"FY-Current": {"revenue": 5000, "cost_of_sales": 3000}},
        multi_period_financials={"FY-Current": {"revenue": 5000, "cost_of_sales": 3000}},
        financials={"FY-Current": {"ebitda": 800}},
        ratios={"FY-Current": {"dscr": 1.5}},
        steps_completed=["spread"],
    )
    saved_path = Path(state_path(company, proposal, date_str=date_str, base_dir=str(cli.workdir)))
    before = saved_path.read_bytes()
    financials = cli.workdir / "forward.json"
    financials.write_text(
        json.dumps({"FY+1": {"revenue": 5400, "cost_of_sales": 3300}}),
        encoding="utf-8",
    )

    result = cli(
        "orchestrator", "--company", company, "--proposal", proposal,
        "--financials", financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "/spread's analyst-supplied mode" in result.stderr
    assert "nothing was written" in result.stderr
    assert "no Anthropic credentials" not in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


def test_spreading_check_no_update_flag_cannot_overwrite_analyst_supplied_history(cli):
    """Decision D3 (#271): the exact scenario /project's flag used to allow - historical raw figures on an
    analyst-supplied deal - is one error line, exit 1, state byte-identical, no multi_period_financials."""
    deal_dir = cli.workdir / "deals" / "Acme" / "Loan_2026-01-01"
    deal_dir.mkdir(parents=True)
    state_path = deal_dir / "state.json"
    state_path.write_text(json.dumps({
        "company": "Acme", "proposal": "Loan", "date": "2026-01-01", "schema_version": "1.1.0",
        "financials_source": "analyst-supplied",
        "analyst_supplied_financials": {"FY-Current": {"revenue": 5000, "cost_of_sales": 4001}},
        "financials": {"FY-Current": {"ebitda": 300, "gross_profit": 999}},
        "ratios": {"FY-Current": {"dscr": 1.3}},
    }), encoding="utf-8")
    before = state_path.read_bytes()
    financials = cli.workdir / "hist.json"
    financials.write_text(json.dumps({"FY-Current": {"revenue": 5000, "cost_of_sales": 3000}}), encoding="utf-8")

    result = cli("spreading_check", "--company", "Acme", "--proposal", "Loan", "--financials", financials,
                 "--no-update-financials-source")

    assert result.returncode == 1 and result.stdout == ""
    assert result.stderr.startswith("error: ") and result.stderr.count("\n") == 1
    assert "historical" in result.stderr and "Traceback" not in result.stderr
    assert state_path.read_bytes() == before
    assert "multi_period_financials" not in state_path.read_text(encoding="utf-8")


@pytest.mark.parametrize(
    "token",
    ['"abc"', '"1e999"', '"nan"', '"1250"', "true", "1e999"],
    ids=["text", "overflow-text", "nan-text", "numeric-text", "boolean", "overflow-number"],
)
def test_spreading_check_rejects_non_numeric_raw_line_items_before_writing(cli, token):
    from state_manager import state_path, write_state

    company, proposal, date_str = "Synthetic Co", "Synthetic Facility", "2026-01-05"
    write_state(
        company, proposal, date_str=date_str, base_dir=str(cli.workdir),
        financials_source="framework-computed",
        financials={"FY-Current": {"raw": {"revenue": 1000}}},
        multi_period_financials={"FY-Current": {"revenue": 1000}},
    )
    saved_path = Path(state_path(company, proposal, date_str=date_str, base_dir=str(cli.workdir)))
    before = saved_path.read_bytes()
    financials = cli.workdir / "financials.json"
    financials.write_text(
        '{"FY-Current":{"revenue":' + token + ',"cost_of_sales":600}}',
        encoding="utf-8",
    )

    result = cli(
        "spreading_check", "--company", company, "--proposal", proposal,
        "--financials", financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "Traceback" not in result.stderr
    assert "FY-Current" in result.stderr and "revenue" in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("flag", ["--financials", "--spread"])
@pytest.mark.parametrize(
    "token",
    ['"abc"', '"1e999"', '"nan"', '"1250"', "true", "false", "1e999"],
    ids=["text", "overflow-text", "nan-text", "numeric-text", "true", "false", "overflow-number"],
)
def test_orchestrator_rejects_non_numeric_raw_line_items_before_model_or_write(cli, flag, token):
    from state_manager import state_path, write_state

    company, proposal, date_str = "Synthetic Co", "Synthetic Facility", "2026-01-05"
    write_state(
        company, proposal, date_str=date_str, base_dir=str(cli.workdir),
        financials_source="framework-computed",
        financials={"FY-Current": {"raw": {"revenue": 1000}}},
        multi_period_financials={"FY-Current": {"revenue": 1000}},
    )
    saved_path = Path(state_path(company, proposal, date_str=date_str, base_dir=str(cli.workdir)))
    before = saved_path.read_bytes()
    financials = cli.workdir / "financials.json"
    financials.write_text(
        '{"FY-Current":{"revenue":' + token + ',"cost_of_sales":600}}',
        encoding="utf-8",
    )

    result = cli(
        "orchestrator", "--company", company, "--proposal", proposal,
        flag, financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "Traceback" not in result.stderr
    assert "financial inputs" in result.stderr
    assert "FY-Current" in result.stderr and "revenue" in result.stderr
    assert saved_path.read_bytes() == before


def test_wiki_publish_generates_the_real_docs_tree(cli, tmp_path):
    out = tmp_path / "wiki"
    result = cli("wiki_publish", "--repo", "example/repo", "--root", REPO_ROOT, "--out", out)
    assert result.returncode == 0, result.stderr
    assert (out / "Home.md").is_file()
    assert (out / "_Sidebar.md").is_file()
    assert not list(out.glob("*.md")) or all(p.suffix == ".md" for p in out.glob("*.md"))


def test_wiki_publish_on_a_broken_corpus_fails_cleanly(cli, tmp_path):
    (tmp_path / "docs").mkdir()
    (tmp_path / "docs" / "README.md").write_text("# Index\n", encoding="utf-8")
    (tmp_path / "docs" / "source.md").write_text("[gone](gone.md)\n", encoding="utf-8")
    result = cli("wiki_publish", "--repo", "example/repo", "--root", tmp_path, "--out", tmp_path / "out")
    assert result.returncode == 1
    assert result.stderr.count("error:") == 1
    assert "gone.md" in result.stderr
    assert not (tmp_path / "out").exists()


def test_wiki_publish_publishes_to_a_local_bare_repository(cli, tmp_path):
    """The --publish mode, run for real as a subprocess: generation, clone, full replacement, one
    commit and a push against a local bare repository standing in for the wiki (no token, no network)."""
    docs = tmp_path / "docs"
    docs.mkdir()
    (docs / "README.md").write_text(
        "# Index\n\n## What is in this folder\n\n| Document | Contents |\n| :--- | :--- |\n"
        "| [source.md](source.md) | x |\n", encoding="utf-8")
    (docs / "source.md").write_text("# S\n", encoding="utf-8")
    seed = tmp_path / "seed"
    seed.mkdir()
    (seed / "Home.md").write_text("placeholder\n", encoding="utf-8")
    env = {k: v for k, v in os.environ.items() if not k.startswith("GIT_")}
    for args in (["-c", "init.defaultBranch=master", "init", "-q"], ["add", "Home.md"],
                 ["commit", "-q", "-m", "Initial Home page"]):
        subprocess.run(["git", "-c", "user.email=t@example.invalid", "-c", "user.name=t", *args],
                       cwd=seed, check=True, capture_output=True, env=env, timeout=60)
    remote = tmp_path / "wiki.git"
    subprocess.run(["git", "clone", "--quiet", "--bare", str(seed), str(remote)], check=True,
                  capture_output=True, env=env, timeout=60)

    result = cli("wiki_publish", "--repo", "example/repo", "--root", str(tmp_path), "--publish",
                 "--source-sha", "1" * 40, "--remote", str(remote))

    assert result.returncode == 0, result.stderr
    tree = subprocess.run(
        ["git", "--git-dir", str(remote), "ls-tree", "-r", "--name-only", "HEAD"],
        check=True, capture_output=True, text=True, env=env, timeout=60)
    files = tree.stdout.splitlines()
    assert "Home.md" in files and "source.md" in files and "_Sidebar.md" in files
    message = subprocess.run(
        ["git", "--git-dir", str(remote), "log", "-1", "--format=%B"],
        check=True, capture_output=True, text=True, env=env, timeout=60)
    assert "docs-source: " + "1" * 40 in message.stdout


STRESS_DOMAIN_CASES = [
    ("revenue_haircut_pct", -10),
    ("opex_increase_pct", -50),
    ("interest_rate_bump_bps", -300),
    ("revenue_haircut_pct", 100),
]


def _seed_stress_deal(cli):
    from state_manager import state_path, write_state

    company, proposal, date_str = "Synthetic Co", "Synthetic Facility", "2026-01-05"
    raw = {
        "revenue": 5000, "cost_of_sales": 3000, "admin_expenses": 1200,
        "interest_paid": 50, "current_debt": 0, "overdraft": 0,
        "long_term_debt": 2400, "loan_notes": 0, "scheduled_principal": 200,
    }
    write_state(
        company, proposal, date_str=date_str, base_dir=str(cli.workdir),
        financials_source="framework-computed",
        financials={"FY+1": {"raw": raw}},
        multi_period_financials={"FY+1": raw},
    )
    return company, proposal, Path(state_path(company, proposal, date_str=date_str, base_dir=str(cli.workdir)))


@pytest.mark.parametrize("key,value", STRESS_DOMAIN_CASES)
def test_spreading_check_refuses_domain_invalid_stress_before_writing(cli, key, value):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    stress = cli.workdir / "stress.json"
    stress.write_text(json.dumps({key: value}), encoding="utf-8")

    result = cli(
        "spreading_check", "--company", company, "--proposal", proposal,
        "--stress-assumptions", stress,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert key in result.stderr and "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("key,value", STRESS_DOMAIN_CASES)
def test_orchestrator_refuses_domain_invalid_stress_before_credentials_or_writing(cli, key, value):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    financials = cli.workdir / "financials.json"
    financials.write_text(json.dumps({"FY+1": {
        "revenue": 5000, "cost_of_sales": 3000, "admin_expenses": 1200,
        "interest_paid": 50, "current_debt": 0, "overdraft": 0,
        "long_term_debt": 2400, "loan_notes": 0, "scheduled_principal": 200,
    }}), encoding="utf-8")
    stress = cli.workdir / "stress.json"
    stress.write_text(json.dumps({key: value}), encoding="utf-8")

    result = cli(
        "orchestrator", "--company", company, "--proposal", proposal,
        "--financials", financials, "--stress-assumptions", stress,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert key in result.stderr and "stress assumptions" in result.stderr
    assert "no Anthropic credentials" not in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


def test_spreading_check_warns_for_percent_value_between_zero_and_one_and_keeps_it(cli):
    company, proposal, saved_path = _seed_stress_deal(cli)
    stress = cli.workdir / "stress.json"
    stress.write_text('{"revenue_haircut_pct":0.5}', encoding="utf-8")

    result = cli(
        "spreading_check", "--company", company, "--proposal", proposal,
        "--stress-assumptions", stress,
    )

    assert result.returncode == 0, result.stderr
    assert result.stderr.startswith("warning:")
    assert "revenue_haircut_pct" in result.stderr and "0.5" in result.stderr
    assert len(result.stderr.strip().splitlines()) == 1
    assert json.loads(result.stdout)["downside_case"] is not None
    assert json.loads(saved_path.read_text(encoding="utf-8"))["stress_assumptions"]["revenue_haircut_pct"] == 0.5


def _seed_analyst_supplied_stress_deal(cli):
    from state_manager import state_path, write_state

    company, proposal, date_str = "Synthetic Co", "Synthetic Facility", "2026-01-05"
    write_state(
        company, proposal, date_str=date_str, base_dir=str(cli.workdir),
        financials_source="analyst-supplied",
        financials_source_note="Synthetic budget convention.",
        analyst_supplied_financials={"FY-Current": {"revenue": 5000, "cost_of_sales": 3000}},
        financials={"FY-Current": {"ebitda": 800}},
    )
    return company, proposal, Path(state_path(company, proposal, date_str=date_str, base_dir=str(cli.workdir)))


@pytest.mark.parametrize("key,value", STRESS_DOMAIN_CASES)
def test_supplied_forecast_refuses_domain_invalid_stress_before_writing(cli, key, value):
    company, proposal, saved_path = _seed_analyst_supplied_stress_deal(cli)
    before = saved_path.read_bytes()
    forecast = cli.workdir / "forecast.json"
    forecast.write_text('{"forecast":{"FY+1":{"subtotals":{"ebitda":300}}}}', encoding="utf-8")
    stress = cli.workdir / "stress.json"
    stress.write_text(json.dumps({key: value}), encoding="utf-8")

    result = cli(
        "supplied_forecast", "--company", company, "--proposal", proposal,
        "--forecast", forecast, "--stress-assumptions", stress, "--note", "Synthetic forecast convention",
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert key in result.stderr and "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


def test_supplied_forecast_warns_and_continues_for_small_positive_percent_stress(cli):
    company, proposal, saved_path = _seed_analyst_supplied_stress_deal(cli)
    forecast = cli.workdir / "forecast.json"
    forecast.write_text('{"forecast":{"FY+1":{"subtotals":{"ebitda":300}}}}', encoding="utf-8")
    stress = cli.workdir / "stress.json"
    stress.write_text('{"opex_increase_pct":0.5}', encoding="utf-8")

    result = cli(
        "supplied_forecast", "--company", company, "--proposal", proposal,
        "--forecast", forecast, "--stress-assumptions", stress, "--note", "Synthetic forecast convention",
    )

    assert result.returncode == 0, result.stderr
    assert result.stderr.startswith("warning:")
    assert "opex_increase_pct" in result.stderr and "0.5" in result.stderr
    assert len(result.stderr.strip().splitlines()) == 1
    saved = json.loads(saved_path.read_text(encoding="utf-8"))
    assert saved["stress_assumptions"]["opex_increase_pct"] == 0.5


# Issue #359: malformed stress values and period blocks must be refused by the real
# command-line entry points before state is written or a model client is created.
STRESS_TYPE_TOKENS = [
    ('"50"', "string"),
    ("true", "boolean"),
    ("null", "null"),
    ("[50]", "list"),
    ('{"value": 50}', "object"),
]
PERIOD_TYPE_TOKENS = [
    ('"abc"', "string"),
    ("7", "number"),
    ("[7]", "list"),
]


def _write_stress_token(cli, token):
    path = cli.workdir / "stress.json"
    path.write_text('{"revenue_haircut_pct":' + token + "}", encoding="utf-8")
    return path


@pytest.mark.parametrize("token,kind", STRESS_TYPE_TOKENS, ids=[kind for _, kind in STRESS_TYPE_TOKENS])
def test_spreading_check_rejects_stress_values_that_are_not_finite_json_numbers(cli, token, kind):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    stress = _write_stress_token(cli, token)

    result = cli(
        "spreading_check", "--company", company, "--proposal", proposal,
        "--stress-assumptions", stress,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "revenue_haircut_pct" in result.stderr
    assert "stress assumptions" in result.stderr and "must be a finite JSON number" in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("token,kind", STRESS_TYPE_TOKENS, ids=[kind for _, kind in STRESS_TYPE_TOKENS])
def test_orchestrator_rejects_stress_values_that_are_not_finite_json_numbers_before_credentials(
        cli, token, kind):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    stress = _write_stress_token(cli, token)

    result = cli(
        "orchestrator", "--company", company, "--proposal", proposal,
        "--stress-assumptions", stress,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "revenue_haircut_pct" in result.stderr
    assert "stress assumptions" in result.stderr and "must be a finite JSON number" in result.stderr
    assert "no Anthropic credentials" not in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("token,kind", STRESS_TYPE_TOKENS, ids=[kind for _, kind in STRESS_TYPE_TOKENS])
def test_supplied_forecast_uses_shared_stress_type_validation_before_writing(cli, token, kind):
    company, proposal, saved_path = _seed_analyst_supplied_stress_deal(cli)
    before = saved_path.read_bytes()
    forecast = cli.workdir / "forecast.json"
    forecast.write_text('{"forecast":{"FY+1":{"subtotals":{"ebitda":300}}}}', encoding="utf-8")
    stress = _write_stress_token(cli, token)

    result = cli(
        "supplied_forecast", "--company", company, "--proposal", proposal,
        "--forecast", forecast, "--stress-assumptions", stress,
        "--note", "Synthetic forecast convention",
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "revenue_haircut_pct" in result.stderr
    assert "stress assumptions" in result.stderr and "must be a finite JSON number" in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("token,kind", PERIOD_TYPE_TOKENS, ids=[kind for _, kind in PERIOD_TYPE_TOKENS])
def test_spreading_check_rejects_non_object_financial_periods_before_writing(cli, token, kind):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    financials = cli.workdir / "financials.json"
    financials.write_text('{"FY+1":' + token + "}", encoding="utf-8")

    result = cli(
        "spreading_check", "--company", company, "--proposal", proposal,
        "--financials", financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "FY+1" in result.stderr
    assert "must be an object of line items" in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("flag", ["--financials", "--spread"])
@pytest.mark.parametrize("token,kind", PERIOD_TYPE_TOKENS, ids=[kind for _, kind in PERIOD_TYPE_TOKENS])
def test_orchestrator_rejects_non_object_financial_periods_before_credentials_or_writing(cli, flag, token, kind):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    financials = cli.workdir / "financials.json"
    financials.write_text('{"FY+1":' + token + "}", encoding="utf-8")

    result = cli(
        "orchestrator", "--company", company, "--proposal", proposal,
        flag, financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "FY+1" in result.stderr
    assert "must be an object of line items" in result.stderr
    assert "no Anthropic credentials" not in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("token,kind", PERIOD_TYPE_TOKENS, ids=[kind for _, kind in PERIOD_TYPE_TOKENS])
def test_supplied_forecast_names_a_period_whose_value_is_not_an_object_of_line_items(cli, token, kind):
    company, proposal, saved_path = _seed_analyst_supplied_stress_deal(cli)
    before = saved_path.read_bytes()
    forecast = cli.workdir / "forecast.json"
    forecast.write_text('{"forecast":{"FY+1":' + token + "}}", encoding="utf-8")

    result = cli(
        "supplied_forecast", "--company", company, "--proposal", proposal,
        "--forecast", forecast, "--note", "Synthetic forecast convention",
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "forecast.FY+1" in result.stderr
    assert "must be an object of line items" in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("token,kind", PERIOD_TYPE_TOKENS, ids=[kind for _, kind in PERIOD_TYPE_TOKENS])
def test_supplied_forecast_names_a_downside_period_whose_value_is_not_an_object(cli, token, kind):
    company, proposal, saved_path = _seed_analyst_supplied_stress_deal(cli)
    before = saved_path.read_bytes()
    forecast = cli.workdir / "forecast.json"
    forecast.write_text(
        '{"forecast":{"FY+1":{"subtotals":{"ebitda":300}}},'
        '"downside":{"description":"Synthetic scenario","periods":{"FY+1":' + token + "}}}",
        encoding="utf-8",
    )

    result = cli(
        "supplied_forecast", "--company", company, "--proposal", proposal,
        "--forecast", forecast, "--note", "Synthetic forecast convention",
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "downside.periods.FY+1" in result.stderr
    assert "must be an object of line items" in result.stderr
    assert "Traceback" not in result.stderr
    assert saved_path.read_bytes() == before


def test_state_manager_find_reports_exact_record_as_one_json_object_without_writing(cli):
    folder = cli.workdir / "deals" / "Synthetic Co" / "Credit Facility_2026-03-15"
    folder.mkdir(parents=True)
    (folder / "state.json").write_text('{"company":"Synthetic Co"}', encoding="utf-8")
    before = sorted(str(p.relative_to(cli.workdir)) for p in cli.workdir.rglob("*"))

    result = cli("state_manager", "--find", "--company", "Synthetic Co", "--proposal", "Credit Facility")

    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    assert json.loads(result.stdout) == {
        "record_found": True,
        "state_path": os.path.relpath(str(folder / "state.json"), str(cli.workdir)),
        "date": "2026-03-15",
    }
    assert result.stdout.endswith("\n")
    assert sorted(str(p.relative_to(cli.workdir)) for p in cli.workdir.rglob("*")) == before





def test_state_manager_find_reports_unsafe_company_as_one_error_without_writing(cli):
    result = cli("state_manager", "--find", "--company", "..", "--proposal", "Synthetic Loan")

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert result.stderr.count("\n") == 1
    assert "Traceback" not in result.stderr
    assert not (cli.workdir / "deals").exists()

def test_state_manager_find_reports_only_normalized_near_miss_and_no_write(cli):
    folder = cli.workdir / "deals" / "Synthetic Co" / "Credit Facility_2026-03-15"
    folder.mkdir(parents=True)
    (folder / "state.json").write_text("{}", encoding="utf-8")
    before = sorted(str(p.relative_to(cli.workdir)) for p in cli.workdir.rglob("*"))

    result = cli("state_manager", "--find", "--company", "Synthetic Co", "--proposal", "credit   facility")

    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    report = json.loads(result.stdout)
    assert report["record_found"] is False
    assert report["company_folder_exists"] is True
    assert [item["proposal"] for item in report["near_misses"]] == ["Credit Facility"]
    assert sorted(str(p.relative_to(cli.workdir)) for p in cli.workdir.rglob("*")) == before


# ---------------------------------------------------------------------------
# Issue #203: real subprocess contracts for review candidate/start modes.
# ---------------------------------------------------------------------------

def test_state_manager_help_lists_both_review_modes_and_all_new_options(cli):
    result = cli("state_manager", "--help")
    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    for flag in ("--review-candidates", "--start-review", "--date", "--prior-date", "--no-prior"):
        assert flag in result.stdout


def test_state_manager_review_candidates_cli_returns_one_read_only_json_object(cli):
    folder = cli.workdir / "deals" / "Synthetic Co" / "Credit Facility_2025-06-01"
    folder.mkdir(parents=True)
    (folder / "state.json").write_text('{"company":"Synthetic Co"}\n', encoding="utf-8")
    before = sorted(str(path.relative_to(cli.workdir)) for path in cli.workdir.rglob("*"))

    result = cli(
        "state_manager", "--review-candidates", "--company", "Synthetic Co",
        "--proposal", "Credit Facility", "--date", "2026-04-15",
    )

    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    assert result.stdout.endswith("\n")
    assert json.loads(result.stdout) == {
        "company": "Synthetic Co",
        "proposal": "Credit Facility",
        "date": "2026-04-15",
        "candidates": [{
            "date": "2025-06-01",
            "state_path": os.path.join(
                "deals", "Synthetic Co", "Credit Facility_2025-06-01", "state.json",
            ),
        }],
        "default_prior_date": "2025-06-01",
        "skipped": [],
    }
    assert sorted(str(path.relative_to(cli.workdir)) for path in cli.workdir.rglob("*")) == before


def test_state_manager_start_review_cli_creates_a_fresh_unlinked_state(cli):
    (cli.workdir / "deals" / "Synthetic Co").mkdir(parents=True)
    result = cli(
        "state_manager", "--start-review", "--company", "Synthetic Co",
        "--proposal", "Credit Facility", "--date", "2026-04-15",
    )

    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    report = json.loads(result.stdout)
    assert report["started"] is True
    assert report["date"] == "2026-04-15"
    assert report["prior_review_date"] is None
    assert report["linked"] is False
    assert "no eligible earlier review was found" in report["message"]
    state = json.loads((cli.workdir / report["state_path"]).read_text(encoding="utf-8"))
    assert set(state) == {"company", "proposal", "date", "schema_version"}


def test_state_manager_start_review_refusal_is_one_error_line_with_no_writes(cli):
    prior = cli.workdir / "deals" / "Synthetic Co" / "Credit Facility_2026-06-01"
    prior.mkdir(parents=True)
    (prior / "state.json").write_text('{"company":"Synthetic Co"}\n', encoding="utf-8")
    before = sorted(str(path.relative_to(cli.workdir)) for path in cli.workdir.rglob("*"))

    result = cli(
        "state_manager", "--start-review", "--company", "Synthetic Co",
        "--proposal", "Credit Facility", "--date", "2026-05-01",
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert result.stderr.count("\n") == 1
    assert "strictly later than the newest existing state-bearing review" in result.stderr
    assert "Traceback" not in result.stderr
    assert sorted(str(path.relative_to(cli.workdir)) for path in cli.workdir.rglob("*")) == before


def test_state_manager_review_candidates_impossible_date_is_one_error_line(cli):
    result = cli(
        "state_manager", "--review-candidates", "--company", "Synthetic Co",
        "--proposal", "Credit Facility", "--date", "2026-02-30",
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error:")
    assert result.stderr.count("\n") == 1
    assert "valid calendar date" in result.stderr
    assert "Traceback" not in result.stderr
    assert not (cli.workdir / "deals").exists()


def test_state_manager_prior_options_are_an_argparse_usage_error(cli):
    result = cli(
        "state_manager", "--start-review", "--company", "Synthetic Co",
        "--proposal", "Credit Facility", "--date", "2026-04-15",
        "--prior-date", "2025-06-01", "--no-prior",
    )

    assert result.returncode == 2
    assert result.stdout == ""
    assert "usage:" in result.stderr
    assert "not allowed with argument" in result.stderr
    assert "Traceback" not in result.stderr


@pytest.mark.parametrize(
    ("mode", "extra_option"),
    [("--find", "--date"), ("--find", "--no-prior"), ("--review-candidates", "--no-prior")],
)
def test_state_manager_review_only_options_refuse_wrong_mode_as_usage_error(cli, mode, extra_option):
    args = [
        "state_manager", mode, "--company", "Synthetic Co", "--proposal", "Credit Facility",
        extra_option,
    ]
    if extra_option == "--date":
        args.append("2026-04-15")
    result = cli(*args)
    assert result.returncode == 2
    assert result.stdout == ""
    assert "usage:" in result.stderr
    assert "Traceback" not in result.stderr



@pytest.mark.parametrize("mode", ["--start-review", "--review-candidates"])
def test_state_manager_review_modes_refuse_an_unknown_company_with_one_error_line_and_no_writes(cli, mode):
    (cli.workdir / "deals" / "Wren Kitchens Limited").mkdir(parents=True)
    before = sorted(str(path.relative_to(cli.workdir)) for path in cli.workdir.rglob("*"))

    result = cli("state_manager", mode, "--company", "Wrem Kitchens Limited", "--proposal", "Finance")

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error: Refusing to start a review: there is no deal folder for company ")
    assert "/triage or /spread" in result.stderr
    assert len(result.stderr.strip().splitlines()) == 1 and "Traceback" not in result.stderr
    assert sorted(str(path.relative_to(cli.workdir)) for path in cli.workdir.rglob("*")) == before


def test_state_manager_find_refuses_an_empty_company_with_one_error_line(cli):
    result = cli("state_manager", "--find", "--company", "", "--proposal", "Finance")

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error: company must not be empty")
    assert len(result.stderr.strip().splitlines()) == 1 and "Traceback" not in result.stderr


def test_review_compare_compares_exact_dated_states_read_only(cli):
    company, proposal = "Synthetic Co", "Asset Facility"
    prior_date, current_date = "2025-10-20", "2026-10-20"
    for date, ebitda in ((prior_date, 100), (current_date, 110)):
        state_dir = cli.workdir / "deals" / company / f"{proposal}_{date}"
        state_dir.mkdir(parents=True)
        state = {
            "company": company,
            "proposal": proposal,
            "date": date,
            "financials_source": "framework-computed",
            "financials": {"FY-Current": {"ebitda": ebitda}},
        }
        (state_dir / "state.json").write_text(json.dumps(state), encoding="utf-8")
    before = {
        date: (cli.workdir / "deals" / company / f"{proposal}_{date}" / "state.json").read_bytes()
        for date in (prior_date, current_date)
    }

    result = cli(
        "review_compare", "--company", company, "--proposal", proposal,
        "--current-date", current_date, "--prior-date", prior_date,
    )

    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    report = json.loads(result.stdout)
    assert report["company"] == company
    assert report["proposal"] == proposal
    assert report["prior_date"] == prior_date
    assert report["current_date"] == current_date
    assert report["summary"]["changed_paths"] == len(report["changes"])
    row = next(change for change in report["changes"]
               if change["path"] == "financials.FY-Current.ebitda")
    assert row["change"] == "value_changed"
    assert row["delta"] == 10
    assert before == {
        date: (cli.workdir / "deals" / company / f"{proposal}_{date}" / "state.json").read_bytes()
        for date in (prior_date, current_date)
    }
    assert sorted(p.relative_to(cli.workdir).as_posix() for p in cli.workdir.rglob("state.json")) == [
        f"deals/{company}/{proposal}_{prior_date}/state.json",
        f"deals/{company}/{proposal}_{current_date}/state.json",
    ]


def test_review_compare_markdown_format_is_a_real_cli_output(cli):
    company, proposal = "Synthetic Co", "Asset Facility"
    prior_date, current_date = "2025-10-20", "2026-10-20"
    for date, ebitda in ((prior_date, 100), (current_date, 110)):
        folder = cli.workdir / "deals" / company / f"{proposal}_{date}"
        folder.mkdir(parents=True)
        (folder / "state.json").write_text(json.dumps({
            "company": company, "proposal": proposal, "date": date,
            "financials": {"FY-Current": {"ebitda": ebitda}},
        }), encoding="utf-8")
    result = cli(
        "review_compare", "--company", company, "--proposal", proposal,
        "--current-date", current_date, "--prior-date", prior_date, "--format", "markdown",
    )
    assert result.returncode == 0, result.stderr
    assert result.stderr == ""
    assert result.stdout.startswith("# Review comparison — Synthetic Co / Asset Facility")
    assert "| Item | Prior | Current | Change |" in result.stdout
    assert "financials.FY-Current.ebitda" in result.stdout
    assert "delta 10" in result.stdout


def test_review_compare_missing_current_state_is_one_cli_error_line(cli):
    company, proposal, prior_date = "Synthetic Co", "Asset Facility", "2025-10-20"
    prior_folder = cli.workdir / "deals" / company / f"{proposal}_{prior_date}"
    prior_folder.mkdir(parents=True)
    (prior_folder / "state.json").write_text(json.dumps({
        "company": company, "proposal": proposal, "date": prior_date,
    }), encoding="utf-8")
    result = cli(
        "review_compare", "--company", "Synthetic Co", "--proposal", "Asset Facility",
        "--current-date", "2026-10-20", "--prior-date", "2025-10-20",
    )
    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr.startswith("error: current state not found: ")
    assert len(result.stderr.strip().splitlines()) == 1
    assert "Traceback" not in result.stderr


def test_review_compare_writes_utf8_when_stdout_was_redirected_as_cp1252(tmp_path):
    """Mirror tests/test_stdio_encoding.py with a synthetic non-cp1252 proposal."""
    work = tmp_path / "review-compare-utf8"
    work.mkdir()
    company, proposal = "Synthetic Co", "Asset 東京 Facility"
    prior_date, current_date = "2025-10-20", "2026-10-20"
    for date, ebitda in ((prior_date, 100), (current_date, 110)):
        folder = work / "deals" / company / f"{proposal}_{date}"
        folder.mkdir(parents=True)
        (folder / "state.json").write_text(json.dumps({
            "company": company, "proposal": proposal, "date": date,
            "financials": {"FY-Current": {"ebitda": ebitda}},
        }, ensure_ascii=False), encoding="utf-8")
    env = child_env()
    env.pop("PYTHONUTF8", None)
    env["PYTHONIOENCODING"] = "cp1252"
    result = subprocess.run(
        [sys.executable, str(SCRIPTS_DIR / "review_compare.py"),
         "--company", company, "--proposal", proposal, "--current-date", current_date,
         "--prior-date", prior_date],
        cwd=work, env=env, capture_output=True, timeout=120, stdin=subprocess.DEVNULL,
    )
    assert result.returncode == 0, result.stderr.decode("utf-8", "replace")
    report = json.loads(result.stdout.decode("utf-8"))
    assert report["proposal"] == proposal
    assert report["summary"]["changed_paths"] == len(report["changes"])
    assert b"\xe6\x9d\xb1\xe4\xba\xac" in result.stdout
    assert result.stderr == b""



# ---------------------------------------------------------------------------
# state_manager --record-group-basis: real-process contract (issue #48, case A)
# ---------------------------------------------------------------------------

def _cli_group_basis_input():
    return {
        "basis": "supplied_consolidated",
        "reporting_entity": "Synthetic Group Ltd",
        "entities_covered": ["Synthetic Group Ltd", "Synthetic Group OpCo"],
        "source": {
            "document": "Borrower-supplied consolidated accounts",
            "accounting_basis": "UK GAAP, consolidated",
            "date": "2026-08-31",
        },
        "period_end": "2026-03-31",
        "periods": ["FY-Current", "FY-1"],
        "figures_basis": "supplied",
    }


def _install_group_basis_state(cli, financials_source="analyst-supplied", group_basis=None):
    folder = cli.workdir / "deals" / "Synthetic Group Ltd" / "Synthetic Group Facility_2026-01-01"
    folder.mkdir(parents=True, exist_ok=True)
    state = {
        "company": "Synthetic Group Ltd",
        "proposal": "Synthetic Group Facility",
        "date": "2026-01-01",
        "schema_version": "1.1.0",
        "financials_source": financials_source,
        "financials": {"FY-Current": {"ebitda": 12}},
        "ratios": {"FY-Current": {"dscr": 2.0}},
        "multi_period_financials": {"FY-Current": {"revenue": 100}},
        "analyst_supplied_financials": {"FY-Current": {"revenue": 101}},
        "marker": "preserve",
    }
    if group_basis is not None:
        state["group_basis"] = group_basis
    path = folder / "state.json"
    path.write_text(json.dumps(state), encoding="utf-8")
    return path, state


def _run_record_group_basis(cli, group_basis_path, *extra):
    return cli(
        "state_manager", "--record-group-basis", "--company", "Synthetic Group Ltd",
        "--proposal", "Synthetic Group Facility", "--group-basis-file", group_basis_path,
        "--perimeter-confirmed", *extra,
    )


def test_state_manager_record_group_basis_subprocess_round_trips_and_preserves_other_keys(cli):
    state_path, original = _install_group_basis_state(cli)
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(_cli_group_basis_input()), encoding="utf-8")
    result = _run_record_group_basis(cli, group_file, "--date", "2026-10-21")
    assert result.returncode == 0, result.stderr
    assert result.stderr == "" and result.stdout.count("\n") == 1
    report = json.loads(result.stdout)
    assert report == {
        "recorded": True,
        "state_path": report["state_path"],
        "figures_basis": "supplied",
    }
    assert report["state_path"].replace("\\", "/") == (
        "deals/Synthetic Group Ltd/Synthetic Group Facility_2026-01-01/state.json"
    )
    written = json.loads(state_path.read_text(encoding="utf-8"))
    assert written["group_basis"]["perimeter_confirmed"] is True
    assert written["group_basis"]["perimeter_confirmed_on"] == "2026-10-21"
    assert written["financials"] == original["financials"]
    assert written["ratios"] == original["ratios"]
    assert written["multi_period_financials"] == original["multi_period_financials"]
    assert written["analyst_supplied_financials"] == original["analyst_supplied_financials"]
    assert written["financials_source"] == "analyst-supplied" and written["marker"] == "preserve"


def test_state_manager_record_group_basis_refuses_without_confirmation_and_does_not_write(cli):
    state_path, _ = _install_group_basis_state(cli)
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(_cli_group_basis_input()), encoding="utf-8")
    before = state_path.read_bytes()
    result = cli(
        "state_manager", "--record-group-basis", "--company", "Synthetic Group Ltd",
        "--proposal", "Synthetic Group Facility", "--group-basis-file", group_file,
    )
    _assert_a_clean_state_error(result, "--perimeter-confirmed")
    assert state_path.read_bytes() == before


def test_state_manager_record_group_basis_refuses_a_missing_deal_without_creating_it(cli):
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(_cli_group_basis_input()), encoding="utf-8")
    result = _run_record_group_basis(cli, group_file)
    _assert_a_clean_state_error(result, "no exact existing state.json")
    assert not (cli.workdir / "deals").exists()


@pytest.mark.parametrize("file_case", ["missing", "directory", "invalid_json", "non_object"])
def test_state_manager_record_group_basis_refuses_unreadable_or_invalid_json_file(cli, file_case):
    state_path, _ = _install_group_basis_state(cli)
    file_path = cli.workdir / "group-basis.json"
    if file_case == "directory":
        file_path.mkdir()
    elif file_case == "invalid_json":
        file_path.write_text("{bad", encoding="utf-8")
    elif file_case == "non_object":
        file_path.write_text(json.dumps(["not", "an", "object"]), encoding="utf-8")
    before = state_path.read_bytes()
    result = _run_record_group_basis(cli, file_path)
    _assert_a_clean_state_error(result)
    assert state_path.read_bytes() == before


def test_state_manager_record_group_basis_refuses_mismatched_financials_source_without_write(cli):
    state_path, _ = _install_group_basis_state(cli, financials_source="framework-computed")
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(_cli_group_basis_input()), encoding="utf-8")
    before = state_path.read_bytes()
    result = _run_record_group_basis(cli, group_file)
    _assert_a_clean_state_error(result, "figures_basis", "financials_source")
    assert state_path.read_bytes() == before


def test_state_manager_record_group_basis_refuses_malformed_existing_metadata(cli):
    state_path, _ = _install_group_basis_state(cli, group_basis={"basis": "not-supported"})
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(_cli_group_basis_input()), encoding="utf-8")
    before = state_path.read_bytes()
    result = _run_record_group_basis(cli, group_file)
    _assert_a_clean_state_error(result, "group_basis")
    assert state_path.read_bytes() == before


def test_state_manager_record_group_basis_help_lists_all_flags(cli):
    result = cli("state_manager", "--help")
    assert result.returncode == 0 and result.stderr == ""
    for flag in ("--record-group-basis", "--company", "--proposal", "--group-basis-file",
                 "--perimeter-confirmed", "--date"):
        assert flag in result.stdout


def test_state_manager_record_group_basis_reports_missing_json_argument_as_one_error(cli):
    _install_group_basis_state(cli)
    result = cli(
        "state_manager", "--record-group-basis", "--company", "Synthetic Group Ltd",
        "--proposal", "Synthetic Group Facility", "--perimeter-confirmed",
    )
    _assert_a_clean_state_error(result, "--group-basis-file is required")


def test_state_manager_record_group_basis_refuses_invalid_utf8_without_writing(cli):
    state_path, _ = _install_group_basis_state(cli)
    group_file = cli.workdir / "group-basis.json"
    group_file.write_bytes(b"\xff")
    before = state_path.read_bytes()
    result = _run_record_group_basis(cli, group_file)
    _assert_a_clean_state_error(result)
    assert state_path.read_bytes() == before


def test_state_manager_record_group_basis_refuses_impossible_confirmation_date_without_writing(cli):
    state_path, _ = _install_group_basis_state(cli)
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(_cli_group_basis_input()), encoding="utf-8")
    before = state_path.read_bytes()
    result = _run_record_group_basis(cli, group_file, "--date", "2026-02-30")
    _assert_a_clean_state_error(result, "group_basis.perimeter_confirmed_on")
    assert state_path.read_bytes() == before


@pytest.mark.parametrize("extra_flags", [
    ("--prior-date", "2025-10-21"),
    ("--required", "spread"),
])
def test_state_manager_record_group_basis_rejects_flags_for_other_modes(cli, extra_flags):
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(_cli_group_basis_input()), encoding="utf-8")
    result = cli(
        "state_manager", "--record-group-basis", "--company", "Synthetic Group Ltd",
        "--proposal", "Synthetic Group Facility", "--group-basis-file", group_file,
        "--perimeter-confirmed", *extra_flags,
    )
    assert result.returncode == 2
    assert result.stdout == "" and result.stderr.startswith("usage:") and "error:" in result.stderr


def test_state_manager_rejects_group_basis_flags_without_record_mode(cli):
    result = cli(
        "state_manager", "--find", "--company", "Synthetic Group Ltd",
        "--proposal", "Synthetic Group Facility", "--group-basis-file", "group-basis.json",
    )
    assert result.returncode == 2
    assert result.stdout == "" and result.stderr.startswith("usage:") and "only valid with --record-group-basis" in result.stderr


def test_state_manager_record_group_basis_rejects_invalid_metadata_before_writing(cli):
    state_path, _ = _install_group_basis_state(cli)
    group_input = _cli_group_basis_input()
    group_input["source"]["date"] = "2026-02-30"
    group_file = cli.workdir / "group-basis.json"
    group_file.write_text(json.dumps(group_input), encoding="utf-8")
    before = state_path.read_bytes()
    result = _run_record_group_basis(cli, group_file)
    _assert_a_clean_state_error(result, "group_basis.source.date")
    assert state_path.read_bytes() == before


ENTITY_KEYED_FINANCIALS = '{"HoldCo": {"FY-Current": {"revenue": 1000}}, "OpCo A": {"FY-Current": {"revenue": 600}}}'


def test_spreading_check_refuses_an_entity_keyed_financials_file_before_writing(cli):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    financials = cli.workdir / "financials.json"
    financials.write_text(ENTITY_KEYED_FINANCIALS, encoding="utf-8")

    result = cli(
        "spreading_check", "--company", company, "--proposal", proposal,
        "--financials", financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr == (
        "error: financial inputs: unexpected period key 'HoldCo'; expected one of FY-2, FY-1, FY-Current, FY+1, FY+2, FY+3\n"
    )
    assert saved_path.read_bytes() == before


@pytest.mark.parametrize("flag", ["--financials", "--spread"])
def test_orchestrator_refuses_an_entity_keyed_financials_file_before_credentials_or_writing(cli, flag):
    company, proposal, saved_path = _seed_stress_deal(cli)
    before = saved_path.read_bytes()
    financials = cli.workdir / "financials.json"
    financials.write_text(ENTITY_KEYED_FINANCIALS, encoding="utf-8")

    result = cli(
        "orchestrator", "--company", company, "--proposal", proposal,
        flag, financials,
    )

    assert result.returncode == 1
    assert result.stdout == ""
    assert result.stderr == (
        "error: financial inputs: unexpected period key 'HoldCo'; expected one of FY-2, FY-1, FY-Current, FY+1, FY+2, FY+3\n"
    )
    assert "no Anthropic credentials" not in result.stderr
    assert saved_path.read_bytes() == before
