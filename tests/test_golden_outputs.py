"""Golden-output regression tests for the exported .docx and .xlsx (issue #145).

The unit tests check individual builder functions; these check the *whole exported document* a reader would see, so
an unintended change to heading levels, table alignment, the monospace ownership tree, a formula, a number format or
the column layout shows up as a reviewable text diff instead of going unnoticed.

Inputs are synthetic only (tests/fixtures/). Outputs are reduced to readable text by tests/snapshot_utils.py -- no
timestamps, no run-splitting detail -- and compared with tests/snapshots/. After an *intended* change:

    pytest tests/test_golden_outputs.py --update-snapshots

then review the diff of tests/snapshots/ like any other code change. A missing snapshot fails rather than being
created, and CI never passes the flag (guarded below), so a regression cannot refresh its own expectation.
"""
import json
import re
import zipfile
from pathlib import Path

import docx
import pytest
import snapshot_utils
from snapshot_utils import (
    SNAPSHOT_DIR,
    compare_to_snapshot,
    forbidden_content,
    normalize_docx,
    normalize_xlsx,
    visible_text_docx,
    visible_text_xlsx,
)

import deal_export
import docx_builder
import pii_scan
import spreading_builder

FIXTURES = Path(__file__).resolve().parent / "fixtures"
COMPANY, PROPOSAL, DATE = "Synthetic Legacy Co", "Synthetic Fleet Loan", "2026-01-01"
REPO = Path(__file__).resolve().parents[1]


def draft_text():
    return (FIXTURES / "snapshots" / "synthetic_cam_draft.md").read_text(encoding="utf-8")


def export(tmp_path, monkeypatch, state_name, extra=None):
    """Export the synthetic draft for a synthetic state; returns (docx_path, xlsx_path).

    Each call works in its own subdirectory, so one test can export two different states."""
    tmp_path = tmp_path / state_name
    tmp_path.mkdir(exist_ok=True)
    monkeypatch.chdir(tmp_path)
    state = json.loads((FIXTURES / "state" / f"{state_name}.json").read_text(encoding="utf-8"))
    state.update(extra or {})
    folder = tmp_path / "deals" / COMPANY / f"{PROPOSAL}_{DATE}"
    folder.mkdir(parents=True)
    (folder / "state.json").write_text(json.dumps(state), encoding="utf-8")
    out = Path(deal_export.export_deal(COMPANY, PROPOSAL, "corporate_credit", draft_text(), date_str=DATE))
    return next(out.glob("*_CAM.docx")), next(out.glob("*_Spreading.xlsx"))


@pytest.fixture
def framework_computed(tmp_path, monkeypatch, capsys):
    paths = export(tmp_path, monkeypatch, "02_pre_spreading_check")   # downside case + collateral
    capsys.readouterr()
    return paths


@pytest.fixture
def analyst_supplied(tmp_path, monkeypatch, capsys):
    paths = export(tmp_path, monkeypatch, "04_pre_analyst_store",
                   {"analyst_supplied_financials": {"FY-Current": {"revenue": 4650, "cost_of_sales": 2900}}})
    capsys.readouterr()
    return paths


# ---------------------------------------------------------------------------
# Golden files
# ---------------------------------------------------------------------------

def test_cam_docx_matches_its_golden_file(framework_computed, update_snapshots):
    compare_to_snapshot("cam_draft.docx.txt", normalize_docx(framework_computed[0]), update_snapshots)


def test_framework_computed_workbook_matches_its_golden_file(framework_computed, update_snapshots):
    compare_to_snapshot("spreading_framework_computed.xlsx.txt", normalize_xlsx(framework_computed[1]),
                        update_snapshots)


def test_analyst_supplied_workbook_matches_its_golden_file(analyst_supplied, update_snapshots):
    compare_to_snapshot("spreading_analyst_supplied.xlsx.txt", normalize_xlsx(analyst_supplied[1]), update_snapshots)


def test_blank_template_workbook_matches_its_golden_file(tmp_path, update_snapshots):
    path = tmp_path / "blank.xlsx"
    spreading_builder.export_to_xlsx(COMPANY, str(path))
    compare_to_snapshot("spreading_blank.xlsx.txt", normalize_xlsx(path), update_snapshots)


# ---------------------------------------------------------------------------
# Forbidden content: what must never reach a client-facing document
# ---------------------------------------------------------------------------

def test_exported_documents_contain_no_leaked_markup_or_error_text(framework_computed, analyst_supplied):
    for docx_path, xlsx_path in (framework_computed, analyst_supplied):
        assert forbidden_content(visible_text_docx(docx_path)) == []
        assert forbidden_content(visible_text_xlsx(xlsx_path)) == []


def test_the_structured_output_block_is_stripped_but_the_ownership_tree_is_kept(framework_computed):
    text = visible_text_docx(framework_computed[0])
    assert "cp_ids_included" not in text and "```" not in text
    assert "├── 100% Synthetic Legacy Co" in text and "└── 60% Synthetic Subsidiary Co" in text


def test_every_forbidden_pattern_is_detected_by_the_check(tmp_path):
    """The check must be able to fail, or the test above proves nothing."""
    samples = ["```", "cp_ids_included", "| --- | --- |", "**bold**", "# Heading", "value nan", "None",
               "Traceback (most recent call last)", "ValueError: bad", "an Exception", "[Image not embedded: x]"]
    for sample in samples:
        assert forbidden_content(sample), sample
    assert forbidden_content("A clean paragraph with 4,650 revenue and a [placeholder] token.") == []
    assert forbidden_content("Exceptional Costs / (Income) and IFERROR(G2/B2,\"N/A\")") == []   # no false positives


def test_golden_files_contain_nothing_that_looks_like_real_data():
    files = sorted(SNAPSHOT_DIR.glob("*.txt"))
    assert files, "no golden files found"
    for path in files:
        findings = pii_scan.scan_for_likely_real_data(path.read_text(encoding="utf-8"))
        assert findings == [], path.name


# ---------------------------------------------------------------------------
# The machinery itself
# ---------------------------------------------------------------------------

def test_a_changed_output_fails_the_comparison_with_a_readable_diff(tmp_path, monkeypatch):
    monkeypatch.setattr(snapshot_utils, "SNAPSHOT_DIR", tmp_path)    # never write into the tracked snapshots/
    path = tmp_path / "out.docx"
    docx_builder.export_to_docx("# Title\n\nSome text.\n", str(path))
    first = normalize_docx(path)
    docx_builder.export_to_docx("# Title\n\nSome other text.\n", str(path))
    with pytest.raises(AssertionError, match=r"(?s)--update-snapshots.*-P\[Normal\] Some text\.\n"
                                              r".*\+P\[Normal\] Some other text\."):
        # compare against a golden file that holds `first`
        compare_to_snapshot("_unit_probe.txt", first, True)
        compare_to_snapshot("_unit_probe.txt", normalize_docx(path), False)


def test_a_missing_snapshot_fails_instead_of_being_created(monkeypatch, tmp_path):
    monkeypatch.setattr(snapshot_utils, "SNAPSHOT_DIR", tmp_path)
    probe = tmp_path / "_never_created.txt"
    assert not probe.exists()
    with pytest.raises(AssertionError, match="no snapshot"):
        compare_to_snapshot(probe.name, "anything\n", False)
    assert not probe.exists()


def test_normalisation_ignores_timestamps_and_run_splitting(tmp_path):
    a, b = tmp_path / "a.docx", tmp_path / "b.docx"
    docx_builder.export_to_docx("A **bold** word.\n", str(a))
    document = docx.Document(str(a))
    document.core_properties.author = "someone else"
    document.core_properties.comments = "changed metadata"
    document.save(str(b))
    assert normalize_docx(a) == normalize_docx(b) == "P[Normal] A **bold** word.\n"


def test_normalisation_merges_adjacent_runs_of_the_same_formatting(tmp_path):
    """How the builder splits a paragraph into runs is an implementation detail, not part of the snapshot. (Only
    formatted runs can show it: two plain runs concatenate the same either way, but `**bo****ld**` would differ
    from `**bold**` without the merge.)"""
    whole, split = tmp_path / "whole.docx", tmp_path / "split.docx"
    document = docx.Document()
    document.add_paragraph().add_run("a bold word").bold = True
    document.save(str(whole))
    document = docx.Document()
    paragraph = document.add_paragraph()
    for piece in ("a b", "old ", "word"):
        paragraph.add_run(piece).bold = True
    assert len(paragraph.runs) == 3
    document.save(str(split))
    assert normalize_docx(whole) == normalize_docx(split) == "P[Normal] **a bold word**\n"


def test_normalisation_shows_formatting_differences(tmp_path):
    plain, styled = tmp_path / "plain.docx", tmp_path / "styled.docx"
    docx_builder.export_to_docx("A word.\n", str(plain))
    docx_builder.export_to_docx("A *word*.\n", str(styled))
    assert normalize_docx(plain) != normalize_docx(styled)


def test_ci_never_passes_the_snapshot_update_flag():
    """A regression must not be able to refresh its own expectation in CI."""
    workflows = sorted((REPO / ".github" / "workflows").glob("*.y*ml"))
    assert workflows
    for workflow in workflows:
        text = workflow.read_text(encoding="utf-8")
        assert "--update-snapshots" not in text, workflow.name
        assert not re.search(r"addopts[^\n]*update-snapshots", text)
    config = (REPO / "pyproject.toml").read_text(encoding="utf-8")
    assert "update-snapshots" not in config


def test_the_exported_files_are_valid_zip_packages(framework_computed):
    for path in framework_computed:
        with zipfile.ZipFile(path) as package:
            assert package.testzip() is None
