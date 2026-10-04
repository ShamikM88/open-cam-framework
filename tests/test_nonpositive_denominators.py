"""A ratio with a zero or negative denominator is N/A and its covenant is UNRESOLVABLE (issue #169).

Before the fix, negative EBITDA or negative equity gave a finite NEGATIVE leverage/gearing figure, which sits below
every maximum covenant threshold, so a loss-making borrower PASSED "leverage <= 3.5x". The decided semantics, all
pinned here from the ratio through the covenant, the downside case, the compliance reasons and the workbook:

- only the DENOMINATOR is tested; a negative numerator over a positive denominator is a real number (DSCR below
  zero, net cash) and is kept;
- a zero or negative denominator gives None (N/A), never a number;
- a covenant on an N/A ratio is UNRESOLVABLE, for a minimum and a maximum covenant alike -- never PASS, never a
  false FAIL -- with a `reason` naming the metric and the denominator ("gross leverage not meaningful: EBITDA is
  negative");
- a covenant that passes in the base case but is UNRESOLVABLE under stress is reported as a downside breach rather
  than silently dropped;
- the exported workbook agrees with the code, cell by cell.
"""
import openpyxl
import pytest
from hypothesis import given
from hypothesis import strategies as st
from test_spreading_builder import _evaluate

import policy_checks
import policy_engine
import spreading_builder
from spreading_builder import RATIO_DENOMINATORS, describe_undefined_ratio, evaluate_downside_case, evaluate_financial_model

HEALTHY = {"revenue": 1000, "cost_of_sales": 400, "admin_expenses": 200, "interest_paid": 50, "scheduled_principal": 100,
           "trade_debtors": 120, "stock": 60, "trade_creditors": 80, "cash": 100, "current_debt": 100,
           "long_term_debt": 400, "share_capital": 100, "retained_profit": 400}
NEGATIVE_EBITDA = dict(HEALTHY, revenue=500)                       # 500 - 400 - 200 = -100
NEGATIVE_EQUITY = dict(HEALTHY, retained_profit=-300)              # 100 - 300 = -200
ZERO_EBITDA = dict(HEALTHY, revenue=600)                           # 600 - 400 - 200 = 0
ZERO_EQUITY = dict(HEALTHY, retained_profit=-100)                  # 100 - 100 = 0

MAX_LEVERAGE = {"metric": "gross_leverage", "type": "maximum", "threshold": 3.5}
MIN_LEVERAGE = {"metric": "gross_leverage", "type": "minimum", "threshold": 0.5}
MAX_GEARING = {"metric": "gearing", "type": "maximum", "threshold": 2.0}
MIN_DSCR = {"metric": "dscr", "type": "minimum", "threshold": 1.25}


def model(raw, period="FY-Current"):
    return evaluate_financial_model({period: raw})


def state_for(raw, covenants):
    out = model(raw)
    return {"financials": out["financials"], "ratios": out["ratios"], "covenants": covenants}


def covenant_results(raw, covenants):
    return policy_engine.evaluate_deal_policy(state_for(raw, covenants))["covenant_results"]


# ---------------------------------------------------------------------------
# The ratios
# ---------------------------------------------------------------------------

def test_a_healthy_borrower_is_unchanged_every_ratio_is_a_number():
    ratios = model(HEALTHY)["ratios"]["FY-Current"]
    assert ratios["gross_leverage"] == pytest.approx(500 / 400)
    assert ratios["gearing"] == pytest.approx(500 / 500)
    assert all(value is not None for key, value in ratios.items())


@pytest.mark.parametrize("raw, ebitda", [(NEGATIVE_EBITDA, -100), (ZERO_EBITDA, 0)])
def test_non_positive_ebitda_makes_every_ebitda_denominated_ratio_na(raw, ebitda):
    out = model(raw)
    assert out["financials"]["FY-Current"]["ebitda"] == ebitda
    ratios = out["ratios"]["FY-Current"]
    assert ratios["gross_leverage"] is None and ratios["net_debt_to_ebitda"] is None
    assert ratios["fcf_conversion_pct"] is None


@pytest.mark.parametrize("raw, equity", [(NEGATIVE_EQUITY, -200), (ZERO_EQUITY, 0)])
def test_non_positive_equity_makes_gearing_na(raw, equity):
    out = model(raw)
    assert out["financials"]["FY-Current"]["total_equity"] == equity
    assert out["ratios"]["FY-Current"]["gearing"] is None


def test_a_negative_numerator_over_a_positive_denominator_is_kept():
    """Only the denominator decides. Negative EBITDA over a positive debt service is a real DSCR below zero (a
    minimum-DSCR covenant correctly FAILs on it); net cash over positive EBITDA is a real negative net debt ratio."""
    loss = model(NEGATIVE_EBITDA)["ratios"]["FY-Current"]
    assert loss["dscr"] == pytest.approx(-100 / 150) and loss["dscr"] < 0
    assert loss["ebitda_interest_cover"] == pytest.approx(-100 / 50)
    cash_rich = model(dict(HEALTHY, cash=2000))["ratios"]["FY-Current"]
    assert cash_rich["net_debt_to_ebitda"] == pytest.approx((500 - 2000) / 400) and cash_rich["net_debt_to_ebitda"] < 0


def test_the_other_denominators_follow_the_same_rule():
    ratios = model(dict(HEALTHY, revenue=-10, cost_of_sales=-5, interest_paid=-1, scheduled_principal=0,
                        current_debt=0, overdraft=0, trade_creditors=-3, other_current_liabilities=0))["ratios"]["FY-Current"]
    for name in ("trade_debtor_days", "trade_creditor_days", "stock_days", "working_capital_cycle_days",
                 "ebit_interest_cover", "ebitda_interest_cover", "dscr", "current_ratio"):
        assert ratios[name] is None, name


# ---------------------------------------------------------------------------
# Covenants: UNRESOLVABLE, never PASS, with a reason
# ---------------------------------------------------------------------------

def test_the_reported_defect_a_loss_making_borrower_no_longer_passes_maximum_covenants():
    results = covenant_results(dict(NEGATIVE_EBITDA, retained_profit=-300), [MAX_LEVERAGE, MAX_GEARING])
    assert [r["status"] for r in results] == ["UNRESOLVABLE", "UNRESOLVABLE"]
    assert all(r["actual"] is None and r["headroom_pct"] is None for r in results)


def test_the_reason_names_the_metric_and_the_denominator():
    leverage, gearing = covenant_results(dict(NEGATIVE_EBITDA, retained_profit=-300), [MAX_LEVERAGE, MAX_GEARING])
    assert leverage["reason"] == "gross leverage not meaningful: EBITDA is negative"
    assert gearing["reason"] == "gearing not meaningful: total equity is negative"
    assert covenant_results(ZERO_EBITDA, [MAX_LEVERAGE])[0]["reason"] == "gross leverage not defined: EBITDA is zero"
    assert covenant_results(ZERO_EQUITY, [MAX_GEARING])[0]["reason"] == "gearing not defined: total equity is zero"


@pytest.mark.parametrize("covenant_type, threshold", [("maximum", 3.5), ("minimum", 0.5), ("maximum", 0), ("minimum", -5)])
def test_minimum_and_maximum_covenants_are_handled_the_same_way(covenant_type, threshold):
    """Whatever the type or threshold (even zero or negative), a ratio that is N/A cannot pass or fail."""
    covenant = {"metric": "gross_leverage", "type": covenant_type, "threshold": threshold}
    result = covenant_results(NEGATIVE_EBITDA, [covenant])[0]
    assert result["status"] == "UNRESOLVABLE" and result["reason"] == "gross leverage not meaningful: EBITDA is negative"


def test_a_covenant_on_a_ratio_with_a_meaningful_negative_value_still_resolves():
    result = covenant_results(NEGATIVE_EBITDA, [MIN_DSCR])[0]
    assert result["status"] == "FAIL" and result["actual"] < 0 and result["reason"] is None


def test_resolved_covenants_carry_a_null_reason():
    results = covenant_results(HEALTHY, [MAX_LEVERAGE, MIN_LEVERAGE])
    assert [(r["status"], r["reason"]) for r in results] == [("PASS", None), ("PASS", None)]


def test_every_unresolvable_cause_has_a_reason():
    ratios = {"dscr": None, "gross_leverage": 1.0}
    cases = [
        ({"metric": "dscr", "type": "sideways", "threshold": 1}, "unrecognized covenant type 'sideways'"),
        ({"metric": "dscr", "type": "minimum"}, "threshold is missing or not a number"),
        ({"metric": "dscr", "type": "minimum", "threshold": "1.25"}, "threshold is missing or not a number"),
        ({"metric": "nope", "type": "minimum", "threshold": 1}, "metric 'nope' is not among"),
        ({"metric": "dscr", "type": "minimum", "threshold": 1}, "dscr has no computed value for this period"),
    ]
    for covenant, expected in cases:
        result = policy_engine._evaluate_covenant(covenant, ratios)
        assert result["status"] == "UNRESOLVABLE" and expected in result["reason"], covenant


def test_an_n_a_ratio_with_a_positive_denominator_gets_a_generic_reason_not_a_wrong_one():
    """If the stored ratio is None but the figures show a positive denominator (hand-edited state), do not invent a
    denominator problem."""
    state = state_for(HEALTHY, [MAX_LEVERAGE])
    state["ratios"]["FY-Current"]["gross_leverage"] = None
    result = policy_engine.evaluate_deal_policy(state)["covenant_results"][0]
    assert result["status"] == "UNRESOLVABLE" and result["reason"] == "gross_leverage has no computed value for this period"


# ---------------------------------------------------------------------------
# describe_undefined_ratio
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("metric", sorted(RATIO_DENOMINATORS))
def test_every_ratio_with_a_denominator_can_explain_itself(metric):
    name, denominators = RATIO_DENOMINATORS[metric]
    healthy = model(HEALTHY)["financials"]["FY-Current"]
    assert describe_undefined_ratio(metric, healthy) is None            # positive denominators: nothing to explain
    for label, _read in denominators:
        broken = {"raw": dict(HEALTHY), **{k: v for k, v in healthy.items() if k != "raw"}}
        raw_key, financial_key, value = _knob(metric, label)
        if raw_key:
            broken["raw"][raw_key] = value
        else:
            broken[financial_key] = value
        text = describe_undefined_ratio(metric, broken)
        assert text.startswith(f"{name} not meaningful: ") and text.endswith(f"{label} is negative"), (metric, text)


def _knob(metric, label):
    """(raw field, financials field, value) to set so that this denominator is negative."""
    return {
        "debt service (interest paid + scheduled principal)": ("interest_paid", None, -1000),   # + 100 principal < 0
        "EBITDA": (None, "ebitda", -1),
        "total equity": (None, "total_equity", -1),
        "current liabilities": (None, "current_liabilities", -1),
        "interest paid": ("interest_paid", None, -1),
        "revenue": ("revenue", None, -1),
        "cost of sales": ("cost_of_sales", None, -1),
    }[label]


def test_the_dscr_denominator_is_interest_plus_scheduled_principal():
    financials = {"raw": {"interest_paid": 30, "scheduled_principal": -50}}
    assert describe_undefined_ratio("dscr", financials) == \
        "DSCR not meaningful: debt service (interest paid + scheduled principal) is negative"
    assert describe_undefined_ratio("dscr", {"raw": {"interest_paid": 30, "scheduled_principal": -30}}) == \
        "DSCR not defined: debt service (interest paid + scheduled principal) is zero"


def test_the_working_capital_cycle_reports_the_first_non_positive_denominator():
    assert "revenue" in describe_undefined_ratio("working_capital_cycle_days", {"raw": {"revenue": 0, "cost_of_sales": -1}})
    assert "cost of sales" in describe_undefined_ratio("working_capital_cycle_days", {"raw": {"revenue": 5, "cost_of_sales": 0}})


def test_aliases_unknown_metrics_and_missing_financials():
    assert describe_undefined_ratio("EBITDA/Interest", {"raw": {"interest_paid": 0}}) == \
        "EBITDA interest cover not defined: interest paid is zero"
    assert describe_undefined_ratio("not_a_ratio", {"ebitda": -1}) is None
    assert describe_undefined_ratio("gross_leverage", None) is None          # nothing recorded: cannot say
    assert describe_undefined_ratio("gross_leverage", "garbage") is None


# ---------------------------------------------------------------------------
# Property: an N/A ratio never passes or fails; a defined one always does
# ---------------------------------------------------------------------------

amount = st.integers(min_value=-5000, max_value=5000)
raw_amounts = st.fixed_dictionaries({k: amount for k in (
    "revenue", "cost_of_sales", "admin_expenses", "interest_paid", "scheduled_principal", "cash", "current_debt",
    "long_term_debt", "share_capital", "retained_profit", "trade_debtors", "stock", "trade_creditors")})
amount = st.one_of(st.integers(min_value=-5000, max_value=5000),
                   st.floats(min_value=-5000, max_value=5000, allow_nan=False, allow_infinity=False))
covenant_strategy = st.builds(
    lambda metric, kind, threshold: {"metric": metric, "type": kind, "threshold": threshold},
    st.sampled_from(["gross_leverage", "net_debt_to_ebitda", "gearing", "dscr", "current_ratio", "fcf_conversion_pct",
                     "ebit_interest_cover", "ebitda_interest_cover", "trade_debtor_days", "trade_creditor_days",
                     "stock_days", "working_capital_cycle_days"]),
    st.sampled_from(["minimum", "maximum"]), st.integers(min_value=-10, max_value=10))


@given(raw_amounts, st.lists(covenant_strategy, min_size=1, max_size=6))
def test_a_covenant_is_unresolvable_exactly_when_its_ratio_is_na(raw, covenants):
    out = model(raw)
    financials, ratios = out["financials"]["FY-Current"], out["ratios"]["FY-Current"]
    results = policy_engine.evaluate_deal_policy(
        {"financials": out["financials"], "ratios": out["ratios"], "covenants": covenants})["covenant_results"]
    for covenant, result in zip(covenants, results, strict=True):
        metric = covenant["metric"]
        if ratios[metric] is None:
            assert result["status"] == "UNRESOLVABLE" and result["actual"] is None
            _, denominators = RATIO_DENOMINATORS[metric]
            assert any(read(financials) <= 0 for _, read in denominators)          # a denominator really is non-positive
            assert result["reason"] == describe_undefined_ratio(metric, financials)  # and the reason is exactly that
        else:
            assert result["status"] in ("PASS", "FAIL") and result["reason"] is None
            assert (result["status"] == "PASS") == (
                ratios[metric] >= covenant["threshold"] if covenant["type"] == "minimum"
                else ratios[metric] <= covenant["threshold"])


# ---------------------------------------------------------------------------
# The downside case
# ---------------------------------------------------------------------------

STRESS = {"revenue_haircut_pct": 60}                                # revenue 1000 -> 400: EBITDA 400 -> -200


def downside_state(base_raw, covenants, stress=STRESS):
    multi = {"FY-Current": base_raw, "FY+1": base_raw}
    base = evaluate_financial_model(multi)
    downside = evaluate_downside_case(multi, stress)
    return {"financials": base["financials"], "ratios": base["ratios"], "covenants": covenants,
            "downside_case": {"financials": downside["financials"], "ratios": downside["ratios"]}}


def test_stress_that_drives_ebitda_negative_is_reported_not_silently_dropped():
    policy = policy_engine.evaluate_deal_policy(downside_state(HEALTHY, [MAX_LEVERAGE, MAX_GEARING]))
    assert [r["status"] for r in policy["covenant_results"]] == ["PASS", "PASS"]       # fine in the base case
    breaches = {b["metric"]: b for b in policy["downside_covenant_breaches"]}
    assert set(breaches) == {"gross_leverage"}        # equity is unchanged by a revenue haircut: gearing stays PASS
    breach = breaches["gross_leverage"]
    assert breach["downside_status"] == "UNRESOLVABLE" and breach["downside_actual"] is None
    assert breach["reason"] == "gross leverage not meaningful: EBITDA is negative"
    assert breach["year"] == "FY+1" and breach["breach_id"] == "DOWNSIDE-FY-1-GROSS-LEVERAGE"
    assert breach["base_actual"] == pytest.approx(500 / 400) and breach["threshold"] == 3.5


def test_an_ordinary_downside_fail_is_still_reported_as_a_fail_with_no_reason():
    state = downside_state(HEALTHY, [{"metric": "dscr", "type": "minimum", "threshold": 1.25}],
                           stress={"revenue_haircut_pct": 30})
    breach = policy_engine.evaluate_deal_policy(state)["downside_covenant_breaches"][0]
    assert breach["downside_status"] == "FAIL" and breach["reason"] is None and breach["downside_actual"] is not None


def test_a_covenant_that_is_unresolvable_in_the_base_case_is_not_a_downside_breach():
    state = downside_state(NEGATIVE_EBITDA, [MAX_LEVERAGE])
    assert policy_engine.evaluate_deal_policy(state)["downside_covenant_breaches"] == []


def test_a_covenant_that_survives_the_stress_is_not_a_breach():
    assert policy_engine.evaluate_deal_policy(downside_state(HEALTHY, [MAX_LEVERAGE], {"revenue_haircut_pct": 5})
                                              )["downside_covenant_breaches"] == []


def test_a_malformed_downside_financials_value_never_adds_a_failure():
    state = downside_state(HEALTHY, [MAX_LEVERAGE])
    for broken in ("garbage", [1, 2], None):
        state["downside_case"]["financials"] = broken
        breach = policy_engine.evaluate_deal_policy(state)["downside_covenant_breaches"][0]
        assert breach["downside_status"] == "UNRESOLVABLE"
        assert breach["reason"] == "gross_leverage has no computed value for this period"   # honest: cannot say why


# ---------------------------------------------------------------------------
# What the Reviewer / orchestrator sees
# ---------------------------------------------------------------------------

def _compliance_reasons(policy_state):
    return policy_checks.check_draft_compliance(None, policy_state, {})


def test_the_code_enforced_covenant_reason_carries_the_explanation():
    policy = policy_engine.evaluate_deal_policy(state_for(NEGATIVE_EBITDA, [MAX_LEVERAGE]))
    reason = next(r for r in _compliance_reasons(policy) if r.startswith("Covenant UNRESOLVABLE"))
    assert reason == ("Covenant UNRESOLVABLE: gross_leverage (maximum 3.5, actual None) "
                      "-- gross leverage not meaningful: EBITDA is negative")


def test_a_policy_state_without_a_reason_still_renders_the_old_message():
    legacy = {"covenant_results": [{"metric": "dscr", "type": "minimum", "threshold": 1.25, "actual": None,
                                    "status": "UNRESOLVABLE", "headroom_pct": None}]}
    assert _compliance_reasons(legacy) == ["Covenant UNRESOLVABLE: dscr (minimum 1.25, actual None)"]


def test_the_undisclosed_downside_breach_message_carries_the_explanation():
    policy = policy_engine.evaluate_deal_policy(downside_state(HEALTHY, [MAX_LEVERAGE]))
    draft = "Narrative.\n\n```json\n{\"downside_breaches_acknowledged\": []}\n```\n"
    reasons = policy_checks.check_draft_compliance(draft, policy, {})
    message = next(r for r in reasons if r.startswith("Undisclosed Downside Breach"))
    assert message == ("Undisclosed Downside Breach DOWNSIDE-FY-1-GROSS-LEVERAGE: gross_leverage cannot be tested in "
                       "FY+1 under stress (gross leverage not meaningful: EBITDA is negative) and is not addressed "
                       "in the draft.")


def test_an_ordinary_downside_fail_keeps_the_breaches_threshold_wording():
    policy = policy_engine.evaluate_deal_policy(downside_state(HEALTHY, [MIN_DSCR], {"revenue_haircut_pct": 30}))
    draft = "Narrative.\n\n```json\n{\"downside_breaches_acknowledged\": []}\n```\n"
    message = next(r for r in policy_checks.check_draft_compliance(draft, policy, {})
                   if r.startswith("Undisclosed Downside Breach"))
    assert message.endswith("dscr breaches threshold in FY+1 under stress but is not addressed in the draft.")


# ---------------------------------------------------------------------------
# The exported workbook says the same thing as the code
# ---------------------------------------------------------------------------

WORKBOOK_ROWS = {
    "DSCR": "dscr", "EBIT/Interest": "ebit_interest_cover", "EBITDA/Interest": "ebitda_interest_cover",
    "FCF Conversion %": "fcf_conversion_pct", "Gearing % (Interest-Bearing Debt / Equity)": "gearing",
    "Current Ratio": "current_ratio", "Gross Leverage": "gross_leverage", "Net Debt / EBITDA": "net_debt_to_ebitda",
    "Trade Debtor Days": "trade_debtor_days", "Trade Creditor Days": "trade_creditor_days", "Stock Days": "stock_days",
    "Working Capital Cycle (days)": "working_capital_cycle_days",
}
GUARDED_ROWS = [label for label in WORKBOOK_ROWS if label != "Working Capital Cycle (days)"]


@pytest.mark.parametrize("raw", [HEALTHY, NEGATIVE_EBITDA, NEGATIVE_EQUITY, ZERO_EBITDA, ZERO_EQUITY,
                                 dict(NEGATIVE_EBITDA, retained_profit=-300),
                                 dict(HEALTHY, revenue=-10, cost_of_sales=-5)],
                         ids=["healthy", "negative-ebitda", "negative-equity", "zero-ebitda", "zero-equity",
                              "both-negative", "negative-revenue-and-cogs"])
def test_every_workbook_ratio_cell_equals_the_code_ratio_na_included(tmp_path, raw):
    path = tmp_path / "w.xlsx"
    spreading_builder.export_to_xlsx("Synthetic Co", str(path), financial_data={"FY-Current": raw})
    sheet = openpyxl.load_workbook(path)["Financial Spreading"]
    row_of = {sheet.cell(row=r, column=1).value: r for r in range(1, sheet.max_row + 1)}
    ratios = model(raw)["ratios"]["FY-Current"]
    memo = {}
    for label, key in WORKBOOK_ROWS.items():
        cell = _evaluate(sheet, "D", row_of[label], memo)           # column D = FY-Current
        if ratios[key] is None:
            assert cell == "N/A", (label, cell)
        else:
            assert cell == pytest.approx(ratios[key]), (label, cell, ratios[key])


def test_the_workbook_formulas_guard_the_denominator_not_only_the_error(tmp_path):
    path = tmp_path / "w.xlsx"
    spreading_builder.export_to_xlsx("Synthetic Co", str(path))
    sheet = openpyxl.load_workbook(path)["Financial Spreading"]
    row_of = {sheet.cell(row=r, column=1).value: r for r in range(1, sheet.max_row + 1)}
    for label in GUARDED_ROWS:
        formula = sheet.cell(row=row_of[label], column=4).value
        assert formula.startswith("=IFERROR(IF((") and ")>0," in formula and formula.endswith('"N/A"),"N/A")'), label
    wcc = sheet.cell(row=row_of["Working Capital Cycle (days)"], column=4).value
    assert wcc.startswith("=IFERROR(") and wcc.endswith(',"N/A")')      # sums three rows that may read "N/A"


# ---------------------------------------------------------------------------
# Review round: malformed or merely missing figures, and ratios recorded by something else (issue #169)
# ---------------------------------------------------------------------------

def _covenant_against(financials, ratios, covenant=MAX_LEVERAGE):
    return policy_engine.evaluate_deal_policy(
        {"financials": {"FY-Current": financials}, "ratios": {"FY-Current": ratios}, "covenants": [covenant]}
    )["covenant_results"][0]


@pytest.mark.parametrize("financials", [
    {"ebitda": "abc"}, {"ebitda": "n/m"}, {"ebitda": True}, {"ebitda": False}, {"ebitda": [1]}, {"ebitda": {"v": 1}}, {"raw": "zzz"},
    {"raw": ["x"]}, {"ebitda": None, "raw": None}, {}, {"total_equity": "?"},
], ids=repr)
def test_malformed_recorded_figures_never_crash_and_give_the_generic_reason(financials):
    """On main these gave a clean UNRESOLVABLE; the explanation must not turn them into a crash."""
    for metric in ("gross_leverage", "gearing", "dscr", "ebit_interest_cover", "stock_days", "working_capital_cycle_days"):
        result = _covenant_against(financials, {metric: None}, {"metric": metric, "type": "maximum", "threshold": 3})
        assert result["status"] == "UNRESOLVABLE" and result["actual"] is None
        assert result["reason"] == f"{metric} has no computed value for this period", (metric, result["reason"])


@pytest.mark.parametrize("metric", ["dscr", "ebit_interest_cover", "ebitda_interest_cover"])
def test_a_text_value_in_a_raw_denominator_item_gives_the_generic_reason(metric):
    """Only the metrics that read `interest_paid` are affected by it being text; the others read other items."""
    result = _covenant_against({"raw": {"interest_paid": "x"}}, {metric: None},
                               {"metric": metric, "type": "minimum", "threshold": 1})
    assert result["status"] == "UNRESOLVABLE" and result["reason"] == f"{metric} has no computed value for this period"


def test_a_figure_that_is_merely_not_recorded_is_never_called_zero():
    """An analyst-supplied deal records its own subtotals and no `raw` block; a missing EBITDA is not 'zero'."""
    for financials in ({}, {"ebitda": None}, {"gross_profit": 5}):
        assert _covenant_against(financials, {"gross_leverage": None})["reason"] == \
            "gross_leverage has no computed value for this period"
    # raw-based metrics need a raw block at all (only a framework-computed period has one)
    assert _covenant_against({"ebitda": 5}, {"dscr": None}, MIN_DSCR)["reason"] == "dscr has no computed value for this period"
    # a recorded zero, and a recorded negative, are still reported as such
    assert "is zero" in _covenant_against({"ebitda": 0}, {"gross_leverage": None})["reason"]
    assert "is negative" in _covenant_against({"ebitda": -1}, {"gross_leverage": None})["reason"]


def test_an_unreadable_component_does_not_hide_a_readable_non_positive_one():
    """The working capital cycle has two denominators; text in one must not stop the other from being reported."""
    reason = _covenant_against({"raw": {"revenue": "n/m", "cost_of_sales": -5}}, {"working_capital_cycle_days": None},
                               {"metric": "working_capital_cycle_days", "type": "maximum", "threshold": 90})["reason"]
    assert reason == "working capital cycle not meaningful: cost of sales is negative"


def test_inside_a_raw_block_a_missing_item_counts_as_zero_as_the_framework_reads_it():
    reason = _covenant_against({"raw": {"revenue": 10}}, {"ebit_interest_cover": None},
                               {"metric": "ebit_interest_cover", "type": "minimum", "threshold": 2})["reason"]
    assert reason == "EBIT interest cover not defined: interest paid is zero"


@pytest.mark.parametrize("covenant, stored", [(MAX_LEVERAGE, -5.0), (MIN_LEVERAGE, -5.0), (MAX_LEVERAGE, 0.0),
                                              (MAX_GEARING, -2.0), (MAX_LEVERAGE, 1.2)])
def test_a_stored_number_cannot_pass_when_the_recorded_denominator_is_not_positive(covenant, stored):
    """Defence in depth: analyst-supplied ratios are recorded 'exactly as given', and a state checkpointed before this
    fix may hold a negative leverage; neither may PASS against a recorded EBITDA / equity that is not positive."""
    metric = covenant["metric"]
    financials = {"ebitda": -100, "total_equity": -200}
    result = _covenant_against(financials, {metric: stored}, covenant)
    assert result["status"] == "UNRESOLVABLE" and result["actual"] is None and result["headroom_pct"] is None
    assert result["reason"].endswith(f"(the recorded value {stored!r} is ignored)")
    assert "not meaningful" in result["reason"] and ("EBITDA" in result["reason"] or "equity" in result["reason"])


def test_a_stored_number_is_trusted_when_the_denominator_is_positive_or_cannot_be_read():
    assert _covenant_against({"ebitda": 400}, {"gross_leverage": 1.25})["status"] == "PASS"          # positive: trusted
    assert _covenant_against({"ebitda": "n/m"}, {"gross_leverage": 1.25})["status"] == "PASS"        # unreadable: trusted
    assert _covenant_against({}, {"gross_leverage": 1.25})["status"] == "PASS"                       # absent: trusted
    # a DSCR from an analyst-supplied deal (no raw block) is theirs; the engine cannot check its denominator
    assert _covenant_against({"ebitda": 100}, {"dscr": 2.0}, MIN_DSCR)["status"] == "PASS"
    # and a meaningful negative numerator over a positive denominator still resolves as a real number
    result = _covenant_against({"raw": {"interest_paid": 50, "scheduled_principal": 100}}, {"dscr": -0.66}, MIN_DSCR)
    assert result["status"] == "FAIL" and result["actual"] == -0.66 and result["reason"] is None


# ---------------------------------------------------------------------------
# End to end through the real CLI-facing functions
# ---------------------------------------------------------------------------

def test_end_to_end_a_loss_making_deal_is_unresolvable_with_the_reason(tmp_path, monkeypatch):
    import policy_check
    import spreading_check
    import state_manager
    monkeypatch.chdir(tmp_path)
    spreading_check.compute("Synthetic Co", "Loan", {"FY-Current": dict(NEGATIVE_EBITDA, retained_profit=-300)})
    state_manager.write_state("Synthetic Co", "Loan", covenants=[MAX_LEVERAGE, MAX_GEARING])
    results = policy_check.compute("Synthetic Co", "Loan")["policy_state"]["covenant_results"]
    assert [r["status"] for r in results] == ["UNRESOLVABLE", "UNRESOLVABLE"]
    assert [r["reason"] for r in results] == ["gross leverage not meaningful: EBITDA is negative",
                                              "gearing not meaningful: total equity is negative"]
