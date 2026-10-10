"""Tests for scripts/supplied_forecast.py: /project's analyst-supplied mode (issue #124).

What is proved: forward-year figures an analyst supplies are recorded as given and labelled (never run through the
framework's formulas to produce a recorded value); incompatible deals and malformed input are refused with nothing
written; the framework's stress shocks reach an analyst-supplied year only when its raw lines reproduce the supplied
figures; a downside that was wanted but cannot be produced is an explicit, UNRESOLVABLE record that the policy engine
and the draft check cannot pass over; the exported workbook shows what was supplied; and the ordinary framework-
computed /project path is unchanged. All data is synthetic.
"""
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook

import spreading_check
import supplied_forecast as sf
from deal_export import export_deal
from policy_check import compute as policy_compute
from policy_checks import REQUIRED_RISK_TAXONOMY, check_draft_compliance
from policy_engine import evaluate_deal_policy
from spreading_builder import (
    FORWARD_PERIOD_KEYS,
    PERIOD_HEADERS,
    SECTIONS,
    SUPPLIED_RATIOS,
    SUPPLIED_ROW_LABELS,
    SUPPLIED_SUBTOTALS,
    evaluate_downside_case,
    evaluate_financial_model,
)
from state_manager import SchemaVersionError, StateError, read_state, write_state

FORECAST = {
    "FY+1": {"subtotals": {"ebitda": 300, "gross_profit": 500}, "ratios": {"dscr": 1.4, "gross_leverage": 2.0}},
    "FY+2": {"subtotals": {"ebitda": 320}, "ratios": {"dscr": 1.5, "gross_leverage": None}},
}
NOTE = "Management budget convention"
COMPANY, PROPOSAL = "Synthetic Co", "Fleet Loan"


@pytest.fixture
def deal(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    return tmp_path


def state_file(root):
    (path,) = root.glob(f"deals/{COMPANY}/{PROPOSAL}_*/state.json")
    return path


def deals_bytes(root):
    return {p.relative_to(root).as_posix(): p.read_bytes() for p in (root / "deals").rglob("*") if p.is_file()}


def record(forecast=None, downside=None, stress=None, note=NOTE):
    data = {"forecast": FORECAST if forecast is None else forecast}
    if downside is not None:
        data["downside"] = downside
    return sf.record(COMPANY, PROPOSAL, sf.validate(data), sf.validate_stress(stress) if stress else None, note)


def state():
    return read_state(COMPANY, PROPOSAL)


# A year whose raw lines the framework's formulas reproduce exactly (to the supplied two decimals). Every line a
# covenant ratio reads is given, the ones that are nil as explicit zeros: an omitted line is not a known zero.
LINES = {"revenue": 1000, "cost_of_sales": 600, "admin_expenses": 100, "interest_paid": 50, "current_debt": 20,
         "overdraft": 0, "long_term_debt": 200, "loan_notes": 0, "cash": 10, "trade_debtors": 100, "stock": 20,
         "tangible_assets": 300, "trade_creditors": 80, "share_capital": 100, "retained_profit": 200,
         "depreciation": 0, "amortisation": 0, "other_income": 0, "interest_received": 0, "scheduled_principal": 0,
         "capex": 0, "tax_paid": 0, "other_current_assets": 0, "other_current_liabilities": 0}
# What the framework-derived downside leaves out for that year: figures that read a line nobody supplied.
UNSUPPORTED_OUTPUTS = {"profit_before_tax", "net_profit", "total_assets", "total_liabilities", "tangible_net_worth"}


def reconciled_year(lines=None, drop=()):
    lines = {k: v for k, v in (lines or LINES).items() if k not in drop}
    derived = evaluate_financial_model({"FY+1": lines})
    return {"subtotals": {k: round(derived["financials"]["FY+1"][k], 2) for k in ("ebitda", "gross_profit")},
            "ratios": {k: None if derived["ratios"]["FY+1"][k] is None else round(derived["ratios"]["FY+1"][k], 2)
               for k in ("dscr", "gross_leverage")},
            "lines": lines}


# ---------------------------------------------------------------------------
# The supplied input is validated
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("mutate, message", [
    (lambda d: d.update(extra=1), "unknown key"),
    (lambda d: d.update(forecast={}), "non-empty object"),
    (lambda d: d.update(forecast=[]), "non-empty object"),
    (lambda d: d["forecast"].update({"FY-1": {"subtotals": {"ebitda": 1}}}), "unknown forecast year"),
    (lambda d: d["forecast"]["FY+1"]["subtotals"].update(ebtida=1), "unknown name"),
    (lambda d: d["forecast"]["FY+1"]["ratios"].update(dscrr=1), "unknown name"),
    (lambda d: d["forecast"]["FY+1"]["subtotals"].update(ebitda="300"), "expected a finite number"),
    (lambda d: d["forecast"]["FY+1"]["subtotals"].update(ebitda=None), "expected a finite number"),
    (lambda d: d["forecast"]["FY+1"]["subtotals"].update(ebitda=True), "expected a finite number"),
    (lambda d: d["forecast"]["FY+1"]["subtotals"].update(ebitda=float("nan")), "expected a finite number"),
    (lambda d: d["forecast"]["FY+1"]["ratios"].update(dscr=float("inf")), "expected a finite number"),
    (lambda d: d["forecast"]["FY+1"]["ratios"].update(dscr="1.4x"), "expected a finite number or null"),
    (lambda d: d["forecast"]["FY+1"].update(colour=1), "unknown key"),
    (lambda d: d["forecast"].update({"FY+3": {"lines": {"revenue": 1}}}), "framework-computed mode's input"),
    (lambda d: d["forecast"].update({"FY+3": "x"}), "must be an object"),
    (lambda d: d["forecast"]["FY+1"].update(lines={"reveneu": 1}), "unknown name"),
    (lambda d: d.update(downside={"periods": {"FY+1": {"ratios": {"dscr": 1}}}}), "downside.description is required"),
    (lambda d: d.update(downside={"description": " ", "periods": {"FY+1": {"ratios": {"dscr": 1}}}}),
     "downside.description is required"),
    (lambda d: d.update(downside={"description": "x"}), "downside.periods must be"),
    (lambda d: d.update(downside={"description": "x", "periods": {"FY-1": {"ratios": {"dscr": 1}}}}),
     "unknown year"),
    (lambda d: d.update(downside={"description": "x", "periods": {"FY+1": {"ratios": {"dscr": 1}, "lines":
                                                                           {"revenue": 1}}}}), "raw lines are not"),
    (lambda d: d.update(downside={"description": "x", "extra": 1, "periods": {}}), "unknown key"),
    (lambda d: d.update(downside="a scenario"), "must be an object"),
])
def test_malformed_or_incomplete_input_is_refused_with_a_message_naming_the_problem(mutate, message):
    data = json.loads(json.dumps({"forecast": FORECAST}))
    mutate(data)
    with pytest.raises(sf.ForecastError, match=message):
        sf.validate(data)


def test_the_top_level_must_be_an_object_and_stress_assumptions_are_checked_too():
    with pytest.raises(sf.ForecastError, match="JSON object"):
        sf.validate([])
    with pytest.raises(sf.ForecastError, match="unknown key"):
        sf.validate_stress({"revenue_haircut": 5})
    with pytest.raises(sf.ForecastError, match="finite number"):
        sf.validate_stress({"revenue_haircut_pct": "five"})
    with pytest.raises(sf.ForecastError, match="JSON object"):
        sf.validate_stress([1])
    assert sf.validate_stress({"revenue_haircut_pct": 5, "interest_rate_bump_bps": 200}) == {
        "revenue_haircut_pct": 5, "interest_rate_bump_bps": 200}


def test_a_refused_input_writes_nothing(deal):
    with pytest.raises(sf.ForecastError):
        record(forecast={"FY+1": {"subtotals": {"ebitda": "x"}}})
    assert not (deal / "deals").exists()


# ---------------------------------------------------------------------------
# Recorded as supplied, labelled, and reloadable
# ---------------------------------------------------------------------------

def test_supplied_subtotals_and_ratios_are_stored_exactly_as_given_and_survive_a_reload(deal):
    record()
    on_disk = json.loads(state_file(deal).read_text(encoding="utf-8"))
    assert on_disk["financials"] == {"FY+1": {"ebitda": 300, "gross_profit": 500}, "FY+2": {"ebitda": 320}}
    assert on_disk["ratios"] == {"FY+1": {"dscr": 1.4, "gross_leverage": 2.0}, "FY+2": {"dscr": 1.5, "gross_leverage": None}}
    assert on_disk["financials_source"] == "analyst-supplied" and on_disk["forecast_source"] == "analyst-supplied"
    assert "multi_period_financials" not in on_disk, "the framework-computed raw store is never written"
    assert state() == on_disk                       # read_state returns exactly what was saved
    assert on_disk["financials_source_note"] == f"Forecast (FY+1 to FY+3) figures are analyst-supplied: {NOTE}"


def test_nothing_supplied_is_ever_recomputed_by_the_frameworks_formulas(deal, monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("an analyst-supplied forecast must not be evaluated by the framework's formulas")

    monkeypatch.setattr(sf, "evaluate_financial_model", forbidden)
    monkeypatch.setattr(sf, "evaluate_downside_case", forbidden)
    monkeypatch.setattr(spreading_check, "evaluate_financial_model", forbidden)
    monkeypatch.setattr(spreading_check, "evaluate_downside_case", forbidden)
    record(downside={"description": "Revenue down 10%", "periods": {"FY+1": {"ratios": {"dscr": 1.1}}}},
           stress={"revenue_haircut_pct": 10})            # stress requested, but there are no raw lines to test
    assert state()["financials"]["FY+1"] == {"ebitda": 300, "gross_profit": 500}


def test_an_optional_raw_breakdown_goes_to_the_existing_analyst_supplied_store_only(deal):
    forecast = json.loads(json.dumps(FORECAST))
    forecast["FY+1"]["lines"] = {"revenue": 1000, "cost_of_sales": 600}
    record(forecast=forecast)
    saved = state()
    assert saved["analyst_supplied_financials"] == {"FY+1": {"revenue": 1000, "cost_of_sales": 600}}
    assert "multi_period_financials" not in saved and "raw" not in saved["financials"]["FY+1"]


def test_the_deal_wide_basis_and_the_caveat_note_are_set_so_the_cam_must_disclose(deal):
    record()
    saved = state()
    assert saved["financials_source"] == "analyst-supplied"
    # The existing code-enforced rule: an analyst-supplied deal's draft must carry the caveat.
    payload = {"cp_ids_included": [], "cs_ids_included": [], "risk_categories_covered": {}, "reported_figures": {},
               "sources": ["x"]}
    draft = "# CAM\n\n```json\n" + json.dumps(payload) + "\n```"
    reasons = check_draft_compliance(draft, evaluate_deal_policy(saved), {}, financials_source="analyst-supplied")
    assert any("Missing Analyst-Supplied Spreading Disclosure" in r for r in reasons)


def test_a_second_call_adds_years_and_merges_without_dropping_what_is_recorded(deal):
    record(forecast={"FY+1": FORECAST["FY+1"]})
    record(forecast={"FY+2": FORECAST["FY+2"], "FY+1": {"ratios": {"dscr": 1.45}}})
    saved = state()
    assert saved["ratios"]["FY+1"] == {"dscr": 1.45, "gross_leverage": 2.0}
    assert saved["financials"]["FY+1"] == {"ebitda": 300, "gross_profit": 500} and "FY+2" in saved["financials"]
    assert saved["financials_source_note"].count("Forecast (FY+1 to FY+3)") == 1, "the note is not repeated"


def test_historical_figures_an_analyst_supplied_deal_already_has_are_kept(deal):
    write_state(COMPANY, PROPOSAL, financials_source="analyst-supplied", financials={"FY-Current": {"ebitda": 250}},
                ratios={"FY-Current": {"dscr": 1.3}}, analyst_supplied_financials={"FY-Current": {"revenue": 900}},
                financials_source_note="Depreciation embedded in Cost of Goods Sold", steps_completed=["spread"])
    record()
    saved = state()
    assert saved["financials"]["FY-Current"] == {"ebitda": 250} and saved["ratios"]["FY-Current"] == {"dscr": 1.3}
    assert saved["analyst_supplied_financials"] == {"FY-Current": {"revenue": 900}}
    assert saved["financials_source_note"].startswith("Depreciation embedded in Cost of Goods Sold Forecast (FY+1")
    assert saved["steps_completed"] == ["spread"]


def test_a_note_is_required(deal):
    with pytest.raises(sf.ForecastError, match="--note is required"):
        record(note=" ")
    assert not (deal / "deals").exists()


# ---------------------------------------------------------------------------
# One basis: what the deal cannot take is refused before anything is written
# ---------------------------------------------------------------------------

def refused(deal, message, error=sf.ForecastError):
    before = deals_bytes(deal)
    with pytest.raises(error, match=message):
        record()
    assert deals_bytes(deal) == before


def test_a_deal_with_framework_computed_forward_years_is_refused_not_blended(deal):
    spreading_check.compute(COMPANY, PROPOSAL, {"FY+1": {"revenue": 100, "cost_of_sales": 60}},
                            update_financials_source=False)
    refused(deal, r"forward years \['FY\+1'\] are already recorded from raw line items")


def test_a_deal_with_framework_computed_history_is_refused(deal):
    spreading_check.compute(COMPANY, PROPOSAL, {"FY-Current": {"revenue": 100, "cost_of_sales": 60}})
    refused(deal, "financials_source is 'framework-computed'.*default /project path")


def test_figures_with_no_recorded_basis_are_not_assumed_to_be_analyst_supplied(deal):
    write_state(COMPANY, PROPOSAL, financials={"FY-Current": {"ebitda": 250}})
    refused(deal, "no financials_source, so their basis is unknown")


def test_forward_figures_with_no_analyst_flag_are_refused_not_overwritten(deal):
    write_state(COMPANY, PROPOSAL, financials_source="analyst-supplied", ratios={"FY+1": {"dscr": 9}})
    refused(deal, "already recorded from raw line items")


def test_a_deal_with_only_non_financial_state_or_none_at_all_can_take_it(deal):
    write_state(COMPANY, PROPOSAL, inputs={"pd": "0.2%"}, steps_completed=["triage"])
    record()
    assert state()["inputs"] == {"pd": "0.2%"} and state()["forecast_source"] == "analyst-supplied"


def test_a_newer_schema_is_refused_and_the_state_is_byte_identical(deal):
    write_state(COMPANY, PROPOSAL, inputs={})
    path = state_file(deal)
    saved = json.loads(path.read_text(encoding="utf-8"))
    saved["schema_version"] = "9.0.0"
    path.write_text(json.dumps(saved), encoding="utf-8")
    refused(deal, "newer framework", error=SchemaVersionError)


@pytest.mark.parametrize("key, value", [("financials", ["x"]), ("ratios", "x"), ("downside_case", [1]),
                                        ("analyst_supplied_financials", 3), ("financials_source", "made-up")])
def test_a_wrongly_shaped_stored_key_is_refused_before_anything_is_written(deal, key, value):
    write_state(COMPANY, PROPOSAL, inputs={}, **{key: value})
    refused(deal, key, error=StateError)


def test_an_unrecognised_forecast_source_is_refused(deal):
    write_state(COMPANY, PROPOSAL, forecast_source="guessed")
    refused(deal, "forecast_source", error=StateError)


# ---------------------------------------------------------------------------
# The framework paths cannot overwrite or relabel an analyst-supplied forecast
# ---------------------------------------------------------------------------

def test_framework_computed_forward_years_cannot_overwrite_an_analyst_supplied_forecast(deal):
    record()
    before = deals_bytes(deal)
    with pytest.raises(ValueError, match="supplied by the analyst and is recorded as given"):
        spreading_check.compute(COMPANY, PROPOSAL, {"FY+1": {"revenue": 100}}, update_financials_source=False)
    assert deals_bytes(deal) == before


def test_a_framework_computed_historical_spread_cannot_relabel_the_whole_deal(deal):
    record()
    before = deals_bytes(deal)
    with pytest.raises(ValueError, match="stamp the whole deal framework-computed"):
        spreading_check.compute(COMPANY, PROPOSAL, {"FY-Current": {"revenue": 100}})
    assert deals_bytes(deal) == before


def test_the_ordinary_framework_computed_project_path_is_unchanged(deal):
    state_after = spreading_check.compute(
        COMPANY, PROPOSAL, {"FY+1": {"revenue": 100, "cost_of_sales": 60, "admin_expenses": 10}},
        stress_assumptions={"revenue_haircut_pct": 10}, update_financials_source=False)
    expected = evaluate_downside_case({"FY+1": {"revenue": 100, "cost_of_sales": 60, "admin_expenses": 10}},
                                      {"revenue_haircut_pct": 10})
    assert state_after["downside_case"] == expected
    assert set(state_after["downside_case"]) == {"financials", "ratios"}, "no basis/unavailable keys appear"
    assert "forecast_source" not in state_after and state_after["multi_period_financials"]["FY+1"]["revenue"] == 100


# ---------------------------------------------------------------------------
# The downside case: supplied, derived only where valid, otherwise unavailable
# ---------------------------------------------------------------------------

def test_an_analysts_own_stressed_forecast_is_recorded_as_analyst_supplied_with_its_description(deal):
    scenario = {"description": "Revenue down 10% for twelve months", "periods": {
        "FY+1": {"subtotals": {"ebitda": 240}, "ratios": {"dscr": 1.1, "gross_leverage": 2.5}}}}
    record(downside=scenario)
    case = state()["downside_case"]
    assert case["financials"] == {"FY+1": {"ebitda": 240}} and case["ratios"]["FY+1"]["dscr"] == 1.1
    assert case["basis"] == {"FY+1": "analyst-supplied"} and case["description"] == scenario["description"]
    assert set(case["unavailable"]) == {"FY+2"}, "a year with no scenario is not silently absent"
    assert "no downside scenario was supplied for this year" in case["unavailable"]["FY+2"]


def test_a_stressed_forecast_for_a_year_with_no_forecast_is_refused(deal):
    with pytest.raises(sf.ForecastError, match="no analyst-supplied forecast for FY\\+3 to stress"):
        record(downside={"description": "x", "periods": {"FY+3": {"ratios": {"dscr": 1}}}})
    assert not (deal / "deals").exists()


def test_an_earlier_analyst_scenario_survives_a_later_call_that_does_not_restate_it(deal):
    record(downside={"description": "Scenario A", "periods": {"FY+1": {"ratios": {"dscr": 1.1}}}})
    record(forecast={"FY+2": FORECAST["FY+2"]})
    case = state()["downside_case"]
    assert case["basis"] == {"FY+1": "analyst-supplied"} and case["description"] == "Scenario A"
    assert case["ratios"]["FY+1"] == {"dscr": 1.1} and set(case["unavailable"]) == {"FY+2"}


def test_the_frameworks_shocks_are_applied_only_to_a_year_whose_lines_reproduce_the_supplied_figures(deal):
    year = reconciled_year()
    record(forecast={"FY+1": year}, stress={"revenue_haircut_pct": 10, "opex_increase_pct": 5})
    case = state()["downside_case"]
    expected = evaluate_downside_case({"FY+1": LINES}, {"revenue_haircut_pct": 10, "opex_increase_pct": 5})
    kept = {k: v for k, v in expected["financials"]["FY+1"].items() if k not in UNSUPPORTED_OUTPUTS}
    assert case["financials"] == {"FY+1": kept} and case["ratios"] == expected["ratios"]
    assert case["basis"] == {"FY+1": sf.FRAMEWORK_DERIVED} and case["unavailable"] == {}
    assert state()["financials"]["FY+1"] == year["subtotals"], "the base case stays exactly as supplied"


def test_a_year_whose_lines_do_not_reproduce_the_supplied_figures_gets_no_shocked_case(deal):
    year = reconciled_year()
    year["subtotals"]["ebitda"] += 25            # the analyst's convention differs from the framework's formula
    record(forecast={"FY+1": year}, stress={"revenue_haircut_pct": 10})
    case = state()["downside_case"]
    assert case["financials"] == {} and case["basis"] == {}
    assert "ebitda: supplied" in case["unavailable"]["FY+1"] and "formulas give" in case["unavailable"]["FY+1"]


def test_a_ratio_that_does_not_reproduce_also_blocks_the_shocks(deal):
    year = reconciled_year()
    year["ratios"]["dscr"] = 9.99
    record(forecast={"FY+1": year}, stress={"revenue_haircut_pct": 10})
    assert "dscr: supplied 9.99" in state()["downside_case"]["unavailable"]["FY+1"]


def test_a_requested_shock_needs_the_lines_it_acts_on(deal):
    year = reconciled_year(drop=("admin_expenses",))
    # Lines that omit admin_expenses still reproduce the figures supplied from those same lines.
    record(forecast={"FY+1": year}, stress={"opex_increase_pct": 5})
    assert "opex_increase_pct acts on ['admin_expenses'], which were not supplied" in \
        state()["downside_case"]["unavailable"]["FY+1"]


def test_a_year_with_no_lines_or_no_non_zero_shock_cannot_be_stressed_by_the_framework(deal):
    record(stress={"revenue_haircut_pct": 10})
    assert "no raw lines were supplied" in state()["downside_case"]["unavailable"]["FY+1"]
    record(forecast={"FY+1": reconciled_year()}, stress={"revenue_haircut_pct": 0})
    assert "none of the stress shocks is non-zero" in state()["downside_case"]["unavailable"]["FY+1"]


def test_an_analyst_scenario_takes_precedence_over_the_frameworks_shocks_for_its_year(deal):
    year = reconciled_year()
    record(forecast={"FY+1": year}, stress={"revenue_haircut_pct": 10},
           downside={"description": "Own scenario", "periods": {"FY+1": {"ratios": {"dscr": 1.0}}}})
    case = state()["downside_case"]
    assert case["basis"] == {"FY+1": "analyst-supplied"} and case["ratios"]["FY+1"] == {"dscr": 1.0}


def test_no_stress_and_no_scenario_means_no_downside_was_asked_for_and_none_is_invented(deal):
    record()
    assert state()["downside_case"] == {}


# ---------------------------------------------------------------------------
# Policy: an unavailable downside is never a pass
# ---------------------------------------------------------------------------

COVENANTS = [{"metric": "dscr", "type": "minimum", "threshold": 1.2}]


def policy(deal, covenants=COVENANTS):
    write_state(COMPANY, PROPOSAL, covenants=covenants)
    return policy_compute(COMPANY, PROPOSAL)["policy_state"]


def test_a_supplied_stressed_ratio_that_fails_a_covenant_is_a_downside_breach(deal):
    record(downside={"description": "Severe", "periods": {"FY+1": {"ratios": {"dscr": 1.1}},
                                                        "FY+2": {"ratios": {"dscr": 1.3}}}})
    breaches = policy(deal)["downside_covenant_breaches"]
    assert [(b["year"], b["downside_status"], b["downside_actual"]) for b in breaches] == [("FY+1", "FAIL", 1.1)]


def test_a_year_with_no_downside_is_an_unresolvable_breach_where_the_base_case_passes(deal):
    record(downside={"description": "Severe", "periods": {"FY+1": {"ratios": {"dscr": 1.3}}}})
    breaches = policy(deal)["downside_covenant_breaches"]
    (breach,) = breaches
    assert (breach["year"], breach["metric"], breach["downside_status"]) == ("FY+2", "dscr", "UNRESOLVABLE")
    assert breach["downside_actual"] is None and breach["base_actual"] == 1.5
    assert breach["breach_id"] == "DOWNSIDE-FY-2-DSCR"
    assert "downside analysis unavailable for FY+2: no downside scenario was supplied" in breach["reason"]


def test_every_covenant_that_passes_in_an_unavailable_year_is_recorded_not_just_the_first(deal):
    record(forecast={"FY+1": {"ratios": {"dscr": 1.4, "gross_leverage": 2.0}}}, stress={"revenue_haircut_pct": 10})
    covenants = COVENANTS + [{"metric": "gross_leverage", "type": "maximum", "threshold": 3.0}]
    ids = [b["breach_id"] for b in policy(deal, covenants)["downside_covenant_breaches"]]
    assert ids == ["DOWNSIDE-FY-1-DSCR", "DOWNSIDE-FY-1-GROSS-LEVERAGE"]


def test_a_year_whose_base_case_already_fails_or_has_no_ratios_adds_no_unavailable_breach(deal):
    record(forecast={"FY+1": {"ratios": {"dscr": 1.0}}, "FY+2": {"subtotals": {"ebitda": 1}}},
           stress={"revenue_haircut_pct": 10})
    assert policy(deal)["downside_covenant_breaches"] == []


def test_an_unavailable_downside_is_not_silently_passed_by_the_draft_check(deal):
    write_state(COMPANY, PROPOSAL, financials_source="analyst-supplied", financials={"FY-Current": {"ebitda": 400}},
                ratios={"FY-Current": {"dscr": 1.6}})
    record(downside={"description": "Severe", "periods": {"FY+1": {"ratios": {"dscr": 1.3}}}})
    saved_policy = policy(deal)
    taxonomy = {category: {"status": "covered"} for category in REQUIRED_RISK_TAXONOMY}

    def reasons_for(acknowledged):
        payload = {"cp_ids_included": [c["cp_id"] for c in saved_policy["required_conditions_precedent"]],
                   "cs_ids_included": [c["cs_id"] for c in saved_policy["required_conditions_subsequent"]],
                   "risk_categories_covered": taxonomy, "reported_figures": {}, "sources": ["x"],
                   "financials_source_disclosed": True, "downside_breaches_acknowledged": acknowledged}
        draft = "# CAM\n\n```json\n" + json.dumps(payload) + "\n```"
        return check_draft_compliance(draft, saved_policy, {}, financials_source="analyst-supplied")

    silent = reasons_for([])
    assert any("Undisclosed Downside Breach DOWNSIDE-FY-2-DSCR" in r and "cannot be tested in FY+2 under stress" in r
               and "downside analysis unavailable" in r for r in silent)
    assert reasons_for(["DOWNSIDE-FY-2-DSCR"]) == []


def test_deals_without_the_unavailable_record_are_evaluated_exactly_as_before():
    state_dict = {"ratios": {"FY-Current": {"dscr": 1.5}, "FY+1": {"dscr": 1.4}}, "covenants": COVENANTS,
                  "downside_case": {"financials": {}, "ratios": {"FY+1": {"dscr": 1.0}}}}
    (breach,) = evaluate_deal_policy(state_dict)["downside_covenant_breaches"]
    assert breach["downside_status"] == "FAIL"
    for tolerated in (None, [], "x", {"FY+1": 3}):
        state_dict["downside_case"] = {"financials": {}, "ratios": {}, "unavailable": tolerated}
        assert isinstance(evaluate_deal_policy(state_dict)["downside_covenant_breaches"], list)


# ---------------------------------------------------------------------------
# Export: the workbook shows what was supplied, never a recomputation
# ---------------------------------------------------------------------------

def export(deal):
    draft = "# CAM\n"
    out = export_deal(COMPANY, PROPOSAL, "corporate_credit", draft, date_str=state_file(deal).parent.name.rsplit("_", 1)[1])
    workbook = load_workbook(next(Path(out).glob("*_Spreading.xlsx")))
    return workbook["Financial Spreading"]


def row_of(sheet, label):
    return next(r for r in sheet.iter_rows(min_row=2) if r[0].value == label)


def test_an_analyst_supplied_forecast_is_exported_as_values_with_labelled_columns(deal):
    record(downside={"description": "Severe", "periods": {"FY+1": {"subtotals": {"ebitda": 240},
                                                                  "ratios": {"dscr": 1.1}}}},
           stress={"revenue_haircut_pct": 10})
    sheet = export(deal)
    headers = [c.value for c in sheet[1]]
    assert headers[4] == "FY+1 (analyst-supplied)" and headers[5] == "FY+2 (analyst-supplied)"
    assert headers[6] == "FY+3", "a year with nothing supplied keeps its ordinary header"
    assert headers[7] == "FY+1 (Downside, analyst-supplied)" and headers[8] == "FY+2 (Downside, unavailable)"
    ebitda, dscr = row_of(sheet, "EBITDA"), row_of(sheet, "DSCR")
    assert ebitda[4].value == 300 and ebitda[5].value == 320 and ebitda[7].value == 240
    assert dscr[4].value == 1.4 and dscr[7].value == 1.1
    assert row_of(sheet, "Gross Leverage")[5].value == "N/A", "a ratio supplied as null is N/A, as everywhere else"
    gross_profit = row_of(sheet, "Gross Profit")
    assert gross_profit[4].value == 500 and gross_profit[5].value is None
    for column in (4, 5, 7, 8):
        assert not str(gross_profit[column].value or "").startswith("="), column
    assert str(gross_profit[6].value).startswith("="), "FY+3 was not supplied, so it keeps its ordinary formula"
    assert row_of(sheet, "EBITDA")[8].value is None, "an unavailable downside column is blank, not a formula"
    assert row_of(sheet, "Net Profit")[4].value is None, "a figure nobody supplied is blank, not computed"


def test_a_framework_derived_downside_still_uses_its_raw_lines_in_the_workbook(deal):
    record(forecast={"FY+1": reconciled_year()}, stress={"revenue_haircut_pct": 10})
    sheet = export(deal)
    assert [c.value for c in sheet[1]][7] == "FY+1 (Downside)"
    assert row_of(sheet, "Revenue")[7].value == 900, "the shocked revenue is written to the raw-input cell"
    assert str(row_of(sheet, "EBITDA")[7].value).startswith("="), "its subtotals stay formulas over those cells"


def test_every_other_deals_workbook_is_unchanged(deal):
    spreading_check.compute(COMPANY, PROPOSAL, {"FY-Current": {"revenue": 100, "cost_of_sales": 60},
                                                "FY+1": {"revenue": 110, "cost_of_sales": 60}},
                            stress_assumptions={"revenue_haircut_pct": 10})
    sheet = export(deal)
    assert [c.value for c in sheet[1]] == PERIOD_HEADERS
    assert str(row_of(sheet, "EBITDA")[4].value).startswith("=")


# ---------------------------------------------------------------------------
# The names the script, the workbook and the commands share
# ---------------------------------------------------------------------------

def test_the_supplied_names_and_their_workbook_rows_are_the_frameworks_own():
    import transcription_check
    assert SUPPLIED_SUBTOTALS == transcription_check.SUBTOTAL_FIELDS and SUPPLIED_RATIOS == transcription_check.RATIO_FIELDS
    labels = {label for _, rows in SECTIONS for label, _ in rows}
    assert set(SUPPLIED_ROW_LABELS.values()) <= labels
    assert set(SUPPLIED_ROW_LABELS) <= set(SUPPLIED_SUBTOTALS) | set(SUPPLIED_RATIOS)
    assert set(SUPPLIED_SUBTOTALS) - set(SUPPLIED_ROW_LABELS) == {"total_debt"}
    derived = evaluate_financial_model({"FY+1": {"revenue": 1}})
    assert set(SUPPLIED_SUBTOTALS) <= set(derived["financials"]["FY+1"]) and set(SUPPLIED_RATIOS) <= set(
        derived["ratios"]["FY+1"])


def test_the_forward_years_are_the_frameworks():
    assert tuple(FORWARD_PERIOD_KEYS) == ("FY+1", "FY+2", "FY+3")


# ---------------------------------------------------------------------------
# main()
# ---------------------------------------------------------------------------

def write_json(path, data):
    path.write_text(json.dumps(data), encoding="utf-8")
    return path


def test_main_prints_a_summary_of_what_was_recorded_and_what_is_unavailable(deal, capsys):
    forecast = write_json(deal / "f.json", {"forecast": FORECAST})
    stress = write_json(deal / "s.json", {"revenue_haircut_pct": 10})
    sf.main(["--company", COMPANY, "--proposal", PROPOSAL, "--forecast", str(forecast), "--stress-assumptions",
             str(stress), "--note", NOTE])
    out = json.loads(capsys.readouterr().out)
    assert out["forecast_years"] == ["FY+1", "FY+2"] and out["financials_source"] == "analyst-supplied"
    assert set(out["downside"]["unavailable"]) == {"FY+1", "FY+2"} and out["downside"]["basis"] == {}


def test_main_turns_every_refusal_into_one_error_line(deal):
    bad = write_json(deal / "bad.json", {"forecast": {"FY+9": {}}})
    with pytest.raises(SystemExit) as exit_info:
        sf.main(["--company", COMPANY, "--proposal", PROPOSAL, "--forecast", str(bad), "--note", NOTE])
    assert str(exit_info.value).startswith("error: unknown forecast year") and "\n" not in str(exit_info.value)
    with pytest.raises(SystemExit) as missing:
        sf.main(["--company", COMPANY, "--proposal", PROPOSAL, "--forecast", str(deal / "none.json"), "--note", NOTE])
    assert str(missing.value).startswith("error: cannot read forecast file")
    broken = deal / "broken.json"
    broken.write_text("{nope", encoding="utf-8")
    with pytest.raises(SystemExit) as invalid:
        sf.main(["--company", COMPANY, "--proposal", PROPOSAL, "--forecast", str(broken), "--note", NOTE])
    assert "is not valid JSON" in str(invalid.value)
    assert not (deal / "deals").exists()


def test_main_reports_a_deal_that_cannot_take_it_as_an_error_line(deal):
    spreading_check.compute(COMPANY, PROPOSAL, {"FY-Current": {"revenue": 100}})
    forecast = write_json(deal / "f.json", {"forecast": FORECAST})
    with pytest.raises(SystemExit) as exit_info:
        sf.main(["--company", COMPANY, "--proposal", PROPOSAL, "--forecast", str(forecast), "--note", NOTE])
    assert str(exit_info.value).startswith("error: refused: this deal's historical figures were computed")


# ---------------------------------------------------------------------------
# The command text routes the analyst-supplied option here, opt-in only
# ---------------------------------------------------------------------------

def test_project_command_offers_the_analyst_supplied_option_explicitly_and_keeps_the_default():
    text = Path(sf.__file__).resolve().parents[1].joinpath(".claude", "commands", "project.md").read_text(
        encoding="utf-8")
    flat = " ".join(text.split())
    for needle in ("supplied_forecast.py", "analyst-supplied", "Ask which applies before assuming", "never infer",
                   "--note", "unavailable", "UNRESOLVABLE", "spreading_check.py", "--no-update-financials-source",
                   "one basis"):
        assert needle in flat, needle


# ---------------------------------------------------------------------------
# The documentation's worked example is the real thing
# ---------------------------------------------------------------------------

def test_the_workflows_page_quotes_what_the_synthetic_example_really_records(deal, capsys):
    repo = Path(sf.__file__).resolve().parents[1]
    examples = repo / "docs" / "examples" / "synthetic_co"
    page = " ".join((repo / "docs" / "workflows.md").read_text(encoding="utf-8").split())
    sf.main(["--company", COMPANY, "--proposal", "Synthetic Budget Loan", "--forecast",
             str(examples / "forecast_supplied.json"), "--stress-assumptions", str(examples / "forecast_stress.json"),
             "--note", "Management budget; depreciation is within cost of sales"])
    summary = json.loads(capsys.readouterr().out)
    basis, unavailable = summary["downside"]["basis"], summary["downside"]["unavailable"]
    assert basis == {"FY+1": sf.FRAMEWORK_DERIVED}
    assert f'"basis": {{"FY+1": "{basis["FY+1"]}"}}' in page
    assert f'"unavailable": {{"FY+2": "{unavailable["FY+2"]}"}}' in page
    saved = read_state(COMPANY, "Synthetic Budget Loan")
    assert saved["downside_case"]["financials"]["FY+1"]["ebitda"] == 200, "the page says stressed EBITDA is 200"
    assert saved["ratios"]["FY+1"] == {"dscr": 6.0, "gross_leverage": 0.67}
    write_state(COMPANY, "Synthetic Budget Loan", covenants=[{"metric": "dscr", "type": "minimum", "threshold": 1.25}])
    (breach,) = policy_compute(COMPANY, "Synthetic Budget Loan")["policy_state"]["downside_covenant_breaches"]
    assert breach["breach_id"] == "DOWNSIDE-FY-2-DSCR" and breach["downside_status"] == "UNRESOLVABLE"
    assert breach["reason"] in page and "DOWNSIDE-FY-2-DSCR" in page


# ---------------------------------------------------------------------------
# Correctness pass (issue #124): incomplete raw lines never count as known zeros
# ---------------------------------------------------------------------------

STRESS = {"revenue_haircut_pct": 10, "opex_increase_pct": 5}


@pytest.mark.parametrize("omitted", sorted(set(sf.REQUIRED_LINES) - set(sf.SHOCK_LINES["revenue_haircut_pct"])
                                           - set(sf.SHOCK_LINES["opex_increase_pct"])))
def test_a_line_a_covenant_ratio_reads_but_that_is_missing_makes_the_year_unavailable_even_if_the_rest_reconciles(
        deal, omitted):
    """The directly shocked lines (revenue, admin_expenses) and the supplied ratios are present and happen to
    reconcile, because the evaluator reads an omitted line as zero. That is not a known zero."""
    year = reconciled_year(drop=(omitted,))
    only = sf.reconciliation_problems(year["lines"], year["subtotals"], year["ratios"], STRESS)
    assert len(only) == 1 and omitted in only[0], "everything supplied does reconcile; only the omission is a problem"
    record(forecast={"FY+1": year}, stress=STRESS)
    case = state()["downside_case"]
    assert case["financials"] == {} and case["ratios"] == {} and case["basis"] == {}
    reason = case["unavailable"]["FY+1"]
    assert omitted in reason and "not supplied" in reason and "explicit 0" in reason
    saved = state()
    assert saved["financials"]["FY+1"] == year["subtotals"] and saved["ratios"]["FY+1"] == year["ratios"]
    assert saved["analyst_supplied_financials"]["FY+1"] == year["lines"], "the base case is exactly as supplied"


def test_an_omitted_line_is_unavailable_but_the_same_line_given_as_an_explicit_zero_is_accepted(deal):
    record(forecast={"FY+1": reconciled_year(drop=("cash",))}, stress=STRESS)
    assert "FY+1" in state()["downside_case"]["unavailable"]
    record(forecast={"FY+1": reconciled_year(lines={**LINES, "cash": 0})}, stress=STRESS)
    case = state()["downside_case"]
    assert case["basis"] == {"FY+1": sf.FRAMEWORK_DERIVED} and case["unavailable"] == {}


def test_an_incomplete_year_is_an_unresolvable_breach_where_the_base_case_passes(deal):
    record(forecast={"FY+1": reconciled_year(drop=("trade_debtors",))}, stress=STRESS)
    (breach,) = policy(deal)["downside_covenant_breaches"]
    assert (breach["year"], breach["downside_status"], breach["downside_actual"]) == ("FY+1", "UNRESOLVABLE", None)
    assert "downside analysis unavailable for FY+1" in breach["reason"] and "trade_debtors" in breach["reason"]


def test_the_derived_downside_carries_only_figures_whose_every_line_was_supplied(deal):
    record(forecast={"FY+1": reconciled_year()}, stress=STRESS)
    figures = state()["downside_case"]["financials"]["FY+1"]
    assert not UNSUPPORTED_OUTPUTS & set(figures), "no profit before tax, total assets, TNW... from absent lines"
    assert {"ebitda", "gross_profit", "operating_profit", "fcf", "total_debt", "total_equity"} <= set(figures)
    assert set(SUPPLIED_RATIOS) <= set(state()["downside_case"]["ratios"]["FY+1"])
    complete = {**LINES, "exceptional_costs": 0, "intangible_assets": 0, "other_fixed_assets": 0,
                "other_long_term_liabilities": 0, "provisions": 0}
    record(forecast={"FY+1": reconciled_year(lines=complete)}, stress=STRESS)
    assert set(figures) | UNSUPPORTED_OUTPUTS <= set(state()["downside_case"]["financials"]["FY+1"])


def test_the_dependency_tables_cover_every_output_and_name_every_line_it_reads():
    """Sound: a line missing from an output's entry never changes that output. Complete: every output is listed."""
    base = {name: 7.0 + 3 * i for i, name in enumerate(sf.RAW_FIELDS)}
    reference = evaluate_financial_model({"p": base})
    assert set(sf.SUBTOTAL_DEPENDS) == set(reference["financials"]["p"]) - {"raw"}
    assert set(sf.RATIO_DEPENDS) == set(reference["ratios"]["p"])
    for name in sf.RAW_FIELDS:
        bumped = evaluate_financial_model({"p": {**base, name: base[name] + 1000}})
        for store, depends in (("financials", sf.SUBTOTAL_DEPENDS), ("ratios", sf.RATIO_DEPENDS)):
            for output, lines in depends.items():
                if name not in lines:
                    assert bumped[store]["p"][output] == reference[store]["p"][output], (name, output)
    assert set(sf.REQUIRED_LINES) == {n for ratio in SUPPLIED_RATIOS for n in sf.RATIO_DEPENDS[ratio]}
    assert {n for lines in sf.RATIO_DEPENDS.values() for n in lines} <= set(sf.REQUIRED_LINES),         "a derived year holds every ratio, so no ratio may read a line that is not required"


def test_non_numeric_stored_figures_make_the_year_unavailable_instead_of_failing():
    problems = sf.reconciliation_problems({**LINES, "cash": "n/m"}, {"ebitda": 300}, {}, STRESS)
    assert problems and "cash" in problems[0] and "not a number" in problems[0]
    problems = sf.reconciliation_problems(LINES, {"ebitda": "n/m"}, {"dscr": "n/m"}, STRESS)
    assert any("ebitda" in p and "not a number" in p for p in problems)
    assert any("dscr" in p and "not a number" in p for p in problems)


# ---------------------------------------------------------------------------
# Correctness pass: a scenario's description stays with the years it describes
# ---------------------------------------------------------------------------

def scenario(description, *years):
    return {"description": description, "periods": {year: {"ratios": {"dscr": 1.1}} for year in years}}


def test_a_different_description_for_another_year_is_refused_while_an_earlier_analyst_year_would_keep_the_old_one(deal):
    record(downside=scenario("Scenario A", "FY+1"))
    before = deals_bytes(deal)
    with pytest.raises(sf.ForecastError) as refused:
        record(forecast={"FY+2": FORECAST["FY+2"]}, downside=scenario("Scenario B", "FY+2"))
    message = str(refused.value)
    assert "Scenario A" in message and "Scenario B" in message and "FY+1" in message and len(message.splitlines()) == 1
    assert deals_bytes(deal) == before, "refused with nothing written"
    assert state()["downside_case"]["description"] == "Scenario A"


def test_restating_every_affected_year_with_the_new_description_is_accepted(deal):
    record(downside=scenario("Scenario A", "FY+1"))
    record(downside=scenario("Scenario B", "FY+1", "FY+2"))
    case = state()["downside_case"]
    assert case["description"] == "Scenario B"
    assert case["basis"] == {"FY+1": "analyst-supplied", "FY+2": "analyst-supplied"}


def test_the_same_description_across_calls_and_years_is_a_normal_update(deal):
    record(downside=scenario("Scenario A", "FY+1"))
    record(downside=scenario("  Scenario A  ", "FY+2"))
    case = state()["downside_case"]
    assert case["description"] == "Scenario A" and set(case["basis"]) == {"FY+1", "FY+2"}
    record(downside=scenario("Scenario A", "FY+1", "FY+2"))
    assert state()["downside_case"]["description"] == "Scenario A"


def test_a_year_that_is_not_an_analyst_scenario_does_not_hold_a_description_back(deal):
    """Only years whose downside is the analyst's own carry a description; a framework-derived or unavailable year
    does not, so a new description for another year conflicts with nothing."""
    record(forecast={"FY+1": reconciled_year()}, stress=STRESS)
    record(forecast={"FY+2": FORECAST["FY+2"]}, downside=scenario("Scenario B", "FY+2"))
    case = state()["downside_case"]
    assert case["description"] == "Scenario B" and case["basis"]["FY+1"] == sf.FRAMEWORK_DERIVED


def test_a_record_with_one_description_keeps_its_shape_and_its_description(deal):
    """Existing records have a single `description`, which describes all of their analyst years."""
    record(downside=scenario("Scenario A", "FY+1"))
    assert state()["downside_case"]["description"] == "Scenario A" and "descriptions" not in state()["downside_case"]
    record(forecast={"FY+2": FORECAST["FY+2"]})
    assert state()["downside_case"]["description"] == "Scenario A"


# ---------------------------------------------------------------------------
# Correctness pass: failures before the state update are one clear error line
# ---------------------------------------------------------------------------

@pytest.mark.parametrize("failure", [TimeoutError("Could not acquire lock 'state.json.lock' within 30s"),
                                     OSError(28, "No space left on device")])
def test_a_failure_writing_state_is_a_state_error_that_keeps_its_cause_and_changes_nothing(deal, monkeypatch, failure):
    record()
    before = deals_bytes(deal)

    def refuse(*args, **kwargs):
        raise failure

    monkeypatch.setattr(sf, "write_state", refuse)
    with pytest.raises(StateError, match="could not write the deal's state") as raised:
        record(forecast={"FY+3": {"ratios": {"dscr": 1.6}}})
    assert raised.value.__cause__ is failure and len(str(raised.value).splitlines()) == 1
    assert deals_bytes(deal) == before


def test_a_failure_at_the_final_replace_leaves_state_and_the_folder_byte_identical(deal, monkeypatch):
    record()
    before = deals_bytes(deal)

    def refuse(*args, **kwargs):
        raise PermissionError(13, "Access is denied")

    with monkeypatch.context() as patched:
        patched.setattr("state_manager.os.replace", refuse)
        with pytest.raises(StateError, match="could not write the deal's state"):
            record(forecast={"FY+3": {"ratios": {"dscr": 1.6}}})
    assert deals_bytes(deal) == before, "no changed state.json, no leftover temporary or lock file"


def test_a_failure_reading_state_is_a_state_error_that_keeps_its_cause(deal, monkeypatch):
    failure = PermissionError(13, "Access is denied")

    def refuse(*args, **kwargs):
        raise failure

    monkeypatch.setattr(sf, "read_state", refuse)
    with pytest.raises(StateError, match="could not read the deal's state") as raised:
        record()
    assert raised.value.__cause__ is failure and not (deal / "deals").exists()


def test_main_reports_a_write_failure_as_one_error_line_and_exits(deal, monkeypatch):
    forecast_file = deal / "forecast.json"
    forecast_file.write_text(json.dumps({"forecast": FORECAST}), encoding="utf-8")

    def refuse(*args, **kwargs):
        raise TimeoutError("Could not acquire lock")

    monkeypatch.setattr(sf, "write_state", refuse)
    with pytest.raises(SystemExit) as exit_info:
        sf.main(["--company", COMPANY, "--proposal", PROPOSAL, "--forecast", str(forecast_file), "--note", NOTE])
    message = exit_info.value.code
    assert isinstance(message, str) and message.startswith("error: could not write the deal's state")
    assert "\n" not in message
    assert not list(deal.glob("deals/**/state.json")), "nothing was recorded"


def set_state_field(deal, name, value):
    path = state_file(deal)
    data = json.loads(path.read_text(encoding="utf-8"))
    data[name] = value
    path.write_text(json.dumps(data), encoding="utf-8")


@pytest.mark.parametrize("bad", ["ten", True, [10], {"pct": 10}, float("nan"), float("inf")])
def test_a_malformed_stored_shock_is_one_clear_error_naming_it_and_changes_nothing(deal, bad):
    record(forecast={"FY+1": reconciled_year()})
    set_state_field(deal, "stress_assumptions", {"revenue_haircut_pct": bad})
    before = deals_bytes(deal)
    with pytest.raises(sf.ForecastError) as refused:
        record(forecast={"FY+2": FORECAST["FY+2"]})
    message = str(refused.value)
    assert "stress_assumptions" in message and "revenue_haircut_pct" in message and "state.json" in message
    assert len(message.splitlines()) == 1 and deals_bytes(deal) == before


def test_a_stored_shock_that_is_null_or_an_unrelated_key_is_tolerated_as_the_calculation_tolerates_it(deal):
    record(forecast={"FY+1": reconciled_year()})
    set_state_field(deal, "stress_assumptions", {"revenue_haircut_pct": None, "comment": "from an older run"})
    record(forecast={"FY+2": FORECAST["FY+2"]})
    assert state()["stress_assumptions"]["comment"] == "from an older run"
