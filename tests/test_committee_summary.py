"""The committee summary (issue #196): what it prints, that it prints only recorded values, and what it writes."""
import json
import os
from pathlib import Path

import docx
import pytest

import committee_summary as cs
from source_manifest import save_source
from spreading_builder import evaluate_downside_case, evaluate_financial_model
from state_manager import StateError, state_path, write_state

COMPANY = "Synthetic Co"
PROPOSAL = "Synthetic Fleet Loan"
DATE = "2026-01-05"
FOLDER = Path("deals") / COMPANY / f"{PROPOSAL}_{DATE}"
MARKDOWN = FOLDER / f"{COMPANY}_{PROPOSAL}_Committee_Summary.md"
DOCX = FOLDER / f"{COMPANY}_{PROPOSAL}_Committee_Summary.docx"

RAW = {
    "revenue": 5000, "cost_of_sales": 3000, "admin_expenses": 1200, "interest_paid": 50,
    "current_debt": 0, "overdraft": 0, "long_term_debt": 2400, "loan_notes": 0, "scheduled_principal": 200,
}
GROUP_BASIS = {
    "basis": "supplied_consolidated", "reporting_entity": "Synthetic Group Ltd",
    "entities_covered": ["Synthetic Group Ltd", "Synthetic Group OpCo"],
    "source": {"document": "Borrower-supplied consolidated accounts", "accounting_basis": "UK GAAP",
               "date": "2026-08-31"},
    "period_end": "2026-03-31", "periods": ["FY-Current"], "figures_basis": "supplied",
    "perimeter_confirmed": True, "perimeter_confirmed_on": "2026-10-21",
}


def record(date_str=DATE, **fields):
    write_state(COMPANY, PROPOSAL, date_str=date_str, **fields)


def rich_fields():
    base = evaluate_financial_model({"FY-Current": RAW, "FY+1": RAW})
    stressed = evaluate_downside_case({"FY+1": RAW}, {"revenue_haircut_pct": 40})
    base_dscr = base["ratios"]["FY-Current"]["dscr"]
    stressed_dscr = stressed["ratios"]["FY+1"]["dscr"]
    assert stressed_dscr < base_dscr
    threshold = round((base_dscr + stressed_dscr) / 2, 2)
    return dict(
        deal_type="corporate_credit", inputs={"pd": "0.20%", "lgd": "LGD 3", "experian_score": 61},
        steps_completed=["triage", "spread", "collateral"], financials_source="framework-computed",
        financials=base["financials"], ratios=base["ratios"], downside_case=stressed,
        covenants=[{"metric": "dscr", "type": "minimum", "threshold": threshold},
                   {"metric": "gross_leverage", "type": "maximum", "threshold": 99}],
        guarantees=[{"provider": "Synthetic Parent Ltd", "type": "corporate", "amount": 1000}],
        collateral=[{"asset_id": "A1", "description": "Synthetic vehicle | fleet\nline two", "value": 100}],
        security_package=[{"secures_asset_id": "A1", "perfection_status": "Perfected", "ranking": "First"}],
        review_verdict="APPROVED",
        review_trail=[{"iteration": 1, "verdict": "REJECTED", "notes": "x", "timestamp": "2026-01-06"},
                      {"iteration": 2, "verdict": "APPROVED", "notes": None, "timestamp": "2026-01-07"}],
    )


@pytest.fixture
def deal(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record(**rich_fields())
    return tmp_path


def summary_text(**kwargs):
    cs.generate(COMPANY, PROPOSAL, **kwargs)
    return MARKDOWN.read_text(encoding="utf-8")


def folder_bytes():
    return {p.as_posix(): p.read_bytes() for p in sorted(Path("deals").rglob("*")) if p.is_file()}


# ---------------------------------------------------------------------------
# What it prints
# ---------------------------------------------------------------------------

def test_the_summary_has_every_section_in_order_and_the_word_file_mirrors_it(deal):
    report = cs.generate(COMPANY, PROPOSAL)

    assert report["status"] == "written" and Path(report["selected_folder"]) == FOLDER
    assert Path(report["markdown_path"]) == MARKDOWN and Path(report["docx_path"]) == DOCX
    assert report["included_sections"] == [
        "Facility and key terms", "Financial snapshot", "Covenants", "Conditions precedent and subsequent",
        "Collateral and security", "Open gaps", "Recorded review status", "Limits and items needing human review"]
    headings = [p.text for p in docx.Document(DOCX).paragraphs if p.style.name.startswith("Heading")]
    assert headings[0] == f"Committee summary: {COMPANY} / {PROPOSAL}"
    assert headings[1:] == report["included_sections"]


def test_recorded_terms_inputs_and_guarantees_are_quoted_and_unrecorded_terms_say_so(deal):
    text = summary_text()

    assert "- **Deal type:** corporate_credit" in text
    assert "- **pd:** 0.20%" in text and "- **experian_score:** 61" in text
    assert f"- **Facility amount, tenor and pricing:** {cs.NOT_RECORDED} in state.json" in text
    assert "| Synthetic Parent Ltd | corporate | 1000 |" in text
    assert "- **Financial basis:** framework-computed" in text


def test_ratios_and_subtotals_are_printed_exactly_as_recorded_with_n_a_for_null(deal):
    state = json.loads(Path(state_path(COMPANY, PROPOSAL, date_str=DATE)).read_text(encoding="utf-8"))
    text = summary_text()

    assert "| Ratio | FY-Current | FY+1 |" in text and "| Item | FY-Current | FY+1 |" in text
    for name, value in state["ratios"]["FY-Current"].items():
        expected = "N/A" if value is None else str(value)
        assert any(line.startswith(f"| {name} |") and expected in line for line in text.splitlines()), name
    assert "| current_ratio | N/A |" in text


def test_a_covenant_that_passes_now_and_fails_stressed_shows_both_results_with_the_policy_status(deal):
    text = summary_text()

    covenant_section = text.split("## Covenants\n\n", 1)[1].split("\n## Conditions precedent", 1)[0]
    assert "| metric | type | threshold | actual | status | headroom_pct | reason |" in covenant_section
    assert "| year | forward_id | metric | type | threshold | actual | status | headroom_pct | reason |" in covenant_section
    assert "| year | metric | base_actual | downside_actual | threshold | downside_status | reason |" in covenant_section
    assert all(" | basis | " not in line for line in covenant_section.splitlines() if line.startswith("|"))

    # Real policy-engine breach records currently omit basis; downside disclosure stays deferred
    # rather than synthesizing basis data in the summary (policy_engine.py is out of scope).
    _, collected = cs.collect(COMPANY, PROPOSAL)
    breaches = collected["policy"]["downside_covenant_breaches"]
    assert breaches
    assert all("basis" not in row for row in breaches)

    current = next(line for line in covenant_section.splitlines() if line.startswith("| dscr | minimum |"))
    assert "| PASS |" in current
    breach = next(line for line in covenant_section.splitlines() if line.startswith("| FY+1 | dscr |"))
    assert "FAIL" in breach
    assert "disclosure only" in text and "Downside breaches" in text


def test_an_unresolvable_covenant_is_listed_and_counted(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    fields = rich_fields()
    fields["ratios"] = {"FY-Current": {"dscr": None, "gross_leverage": 3.0}}
    record(**fields)

    text = summary_text()

    assert "UNRESOLVABLE" in text and "- **UNRESOLVABLE covenant results:** 1" in text


def test_conditions_collateral_and_security_gaps_come_from_the_policy_check(deal):
    text = summary_text()

    assert "- KYC-AML:" in text and "- FACILITY-EXECUTION:" in text
    assert "- GUARANTEE-SYNTHETIC-PARENT-LTD:" in text and "SEC-PERFECT-A1" not in text
    assert "- CS-COVENANT-COMPLIANCE:" in text
    assert "| A1 | Synthetic vehicle / fleet line two | 100 |" in text, "pipes and newlines must not break a table"
    assert "| A1 | Perfected | First |" in text
    assert "None reported." in text


def test_a_security_gap_in_the_state_is_reported(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    fields = rich_fields()
    fields["security_package"] = []
    record(**fields)

    text = summary_text()

    assert "**Security gaps reported by the policy check**" in text and "None reported." not in text


def test_open_gaps_name_the_missing_steps_and_the_hard_prerequisite(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record(steps_completed=["triage"])

    text = summary_text()

    assert "- **Required step missing:** spread" in text
    assert "- **Steps not recorded as completed:** spread, commercial, collateral, project, review, assemble" in text


def test_review_status_quotes_the_record_and_calls_it_historical(deal):
    text = summary_text()

    assert "- **Latest recorded verdict:** APPROVED (a historical record, not a current approval)" in text
    assert "- **Recorded review iterations:** 2; latest APPROVED at 2026-01-07" in text


def test_a_nearly_empty_deal_prints_only_not_recorded_statements(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record()

    text = summary_text()

    for line in ("No guarantees are recorded.", "No financial figures or ratios are recorded.",
                 "No covenants are recorded.", "No collateral is recorded.",
                 "No downside case is recorded; that is not evidence that downside risk is absent.",
                 "- **Deal type:** not recorded", "- **Latest recorded verdict:** not recorded"):
        assert line in text, line
    assert "- **Financial basis:** not recorded (read as framework-computed)" in text


def test_the_group_basis_section_appears_only_when_recorded_and_carries_the_disclosure(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record(financials_source="analyst-supplied")
    assert "Group accounts basis" not in summary_text()

    record(financials_source="analyst-supplied", group_basis=GROUP_BASIS)
    text = summary_text()

    assert "## Group accounts basis" in text
    assert cs.GROUP_DISCLOSURE in text
    assert "- **Reporting entity:** Synthetic Group Ltd" in text and "- **Common period end:** 2026-03-31" in text
    assert "statutory consolidation prepared by this framework" in text


def test_the_basis_note_and_forecast_basis_are_shown_when_recorded(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record(financials_source="analyst-supplied", forecast_source="analyst-supplied",
           financials_source_note="Depreciation sits inside cost of sales.")

    text = summary_text()

    assert "- **Financial basis:** analyst-supplied; forecast: analyst-supplied" in text
    assert "- **Basis note:** Depreciation sits inside cost of sales." in text


def test_the_text_contains_no_verdict_or_recommendation_wording(deal):
    text = summary_text().lower()

    for phrase in ("we recommend", "recommended", "should lend", "should decline", "approve this", "fit to lend"):
        assert phrase not in text
    assert "not a credit decision, recommendation or approval" in text


def test_saved_source_wording_distinguishes_unknown_missing_and_fine(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record(steps_completed=["triage"])
    assert ("citation declaration is unknown, so the saved-source check could not be completed"
            in summary_text())

    record(steps_completed=["triage"], triage={"verdict": "Go", "sources": ["https://example.invalid/a"]})
    assert "; declared citations have no saved source" in summary_text()

    source = tmp_path / "filing.txt"
    source.write_text("filing", encoding="utf-8")
    save_source(COMPANY, PROPOSAL, step="triage", claim="registry entry", source_path=str(source),
                url="https://example.invalid/a", date_str=DATE)
    text = summary_text()
    assert "- **Saved sources:** 1 in the manifest; citations declared: declared" in text
    assert "no saved source" not in text


# ---------------------------------------------------------------------------
# Which record
# ---------------------------------------------------------------------------

def test_the_named_date_selects_that_folder_and_the_default_is_the_newest(deal):
    record(date_str="2026-06-01", deal_type="another_type", steps_completed=["spread"])

    newest = cs.generate(COMPANY, PROPOSAL)
    older = cs.generate(COMPANY, PROPOSAL, date_str=DATE)

    assert Path(newest["selected_folder"]).name == f"{PROPOSAL}_2026-06-01"
    assert Path(older["selected_folder"]) == FOLDER
    assert "another_type" in Path(newest["markdown_path"]).read_text(encoding="utf-8")
    assert "corporate_credit" in MARKDOWN.read_text(encoding="utf-8")


def test_the_policy_check_runs_on_the_selected_state_not_the_newest(deal):
    record(date_str="2026-06-01", covenants=[], steps_completed=["spread"])

    cs.generate(COMPANY, PROPOSAL, date_str=DATE)

    assert "No covenants are recorded." not in MARKDOWN.read_text(encoding="utf-8")


# ---------------------------------------------------------------------------
# Refusals
# ---------------------------------------------------------------------------

def test_a_deal_with_no_state_or_a_named_date_without_state_is_refused_and_nothing_is_written(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    with pytest.raises(StateError, match="no existing state.json"):
        cs.generate(COMPANY, PROPOSAL)
    assert not Path("deals").exists()

    record()
    with pytest.raises(StateError, match="no state.json in the dated folder"):
        cs.generate(COMPANY, PROPOSAL, date_str="2025-12-31")
    with pytest.raises(ValueError, match="--date-str"):
        cs.generate(COMPANY, PROPOSAL, date_str="05-01-2026")
    assert not MARKDOWN.exists() and not DOCX.exists()


@pytest.mark.parametrize("fields, message", [
    ({"financials_source": 5}, "financials_source"),
    ({"covenants": "not a list"}, "covenants"),
    ({"group_basis": {"basis": "nonsense"}}, "group_basis"),
])
def test_malformed_state_is_refused_naming_the_key_and_nothing_is_written(tmp_path, monkeypatch, fields, message):
    monkeypatch.chdir(tmp_path)
    record(**fields)
    before = folder_bytes()

    with pytest.raises(StateError, match=message):
        cs.generate(COMPANY, PROPOSAL)

    assert folder_bytes() == before


def test_an_unusable_source_manifest_is_refused(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record()
    manifest = FOLDER / "sources" / "manifest.json"
    manifest.parent.mkdir()
    manifest.write_text("{not json", encoding="utf-8")

    with pytest.raises(StateError):
        cs.generate(COMPANY, PROPOSAL)

    assert not MARKDOWN.exists()


# ---------------------------------------------------------------------------
# What it writes
# ---------------------------------------------------------------------------

def test_only_the_two_report_files_are_created_and_the_record_is_untouched(deal):
    before = folder_bytes()

    cs.generate(COMPANY, PROPOSAL)

    after = folder_bytes()
    assert set(after) - set(before) == {MARKDOWN.as_posix(), DOCX.as_posix()}
    assert {k: v for k, v in after.items() if k in before} == before
    assert not list(FOLDER.glob(".committee_summary_*"))


def test_running_again_replaces_the_reports_with_identical_content(deal):
    cs.generate(COMPANY, PROPOSAL)
    first = MARKDOWN.read_bytes()

    cs.generate(COMPANY, PROPOSAL)

    assert MARKDOWN.read_bytes() == first and first.endswith(b"\n") and b"\r" not in first


def test_a_word_failure_is_reported_truthfully_and_leaves_no_temporary_file(deal, monkeypatch):
    def fail(markdown, path):
        raise OSError("disk full")

    monkeypatch.setattr(cs, "export_to_docx", fail)

    with pytest.raises(OSError, match="Markdown summary was written to .* but the Word summary failed: disk full"):
        cs.generate(COMPANY, PROPOSAL)

    assert MARKDOWN.exists() and not DOCX.exists()
    assert not list(FOLDER.glob(".committee_summary_*"))


def test_a_failed_markdown_write_leaves_no_partial_file(deal, monkeypatch):
    def broken_replace(source, destination):
        raise OSError("replace failed")

    monkeypatch.setattr(os, "replace", broken_replace)

    with pytest.raises(OSError, match="replace failed"):
        cs.generate(COMPANY, PROPOSAL)

    assert not MARKDOWN.exists() and not list(FOLDER.glob(".committee_summary_*"))


def test_rendering_is_pure_and_needs_no_files(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    data = {"company": COMPANY, "proposal": PROPOSAL, "date": DATE, "state": {},
            "policy": {"covenant_results": [], "forward_covenant_results": [], "downside_covenant_breaches": [],
                       "required_conditions_precedent": [], "required_conditions_subsequent": [],
                       "security_gaps": []},
            "manifest_count": 0, "sources_declared": "not_declared", "sources_missing": False}

    text = cs.render_markdown(data)

    assert text.startswith(f"# Committee summary: {COMPANY} / {PROPOSAL}") and not list(tmp_path.iterdir())
    assert "- **Saved sources:** 0 in the manifest; citations declared: not_declared" in text


# ---------------------------------------------------------------------------
# The command line, in process
# ---------------------------------------------------------------------------

def test_main_prints_the_report_as_json(deal, capsys):
    cs.main(["--company", COMPANY, "--proposal", PROPOSAL, "--date-str", DATE])

    report = json.loads(capsys.readouterr().out)
    assert report["status"] == "written" and Path(report["markdown_path"]) == MARKDOWN


def test_main_turns_a_refusal_into_one_error_line(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)

    with pytest.raises(SystemExit) as caught:
        cs.main(["--company", COMPANY, "--proposal", PROPOSAL])

    assert caught.value.code == f"error: no existing state.json for company '{COMPANY}' and proposal '{PROPOSAL}'"


def test_main_requires_company_and_proposal(capsys):
    with pytest.raises(SystemExit) as caught:
        cs.main(["--proposal", PROPOSAL])

    assert caught.value.code == 2 and "--company" in capsys.readouterr().err


def test_a_snapshot_with_only_subtotals_or_only_ratios_prints_just_that_table(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    record(financials={"FY-Current": {"ebitda": 800}})
    only_subtotals = summary_text()
    assert "**Recorded subtotals**" in only_subtotals and "**Recorded ratios**" not in only_subtotals

    record(financials={}, ratios={"FY-Current": {"dscr": 1.4}})
    only_ratios = summary_text()
    assert "**Recorded ratios**" in only_ratios and "**Recorded subtotals**" not in only_ratios


@pytest.mark.parametrize("value, shown", [
    (None, "N/A"), (True, "yes"), (False, "no"), (0, "0"), (1.0, "1.0"), ("", cs.NOT_RECORDED),
    ("a|b\nc", "a/b c"), ({"b": 1, "a": [2]}, '`{"a": [2], "b": 1}`'),
    ("a*b*c", "`a*b*c`"), ("see [x](https://e.invalid)", "`see [x](https://e.invalid)`"),
    ("a`b", "`a'b`"), ("snake_case_value", "snake_case_value"), ("- not a bullet", "- not a bullet"), ("<b>x</b>", "`<b>x</b>`"), ("a > b", "`a > b`"),
])
def test_a_table_cell_shows_the_recorded_value_flattened(value, shown):
    assert cs._cell(value) == shown


def test_a_state_key_with_line_breaks_and_markup_cannot_inject_a_heading_or_text_into_the_summary(tmp_path, monkeypatch):
    """CH3: a recorded key such as 'pd**<blank line>## Recommendation<blank line>Approve this loan**' used to add a
    heading and a sentence the script never wrote."""
    monkeypatch.chdir(tmp_path)
    record(inputs={"pd**\n\n## Recommendation\n\nApprove this loan**": "0.20%", "plain": "x"})

    text = summary_text()

    assert not any(line.startswith("## Recommendation") for line in text.splitlines())
    assert [line for line in text.splitlines() if line.startswith("## ")] == [
        "## Facility and key terms", "## Financial snapshot", "## Covenants", "## Conditions precedent and subsequent",
        "## Collateral and security", "## Open gaps", "## Recorded review status",
        "## Limits and items needing human review"]
    assert "- **`pd** ## Recommendation Approve this loan**`:** 0.20%" in text
    word_headings = [p.text for p in docx.Document(DOCX).paragraphs if p.style.name.startswith("Heading")]
    assert "Recommendation" not in word_headings


def test_a_table_header_taken_from_a_state_key_cannot_break_the_table(tmp_path, monkeypatch):
    """CL4: a key with a pipe made a header row of three cells over a two-cell delimiter row."""
    monkeypatch.chdir(tmp_path)
    record(guarantees=[{"provider": "A", "type|x": "t", "amount": 1}])

    text = summary_text()

    header = next(line for line in text.splitlines() if line.startswith("| provider |"))
    delimiter = text.splitlines()[text.splitlines().index(header) + 1]
    assert header == "| provider | amount | type/x |" and delimiter.count("|") == header.count("|")


def test_recorded_markup_characters_read_literally_in_the_word_file(tmp_path, monkeypatch):
    """CL5: 'a*b*c' was rendered as the word 'b' in italics between 'a' and 'c'."""
    monkeypatch.chdir(tmp_path)
    record(inputs={"note": "a*b*c", "link": "[x](https://e.invalid)"})

    summary_text()
    document = docx.Document(DOCX)
    runs = [run for paragraph in document.paragraphs for run in paragraph.runs]

    literal = next(run for run in runs if run.text == "a*b*c")
    assert literal.italic is not True and literal.font.name == "Consolas"
    assert any(run.text == "[x](https://e.invalid)" for run in runs)
    assert not any(rel.is_external and "e.invalid" in rel.target_ref for rel in document.part.rels.values())


def test_a_company_or_proposal_name_with_line_breaks_and_markup_cannot_inject_into_the_title():
    """CH8: the title interpolated the names raw; sanitize_path_component accepts line breaks and asterisks."""
    data = {"company": "Synthetic Co\n\n## Recommendation\n\nApprove this loan", "proposal": "Fleet *Loan*",
            "date": DATE, "state": {},
            "policy": {"covenant_results": [], "forward_covenant_results": [], "downside_covenant_breaches": [],
                       "required_conditions_precedent": [], "required_conditions_subsequent": [],
                       "security_gaps": []},
            "manifest_count": 0, "sources_declared": "not_declared", "sources_missing": False}

    text = cs.render_markdown(data)

    assert text.splitlines()[0] == (
        "# Committee summary: Synthetic Co ## Recommendation Approve this loan / `Fleet *Loan*`")
    assert [line for line in text.splitlines() if line.startswith("#")][:2] == [
        text.splitlines()[0], "## Facility and key terms"]
    assert not any(line.startswith("## Recommendation") for line in text.splitlines())


def test_covenant_basis_is_shown_only_when_a_row_carries_it_and_word_follows_markdown(tmp_path):
    """Basis is conditional, a row without it reads borrower, a recorded value is escaped, and Word matches Markdown."""
    data = {"company": "Synthetic Co", "proposal": "Fleet Loan",
            "date": DATE, "state": {},
            "policy": {"covenant_results": [
                           {"metric": "dscr", "basis": "consolidated", "type": "minimum", "threshold": 1.25,
                            "actual": None, "status": "UNRESOLVABLE", "headroom_pct": None,
                            "reason": "covenant basis differs from recorded view"},
                           {"metric": "gross_leverage", "type": "maximum", "threshold": 4.0,
                            "actual": 3.0, "status": "PASS", "headroom_pct": 0.25, "reason": None}],
                       "forward_covenant_results": [
                           {"year": "FY+1", "forward_id": "FORWARD-FY-1-DSCR", "metric": "dscr",
                            "basis": "guarantor_group", "type": "minimum", "threshold": 1.25, "actual": 1.4,
                            "status": "PASS", "headroom_pct": 0.12, "reason": None}],
                       "downside_covenant_breaches": [
                           {"year": "FY+1", "metric": "dscr", "basis": "guarantor*<group>", "base_actual": 1.4,
                            "downside_actual": 0.8, "threshold": 1.25, "downside_status": "FAIL", "reason": None}],
                       "required_conditions_precedent": [], "required_conditions_subsequent": [],
                       "security_gaps": []},
            "manifest_count": 0, "sources_declared": "not_declared", "sources_missing": False}

    text = cs.render_markdown(data)

    assert "| metric | basis | type | threshold | actual | status | headroom_pct | reason |" in text
    assert "| dscr | consolidated | minimum | 1.25 | N/A | UNRESOLVABLE | N/A | covenant basis differs from recorded view |" in text
    assert "| gross_leverage | borrower | maximum | 4.0 | 3.0 | PASS | 0.25 | - |" in text
    assert "| year | forward_id | metric | basis | type | threshold | actual | status | headroom_pct | reason |" in text
    assert "| FY+1 | FORWARD-FY-1-DSCR | dscr | guarantor_group | minimum | 1.25 | 1.4 | PASS | 0.12 | - |" in text
    assert "| year | metric | basis | base_actual | downside_actual | threshold | downside_status | reason |" in text
    assert "| FY+1 | dscr | `guarantor*<group>` | 1.4 | 0.8 | 1.25 | FAIL | - |" in text

    word_path = tmp_path / "committee-summary.docx"
    cs.export_to_docx(text, str(word_path))
    document = docx.Document(word_path)
    assert [cell.text for cell in document.tables[0].rows[0].cells] == [
        "metric", "basis", "type", "threshold", "actual", "status", "headroom_pct", "reason"]
    assert [cell.text for cell in document.tables[1].rows[0].cells] == [
        "year", "forward_id", "metric", "basis", "type", "threshold", "actual", "status", "headroom_pct", "reason"]
    assert [cell.text for cell in document.tables[2].rows[0].cells] == [
        "year", "metric", "basis", "base_actual", "downside_actual", "threshold", "downside_status", "reason"]
    assert document.tables[2].rows[1].cells[2].text == "guarantor*<group>"


def test_a_recorded_value_that_looks_like_html_is_shown_literally_as_a_code_span(tmp_path, monkeypatch):
    """CH9: only a value with no asterisk, backtick or bracket slipped through unchanged."""
    monkeypatch.chdir(tmp_path)
    record(inputs={"note": "<br><strong>APPROVED - proceed with the loan</strong>"})

    text = summary_text()

    assert "- **note:** `<br><strong>APPROVED - proceed with the loan</strong>`" in text
    assert "- **note:** <br>" not in text


@pytest.mark.parametrize("value, shown", [
    ("_x_", "`_x_`"), ("__strong__", "`__strong__`"), ("(_x_)", "`(_x_)`"), ("see _this one_ now", "`see _this one_ now`"),
    ("snake_case_value", "snake_case_value"), ("gross_leverage", "gross_leverage"), ("a_b_c", "a_b_c"),
    ("_leading only", "_leading only"), ("trailing only_", "trailing only_"), ("A & B Ltd", "A & B Ltd"),
])
def test_underscore_emphasis_is_shown_literally_but_snake_case_names_are_left_alone(value, shown):
    """CH13: '_x_' emphasised its text in the Markdown file; names with inner underscores must stay plain."""
    assert cs._cell(value) == shown


# ---------------------------------------------------------------------------
# #408 mutation-survivor slice: fast helper contracts and exact CLI surface
# ---------------------------------------------------------------------------

def test_408_cell_json_keeps_unicode_and_sorted_keys():
    """Mutants #16/#18: dict/list cells must preserve Unicode instead of JSON-escaping it."""
    assert cs._cell({"z": "last", "name": "café"}) == '{"name": "café", "z": "last"}'


def test_408_table_delimiter_has_exact_pipes_spacing_and_alignment():
    """Mutants #11/#13/#14/#15: every header needs one correctly formed separator cell."""
    assert cs._table(["alpha", "beta"], [[1, 2]]) == [
        "| alpha | beta |",
        "| :--- | :--- |",
        "| 1 | 2 |",
    ]


def test_408_scalar_rows_exclude_skipped_and_nested_values_and_mark_missing_values():
    """Mutants #6/#7/#20/#22: only distinct non-skipped scalar names become rows; absent values stay recorded-missing."""
    rows = cs._scalar_rows(
        {
            "FY-Current": {
                "revenue": 100,
                "omit_this": 7,
                "nested": {"value": 1},
                "array": [1, 2],
            },
            "FY+1": {"revenue": 110, "new_metric": 3},
        },
        ["FY-Current", "FY+1"],
        skip=("omit_this",),
    )
    assert rows == [
        ["revenue", 100, 110],
        ["new_metric", cs.NOT_RECORDED, 3],
    ]


def test_408_records_table_marks_a_missing_preferred_cell_as_not_recorded():
    """Mutants #14/#16: a missing value in a real column must not become None or an empty cell."""
    assert cs._records_table(
        [{"metric": "dscr", "threshold": 1.2}, {"metric": "leverage"}],
        ["metric", "threshold"],
    ) == [
        "| metric | threshold |",
        "| :--- | :--- |",
        "| dscr | 1.2 |",
        f"| leverage | {cs.NOT_RECORDED} |",
    ]


def test_408_collect_passes_names_date_base_dir_state_keys_and_path_to_readers(tmp_path, monkeypatch):
    """Pin every call boundary in collect without exporting either report file."""
    from pathlib import Path

    state_file = tmp_path / "Synthetic Co" / f"Synthetic Fleet Loan_{DATE}" / "state.json"
    state_file.parent.mkdir(parents=True)
    state_file.write_text("{}", encoding="utf-8")
    base_dir = str(tmp_path)
    calls = {"sanitize": [], "dates": [], "state_path": [], "read_state": [],
             "validate_state": [], "manifest": [], "find_deal": []}

    def sanitize(value, field_name):
        calls["sanitize"].append((value, field_name))
        return value

    def validate_date(value, field_name):
        calls["dates"].append((value, field_name))

    def path_for(company, proposal, **kwargs):
        calls["state_path"].append((company, proposal, kwargs.copy()))
        return str(state_file)

    def read_state(company, proposal, **kwargs):
        calls["read_state"].append((company, proposal, kwargs.copy()))
        return {"synthetic": "record"}

    def validate_state(state, **kwargs):
        calls["validate_state"].append((state, kwargs.copy()))

    def read_manifest(company, proposal, **kwargs):
        calls["manifest"].append((company, proposal, kwargs.copy()))
        return []

    def find_deal(company, proposal, **kwargs):
        calls["find_deal"].append((company, proposal, kwargs.copy()))
        return {"record_found": True, "date": DATE}

    monkeypatch.setattr(cs, "sanitize_path_component", sanitize)
    monkeypatch.setattr(cs, "validate_iso_date_str", validate_date)
    monkeypatch.setattr(cs, "state_path", path_for)
    monkeypatch.setattr(cs, "read_state", read_state)
    monkeypatch.setattr(cs, "validate_state", validate_state)
    monkeypatch.setattr(cs, "read_manifest", read_manifest)
    monkeypatch.setattr(cs, "find_deal", find_deal)
    monkeypatch.setattr(cs, "compute_from_state", lambda state: {"policy_state": {"from": state["synthetic"]}})
    monkeypatch.setattr(cs, "sources_declared", lambda state: "not_declared")
    monkeypatch.setattr(cs, "missing_saved_sources", lambda state, manifest: False)

    folder, data = cs.collect("Synthetic Co", "Synthetic Fleet Loan", date_str=DATE, base_dir=base_dir)

    assert folder == str(state_file.parent)
    assert data["date"] == DATE and data["policy"] == {"from": "record"}
    assert calls["sanitize"][:2] == [("Synthetic Co", "company"), ("Synthetic Fleet Loan", "proposal")]
    assert calls["dates"] == [(DATE, "--date-str")]
    assert calls["state_path"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
    ]
    assert calls["read_state"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir,
                                                   "keys": cs.STATE_KEYS_READ})
    ]
    assert calls["validate_state"] == [
        ({"synthetic": "record"}, {"keys": ("group_basis",), "path": str(state_file)})
    ]
    assert calls["manifest"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir})
    ]

    # The implicit-date branch must use find_deal's returned date, and the same base directory for every read.
    calls.update({"state_path": [], "read_state": [], "validate_state": [], "manifest": [], "find_deal": []})
    folder, data = cs.collect("Synthetic Co", "Synthetic Fleet Loan", base_dir=base_dir)
    assert data["date"] == DATE and folder == str(state_file.parent)
    assert calls["find_deal"] == [("Synthetic Co", "Synthetic Fleet Loan", {"base_dir": base_dir})]
    assert calls["state_path"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir}),
    ]
    assert calls["read_state"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir,
                                                   "keys": cs.STATE_KEYS_READ})
    ]
    assert calls["manifest"] == [
        ("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE, "base_dir": base_dir})
    ]


def test_408_atomic_replace_uses_a_same_directory_uniquely_named_temp_file(tmp_path, monkeypatch):
    """Mutants #2-#12: preserve destination directory and the chosen prefix/suffix."""
    import tempfile

    real_mkstemp = tempfile.mkstemp
    recorded = {}

    def recording_mkstemp(*args, **kwargs):
        recorded.update(kwargs)
        return real_mkstemp(*args, **kwargs)

    monkeypatch.setattr(cs.tempfile, "mkstemp", recording_mkstemp)
    target = tmp_path / "summary.md"

    def write(path):
        from pathlib import Path
        Path(path).write_text("synthetic payload", encoding="utf-8")

    cs._replace_atomically(str(target), write)

    assert target.read_text(encoding="utf-8") == "synthetic payload"
    assert recorded == {"dir": str(tmp_path), "prefix": ".committee_summary_", "suffix": ".tmp"}
    assert [p.name for p in tmp_path.iterdir()] == ["summary.md"]


def test_408_generate_keeps_utf8_lf_write_base_dir_and_policy_scope(tmp_path, monkeypatch):
    """Mutants in generate: exact collection kwargs, open encoding/newline, and report key/value."""
    import builtins

    data = {
        "company": "Synthetic Co", "proposal": "Synthetic Fleet Loan", "date": DATE, "state": {},
        "policy": {}, "manifest_count": 0, "sources_declared": "not_declared", "sources_missing": False,
    }
    collect_calls = []
    temp_writes = []
    open_calls = []
    real_open = builtins.open

    def collect(company, proposal, **kwargs):
        collect_calls.append((company, proposal, kwargs.copy()))
        return str(tmp_path), data

    def replace(path, write):
        temp = tmp_path / f"synthetic-{len(temp_writes)}.tmp"
        temp_writes.append((path, write))
        write(str(temp))

    def recording_open(*args, **kwargs):
        open_calls.append(kwargs.copy())
        return real_open(*args, **kwargs)

    monkeypatch.setattr(cs, "collect", collect)
    monkeypatch.setattr(cs, "render_markdown", lambda data: "## Synthetic section\n")
    monkeypatch.setattr(cs, "_replace_atomically", replace)
    monkeypatch.setattr(cs, "export_to_docx", lambda markdown, path: None)
    monkeypatch.setattr(builtins, "open", recording_open)

    report = cs.generate("Synthetic Co", "Synthetic Fleet Loan", date_str=DATE, base_dir=str(tmp_path))

    assert collect_calls == [("Synthetic Co", "Synthetic Fleet Loan",
                              {"date_str": DATE, "base_dir": str(tmp_path)})]
    assert open_calls == [{"encoding": "utf-8", "newline": "\n"}]
    assert report == {
        "status": "written",
        "selected_folder": str(tmp_path),
        "markdown_path": str(tmp_path / "Synthetic Co_Synthetic Fleet Loan_Committee_Summary.md"),
        "docx_path": str(tmp_path / "Synthetic Co_Synthetic Fleet Loan_Committee_Summary.docx"),
        "included_sections": ["Synthetic section"],
        "policy_scope": "deal-structural checks on the recorded state; draft-level findings not assessed",
    }


def test_408_main_help_and_required_proposal_are_exact(capsys, monkeypatch):
    """Pin help invocation and make --proposal genuinely required, without pinning prose."""
    with pytest.raises(SystemExit) as help_exit:
        cs.main(["--help"])
    assert help_exit.value.code == 0
    capsys.readouterr()  # Help wording is prose; only the CLI contract is pinned here.

    with pytest.raises(SystemExit) as required_exit:
        cs.main(["--company", "Synthetic Co"])
    assert required_exit.value.code == 2
    assert "the following arguments are required: --proposal" in capsys.readouterr().err


def test_408_main_passes_date_str_and_prints_two_space_json(capsys, monkeypatch):
    """Mutants #43/#46/#50/#52/#53: forward the selected date and preserve JSON indentation."""
    calls = []
    monkeypatch.setattr(cs, "generate", lambda company, proposal, **kwargs: (
        calls.append((company, proposal, kwargs.copy())) or {"status": "written"}
    ))
    cs.main(["--company", "Synthetic Co", "--proposal", "Synthetic Fleet Loan", "--date-str", DATE])
    assert calls == [("Synthetic Co", "Synthetic Fleet Loan", {"date_str": DATE})]
    assert capsys.readouterr().out == '{\n  "status": "written"\n}\n'
